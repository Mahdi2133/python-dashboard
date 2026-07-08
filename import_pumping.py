"""Import صورت‌وضعیت آزمایش پمپاژ از فایل‌های «total» پوشه‌ی «میلاد عزیز/pumping test».

سه قرارداد پمپاژ با سه ساختار متفاوت. هر فایل «total» یک جدول فهرست‌بها دارد
با ستون‌های سمت راست = سهم هر چاه. برای هر فایل یک سند «آزمایش پمپاژ» ساخته
می‌شود، آیتم‌ها وارد و سهم هر چاه در finance_item_allocations ثبت می‌شود.

فایل‌ها را می‌توانی با نام‌های استاندارد زیر در ریشه‌ی پروژه بگذاری، یا اگر
پوشه‌ی «میلاد عزیز/pumping test» سر جایش باشد، بدون تغییر نام هم پیدا می‌شوند.

    pump_10457.xlsx      ←  میلاد عزیز/pumping test/10457-pumping test/total.xlsx
    pump_beabsaz.xlsx    ←  میلاد عزیز/pumping test/به آب ساز پمپاژ/total صورتجلسات به آب ساز پیروز 2.xlsx
    pump_safarpour.xlsx  ←  میلاد عزیز/pumping test/پمپاژ صفر پور/قابل پرداخت صفر پور/total 1+2 -قابل پرداخت.xlsx

اجرا:  python import_pumping.py
"""
import os
import re

from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models.finance import FinanceStatement, FinanceItem, FinanceItemAllocation
from app.models.well import Well

OP_TYPE = "آزمایش پمپاژ"

# کانفیگ هر فایل: (نام‌های جایگزین برای پیدا کردن), شیت, ردیف سرستون,
# اندیس ستون‌ها (row_no/desc/unit/unit_price/qty/total/contract_qty),
# بازه‌ی ستون‌های چاه, شماره قرارداد, پیمانکار.
FILES = [
    {
        "paths": ["pump_10457.xlsx",
                  os.path.join("میلاد عزیز", "pumping test", "10457-pumping test", "total.xlsx")],
        "sheet": "1404159725 (2)", "hdr": 2,
        "row_no": 1, "desc": 2, "unit": 4, "unit_price": 6, "qty": 5,
        "total": 7, "contract_qty": 3, "wells": range(14, 23),
        "contract_no": "140510457", "contractor": "آب اندیشان دشت توس آسیا",
    },
    {
        "paths": ["pump_beabsaz.xlsx",
                  os.path.join("میلاد عزیز", "pumping test", "به آب ساز پمپاژ",
                               "total صورتجلسات به آب ساز پیروز 2.xlsx")],
        "sheet": "total", "hdr": 0,
        "row_no": 1, "desc": 2, "unit": 3, "unit_price": 4, "qty": 5,
        "total": 6, "contract_qty": None, "wells": range(7, 17),
        "contract_no": "پیروز", "contractor": "به آب ساز پیروز",
    },
    {
        "paths": ["pump_safarpour.xlsx",
                  os.path.join("میلاد عزیز", "pumping test", "پمپاژ صفر پور",
                               "قابل پرداخت صفر پور", "total 1+2 -قابل پرداخت.xlsx")],
        "sheet": "total", "hdr": 1,
        "row_no": 3, "desc": 4, "unit": 5, "unit_price": 6, "qty": 7,
        "total": 8, "contract_qty": None, "wells": range(9, 14),
        "contract_no": "1404711", "contractor": "سیدمحمد حسینی - حامد صفرپور",
    },
]

# نگاشت نام خام فایل → نام سامانه (نام خام همیشه در well_name ذخیره می‌شود)
ALIAS = {
    "اسلام 4": "اسلام اباد 4",
    "اسلام 5": "اسلام اباد 5",
    "باغون 2": "باغون آباد 2",
    "رضا 2": "امام رضا 2",
    "رضا 46": "امام رضا 46",
    "منزل 37": "منزل آباد 37",
    "شهرک راه آهن (مجدد)": "شهرک راه آهن",
    "فناوری 1 مجدد": "فناوری 1",
}

FA = "۰۱۲۳۴۵۶۷۸۹"
def to_eng(s):
    return "".join("0123456789"[FA.index(c)] if c in FA else c for c in str(s))

def match_key(v):
    s = to_eng("" if v is None else str(v)).replace("ي", "ی").replace("ك", "ک")
    s = " ".join(s.split()).strip()
    if not s:
        return ""
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ؤ", "و")
    return re.sub(r"[\s\u200c\u200f\u200e]+", "", s).lower()

def parse_num(v):
    try:
        s = to_eng(str(v)).strip()
        return float(s) if s and s not in ("-", "—", "") else None
    except (ValueError, TypeError):
        return None

def find_path(cfg):
    for p in cfg["paths"]:
        if os.path.exists(p):
            return p
    return None


def run():
    app = create_app()
    with app.app_context():
        db.create_all()   # اگر جدول تخصیص نبود، بساز

        well_by_key = {}
        for w in db.session.scalars(db.select(Well)).all():
            if w.name:
                well_by_key[match_key(w.name)] = w.id

        def resolve(raw):
            name = ALIAS.get(raw.strip(), raw.strip())
            return well_by_key.get(match_key(name))

        # پاک کردن import پمپاژِ قبلی
        old = db.session.query(FinanceStatement).filter_by(
            source="import", op_type=OP_TYPE).all()
        for st in old:
            db.session.delete(st)
        db.session.commit()
        if old:
            print(f"  {len(old)} سند پمپاژِ import قبلی پاک شد.")

        unmatched = set()
        gd = gi = ga = 0
        for cfg in FILES:
            path = find_path(cfg)
            if not path:
                print(f"  ⚠ فایل پیدا نشد (رد شد): {cfg['paths'][0]}")
                continue
            wb = load_workbook(path, read_only=True, data_only=True)
            if cfg["sheet"] not in wb.sheetnames:
                print(f"  ⚠ شیت «{cfg['sheet']}» در {path} نبود. شیت‌ها: {wb.sheetnames}")
                wb.close()
                continue
            ws = wb[cfg["sheet"]]
            rows = list(ws.iter_rows(values_only=True))
            hdr = rows[cfg["hdr"]]

            st = FinanceStatement(
                well_id=None, source="import", op_type=OP_TYPE,
                contractor=cfg["contractor"], contract_no=cfg["contract_no"],
                statement_no=cfg["contract_no"], statement_kind="قطعی",
                op_description=f"آزمایش پمپاژ — {cfg['contractor']}",
            )
            db.session.add(st)
            db.session.flush()

            ic = ac = 0
            total_sum = 0.0
            for i in range(cfg["hdr"] + 1, len(rows)):
                r = rows[i]
                def cell(k):
                    idx = cfg[k]
                    return r[idx] if (idx is not None and idx < len(r)) else None
                desc = cell("desc")
                row_no = cell("row_no")
                if not desc or not str(desc).strip():
                    continue
                d = str(desc).strip()
                if any(t in d for t in ("جمع", "مبلغ کل", "total", "مبلغ ناویژه")):
                    continue
                if parse_num(row_no) is None and (not row_no or not str(row_no).strip()):
                    continue
                total_amt = parse_num(cell("total"))
                qty = parse_num(cell("qty"))
                if not total_amt and not qty:
                    continue

                it = FinanceItem(
                    statement_id=st.id, category=OP_TYPE,
                    row_no=str(row_no).strip() if row_no else None,
                    description=d,
                    unit=str(cell("unit")).strip() if cell("unit") else None,
                    contract_qty=parse_num(cell("contract_qty")),
                    quantity=qty, unit_price=parse_num(cell("unit_price")),
                    coefficient=None, total_amount=total_amt,
                )
                db.session.add(it)
                db.session.flush()
                ic += 1
                if total_amt:
                    total_sum += total_amt

                for c in cfg["wells"]:
                    raw = hdr[c] if c < len(hdr) else None
                    if not raw or not str(raw).strip():
                        continue
                    raw = str(raw).strip()
                    q = parse_num(r[c]) if c < len(r) else None
                    if not q:
                        continue
                    wid = resolve(raw)
                    if wid is None:
                        unmatched.add(raw)
                    db.session.add(FinanceItemAllocation(
                        item_id=it.id, well_id=wid, well_name=raw, quantity=q))
                    ac += 1

            st.total_before = round(total_sum, 2)
            st.total_after = round(total_sum, 2)
            db.session.commit()
            wb.close()
            gd += 1; gi += ic; ga += ac
            print(f"  {cfg['contractor']}: {ic} آیتم | {ac} تخصیص | {total_sum:,.0f} ریال")

        print(f"\n✓ تمام شد: {gd} سند پمپاژ، {gi} آیتم، {ga} تخصیص چاه.")
        if unmatched:
            print("\n⚠ چاه‌های تطبیق‌نخورده (نامشان ذخیره شد):")
            for n in sorted(unmatched):
                print(f"    - {n}")


if __name__ == "__main__":
    run()