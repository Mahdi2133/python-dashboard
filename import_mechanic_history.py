"""Import تاریخی کشیدن/نصب — نسخه‌ی ۳: نگاشت کامل و دقیق همه‌ی فیلدها.

اصلاحات: نگاشت درست ستون‌های ۱۴۰۴/۱۴۰۵ (که یک ستون خطا داشتند)، خواندن موتور و
پمپ در همه‌ی سال‌ها (تیپ/سازنده/نو-تعمیری)، پیمانکار درست ۱۴۰۱-۱۴۰۲، نظر کارگاه،
عمق/سطوح/هد، آزمایش پمپاژ، راه‌انداز، کابل، ماه کارکرد و توضیحات.

اجرا:  python import_mechanic_history.py
"""
import os
import re

from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models.mechanic import MechanicEvent
from app.models.well import Well

XLS_PATH = "گزارش عملکرد کشیدن و نصب.xlsx"

FA = "۰۱۲۳۴۵۶۷۸۹"
def to_eng(s):
    return "".join("0123456789"[FA.index(c)] if c in FA else c for c in str(s))

def match_key(v):
    s = to_eng("" if v is None else str(v)).replace("ي", "ی").replace("ك", "ک")
    s = " ".join(s.split()).strip()
    if not s:
        return ""
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ؤ", "و")
    return re.sub(r"[\s\u200c\u200f\u200e]+", "", s).lower()

def parse_num(v):
    try:
        s = to_eng(str(v)).strip()
        if not s or s in ("-", "—", "#VALUE!"):
            return None
        return float(s)
    except (ValueError, TypeError):
        return None

def parse_int(v):
    n = parse_num(v)
    return int(n) if n is not None else None

def norm_year(y):
    y = int(y)
    if y < 100:
        return 1300 + y if y >= 50 else 1400 + y
    if y < 1000:
        return 1000 + y
    return y

MONTHS_FA = {"فروردین": 1, "اردیبهشت": 2, "خرداد": 3, "تیر": 4, "مرداد": 5,
             "شهریور": 6, "مهر": 7, "آبان": 8, "آذر": 9, "دی": 10, "بهمن": 11, "اسفند": 12}

def to_greg(y, mo, d):
    try:
        import jdatetime
        if 1350 < y < 1410 and 1 <= mo <= 12 and 1 <= d <= 31:
            return jdatetime.date(y, mo, d).togregorian()
    except Exception:
        pass
    return None

def parse_full_date(v):
    if not v:
        return None
    s = to_eng(str(v)).strip()
    m = re.match(r"(\d{2,4})[/\-.](\d{1,2})[/\-.](\d{1,2})", s)
    if not m:
        return None
    return to_greg(norm_year(m.group(1)), int(m.group(2)), int(m.group(3)))

def parse_dmy(day, mon, yr):
    try:
        if mon and str(mon).strip() in MONTHS_FA:
            mon = MONTHS_FA[str(mon).strip()]
        return to_greg(norm_year(to_eng(str(yr)).strip()),
                       int(to_eng(str(mon))), int(to_eng(str(day))))
    except (ValueError, TypeError):
        return None

# نگاشت کامل هر شیت (کلید: نام فیلد، مقدار: شماره ستون)
SHEETS = {
 "1398": {"start": 2, "date": ("full", 1), "name": 2, "fault": 4, "optype": 5, "contractor": 6, "pm": 7,
          "m_prev": 8, "m_maker": 9, "m_cond": 10, "p_prev": 11, "p_maker": 12, "p_cond": 13,
          "prev_date": 14, "notes": 19},
 "1399": {"start": 2, "date": ("full", 1), "name": 2, "fault": 4, "optype": 5, "contractor": 6,
          "m_prev": 7, "m_maker": 8, "m_cond": 9, "p_prev": 10, "p_maker": 11, "p_cond": 12,
          "prev_date": 13, "months": 18},
 "1400": {"start": 2, "date": ("full", 1), "name": 2, "fault": 4, "optype": 5, "contractor": 6,
          "m_prev": 7, "m_maker": 8, "m_cond": 9, "p_prev": 10, "p_maker": 11, "p_cond": 12,
          "prev_date": 13, "tip": 15, "pt_pr": 16, "design": 17, "pt_date": 18, "starter": 19,
          "months": 24, "notes": 25},
 "1401": {"start": 6, "date": ("full", 1), "name": 3, "fault": 5, "opinion": 6, "optype": 7, "pm": 8,
          "contractor": 14, "m_prev": 15, "m_curr": 16, "m_plate": 17, "m_maker": 18, "m_cond": 19,
          "p_prev": 20, "p_curr": 21, "p_plate": 22, "p_maker": 23, "p_cond": 24,
          "prev_date": 25, "depth": 27, "prev_ins": 28, "curr_ins": 29},
 "1402": {"start": 6, "date": ("full", 1), "name": 3, "fault": 5, "opinion": 6, "optype": 7, "pm": 8,
          "contractor": 14, "m_prev": 15, "m_curr": 16, "m_plate": 17, "m_maker": 18, "m_cond": 19,
          "p_prev": 20, "p_curr": 21, "p_stage": 22, "p_plate": 23, "p_maker": 24, "p_cond": 25,
          "prev_date": 26, "depth": 28, "prev_ins": 29, "curr_ins": 30},
 "1403": {"start": 6, "date": ("full", 1), "name": 3, "fault": 5, "opinion": 6, "optype": 7,
          "contractor": 8, "m_prev": 9, "m_curr": 10, "m_plate": 11, "m_maker": 12, "m_cond": 13,
          "p_prev": 14, "p_curr": 15, "p_stage": 16, "p_plate": 17, "p_maker": 18, "p_cond": 19,
          "prev_date": 20, "depth": 22, "prev_ins": 23, "curr_ins": 24, "static": 25, "dynamic": 26,
          "path": 27, "network": 28, "head": 29, "pt_date": 34},
 "1404": {"start": 2, "date": ("dmy", 1, 2, 3), "name": 4, "fault": 6, "optype": 7, "pm": 8,
          "contractor": 9, "m_prev": 10, "m_curr": 11, "m_plate": 12, "m_maker": 13, "m_cond": 14,
          "p_prev": 17, "p_curr": 18, "p_stage": 19, "p_plate": 20, "p_maker": 21, "p_cond": 22,
          "tip": 23, "prev_date": 24, "depth": 25, "prev_ins": 26, "curr_ins": 27, "static": 28,
          "dynamic": 29, "path": 30, "network": 31, "head": 32, "design": 33, "pipe": 34,
          "pt_date": 36, "pt_pr": 37, "pt_fl": 38, "starter": 41, "cable": 42, "months": 47,
          "opinion": 49, "notes": 50},
 "1405": {"start": 2, "date": ("dmy", 1, 2, 3), "name": 4, "fault": 7, "optype": 8, "pm": 9,
          "contractor": 10, "m_prev": 11, "m_curr": 12, "m_plate": 13, "m_maker": 14, "m_cond": 15,
          "p_prev": 18, "p_curr": 19, "p_stage": 20, "p_plate": 21, "p_maker": 22, "p_cond": 23,
          "tip": 24, "prev_date": 25, "depth": 26, "prev_ins": 27, "curr_ins": 28, "static": 29,
          "dynamic": 30, "path": 31, "network": 32, "head": 33, "design": 34, "pipe": 35,
          "pt_date": 37, "pt_pr": 38, "pt_fl": 39, "starter": 42, "cable": 43, "months": 48,
          "opinion": 50, "notes": 51},
}


def g(row, m, key):
    ci = m.get(key)
    if ci is None or ci >= len(row):
        return None
    v = row[ci]
    if v in (None, "", " "):
        return None
    s = str(v).strip()
    return s if s and s != "#VALUE!" else None


def build_desc(prev, curr, plate, maker, stage=None):
    """توضیح موتور/پمپ از اجزا: تیپ قبلی→فعلی، طبقه، پلاک، سازنده."""
    parts = []
    if curr and prev and str(curr) != str(prev):
        parts.append(f"تیپ {prev}→{curr}")
    elif curr or prev:
        parts.append(f"تیپ {curr or prev}")
    if stage:
        parts.append(f"طبقه {stage}")
    if plate:
        parts.append(f"پلاک {plate}")
    if maker:
        parts.append(str(maker))
    return " | ".join(parts) or None


def run():
    if not os.path.exists(XLS_PATH):
        raise SystemExit(f"فایل پیدا نشد: {XLS_PATH}")

    app = create_app()
    with app.app_context():
        wells = {}
        for w in db.session.scalars(db.select(Well)).all():
            if w.name:
                wells.setdefault(match_key(w.name), w.id)

        old = db.session.query(MechanicEvent).filter_by(source="import").delete()
        db.session.commit()
        if old:
            print(f"  {old} رکورد import قبلی پاک شد.")

        wb = load_workbook(XLS_PATH, read_only=True, data_only=True)
        total_added = 0
        total_skip = 0
        for sn, m in SHEETS.items():
            if sn not in wb.sheetnames:
                continue
            rows = list(wb[sn].iter_rows(values_only=True))
            added = 0
            skip = 0
            for row in rows[m["start"]:]:
                name = g(row, m, "name")
                if not name or not match_key(name):
                    continue
                wid = wells.get(match_key(name))
                if wid is None:
                    skip += 1
                    continue
                ds = m["date"]
                if ds[0] == "full":
                    op_date = parse_full_date(g(row, m, "date") if False else (row[ds[1]] if ds[1] < len(row) else None))
                else:
                    op_date = parse_dmy(row[ds[1]] if ds[1] < len(row) else None,
                                        row[ds[2]] if ds[2] < len(row) else None,
                                        row[ds[3]] if ds[3] < len(row) else None)

                rec = MechanicEvent(
                    well_id=wid, source="import", stage="پایان‌یافته",
                    op_date=op_date,
                    op_type=g(row, m, "optype"),
                    fault_description=g(row, m, "fault"),
                    contractor=g(row, m, "contractor"),
                    pm_form_no=g(row, m, "pm"),
                    motor_desc=build_desc(g(row, m, "m_prev"), g(row, m, "m_curr"),
                                          g(row, m, "m_plate"), g(row, m, "m_maker")),
                    motor_condition=g(row, m, "m_cond"),
                    pump_desc=build_desc(g(row, m, "p_prev"), g(row, m, "p_curr"),
                                         g(row, m, "p_plate"), g(row, m, "p_maker"),
                                         g(row, m, "p_stage")),
                    pump_condition=g(row, m, "p_cond"),
                    tip_change=g(row, m, "tip"),
                    prev_install_date=parse_full_date(g(row, m, "prev_date")),
                    well_depth=parse_num(g(row, m, "depth")),
                    prev_install_depth=parse_num(g(row, m, "prev_ins")),
                    curr_install_depth=parse_num(g(row, m, "curr_ins")),
                    static_level=parse_num(g(row, m, "static")),
                    dynamic_level=parse_num(g(row, m, "dynamic")),
                    path_loss=parse_num(g(row, m, "path")),
                    network_pressure=parse_num(g(row, m, "network")),
                    total_head=parse_num(g(row, m, "head")),
                    design_flow=parse_num(g(row, m, "design")),
                    pipe_diameter=parse_num(g(row, m, "pipe")),
                    pt_date=parse_full_date(g(row, m, "pt_date")),
                    pt_pressure=parse_num(g(row, m, "pt_pr")),
                    pt_flow=parse_num(g(row, m, "pt_fl")),
                    starter=g(row, m, "starter"),
                    cable_size=g(row, m, "cable"),
                    months_worked=parse_int(g(row, m, "months")),
                    mechanic_opinion=g(row, m, "opinion"),
                    notes=g(row, m, "notes"),
                )
                db.session.add(rec)
                added += 1
            db.session.commit()
            total_added += added
            total_skip += skip
            print(f"  سال {sn}: {added} رکورد، {skip} رد.")

        print(f"\n✓ تمام شد: {total_added} رکورد با نگاشت کامل وارد شد، {total_skip} رد.")


if __name__ == "__main__":
    run()