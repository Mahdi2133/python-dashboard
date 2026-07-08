"""Import صورت‌وضعیت مالی بهسازی از فایل اکسل (نسخه‌ی چند-چاهی).

ساختار فایل «صورت‌وضعیت_بهسازی.xlsx» (شیت 1404159725 (3)):

    یک قرارداد بهسازی (۱۴۰۴۱۵۹۷۲۵) با ۴ «وضعیت» (صورت‌وضعیت موقت پیاپی).
    هر «وضعیت» یک سند مالی است که چند چاه را پوشش می‌دهد.
    هر ردیف = یک ردیفِ فهرست‌بها (ویدئومتری، احیا، اضافه‌بها، ...).
    برای هر ردیف، مقدارِ هر وضعیت بین چاه‌های همان وضعیت تقسیم شده
    (ستون‌های سمت راست: هر ستون = یک چاه).

بنابراین این اسکریپت ۴ سند (وضعیت ۱ تا ۴) می‌سازد، برای هرکدام آیتم‌های
فهرست‌بها را با مقدار/مبلغِ همان وضعیت وارد می‌کند، و «سهم هر چاه» را در
جدول finance_item_allocations ثبت می‌کند.

پیش‌نیاز: جدول finance_item_allocations باید ساخته شده باشد. اگر نشده:
    python launcher.py db upgrade      (یا)  flask db upgrade

اجرا:  python import_finance.py
"""
import os
import re

from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models.finance import FinanceStatement, FinanceItem, FinanceItemAllocation
from app.models.well import Well

CANDIDATES = [
    os.path.join("میلاد عزیز", "159725-rehabilitation", "صورت‌وضعیت_بهسازی.xlsx"),
    "صورت‌وضعیت_بهسازی.xlsx",
]
SHEET = "1404159725 (3)"
OP_TYPE = "بهسازی"
CONTRACT_NO = "1404159725"

VAZIYAT = {
    "وضعیت ۱": {"qty_col": 9,  "amt_col": 10, "wells": range(19, 38), "video_col": 38},
    "وضعیت ۲": {"qty_col": 11, "amt_col": 12, "wells": range(39, 49), "video_col": 49},
    "وضعیت ۳": {"qty_col": 13, "amt_col": 14, "wells": range(50, 66), "video_col": 66},
    "وضعیت ۴": {"qty_col": 15, "amt_col": 16, "wells": range(67, 75), "video_col": 75},
}
ROW_TOTAL_BEFORE_OVERHEAD = 13
ROW_TOTAL_AFTER_OVERHEAD = 14
ROW_WORKSHOP = 15
ROW_VAZIYAT_TOTAL = 16
FIRST_ITEM_ROW = 2
LAST_ITEM_ROW = 12

ALIAS = {
    "شهرک راه آهن مجدد": "شهرک راه آهن",
    "درآبد 6 رزرو": "درآبد6",
    "رضا 32": "امام رضا 32",
}

FA = "۰۱۲۳۴۵۶۷۸۹"
def to_eng(s):
    return "".join("0123456789"[FA.index(c)] if c in FA else c for c in str(s))

def match_key(v):
    s = to_eng("" if v is None else str(v)).replace("ي", "ی").replace("ك", "ک")
    s = " ".join(s.split()).strip()
    if not s:
        return ""
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا").replace("ة", "ه").replace("ؤ", "و")
    return re.sub(r"[\s\u200c\u200f\u200e]+", "", s).lower()

def parse_num(v):
    try:
        s = to_eng(str(v)).strip()
        return float(s) if s and s not in ("-", "—", "") else None
    except (ValueError, TypeError):
        return None

def find_xls():
    for p in CANDIDATES:
        if os.path.exists(p):
            return p
    raise SystemExit("فایل «صورت‌وضعیت_بهسازی.xlsx» پیدا نشد.")

def run():
    xls = find_xls()
    print(f"فایل: {xls}")
    app = create_app()
    with app.app_context():
        # اگر جدول finance_item_allocations هنوز ساخته نشده (مهاجرت اجرا نشده)،
        # این خط فقط جدول‌های نبوده را می‌سازد و به جدول‌های موجود دست نمی‌زند.
        db.create_all()

        well_by_key = {}
        for w in db.session.scalars(db.select(Well)).all():
            if w.name:
                well_by_key[match_key(w.name)] = w.id

        def resolve_well(raw_name):
            name = ALIAS.get(raw_name.strip(), raw_name.strip())
            return well_by_key.get(match_key(name))

        old = db.session.query(FinanceStatement).filter_by(source="import").all()
        for st in old:
            db.session.delete(st)
        db.session.commit()
        if old:
            print(f"  {len(old)} سند import قبلی پاک شد.")

        wb = load_workbook(xls, read_only=True, data_only=True)
        if SHEET not in wb.sheetnames:
            raise SystemExit(f"شیت «{SHEET}» نیست. شیت‌ها: {wb.sheetnames}")
        ws = wb[SHEET]
        rows = list(ws.iter_rows(values_only=True))
        hdr = rows[1]

        unmatched = set()
        gd = gi = ga = 0
        for vname, cfg in VAZIYAT.items():
            qcol, acol = cfg["qty_col"], cfg["amt_col"]
            well_cols = list(cfg["wells"])
            total_before = parse_num(rows[ROW_TOTAL_BEFORE_OVERHEAD][acol])
            total_after_ovh = parse_num(rows[ROW_TOTAL_AFTER_OVERHEAD][acol])
            vaziyat_total = parse_num(rows[ROW_VAZIYAT_TOTAL][acol])
            workshop = parse_num(rows[ROW_WORKSHOP][acol]) if vname == "وضعیت ۲" else None
            coef_ovh = round(total_after_ovh / total_before, 4) if total_before else None

            st = FinanceStatement(
                well_id=None, source="import", op_type=OP_TYPE,
                contract_no=CONTRACT_NO, statement_no=CONTRACT_NO,
                statement_kind=vname,
                op_description=f"صورت‌وضعیت بهسازی قرارداد {CONTRACT_NO} — {vname}",
                coef_overhead=coef_ovh, workshop_setup=workshop,
                total_before=total_before,
                total_after=vaziyat_total if vaziyat_total else total_after_ovh,
            )
            db.session.add(st)
            db.session.flush()

            ic = ac = 0
            current_cat = ""
            coef_seen = None
            for i in range(FIRST_ITEM_ROW, LAST_ITEM_ROW + 1):
                r = rows[i]
                cat = str(r[0]).strip() if r[0] else ""
                if cat:
                    current_cat = cat
                row_no, desc = r[1], r[2]
                if not desc or not str(desc).strip() or not row_no:
                    continue
                qty = parse_num(r[qcol])
                amt = parse_num(r[acol])
                if not qty and not amt:
                    continue
                coef = parse_num(r[7])
                if coef and coef_seen is None:
                    coef_seen = coef
                it = FinanceItem(
                    statement_id=st.id, category=current_cat,
                    row_no=str(row_no), description=str(desc).strip(),
                    unit=str(r[4]).strip() if r[4] else None,
                    contract_qty=parse_num(r[3]), quantity=qty,
                    unit_price=parse_num(r[6]), coefficient=coef, total_amount=amt,
                )
                db.session.add(it)
                db.session.flush()
                ic += 1
                for c in well_cols:
                    raw = hdr[c]
                    if not raw or not str(raw).strip():
                        continue
                    raw = str(raw).strip()
                    if "ویدئومتری" in raw:
                        continue
                    q = parse_num(r[c]) if c < len(r) else None
                    if not q:
                        continue
                    wid = resolve_well(raw)
                    if wid is None:
                        unmatched.add(raw)
                    db.session.add(FinanceItemAllocation(
                        item_id=it.id, well_id=wid, well_name=raw, quantity=q))
                    ac += 1
            st.coef_contract = coef_seen
            db.session.commit()
            gd += 1; gi += ic; ga += ac
            print(f"  {vname}: {ic} آیتم | {ac} تخصیص | مبلغ: {(st.total_after or 0):,.0f} ریال")
        wb.close()
        print(f"\n✓ تمام شد: {gd} سند، {gi} آیتم، {ga} تخصیص چاه.")
        if unmatched:
            print("\n⚠ چاه‌های تطبیق‌نخورده (نامشان ذخیره شد):")
            for n in sorted(unmatched):
                print(f"    - {n}")

if __name__ == "__main__":
    run()