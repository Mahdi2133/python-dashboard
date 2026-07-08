"""به‌روزرسانی انبار قطعات کارگاه مکانیک از فایل «قطعات از سال ۹۸ تا ۰۵»."""
import os
from collections import defaultdict

from openpyxl import load_workbook
from sqlalchemy import text

from app import create_app
from app.extensions import db
from app.models.mechanic_part import MechanicPart

XLS = "parts_98_05.xlsx"
SHEET = "خلاصه عملکرد قطعات مورد استفاده"
EQUIP = {"الکتروموتور شناور": "موتور", "پمپ شناور": "پمپ"}

NEW_COLS = {
    "inventory_code": "VARCHAR(60)", "part_type": "VARCHAR(120)",
    "manufacturer": "VARCHAR(120)", "installed_count": "INTEGER",
    "collected_count": "INTEGER", "source": "VARCHAR(40)",
}


def ensure_columns():
    existing = {row[1] for row in db.session.execute(
        text("PRAGMA table_info(mechanic_parts)"))}
    for col, typ in NEW_COLS.items():
        if col not in existing:
            db.session.execute(text(
                f"ALTER TABLE mechanic_parts ADD COLUMN {col} {typ}"))
    db.session.commit()


def s(v):
    return str(v).strip() if v not in (None, "") else None


def run():
    if not os.path.exists(XLS):
        raise SystemExit(f"فایل پیدا نشد: {XLS} (باید در ریشه‌ی پروژه باشد)")
    app = create_app()
    with app.app_context():
        ensure_columns()
        wb = load_workbook(XLS, read_only=True, data_only=True)
        ws = wb[SHEET]
        it = ws.iter_rows(values_only=True)
        next(it)
        agg = defaultdict(lambda: dict(total=0, new=0, old=0, usable=0, scrap=0,
                                       inst=0, coll=0, code=None, ptype=None, maker=None))
        for r in it:
            equip = EQUIP.get(s(r[6]) or "", "نامشخص")
            name = s(r[30])
            if not name:
                continue
            a = agg[(equip, name)]
            a["total"] += 1
            st = s(r[26])
            if st == "نو":
                a["new"] += 1
            elif st == "کهنه":
                a["old"] += 1
            ru = s(r[27])
            if ru == "بلی":
                a["usable"] += 1
            elif ru == "خیر":
                a["scrap"] += 1
            act = s(r[22])
            if act == "نصب شد":
                a["inst"] += 1
            elif act == "جمع آوری شد":
                a["coll"] += 1
            if a["code"] is None:
                a["code"] = s(r[29])
            if a["ptype"] is None:
                a["ptype"] = s(r[23])
            if a["maker"] is None:
                a["maker"] = s(r[25])
        wb.close()

        old = db.session.query(MechanicPart).filter_by(source="import").delete()
        db.session.commit()
        if old:
            print(f"  {old} قطعه‌ی import قبلی پاک شد.")

        n = 0
        for (equip, name), a in agg.items():
            db.session.add(MechanicPart(
                equipment=equip, part_name=name, source="import",
                total_count=a["total"], new_count=a["new"], repaired=a["old"],
                usable=a["usable"], scrap=a["scrap"],
                installed_count=a["inst"], collected_count=a["coll"],
                inventory_code=a["code"], part_type=a["ptype"], manufacturer=a["maker"],
            ))
            n += 1
        db.session.commit()
        mo = sum(1 for k in agg if k[0] == "موتور")
        pu = sum(1 for k in agg if k[0] == "پمپ")
        print(f"\n✓ تمام شد: {n} قطعه وارد شد (موتور: {mo} | پمپ: {pu}).")


if __name__ == "__main__":
    run()