# -*- coding: utf-8 -*-
"""Export a report or the raw record table (requirements 21 & 22).

Excel  — openpyxl, sheet forced right-to-left, Persian headers, frozen header.
CSV    — UTF-8 **with BOM** so Excel on Windows opens Persian correctly.
JSON   — ensure_ascii off.
PDF    — reportlab with the bundled Vazirmatn TTF plus arabic-reshaper and
         python-bidi, because raw Persian in reportlab renders disconnected
         and left-to-right otherwise.
"""
from __future__ import annotations

import csv
import io
import json
import logging

from ..paths import resource_dir

log = logging.getLogger(__name__)

_PDF_FONT_REGISTERED = False


# ── Excel ────────────────────────────────────────────────────────────────────
def to_xlsx(columns, rows, sheet_title="گزارش", meta=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "گزارش")[:31]
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill("solid", fgColor="0A3D62")
    header_font = Font(name="Tahoma", bold=True, color="FFFFFF", size=10)
    body_font = Font(name="Tahoma", size=10)
    thin = Side(style="thin", color="D1D9E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_no = 1
    if meta:
        for key, value in meta.items():
            ws.cell(row=row_no, column=1, value=f"{key}: {value}").font = Font(
                name="Tahoma", size=9, italic=True)
            row_no += 1
        row_no += 1

    header_row = row_no
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=idx, value=col["label"])
        cell.fill, cell.font, cell.border, cell.alignment = (
            header_fill, header_font, border, center)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, col in enumerate(columns, start=1):
            value = row.get(col["key"])
            if isinstance(value, (list, tuple)):
                value = "، ".join(str(v) for v in value)
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font, cell.border, cell.alignment = body_font, border, center

    for idx, col in enumerate(columns, start=1):
        widest = max([len(str(col["label"]))]
                     + [len(str(row.get(col["key"], ""))) for row in rows[:400]] or [10])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(widest + 4, 12), 45)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── CSV / JSON ───────────────────────────────────────────────────────────────
def to_csv(columns, rows) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_ALL)
    writer.writerow([c["label"] for c in columns])
    for row in rows:
        writer.writerow([
            "، ".join(str(v) for v in row.get(c["key"]))
            if isinstance(row.get(c["key"]), (list, tuple))
            else ("" if row.get(c["key"]) is None else row.get(c["key"]))
            for c in columns])
    # BOM: without it Excel for Windows shows Persian as mojibake.
    return "﻿".encode("utf-8") + buf.getvalue().encode("utf-8")


def to_json(columns, rows, meta=None) -> bytes:
    payload = {"columns": columns, "rows": rows}
    if meta:
        payload["meta"] = meta
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


# ── PDF ──────────────────────────────────────────────────────────────────────
def _register_pdf_font():
    global _PDF_FONT_REGISTERED
    if _PDF_FONT_REGISTERED:
        return True
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    fonts_dir = resource_dir() / "app" / "static" / "fonts"
    regular = fonts_dir / "Vazirmatn-Regular.ttf"
    bold = fonts_dir / "Vazirmatn-Bold.ttf"
    if not regular.exists():
        log.error("Persian PDF font missing at %s", regular)
        return False
    pdfmetrics.registerFont(TTFont("Vazir", str(regular)))
    pdfmetrics.registerFont(TTFont("Vazir-Bold", str(bold if bold.exists() else regular)))
    pdfmetrics.registerFontFamily("Vazir", normal="Vazir", bold="Vazir-Bold")
    _PDF_FONT_REGISTERED = True
    return True


def shape_rtl(text) -> str:
    """Join Arabic letterforms and reorder for visual RTL output."""
    if text is None:
        return ""
    text = str(text)
    if not text.strip():
        return text
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        log.warning("RTL shaping unavailable; PDF text may render disconnected")
        return text


def to_pdf(columns, rows, title="گزارش", meta=None, landscape_mode=True) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)

    if not _register_pdf_font():
        raise RuntimeError("فونت فارسی برای تولید PDF یافت نشد.")

    page = landscape(A4) if landscape_mode else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, rightMargin=10 * mm, leftMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=title)

    title_style = ParagraphStyle("t", fontName="Vazir-Bold", fontSize=14, alignment=1,
                                 spaceAfter=6, leading=20)
    meta_style = ParagraphStyle("m", fontName="Vazir", fontSize=8, alignment=1,
                                textColor=colors.HexColor("#6b7f8e"), leading=12)

    story = [Paragraph(shape_rtl(title), title_style)]
    if meta:
        story.append(Paragraph(
            shape_rtl(" | ".join(f"{k}: {v}" for k, v in meta.items())), meta_style))
    story.append(Spacer(1, 6))

    # RTL tables read right-to-left, so the column order is reversed and the
    # header row is reversed with it.
    header = [shape_rtl(c["label"]) for c in columns][::-1]
    data = [header]
    for row in rows[:3000]:
        line = []
        for col in columns:
            value = row.get(col["key"])
            if isinstance(value, (list, tuple)):
                value = "، ".join(str(v) for v in value)
            text = "" if value is None else str(value)
            line.append(shape_rtl(text[:60]))
        data.append(line[::-1])

    usable = page[0] - 20 * mm
    col_width = usable / max(len(columns), 1)
    table = Table(data, colWidths=[col_width] * len(columns), repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Vazir"),
        ("FONTNAME", (0, 0), (-1, 0), "Vazir-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0a3d62")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d9e0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f6f9fc")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    if len(rows) > 3000:
        story.append(Spacer(1, 6))
        story.append(Paragraph(
            shape_rtl(f"نمایش ۳۰۰۰ سطر نخست از {len(rows)} سطر. "
                      "برای خروجی کامل از اکسل استفاده کنید."), meta_style))
    doc.build(story)
    return buf.getvalue()


FORMATS = {
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "csv": ("text/csv; charset=utf-8", "csv"),
    "json": ("application/json; charset=utf-8", "json"),
    "pdf": ("application/pdf", "pdf"),
}


def render(fmt, columns, rows, title="گزارش", meta=None):
    fmt = (fmt or "xlsx").lower()
    if fmt == "xlsx":
        return to_xlsx(columns, rows, title, meta), *FORMATS["xlsx"]
    if fmt == "csv":
        return to_csv(columns, rows), *FORMATS["csv"]
    if fmt == "json":
        return to_json(columns, rows, meta), *FORMATS["json"]
    if fmt == "pdf":
        return to_pdf(columns, rows, title, meta), *FORMATS["pdf"]
    raise ValueError(f"قالب خروجی پشتیبانی نمی‌شود: {fmt}")
