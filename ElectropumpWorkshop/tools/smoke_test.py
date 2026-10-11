# -*- coding: utf-8 -*-
"""End-to-end smoke test. Run with: python tools/smoke_test.py

Exercises every subsystem against a throwaway database so it can be run
before a release without touching instance/wells.db.
"""
import json
import os
import io
import shutil
import sys
import tempfile
from urllib.parse import quote

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

PASS, FAIL = [], []


def check(name, condition, detail=""):
    (PASS if condition else FAIL).append(name)
    print(f"  [{'OK ' if condition else 'FAIL'}] {name}" + (f"  {detail}" if detail else ""))


def _export_headers(client):
    """The column list the raw export produces, without downloading the file."""
    from app.routes.api_transfer import _export_columns
    return _export_columns()


def main():
    tmp = tempfile.mkdtemp(prefix="electropump_test_")
    db = os.path.join(tmp, "wells.db")
    from app import create_app
    app = create_app({"SQLALCHEMY_DATABASE_URI": f"sqlite:///{db}",
                      "WTF_CSRF_ENABLED": False})
    c = app.test_client()
    login = c.post("/api/login", json={"username": "admin", "password": "admin"})
    check("ورود مدیر پیش‌فرض", login.status_code == 200)

    print("\n— پایگاه داده و داده‌های مرجع —")
    sysinfo = c.get("/api/system").get_json()["data"]
    check("پایگاه داده ساخته و متصل شد", sysinfo["database"]["connected"])
    check("حالت WAL فعال است", sysinfo["database"]["journal_mode"] == "wal",
          sysinfo["database"]["journal_mode"])
    check("چاه‌ها از رفرنس درج شدند", sysinfo["counts"]["wells"] == 864,
          str(sysinfo["counts"]["wells"]))
    # Counted with a floor rather than an exact number: the admin adds fields
    # and options from the builders, and a release that seeds more of either
    # is not a regression. What matters is that the seed ran at all.
    check("گزینه‌های پایه درج شدند", sysinfo["counts"]["lookup_items"] >= 213,
          str(sysinfo["counts"]["lookup_items"]))
    check("فیلدهای فرم درج شدند", sysinfo["counts"]["form_fields"] >= 68,
          str(sysinfo["counts"]["form_fields"]))

    print("\n— تقویم شمسی —")
    from app.services.jalali import (gregorian_to_jalali, is_leap_jalali,
                                     jalali_month_days, jalali_to_gregorian,
                                     parse_jalali)
    check("۱۴۰۴/۰۱/۰۱ = 2025-03-21",
          str(jalali_to_gregorian(1404, 1, 1)) == "2025-03-21")
    check("۱۴۰۳ کبیسه است و ۳۰ اسفند دارد",
          is_leap_jalali(1403) and jalali_month_days(1403, 12) == 30)
    bad = sum(1 for jy in range(1380, 1440) for jm in range(1, 13)
              for jd in range(1, jalali_month_days(jy, jm) + 1)
              if gregorian_to_jalali(jalali_to_gregorian(jy, jm, jd)) != (jy, jm, jd))
    check("رفت‌وبرگشت ۱۳۸۰ تا ۱۴۳۹ بدون خطا", bad == 0, f"{bad} mismatch")
    check("خواندن '99/01/01' → ۱۳۹۹", parse_jalali("99/01/01") == (1399, 1, 1))
    check("خواندن '400/01/06' → ۱۴۰۰", parse_jalali("400/01/06") == (1400, 1, 6))
    check("خواندن ارقام فارسی", parse_jalali("۱۴۰۵/۰۶/۲۲") == (1405, 6, 22))

    print("\n— چرخه‌ی رکورد —")
    payload = {"op_jdate": "1405/06/22", "well": "امام رضا 11", "center": "سوران",
               "operation": "نصب", "contractor": "جوادی", "motor_curr": "73",
               "pump_curr": "384", "pump_stages": 10, "design_flow": 31,
               "test_flow": 29.5, "failure": ["شولات", "اهم دار"],
               "workshop_opinion": ["شولاتی"], "desc_tags": ["ورود شولات"],
               "description": "smoke test"}
    r = c.post("/api/records", json=payload)
    check("ایجاد رکورد", r.status_code == 200)
    rec = r.get_json()["data"]
    rid = rec["id"]
    check("تاریخ میلادی محاسبه شد", rec["op_date"] == "2026-09-13", rec["op_date"])
    check("چندانتخابی‌ها به جدول تگ رفتند", sorted(rec["failure"]) == sorted(["شولات", "اهم دار"]))
    check("تیپ پمپ به‌عنوان مقدار (نه شناسه) ذخیره شد", rec["pump_curr"] == "384",
          str(rec["pump_curr"]))
    r = c.put(f"/api/records/{rid}", json={"description": "ویرایش شد"})
    check("ویرایش جزئی بدون ارسال دوباره‌ی تاریخ", r.status_code == 200)
    check("مقدار ویرایش‌شده خوانده شد",
          c.get(f"/api/records/{rid}").get_json()["data"]["description"] == "ویرایش شد")
    r = c.post("/api/records", json={"well": ""})
    check("اعتبارسنجی فرم خالی", r.status_code == 422 and r.get_json()["fields"])
    c.delete(f"/api/records/{rid}")
    check("حذف نرم (رکورد باقی می‌ماند)",
          c.get(f"/api/records/{rid}").get_json()["data"]["is_active"] is False)
    c.post(f"/api/records/{rid}/restore")
    check("بازیابی رکورد",
          c.get(f"/api/records/{rid}").get_json()["data"]["is_active"] is True)

    print("\n— نرمال‌سازی املا (نام مستعار) —")
    r = c.post("/api/records", json={**payload, "contractor": "سعدابادی",
                                     "executor": "پیمانی",
                                     "center": "منزل اباد", "starter": "سفت"})
    d = r.get_json()["data"]
    check("«سعدابادی» → «سعدآبادی»", d["contractor"] == "سعدآبادی", d["contractor"])
    check("«منزل اباد» → «منزل آباد»", d["center"] == "منزل آباد", d["center"])
    check("«سفت» → «سافت»", d["starter"] == "سافت", str(d["starter"]))
    c.delete(f"/api/records/{d['id']}?hard=1")

    print("\n— ورود از اکسل —")
    xls = os.environ.get("ELECTROPUMP_TEST_XLS")
    if xls and os.path.exists(xls):
        from app.services.importer import commit_import, preview
        with app.app_context():
            info = preview(xls, "1405", limit=1)
            mapped = sum(1 for v in info["mapping"].values() if v)
            check("تشخیص سرستون شیت ۱۴۰۵", mapped >= 45, f"{mapped} ستون")
            res = commit_import(xls, "1405", info["mapping"], info["data_start"],
                                dry_run=True, filename="test.xls")
            check("بررسی آزمایشی بدون خطا", res["failed"] == 0,
                  f"ins={res['inserted']} fail={res['failed']}")
    else:
        print("  [skip] ELECTROPUMP_TEST_XLS تنظیم نشده")

    print("\n— گزارش‌ها —")
    from app.reports import REPORTS
    for key in REPORTS:
        r = c.get(f"/api/reports/{key}")
        check(f"گزارش {key}", r.status_code == 200 and r.get_json()["ok"])
    r = c.post("/api/reports/builder", json={
        "dataset": "records", "group_by": ["center"],
        "aggregations": [{"fn": "count"}, {"fn": "avg", "field": "test_flow"}]})
    check("گزارش‌ساز (گروه‌بندی)", r.status_code == 200 and r.get_json()["data"]["rows"])
    r = c.post("/api/reports/builder", json={
        "dataset": "records", "fields": ["well", "op_date", "operation"]})
    check("گزارش‌ساز (فهرست)", r.status_code == 200)
    r = c.post("/api/reports/builder", json={"dataset": "records",
                                             "group_by": ["'; DROP TABLE records;--"]})
    check("گزارش‌ساز فیلد ناشناخته را نادیده می‌گیرد", r.status_code == 200)
    check("جدول رکوردها سالم است", c.get("/api/records").status_code == 200)

    print("\n— خروجی‌ها —")
    for fmt, magic in (("xlsx", b"PK"), ("csv", b"\xef\xbb\xbf"), ("json", b"{"),
                       ("pdf", b"%PDF")):
        r = c.get(f"/api/export.{fmt}")
        check(f"خروجی خام {fmt}", r.status_code == 200 and r.data.startswith(magic),
              f"{len(r.data)}B")
    for fmt in ("xlsx", "csv", "json", "pdf"):
        r = c.get(f"/api/reports/overall/export.{fmt}")
        check(f"خروجی گزارش {fmt}", r.status_code == 200 and len(r.data) > 200)

    print("\n— RTL در PDF —")
    from app.services.exporter import shape_rtl
    check("حروف فارسی به شکل متصل درمی‌آیند", shape_rtl("شاخص") != "شاخص")
    check("ترتیب بصری راست‌به‌چپ می‌شود", shape_rtl("شاخص")[::-1][0] == "ﺷ")
    check("تاریخ با اسلش وارونه نمی‌شود", "1405/06/22" in shape_rtl("تاریخ: 1405/06/22"))

    print("\n— فرم‌ساز —")
    r = c.post("/api/form-builder/sections", json={"code": "t", "title": "تست"})
    sid = r.get_json()["data"]["id"]
    r = c.post("/api/form-builder/fields",
               json={"field_name": "t_num", "label": "عدد تست",
                     "field_type": "number", "section_id": sid})
    fid = r.get_json()["data"]["id"]
    check("افزودن بخش و فیلد", bool(sid and fid))
    r = c.post("/api/records", json={**payload, "dynamic": {"t_num": "42"}})
    drid = r.get_json()["data"]["id"]
    check("مقدار فیلد پویا ذخیره شد",
          c.get(f"/api/records/{drid}").get_json()["data"]["dynamic"].get("t_num") == 42.0)
    c.delete(f"/api/form-builder/sections/{sid}")
    check("بخش دارای داده حذف نمی‌شود (فقط غیرفعال)",
          c.get(f"/api/records/{drid}").get_json()["data"]["dynamic"].get("t_num") == 42.0)
    builtin = next(f for s in c.get("/api/form-builder?all=1").get_json()["data"]["sections"]
                   for f in s["fields"] if f["field_name"] == "well")
    c.put(f"/api/form-builder/fields/{builtin['id']}",
          json={"field_name": "hacked", "field_type": "text"})
    still = any(f["field_name"] == "well"
                for s in c.get("/api/form-builder?all=1").get_json()["data"]["sections"]
                for f in s["fields"])
    check("نام فنی فیلد پایه تغییرناپذیر است", still)
    r = c.post("/api/form-builder/fields",
               json={"field_name": "bad", "label": "x", "field_type": "nope",
                     "section_id": sid})
    check("نوع فیلد نامعتبر رد می‌شود", r.status_code == 422)

    print("\n— گزینه‌ها —")
    r = c.post("/api/lookups/starter", json={"value": "اینورتر"})
    oid = r.get_json()["data"]["id"]
    r = c.delete(f"/api/lookups/item/{oid}")
    check("گزینه‌ی بدون استفاده حذف می‌شود", "حذف شد" in r.get_json()["message"])
    used = next(i for i in c.get("/api/lookups/operation").get_json()["data"]["items"]
                if i["value"] == "نصب")
    r = c.delete(f"/api/lookups/item/{used['id']}")
    check("گزینه‌ی در حال استفاده فقط غیرفعال می‌شود",
          "غیرفعال" in r.get_json()["message"])
    c.put(f"/api/lookups/item/{used['id']}", json={"is_active": True})

    print("\n— چاه‌ها —")
    r = c.post("/api/wells", json={"name": "چاه تست الف"})
    w1 = r.get_json()["data"]["id"]
    r = c.post("/api/wells", json={"name": "چاه تست ب"})
    w2 = r.get_json()["data"]["id"]
    c.post("/api/records", json={**payload, "well": "چاه تست ب"})
    r = c.post(f"/api/wells/{w1}/merge", json={"source_id": w2})
    check("ادغام چاه تکراری رکوردها را منتقل می‌کند",
          r.status_code == 200 and "رکورد" in r.get_json()["message"],
          r.get_json().get("message", ""))

    print("\n— انتقال از localStorage —")
    r = c.post("/api/import/localstorage", json={"records": json.dumps([{
        "id": "1", "day": "5", "month": "3", "year": "1404", "wellName": "گلشهر یک",
        "center": "گلشهر", "operation": "کشیدن", "contractor": "سعدآبادی",
        "motorCurr": "55", "pumpCurr": "345",
        "failure": "شولات, اهم دار", "description": "برچسب الف | متن آزاد"}])})
    check("انتقال از localStorage", r.status_code == 200
          and r.get_json()["data"]["inserted"] == 1,
          json.dumps(r.get_json().get("data", {}), ensure_ascii=False)[:80])

    print("\n— پشتیبان و بازیابی —")
    from app.services.backup import create_backup, restore_backup, validate_backup
    with app.app_context():
        info = create_backup(note="smoke")
        check("تهیه پشتیبان", os.path.exists(info["path"]))
        check("پشتیبان تک‌فایل است (بدون wal/shm)",
              not os.path.exists(info["path"] + "-wal"))
        check("اعتبارسنجی پشتیبان", validate_backup(info["path"])["valid"])
        check("پشتیبان، پایگاه‌های جداگانه (انبار و بانک‌ها) را هم می‌گیرد",
              "warehouse.db" in info.get("side_databases", [])
              and "flowtest.db" in info.get("side_databases", []), str(info.get("side_databases")))
        _side_dir = os.path.join(os.path.dirname(info["path"]),
                                 "refdata_" + info["filename"][len("wells_backup_"):-3])
        bogus = os.path.join(tmp, "bogus.db")
        open(bogus, "wb").write(b"not a database")
        check("فایل نامعتبر رد می‌شود", not validate_backup(bogus)["valid"])
        before = c.get("/api/records?page_size=1").get_json()["total"]
        c.delete(f"/api/records/{rid}?hard=1")
        restore_backup(info["path"])
        check("بازیابی داده‌ها را برمی‌گرداند",
              c.get("/api/records?page_size=1").get_json()["total"] == before,
              f"{before} → {c.get('/api/records?page_size=1').get_json()['total']}")
        os.remove(info["path"])
        shutil.rmtree(_side_dir, ignore_errors=True)

    print("\n— گزارش تغییرات و صفحات —")
    check("Audit log ثبت شده", c.get("/api/audit").get_json()["total"] > 0)
    for path in ("/", "/entry", "/dashboard", "/records", "/wells", "/reports",
                 "/report-builder", "/form-builder", "/options", "/transfer",
                 "/users", "/settings", "/print/report/overall"):
        # "/" redirects to the first page the signed-in user may open.
        check(f"صفحه {path}",
              c.get(path, follow_redirects=True).status_code == 200)
    check("صفحه ۴۰۴ فارسی", "یافت نشد" in c.get("/nope").get_data(as_text=True))

    print("\n— ورود و سطح دسترسی —")
    r = c.post("/api/users", json={
        "username": "op_test", "password": "pass1234", "first_name": "کاربر",
        "last_name": "آزمایشی", "personnel_code": "T-1", "role": "operator"})
    check("ایجاد کاربر با نقش «ثبت اطلاعات»", r.status_code == 200)
    op_id = r.get_json()["data"]["id"]
    check("مجوزهای نقش اعمال شد",
          set(r.get_json()["data"]["permissions"]) ==
          {"record.create", "well.view", "report.view", "workflow.act"},
          str(sorted(r.get_json()["data"]["permissions"])))

    op = app.test_client()
    r = op.post("/api/login", json={"username": "op_test", "password": "pass1234"})
    check("ورود کاربر جدید", r.status_code == 200)
    check("کاربر می‌تواند رکورد ثبت کند",
          op.post("/api/records", json=payload).status_code == 200)
    check("کاربر نمی‌تواند ویرایش کند", op.put(f"/api/records/{rid}",
                                              json={"description": "x"}).status_code == 403)
    check("کاربر نمی‌تواند حذف کند",
          op.delete(f"/api/records/{rid}").status_code == 403)
    check("کاربر به مدیریت کاربران دسترسی ندارد",
          op.get("/api/users").status_code == 403)
    check("کاربر به تنظیمات دسترسی ندارد", op.get("/api/system").status_code == 403)
    check("تب‌های غیرمجاز باز نمی‌شوند", op.get("/settings").status_code == 403)
    check("تب مجاز باز می‌شود", op.get("/entry").status_code == 200)

    c.put(f"/api/users/{op_id}", json={"permissions": ["record.create"]})
    check("مجوز دستی جایگزین نقش می‌شود",
          op.get("/api/me").get_json()["data"]["permissions"] == ["record.create"])
    c.put(f"/api/users/{op_id}", json={"is_active": False})
    check("غیرفعال‌سازی بلافاصله اثر می‌کند", op.get("/api/records").status_code == 401)
    check("کاربر غیرفعال نمی‌تواند وارد شود",
          op.post("/api/login", json={"username": "op_test",
                                      "password": "pass1234"}).status_code == 401)

    activity = c.get(f"/api/users/{op_id}/activity").get_json()["data"]
    check("گزارش کارکرد: رکورد ثبت‌شده", activity["summary"]["records_created"] >= 1)
    check("گزارش کارکرد: نشست ثبت شده", len(activity["sessions"]) >= 1)
    check("گزارش کارکرد: اقدامات ثبت شده", len(activity["audits"]) >= 1)
    check("رمز نادرست پذیرفته نمی‌شود",
          app.test_client().post("/api/login",
                                 json={"username": "admin",
                                       "password": "x"}).status_code == 401)
    check("بدون ورود، API بسته است",
          app.test_client().get("/api/records").status_code == 401)

    print("\n— مهلت ویرایش رکورد —")
    # A recorder who may fix their own work, but only for a day.
    r = c.post("/api/users", json={
        "username": "win_op", "password": "window123", "role": "operator",
        "first_name": "مهلت", "last_name": "دار", "edit_window_hours": 24,
        "permissions": ["record.create", "record.view", "record.edit",
                        "record.delete", "well.view"]})
    win_id = r.get_json()["data"]["id"]
    check("مهلت ویرایش ذخیره می‌شود",
          r.get_json()["data"]["edit_window_hours"] == 24)
    check("برچسب فارسی مهلت", r.get_json()["data"]["edit_window_label"] == "۱ روز",
          r.get_json()["data"]["edit_window_label"])
    r = c.put(f"/api/users/{win_id}", json={"edit_window_hours": "abc"})
    check("مهلت نامعتبر رد می‌شود", r.status_code == 422)
    r = c.put(f"/api/users/{win_id}", json={"edit_window_hours": 99999})
    check("مهلت خارج از بازه رد می‌شود", r.status_code == 422)
    c.put(f"/api/users/{win_id}", json={"edit_window_hours": 24})

    c.post("/api/logout")
    c.post("/api/login", json={"username": "win_op", "password": "window123"})
    r = c.post("/api/records", json=payload)
    own = r.get_json()["data"]
    check("کاربر مهلت‌دار می‌تواند ثبت کند", r.status_code == 200)
    check("رکورد تازه قابل ویرایش است", own["can_edit"] is True, str(own["can_edit"]))
    check("تاریخ پایان مهلت برگردانده می‌شود", bool(own["edit_deadline_j"]),
          str(own["edit_deadline_j"]))
    r = c.put(f"/api/records/{own['id']}", json={"description": "داخل مهلت"})
    check("ویرایش داخل مهلت انجام می‌شود", r.status_code == 200)

    # Age the record past the window; the clock is the only thing that changed.
    with app.app_context():
        from app.models import Record
        from app.extensions import db as _db
        from app.services.jalali import local_now
        import datetime as _d
        rec = _db.session.get(Record, own["id"])
        rec.created_at = local_now() - _d.timedelta(hours=25)
        _db.session.commit()
    r = c.get(f"/api/records/{own['id']}")
    check("پس از مهلت، can_edit خاموش می‌شود",
          r.get_json()["data"]["can_edit"] is False)
    r = c.put(f"/api/records/{own['id']}", json={"description": "بعد از مهلت"})
    check("ویرایش پس از مهلت رد می‌شود", r.status_code == 403)
    check("پیام، کاربر را به مدیر ارجاع می‌دهد",
          "مدیر سیستم" in r.get_json()["error"], r.get_json()["error"])
    check("مقدار قبلی دست‌نخورده مانده",
          c.get(f"/api/records/{own['id']}").get_json()["data"]["description"]
          == "داخل مهلت")
    # Delete must be shut too, or the window is trivially sidestepped.
    r = c.delete(f"/api/records/{own['id']}")
    check("حذف پس از مهلت هم رد می‌شود", r.status_code == 403)
    r = c.post(f"/api/records/{own['id']}/restore")
    check("بازیابی پس از مهلت هم رد می‌شود", r.status_code == 403)

    c.post("/api/logout")
    c.post("/api/login", json={"username": "admin", "password": "admin"})
    r = c.put(f"/api/records/{own['id']}", json={"description": "مدیر آزاد است"})
    check("مدیر سیستم محدود نمی‌شود", r.status_code == 200)
    check("مدیر همیشه can_edit دارد",
          c.get(f"/api/records/{own['id']}").get_json()["data"]["can_edit"] is True)
    # Lifting the limit hands the record back without the admin touching it.
    c.put(f"/api/users/{win_id}", json={"edit_window_hours": None})
    c.post("/api/logout")
    c.post("/api/login", json={"username": "win_op", "password": "window123"})
    r = c.put(f"/api/records/{own['id']}", json={"description": "مهلت برداشته شد"})
    check("با برداشتن مهلت، کاربر دوباره می‌تواند ویرایش کند", r.status_code == 200)
    c.post("/api/logout")
    c.post("/api/login", json={"username": "admin", "password": "admin"})
    c.delete(f"/api/records/{own['id']}?hard=1")
    c.delete(f"/api/users/{win_id}")

    print("\n— کد PM و کلاسه چاه —")
    wells = c.get("/api/wells?limit=1&q=10/21/1").get_json()["data"]
    check("جستجوی چاه با کد PM", bool(wells) and wells[0]["pm_code"] == "10/21/1",
          str([w["name"] for w in wells]))
    check("نمایش ترکیبی نام و کد", "PM" in (wells[0]["display"] if wells else ""))
    exact = c.get("/api/wells?limit=3&q=" + quote("کورده 1")).get_json()["data"]
    check("تطابق دقیق در صدر نتایج", bool(exact) and exact[0]["name"] == "کورده 1",
          str([w["name"] for w in exact[:3]]))
    with app.app_context():
        from app.models import Well
        from app.services.lookups import fold_persian, well_key
        actives = Well.query.filter_by(is_active=True).all()
        folded = [fold_persian(w.name) for w in actives]
        keys = [well_key(w.name) for w in actives]
        codes = [w.pm_code for w in actives if w.pm_code]
        check("نام چاه تکراری وجود ندارد", len(folded) == len(set(folded)),
              f"{len(folded)-len(set(folded))} تکراری")
        from collections import Counter as _C
        _dupes = [k for k, n in _C(keys).items() if n > 1]
        check("نام چاه با ایندکس حرفی هم تکراری نیست", not _dupes,
              str([[w.name for w in actives if well_key(w.name) == k]
                   for k in _dupes]))
        check("کد PM تکراری وجود ندارد", len(codes) == len(set(codes)))
        check("چاه‌ها مرکز دارند",
              Well.query.filter(Well.center_id.isnot(None)).count() > 800)
        # The index may be written as a digit or as a word — one well, not two.
        check("«یک» و «1» یک چاه‌اند",
              well_key("ده غیبی یک") == well_key("ده غیبی 1"))
        check("«دو» و «2» یک چاه‌اند",
              well_key("مرکز تحقیقات دو") == well_key("مرکز تحقیقات 2"))
        check("ارقام فارسی هم یکی می‌شوند",
              well_key("الهیه ۳") == well_key("الهیه 3"))
        check("پرانتز و نقطه‌گذاری نادیده گرفته می‌شود",
              well_key("امرغان توس ( قدیم )") == well_key("امرغان توس قدیم"))
        check("ایندکس پیش از «قدیم» هم شناخته می‌شود",
              well_key("ابوطالب یک قدیم") == well_key("ابوطالب 1 قدیم"))
        # A number word that is part of the name must NOT be folded away.
        for a, b in (("ده سرخ", "10 سرخ"), ("چهار فصل", "4 فصل"),
                     ("سه راه دانش", "3 راه دانش"), ("جمال ده", "جمال 10")):
            check(f"«{a}» با «{b}» یکی نمی‌شود", well_key(a) != well_key(b))

    print("\n— کد PM با دو املا —")
    with app.app_context():
        from app.models import Well, Record
        from app.extensions import db as _db
        from app.services.seed import deduplicate_wells
        from app.services.lookups import pm_digits
        check("ارقام کد PM یکسان می‌شوند",
              pm_digits("10/24/41") == pm_digits("102441") == "102441")
        a = Well(name="چاه دوقلو", pm_code="10/99/7", well_class="900001",
                 is_active=True, is_verified=True)
        b = Well(name="چاه دوقلو (BOT)", pm_code="10997", well_class="900001",
                 is_active=True)
        # Same digits but a different کلاسه: a real second well, never merged.
        d = Well(name="چاه ناهمسان", pm_code="10997", well_class="900002",
                 is_active=True)
        _db.session.add_all([a, b, d])
        _db.session.commit()
        res = deduplicate_wells()
        check("دو املای یک کد PM ادغام می‌شوند",
              not Well.query.filter_by(name="چاه دوقلو (BOT)").one().is_active,
              str(res))
        check("املای رسمی «10/99/7» باقی می‌ماند",
              Well.query.filter_by(name="چاه دوقلو").one().is_active)
        check("املای قدیمی به‌عنوان نام مستعار می‌ماند",
              "چاه دوقلو (BOT)" in [al.alias for al in
                                    Well.query.filter_by(name="چاه دوقلو").one().aliases])
        check("کلاسه متفاوت مانع ادغام می‌شود",
              Well.query.filter_by(name="چاه ناهمسان").one().is_active)

    print("\n— «قدیم» و «جدید» —")
    with app.app_context():
        from app.models import Well
        from app.extensions import db as _db
        from app.services.seed import deduplicate_wells
        from app.services.lookups import qualifier_base_key
        check("پسوند «جدید» شناسایی می‌شود",
              qualifier_base_key("امامیه 17( جدید )") == qualifier_base_key("امامیه 17( قدیم )")
              != "", qualifier_base_key("امامیه 17( جدید )"))
        check("نام بدون پسوند، پایه ندارد", qualifier_base_key("امامیه 17") == "")
        base = Well.query.filter_by(name="امامیه 17").one()
        _db.session.add_all([
            Well(name="امامیه 17( جدید )", is_active=True, is_verified=False),
            Well(name="امامیه 17( قدیم )", is_active=True, is_verified=False),
            # Registered under a qualified name: that IS its name, keep it.
            Well(name="امامیه 90 جدید", pm_code="10/26/990", well_class="900090",
                 is_active=True, is_verified=True),
        ])
        _db.session.commit()
        deduplicate_wells()
        for variant in ("امامیه 17( جدید )", "امامیه 17( قدیم )"):
            check(f"«{variant}» حذف شد",
                  not Well.query.filter_by(name=variant).one().is_active)
        check("رکوردها به چاه اصلی رسیدند", Well.query.filter_by(name="امامیه 17")
              .one().is_active)
        check("نام‌های قبلی به‌عنوان نام مستعار ماندند",
              {"امامیه 17( جدید )", "امامیه 17( قدیم )"} <=
              {a.alias for a in Well.query.filter_by(name="امامیه 17").one().aliases})
        check("چاهِ ثبت‌شده با نام «جدید» دست‌نخورده می‌ماند",
              Well.query.filter_by(name="امامیه 90 جدید").one().is_active)

    print("\n— جستجوی چاه بدون تکرار —")
    with app.app_context():
        from app.models import Well
        from app.extensions import db as _db
        _db.session.add(Well(name="آماده سازی یک", is_active=True))
        _db.session.commit()
    hits = c.get("/api/wells?limit=20&q=" + quote("آماده سازی")).get_json()["data"]
    names = [w["name"] for w in hits]
    check("چاه هم‌نام دوباره در نتایج نمی‌آید",
          "آماده سازی یک" not in names and "آماده سازی 1" in names, str(names))
    r = c.post("/api/wells", json={"name": "آماده سازی 1"})
    check("افزودن چاه تکراری رد می‌شود", r.status_code == 409, r.get_json().get("error"))
    r = c.post("/api/wells", json={"name": "آماده سازی هفت"})
    check("چاه واقعاً جدید پذیرفته می‌شود", r.status_code == 200)

    print("\n— مرکز از روی چاه —")
    w = c.get("/api/wells?limit=1&q=" + quote("کورده 1")).get_json()["data"][0]
    check("چاه، مرکز خود را برمی‌گرداند", w["center_value"] == "سوران",
          str(w["center_value"]))

    print("\n— پیمانکار مشروط به مجری —")
    rules = c.get("/api/form-builder").get_json()["data"]["conditional"]
    check("قاعده نمایش پیمانکار تعریف شده",
          any(r.get("field") == "contractor" and r.get("on") == "executor"
              and r.get("value") == "پیمانی" for r in rules),
          str(rules)[:200])
    base = dict(payload)
    base.pop("dynamic", None)
    r = c.post("/api/records", json=dict(base, executor="امانی", contractor="جوادی"))
    check("با «امانی» پیمانکار ذخیره نمی‌شود",
          r.get_json()["data"]["contractor"] is None)
    c.delete(f"/api/records/{r.get_json()['data']['id']}?hard=1")
    r = c.post("/api/records", json=dict(base, executor="پیمانی", contractor="جوادی"))
    check("با «پیمانی» پیمانکار ذخیره می‌شود",
          r.get_json()["data"]["contractor"] == "جوادی")
    c.delete(f"/api/records/{r.get_json()['data']['id']}?hard=1")
    r = c.post("/api/records", json=dict(base, executor="امانی"))
    check("پیمانکار دیگر الزامی نیست", r.status_code == 200)
    if r.get_json().get("ok"):
        c.delete(f"/api/records/{r.get_json()['data']['id']}?hard=1")

    print("\n— ساعت محلی —")
    from app.services.jalali import (local_now, tehran_time_str, to_jalali_str,
                                     to_tehran)
    import datetime as _dt
    # Timestamps are stored on the wall clock, so what is written is what is
    # shown — no offset arithmetic that a mis-set Windows timezone can break.
    stamp = _dt.datetime(2026, 9, 13, 21, 15, 30)
    check("ساعت همان‌طور که ثبت شده نمایش داده می‌شود",
          tehran_time_str(stamp, with_seconds=False) == "21:15",
          tehran_time_str(stamp))
    check("تاریخ شمسی با ساعت محلی می‌خواند",
          to_jalali_str(stamp) == "1405/06/22", to_jalali_str(stamp))
    check("ساعت ثبت‌شده با ساعت سیستم یکی است",
          abs((local_now() - _dt.datetime.now()).total_seconds()) < 2)
    aware = stamp.replace(tzinfo=_dt.timezone.utc)
    check("مقدار دارای منطقه‌زمانی به وقت تهران می‌آید",
          to_tehran(aware).strftime("%H:%M") == "00:45",
          to_tehran(aware).strftime("%H:%M"))
    with app.app_context():
        cols = [col["label"] for col in _export_headers(c)]
    check("ستون «کد PM» در خروجی", "کد PM" in cols)
    check("ستون «کلاسه چاه» در خروجی", "کلاسه چاه" in cols)

    print("\n— ویرایش گزینه‌های فرم —")
    r = c.post("/api/form-builder/sections", json={"code": "opt_t", "title": "تست"})
    sid2 = r.get_json()["data"]["id"]
    r = c.post("/api/form-builder/fields",
               json={"field_name": "opt_pick", "label": "انتخاب", "field_type": "radio",
                     "section_id": sid2,
                     "options": [{"value": "الف", "label": "الف"},
                                 {"value": "ب", "label": "ب"}]})
    fid2 = r.get_json()["data"]["id"]
    r = c.put(f"/api/form-builder/fields/{fid2}",
              json={"options": [{"value": "الف", "label": "الف ویرایش‌شده"},
                                {"value": "ج", "label": "ج"}]})
    opts = {o["value"]: o for o in r.get_json()["data"]["own_options"]}
    check("ویرایش گزینه‌های فیلد سفارشی ذخیره می‌شود",
          opts.get("الف", {}).get("label") == "الف ویرایش‌شده")
    check("گزینه جدید افزوده می‌شود", "ج" in opts)
    check("گزینه حذف‌شده فقط غیرفعال می‌شود",
          "ب" in opts and opts["ب"]["is_active"] is False)

    builtin = next(f for s in c.get("/api/form-builder?all=1").get_json()["data"]["sections"]
                   for f in s["fields"] if f["field_name"] == "center")
    r = c.get(f"/api/form-builder/fields/{builtin['id']}/options").get_json()["data"]
    check("گزینه‌های فیلد پایه خوانده می‌شود", r["source"] == "lookup" and r["options"])
    edited = [{"value": o["value"], "label": o["label"], "is_active": o["is_active"]}
              for o in r["options"]]
    edited[0]["label"] = "سوران ✓"
    r = c.put(f"/api/form-builder/fields/{builtin['id']}/options",
              json={"options": edited})
    check("ویرایش گزینه‌های فیلد پایه ذخیره می‌شود", r.status_code == 200)
    after = c.get(f"/api/form-builder/fields/{builtin['id']}/options").get_json()["data"]
    check("برچسب جدید اعمال شد", after["options"][0]["label"] == "سوران ✓")

    print("\n— ستون‌های جدول و خروجی —")
    with app.app_context():
        labels = [col["label"] for col in _export_headers(c)]
    check("ستون «کد PM» در خروجی", "کد PM" in labels)
    check("ستون «کلاسه چاه» در خروجی", "کلاسه چاه" in labels)
    rec = c.get("/api/records?page_size=1").get_json()["data"]
    check("رکورد کد PM چاه را حمل می‌کند",
          bool(rec) and "well_pm_code" in rec[0])
    check("رکورد کلاسه چاه را حمل می‌کند", bool(rec) and "well_class" in rec[0])

    print("\n— تاریخ در حالت ویرایش —")
    r = c.post("/api/records", json={**payload, "prev_install_date": "1403/07/14",
                                     "test_date": "1405/06/01"})
    dated = r.get_json()["data"]
    check("تاریخ میلادی برای ماشین نگه داشته می‌شود",
          dated["prev_install_date"] == "2024-10-05", str(dated["prev_install_date"]))
    check("تاریخ شمسی برای فرم برگردانده می‌شود",
          dated["prev_install_date_j"] == "1403/07/14",
          str(dated.get("prev_install_date_j")))
    check("تاریخ آزمایش هم شمسی برمی‌گردد",
          dated["test_date_j"] == "1405/06/01", str(dated.get("test_date_j")))
    # A dynamic (admin-built) date field must round-trip in Jalali too, or
    # reopening the record shows 2024-10-05 where 1403/07/14 was typed.
    r = c.post("/api/form-builder/sections", json={"code": "dt_s", "title": "تاریخ"})
    dsid = r.get_json()["data"]["id"]
    r = c.post("/api/form-builder/fields",
               json={"field_name": "extra_date", "label": "تاریخ اضافه",
                     "field_type": "jalali_date", "section_id": dsid})
    check("فیلد تاریخ سفارشی ساخته شد", r.status_code == 200)
    r = c.put(f"/api/records/{dated['id']}",
              json={"dynamic": {"extra_date": "1404/02/03"}})
    back = c.get(f"/api/records/{dated['id']}").get_json()["data"]
    check("تاریخ فیلد سفارشی شمسی برمی‌گردد",
          back["dynamic"].get("extra_date") == "1404/02/03",
          str(back["dynamic"].get("extra_date")))
    c.delete(f"/api/records/{dated['id']}?hard=1")

    print("\n— فرایند: تعریف و مسیر —")
    wf = c.get("/api/workflow/definition").get_json()["data"]
    stages = {s["stage_number"]: s for s in wf["workflow"]["stages"]}
    check("شش مرحله تعریف شده", len(stages) == 6, str(sorted(stages)))
    check("مرحله ۱ فقط برای کشیدن", stages[1]["applies_to"] == "pull")
    check("مرحله ۳ برای هر دو عملیات", stages[3]["applies_to"] == "both")
    check("مرحله ۱ فقط «علت خرابی» را دارد",
          [i["code"] for i in stages[1]["items"] if i["kind"] == "field"]
          == ["failure"])
    check("مرحله ۵ سه بخش دارد", len(stages[5]["items"]) == 3)
    check("پالت شامل بخش‌ها و فیلدهاست",
          wf["palette"]["sections"] and wf["palette"]["fields"])

    # Real accounts, one per person, exactly as the workshop will have them.
    owners, made = {}, {}
    for uname, first, last in (("markaz", "مرکز", "آبرسانی"),
                               ("bozorg", "امین", "بزرگمهر"),
                               ("yaghouti", "مهدی", "یاقوتی‌نیا"),
                               ("kahani", "علی", "کاهانی")):
        rr = c.post("/api/users", json={
            "username": uname, "password": "process123", "role": "stage_owner",
            "first_name": first, "last_name": last})
        made[uname] = rr.get_json()["data"]
        owners[uname] = made[uname]["id"]
    check("نقش «متولی مرحله»: کارتابل، مشاهده رکوردها و خروجی",
          set(made["kahani"]["permissions"])
          == {"workflow.act", "record.view", "record.export"},
          str(sorted(made["kahani"]["permissions"])))
    for number, uname in ((0, "markaz"), (1, "markaz"), (2, "bozorg"),
                          (3, "yaghouti"), (4, "bozorg"), (5, "kahani")):
        rr = c.put(f"/api/workflow/stages/{stages[number]['id']}",
                   json={"assignee_id": owners[uname]})
        if number == 0:
            check("انتساب متولی به مرحله", rr.status_code == 200)

    def who(username):
        cc = app.test_client()
        cc.post("/api/login", json={"username": username, "password": "process123"})
        return cc

    markaz, bozorg, yaghouti, kahani = (who("markaz"), who("bozorg"),
                                        who("yaghouti"), who("kahani"))

    print("\n— فرایند: متولی بدون مجوز ثبت رکورد کار می‌کند —")
    # A stage owner must be able to run their stage with workflow.act alone.
    # Needing record.create for it would hand them the whole data-entry tab.
    for label, path in (("صفحه کارتابل", "/inbox"),
                        ("ساختار فرم", "/api/form-builder"),
                        ("فهرست گزینه‌ها", "/api/lookups"),
                        ("جستجوی چاه", "/api/wells?q=" + quote("امام")),
                        ("کارتابل", "/api/workflow/inbox")):
        check(f"متولی به {label} دسترسی دارد",
              markaz.get(path).status_code == 200, path)
    for label, path in (("جدول رکوردها", "/records"),
                        ("مستندات", "/documents")):
        check(f"متولی به {label} دسترسی دارد",
              markaz.get(path).status_code == 200, path)
    check("متولی خروجی اکسل می‌گیرد",
          markaz.get("/api/export.xlsx").status_code == 200)
    check("متولی خروجی CSV هم می‌گیرد",
          markaz.get("/api/export.csv").status_code == 200)
    for label, path in (("ثبت اطلاعات", "/entry"),
                        ("فرم‌ساز", "/form-builder"), ("کاربران", "/users"),
                        ("فرایندساز", "/workflow")):
        check(f"متولی به {label} دسترسی ندارد",
              markaz.get(path).status_code == 403, path)
    check("متولی نمی‌تواند رکورد را ویرایش کند",
          markaz.put("/api/records/1", json={"description": "x"}).status_code == 403)
    check("متولی نمی‌تواند رکورد مستقیم ثبت کند",
          markaz.post("/api/records", json={"well": "امام رضا 11"}).status_code == 403)
    check("متولی مرحله صفر می‌تواند فرایند را شروع کند",
          markaz.post("/api/workflow/instances",
                      json={"operation_kind": "کشیدن",
                            "well": "امام رضا 11"}).status_code == 200)
    check("و کارتابلش دکمه شروع را نشان می‌دهد",
          markaz.get("/api/workflow/inbox").get_json().get("may_start") is True)
    # Starting is step zero's job; the rest of the chain receives work.
    rr = kahani.post("/api/workflow/instances",
                     json={"operation_kind": "کشیدن", "well": "امام رضا 11"})
    check("متولی مرحله‌های بعد نمی‌تواند فرایند شروع کند", rr.status_code == 422,
          str(rr.get_json().get("error"))[:60])
    check("و دکمه شروع برایش پنهان است",
          kahani.get("/api/workflow/inbox").get_json().get("may_start") is False)
    rr = markaz.post("/api/workflow/instances", json={"operation_kind": "کشیدن"})
    check("بدون نام چاه فرایند شروع نمی‌شود", rr.status_code == 422,
          str(rr.get_json().get("error"))[:50])
    rr = markaz.post("/api/workflow/instances",
                     json={"operation_kind": "کشیدن", "well": "چاه ناموجود ۹۹"})
    check("چاه خارج از فهرست پذیرفته نمی‌شود", rr.status_code == 422)

    print("\n— فرایند: هر عملیات از مرحله‌ی خودش آغاز می‌شود —")
    # کشیدن opens at مرکز آبرسانی's stage, نصب at کارگاه مکانیک's. Neither is
    # a rule in the code any more: both are the ``can_start`` flag on a stage,
    # and «شامل» says which operation each door admits.
    rr = markaz.post("/api/workflow/instances",
                     json={"operation_kind": "نصب", "well": "کورده 1"})
    check("مرکز آبرسانی نمی‌تواند «نصب» را شروع کند", rr.status_code == 422,
          str(rr.get_json().get("error"))[:70])
    check("و کارتابلش فقط «کشیدن» را پیشنهاد می‌دهد",
          [k["value"] for k in
           markaz.get("/api/workflow/inbox").get_json().get("startable", [])]
          == ["pull"])
    check("کارگاه مکانیک فقط «نصب» را می‌تواند شروع کند",
          [k["value"] for k in
           yaghouti.get("/api/workflow/inbox").get_json().get("startable", [])]
          == ["install"])
    # What the کارتابل itself posts. It builds its radios from the list the
    # inbox hands back, so both forms of the answer have to be accepted — a
    # release that took only the label left the start button dead.
    for kind in ("نصب", "install"):
        rr = yaghouti.post("/api/workflow/instances",
                           json={"operation_kind": kind, "well": "کورده 1"})
        check(f"شروع فرایند با «{kind}» پذیرفته می‌شود", rr.status_code == 200,
              str(rr.get_json().get("error") or "")[:60])
    iid = yaghouti.post("/api/workflow/instances",
                        json={"operation_kind": "نصب", "well": "کورده 1"}
                        ).get_json()["data"]["id"]
    det = yaghouti.get(f"/api/workflow/instances/{iid}").get_json()["data"]
    check("نصب مستقیم به مرحله ۳ می‌رود", det["current_stage"] == 3,
          str(det["current_stage"]))
    check("مرحله‌های ۱ و ۲ طی نمی‌شوند",
          [e["status"] for e in det["entries"]
           if e["stage_number"] in (1, 2)] == ["skipped", "skipped"])
    check("در نصب، «علت خرابی» پرسیده نمی‌شود",
          "failure" not in [f["field_name"] for s in det["form"]["sections"]
                            for f in s["fields"]])
    check("در نصب، «نصب مرتبط با…» پرسیده می‌شود",
          "install_relates_to" in [f["field_name"] for s in det["form"]["sections"]
                                   for f in s["fields"]])

    print("\n— فرایند: مرحله‌ها به هم وابسته نیستند —")
    pid = markaz.post("/api/workflow/instances",
                      json={"operation_kind": "کشیدن", "well": "امام رضا 11"}
                      ).get_json()["data"]["id"]
    box = yaghouti.get("/api/workflow/inbox").get_json()["data"]
    check("مرحله ۳ بی‌درنگ در کارتابل می‌آید",
          any(b["id"] == pid and b["stage_number"] == 3 for b in box))
    check("بزرگمهر هر دو مرحله ۲ و ۴ را می‌بیند",
          {b["stage_number"] for b in bozorg.get("/api/workflow/inbox")
           .get_json()["data"] if b["id"] == pid} == {2, 4})
    det = yaghouti.get(f"/api/workflow/instances/{pid}?stage=3").get_json()["data"]
    check("هشدار مرحله‌های ثبت‌نشده داده می‌شود",
          {w["stage_number"] for w in det["waiting_on"]} == {1, 2},
          str(det["waiting_on"]))
    check("ولی مرحله ۳ مسدود نیست", det["may_act"] is True)
    rr = yaghouti.post(f"/api/workflow/instances/{pid}/submit", json={
        "stage_number": 3,
        "data": {"op_jdate": "1405/06/22", "well": "امام رضا 11",
                 "center": "سوران", "operation": "کشیدن",
                 "motor_curr": "73", "pump_curr": "384"}})
    check("مرحله ۳ پیش از مرحله ۲ ثبت می‌شود", rr.status_code == 200,
          str(rr.get_json().get("error")))
    det = markaz.get(f"/api/workflow/instances/{pid}?stage=1").get_json()["data"]
    # «علت خرابی» is still owed, and the cause forms travel with it (closed
    # until a cause is ticked); «اطلاعات پایه» must not come back.
    codes = [s["code"] for s in det["form"]["sections"]]
    check("«اطلاعات پایه» دوباره پرسیده نمی‌شود",
          codes[0] == "field_failure" and "basic" not in codes
          and all(c.startswith("fail_") for c in codes[1:]),
          str(codes[:3]))
    rr = kahani.post(f"/api/workflow/instances/{pid}/submit",
                     json={"stage_number": 3, "data": {}})
    check("متولی دیگری نمی‌تواند مرحله را ثبت کند", rr.status_code == 422)

    print("\n— فرایند: دیدن کار مرحله‌های قبل و قفل بودن چاه —")
    seen = markaz.post("/api/workflow/instances",
                       json={"operation_kind": "کشیدن", "well": "امام رضا 11"}
                       ).get_json()["data"]["id"]
    det = markaz.get(f"/api/workflow/instances/{seen}?stage=1").get_json()["data"]
    wells = [f for s in det["form"]["sections"] for f in s["fields"]
             if f["field_name"] == "well"]
    check("چاه پس از مرحله صفر دوباره جستجو نمی‌شود",
          bool(wells) and wells[0].get("read_only") is True,
          str(wells and wells[0].get("read_only")))
    check("و مقدارش نمایش داده می‌شود",
          wells[0].get("read_only_value") == "امام رضا 11",
          str(wells[0].get("read_only_value")))
    # Two causes ticked means both of their reading forms are owed.
    rr = markaz.post(f"/api/workflow/instances/{seen}/submit", json={
        "stage_number": 1, "data": {
            "op_jdate": "1405/06/22", "center": "سوران",
            "failure": ["شولات", "اهم دار"],
            "fail_sanding_p01": "تخلیه شده", "fail_sanding_p02": "انجام شده",
            "fail_sanding_p03": "نشده", "fail_sanding_p04": "دارد",
            "fail_sanding_p05": 12, "fail_sanding_p06": 30,
            "fail_ohmdrop_p01": 1.5, "fail_ohmdrop_p02": 1.2,
            "fail_ohmdrop_p03": 0.8, "fail_ohmdrop_p04": 0.6}})
    check("مرحله ۱ با پارامترهای دو علت ثبت می‌شود", rr.status_code == 200,
          str(rr.get_json().get("fields") or rr.get_json().get("error"))[:90])
    det = bozorg.get(f"/api/workflow/instances/{seen}?stage=2").get_json()["data"]
    summary = det.get("summary") or []
    check("مرحله ۲ کار مرحله ۱ را می‌بیند",
          any(b["stage_number"] == 1 for b in summary), str(summary))
    values = {v["label"]: v["value"] for b in summary for v in b["values"]}
    check("علت خرابی ثبت‌شده را می‌بیند",
          values.get("علت خرابی") == "شولات، اهم دار", str(values))
    check("مرحله‌ی خودش در خلاصه تکرار نمی‌شود",
          all(b["stage_number"] != 2 for b in summary))
    check("خلاصه فقط خواندنی است (فیلد نیست)",
          all("field_name" not in v for b in summary for v in b["values"]))
    # The page can rebuild the same summary from the entries alone, so a
    # browser running ahead of its server still shows the earlier stages.
    entries = {e["stage_number"]: e for e in det["entries"]}
    check("payload هر مرحله برای بازسازی خلاصه در دسترس است",
          entries[1]["status"] == "submitted"
          and entries[1]["payload"].get("failure") == ["شولات", "اهم دار"],
          str(entries[1].get("payload"))[:70])
    check("عنوان و ثبت‌کننده‌ی مرحله هم همراه payload می‌آید",
          bool(entries[1].get("title")) and bool(entries[1].get("user_name")),
          f"{entries[1].get('title')} / {entries[1].get('user_name')}")
    bozorg.post(f"/api/workflow/instances/{seen}/submit", json={
        "stage_number": 2, "data": {"review_decision": "نیاز به کشیدن دارد"}})
    det = yaghouti.get(f"/api/workflow/instances/{seen}?stage=3").get_json()["data"]
    stages_seen = {b["stage_number"] for b in det.get("summary") or []}
    check("مرحله ۳ کار مرحله‌های ۱ و ۲ را می‌بیند",
          {1, 2} <= stages_seen, str(sorted(stages_seen)))

    print("\n— فرایند: قواعد مرحله ۴ —")
    def run_to_stage4(action, pump_now=None):
        i = markaz.post("/api/workflow/instances",
                        json={"operation_kind": "کشیدن", "well": "امام رضا 11"}
                        ).get_json()["data"]["id"]
        markaz.post(f"/api/workflow/instances/{i}/submit", json={
            "stage_number": 1, "data": {"op_jdate": "1405/06/22",
            "well": "امام رضا 11", "center": "سوران", "failure": ["شولات"],
            # «شولات» ticked, so its readings are owed at this stage.
            "fail_sanding_p01": "تخلیه شده", "fail_sanding_p02": "انجام شده",
            "fail_sanding_p03": "نشده", "fail_sanding_p04": "دارد",
            "fail_sanding_p05": 12, "fail_sanding_p06": 30}})
        bozorg.post(f"/api/workflow/instances/{i}/submit", json={
            "stage_number": 2, "data": {"review_decision": "نیاز به کشیدن دارد"}})
        yaghouti.post(f"/api/workflow/instances/{i}/submit", json={
            "stage_number": 3, "data": {"operation": "کشیدن",
            "motor_curr": "73", "pump_curr": "384"}})
        data = {"required_action": action}
        if pump_now:
            data["pump_type_now"] = pump_now
        bozorg.post(f"/api/workflow/instances/{i}/submit",
                    json={"stage_number": 4, "data": data})
        d = bozorg.get(f"/api/workflow/instances/{i}").get_json()["data"]
        return i, next(e for e in d["entries"] if e["stage_number"] == 4)

    _i, e4 = run_to_stage4("ویدئومتری")
    check("«ویدئومتری» فرم نصب را بایگانی می‌کند", e4["status"] == "archived",
          e4["status"])
    _i, e4 = run_to_stage4("بهسازی")
    check("«بهسازی» هم بایگانی می‌کند", e4["status"] == "archived")
    _i, e4 = run_to_stage4("نصب الکتروپمپ جدید", "خیر")
    check("«تیپ در این مرحله نه» یعنی موکول به بعد",
          e4["status"] == "deferred", e4["status"])
    last, e4 = run_to_stage4("نصب الکتروپمپ جدید", "بله")
    check("«تیپ در این مرحله بله» یعنی فرم پر می‌شود",
          e4["status"] == "submitted", e4["status"])
    rr = bozorg.post(f"/api/workflow/instances/{last}/submit",
                     json={"stage_number": 4, "data": {}})
    check("بدون «اقدام مورد نیاز» مرحله ۴ ثبت نمی‌شود",
          rr.status_code == 422 or rr.get_json().get("ok"))

    print("\n— فرایند: ثبت نهایی و مستندات —")
    det = kahani.get(f"/api/workflow/instances/{last}?stage=5").get_json()["data"]
    check("مرحله ۵ در کشیدن سه بخش دارد",
          len(det["form"]["sections"]) == 3,
          str([s["code"] for s in det["form"]["sections"]]))
    import io as _io
    up = kahani.post(f"/api/workflow/instances/{last}/attachments",
                     data={"file": (_io.BytesIO(b"%PDF-1.4 test"), "gozaresh.pdf")},
                     content_type="multipart/form-data")
    check("بارگذاری مستند PDF", up.status_code == 200,
          str(up.get_json().get("error")))
    doc_id = up.get_json()["data"]["id"]
    bad = kahani.post(f"/api/workflow/instances/{last}/attachments",
                      data={"file": (_io.BytesIO(b"MZ"), "virus.exe")},
                      content_type="multipart/form-data")
    check("فایل اجرایی رد می‌شود", bad.status_code == 415)
    check("دانلود مستند", kahani.get(f"/api/workflow/attachments/{doc_id}")
          .status_code == 200)
    rr = kahani.post(f"/api/workflow/instances/{last}/submit", json={
        "stage_number": 5, "data": {"test_flow": 30, "starter": "سافت",
                                    "workshop_opinion": ["شولاتی"]}})
    body_ = rr.get_json()
    d_ = kahani.get(f"/api/workflow/instances/{last}").get_json()["data"]
    check("با ثبت مرحله ۵ رکورد ساخته می‌شود",
          bool(body_.get("data", {}).get("record_id")),
          str(body_.get("error")) + " | " + str(body_.get("fields"))[:120]
          + " | " + str([(e["stage_number"], e["status"]) for e in d_["entries"]]))
    new_id = body_["data"]["record_id"]
    rec = c.get(f"/api/records/{new_id}").get_json()["data"]
    check("رکورد نوع عملیات را دارد",
          rec["dynamic"].get("operation_kind") == "کشیدن")
    check("رکورد علت خرابی را دارد", rec["failure"] == ["شولات"])
    check("رکورد نام چاه را دارد", rec["well"] == "امام رضا 11", str(rec["well"]))
    check("فرایند تکمیل‌شده علامت خورد",
          c.get(f"/api/workflow/instances/{last}").get_json()["data"]["status"]
          == "completed")

    print("\n— فرایند: مرکز مستندات —")
    listing = kahani.get("/api/workflow/attachments").get_json()
    check("فهرست مستندات برای متولی باز است", listing.get("ok") is True)
    docs = listing.get("data") or []
    check("مستند بارگذاری‌شده در فهرست هست",
          any(d["filename"] == "gozaresh.pdf" for d in docs), str(len(docs)))
    if docs:
        one = next(d for d in docs if d["filename"] == "gozaresh.pdf")
        check("فرایند و چاه مستند مشخص است",
              bool(one.get("instance_id")) and bool(one.get("well")),
              f"{one.get('instance_id')} / {one.get('well')}")
        check("مرحله‌ی بارگذاری مشخص است", one.get("stage_title") is not None,
              str(one.get("stage_title")))
        check("دانلود از مرکز مستندات کار می‌کند",
              kahani.get(one["url"]).status_code == 200)
    check("جستجوی مستند با نام فایل",
          len((kahani.get("/api/workflow/attachments?q=gozaresh")
               .get_json().get("data") or [])) >= 1)
    check("مدیر هم همه‌ی مستندات را می‌بیند",
          c.get("/api/workflow/attachments").get_json().get("ok") is True)

    print("\n— فرایند: مقادیر «قبلی» —")
    prev = c.get("/api/workflow/previous?well=" + quote("امام رضا 11")).get_json()["data"]
    check("مقادیر قبلی از آخرین عملیات خوانده می‌شود",
          prev["values"].get("motor_prev") == "73", str(prev["values"]))
    check("تاریخ نصب قبلی شمسی است",
          str(prev["values"].get("prev_install_date", "")).startswith("14"),
          str(prev["values"].get("prev_install_date")))
    _pv = prev["values"]
    check("ماه نصب قبلی با فهرست ماه‌ها جور است (شماره‌ی ماه) و با تاریخ نصب قبلی یکی است",
          str(_pv.get("old_install_month")) in [str(n) for n in range(1, 13)]
          and str(_pv.get("prev_install_date", "")).split("/")[1:2]
          == [f"{int(_pv['old_install_month']):02d}"],
          f"{_pv.get('old_install_month')} / {_pv.get('prev_install_date')}")
    check("منبع مقادیر قبلی اعلام می‌شود", bool(prev["source"]))
    # A well with no history is not a broken prefill; the page must be able to
    # tell the two apart, which it does from an empty values map.
    with app.app_context():
        from app.models import Well
        from app.extensions import db as _db
        fresh = Well(name="چاه بدون سابقه", is_active=True, is_verified=True)
        _db.session.add(fresh); _db.session.commit()
        fresh_id = fresh.id
    empty = c.get(f"/api/workflow/previous?well_id={fresh_id}").get_json()["data"]
    check("چاه بدون سابقه مقادیر قبلی ندارد", not (empty.get("values") or {}),
          str(empty))
    check("و منبعی هم اعلام نمی‌شود", not empty.get("source"))

    print("\n— فرایند: گزینه‌های قفل‌شده —")
    with app.app_context():
        from app.models import LookupItem
        from app.services.lookups import get_category
        cat = get_category("failure_reason")
        locked = LookupItem.query.filter_by(category_id=cat.id,
                                            value="جمع آوری").one()
    check("«جمع آوری» به علت خرابی اضافه شده و قفل است", locked.is_locked)
    rr = c.delete(f"/api/lookups/item/{locked.id}")
    check("گزینه قفل‌شده حذف نمی‌شود", rr.status_code == 409)
    rr = c.put(f"/api/lookups/item/{locked.id}", json={"is_active": False})
    check("گزینه قفل‌شده غیرفعال نمی‌شود", rr.status_code == 409)
    rr = c.put(f"/api/lookups/item/{locked.id}", json={"label": "جمع‌آوری چاه"})
    check("ولی برچسبش قابل تغییر است", rr.status_code == 200)

    print("\n— فرم علت خرابی در کارتابل باز می‌شود —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser
        from app.services.workflow import (WorkflowError, active_workflow,
                                           stage_by_number, stage_form,
                                           start_instance, submit_stage,
                                           sync_entries)
        boss = AppUser.query.filter_by(role="admin").first()
        inst = start_instance({"operation_kind": "کشیدن",
                               "well": "امام رضا 11"}, boss)
        sync_entries(inst)
        _db.session.commit()
        st1 = stage_by_number(inst, 1)
        codes = [b["code"] for b in stage_form(inst, st1)["sections"]]
        check("مرحله‌ای که «علت خرابی» می‌پرسد فرم علت‌ها را هم دارد",
              sum(1 for c in codes if c.startswith("fail_")) == 10,
              f"{sum(1 for c in codes if c.startswith('fail_'))} فرم")
        # A cause ticked with its readings blank is refused at this stage —
        # not five stages later in front of somebody who cannot fill them.
        try:
            submit_stage(inst, {"op_jdate": "1405/07/03", "center": "سوران",
                                "failure": ["هوادهی"]}, boss, stage_number=1)
            refused, fields = False, {}
        except WorkflowError as exc:
            refused, fields = True, exc.fields
        check("پارامترهای علت انتخاب‌شده همان‌جا الزامی‌اند",
              refused and any(k.startswith("fail_aeration") for k in fields),
              "، ".join(sorted(fields))[:80])
        check("ولی پارامترهای علت‌های انتخاب‌نشده خواسته نمی‌شوند",
              not any(k.startswith("fail_burn") for k in fields))

    print("\n— شروع هر عملیات از درِ خود متولی —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser
        from app.services.workflow import (active_workflow, applicable_stages,
                                           start_instance, startable_kinds,
                                           sync_entries)
        wf = active_workflow()
        s1 = next(x for x in wf.stages if x.stage_number == 1)
        s3 = next(x for x in wf.stages if x.stage_number == 3)
        # The reported setup: stage 1 is passed through by both operations,
        # stage 3 is the نصب door.
        s1.applies_to, s1.can_start, s1.start_kind = "both", True, "pull"
        s3.can_start, s3.start_kind = True, "install"
        _db.session.commit()
        mk = AppUser.query.filter_by(username="markaz").one()
        yq = AppUser.query.filter_by(username="yaghouti").one()
        check("مرکز آبرسانی فقط «کشیدن» را شروع می‌کند",
              startable_kinds(mk) == ["pull"], str(startable_kinds(mk)))
        check("متولی مرحله ۳ «نصب» را شروع می‌کند",
              startable_kinds(yq) == ["install"], str(startable_kinds(yq)))
        inst = start_instance({"operation_kind": "نصب",
                               "well": "امام رضا 11"}, yq)
        sync_entries(inst)
        _db.session.commit()
        check("و فرایند نصب از مرحله ۳ آغاز می‌شود",
              inst.entry_stage == 3
              and [x.stage_number for x in applicable_stages(inst)][0] == 3,
              str([x.stage_number for x in applicable_stages(inst)]))

    print("\n— ارجاع همزمان به چند نفر —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser
        from app.models.workflow import ENTRY_DONE, REFER_CHOOSE
        from app.services.workflow import (active_workflow, owners_of,
                                           stage_by_number, start_instance,
                                           submit_stage, sync_entries)
        wf = active_workflow()
        boss = AppUser.query.filter_by(role="admin").first()
        bz = AppUser.query.filter_by(username="bozorg").one()
        kh = AppUser.query.filter_by(username="kahani").one()
        s3 = next(x for x in wf.stages if x.stage_number == 3)
        s3.referral_mode = REFER_CHOOSE
        s3.refer_all = True
        _db.session.commit()

        inst = start_instance({"operation_kind": "نصب",
                               "well": "امام رضا 11"}, boss)
        sync_entries(inst)
        _db.session.commit()
        stage4 = stage_by_number(inst, 4)
        # Stage 3 hands stage 4 to two people at once, both of whom must record.
        from app.services.workflow import _refer_onward
        _refer_onward(inst, s3, boss, refer_to=[bz.id, kh.id])
        _db.session.commit()
        check("کار همزمان در کارتابل هر دو نفر است",
              sorted(owners_of(inst, stage4)) == sorted([bz.id, kh.id]))
        e4 = next(e for e in inst.entries if e.stage_number == 4)
        submit_stage(inst, {"required_action": "ویدئومتری"}, bz, stage_number=4)
        check("با ثبت نفر اول، مرحله هنوز باز می‌ماند",
              e4.status not in ENTRY_DONE, e4.status)
        check("و فقط در کارتابل نفر دوم می‌ماند",
              owners_of(inst, stage4) == [kh.id])
        submit_stage(inst, {}, kh, stage_number=4)
        check("با ثبت نفر دوم، مرحله بسته می‌شود",
              e4.status in ENTRY_DONE, e4.status)
        check("و پاسخ نفر اول نگه داشته شده",
              (e4.payload or {}).get("required_action") == "ویدئومتری")

        # «یکی کافی است»: the first to record closes it for everybody.
        s3.refer_all = False
        _db.session.commit()
        inst2 = start_instance({"operation_kind": "نصب",
                                "well": "امام رضا 11"}, boss)
        sync_entries(inst2)
        _db.session.commit()
        _refer_onward(inst2, s3, boss, refer_to=[bz.id, kh.id])
        _db.session.commit()
        submit_stage(inst2, {"required_action": "ویدئومتری"}, kh, stage_number=4)
        e4b = next(e for e in inst2.entries if e.stage_number == 4)
        check("وقتی یکی کافی است، ثبت یک نفر مرحله را می‌بندد",
              e4b.status in ENTRY_DONE, e4b.status)
        s3.referral_mode = "next"
        _db.session.commit()

    print("\n— اتصال علت خرابی به فرم، قابل ویرایش —")
    links = c.get("/api/form-builder/cause-links?field=failure").get_json()["data"]
    check("فهرست علت‌ها و فرم‌هایشان خوانده می‌شود",
          len(links["causes"]) > 10 and len(links["forms"]) == 10,
          f"{len(links['causes'])} علت، {len(links['forms'])} فرم")
    empty = [x["value"] for x in links["causes"] if not x["forms"]]
    check("علت‌های بی‌فرم مشخص‌اند", len(empty) > 0, f"{len(empty)} علت")
    # «سوختن الکتروموتور» joins «سوختن الکتروپمپ» on the burn form, and
    # «گیرپاژ» opens two forms at once.
    rr = c.put("/api/form-builder/cause-links", json={"field": "failure", "links": {
        "fail_burn": ["سوختن الکتروپمپ", "سوختن الکتروموتور"],
        "fail_vibration": ["صدا و لرزش", "گیرپاژ"],
        "fail_noflow": ["عدم آبدهی", "گیرپاژ"]}})
    check("اتصال‌ها ذخیره می‌شود", rr.status_code == 200,
          str(rr.get_json().get("error")))
    with app.app_context():
        from app.services.records import _hidden_by_condition
        from app.models import FormSection
        burn = FormSection.query.filter_by(code="fail_burn").one()
        vib = FormSection.query.filter_by(code="fail_vibration").one()
        noflow = FormSection.query.filter_by(code="fail_noflow").one()
        hid = _hidden_by_condition({"failure": ["سوختن الکتروموتور"]})
        check("علت تازه‌وصل‌شده فرم را باز می‌کند",
              all(f.field_name not in hid for f in burn.fields))
        hid = _hidden_by_condition({"failure": ["گیرپاژ"]})
        check("یک علت می‌تواند دو فرم را با هم باز کند",
              all(f.field_name not in hid for f in vib.fields + noflow.fields))
    rr = c.put("/api/form-builder/cause-links", json={"field": "failure",
                                                      "links": {"fail_reeng": []}})
    with app.app_context():
        from app.services.records import _hidden_by_condition
        from app.models import FormSection
        reeng = FormSection.query.filter_by(code="fail_reeng").one()
        hid = _hidden_by_condition({"failure": ["مهندسی مجدد"]})
        check("فرمی که به هیچ علتی وصل نیست برای کسی باز نمی‌شود",
              reeng.visible_when == "failure="
              and all(f.field_name in hid for f in reeng.fields),
              str(reeng.visible_when))
    c.put("/api/form-builder/cause-links", json={"field": "failure",
          "links": {"fail_reeng": ["مهندسی مجدد"]}})

    print("\n— پر شدن خودکار از سوابق چاه، تنظیم‌شدنی در فرم‌ساز —")
    with app.app_context():
        from app.models import FormField
        pre = {f.field_name: f.prefill_from for f in
               FormField.query.filter(FormField.prefill_from.isnot(None)).all()}
        check("فیلدهای «قبلی» منبعشان را دارند",
              pre.get("prev_install_date") == "@op_date"
              and pre.get("motor_prev") == "motor_curr", str(pre)[:80])
    fid = next(f["id"] for s_ in c.get("/api/form-builder").get_json()["data"]["sections"]
               for f in s_["fields"] if f["field_name"] == "flow_before_pull")
    rr = c.put(f"/api/form-builder/fields/{fid}",
               json={"prefill_from": "flow_after_install"})
    check("مدیر منبع هر فیلد را تعیین می‌کند", rr.status_code == 200)
    # A known last operation on a well, so the answer is checked against a
    # value rather than against whatever the fixture happens to hold.
    made = c.post("/api/records", json={
        "op_jdate": "1405/05/10", "well": "کورده 1", "center": "سوران",
        "operation": "نصب", "motor_curr": "18.5", "pump_curr": "233",
        "flow_after_install": 27}).get_json()
    check("عملیات آزمایشی روی چاه ثبت شد", made.get("ok") is True,
          str(made.get("fields") or made.get("error"))[:80])
    with app.app_context():
        from app.models import Record
        from app.services.workflow import previous_values_for
        rec = Record.query.get(made["data"]["id"])
        vals = previous_values_for(rec.well_id)["values"]
        check("و مقدار از آخرین عملیات همان چاه خوانده می‌شود",
              str(vals.get("flow_before_pull")) in ("27", "27.0"),
              str(vals.get("flow_before_pull")))
        check("تاریخ نصب قبلی از تاریخ آخرین عملیات پر می‌شود",
              vals.get("prev_install_date") == "1405/05/10",
              str(vals.get("prev_install_date")))
        check("تیپ موتور قبلی از تیپ موتور فعلی همان عملیات",
              str(vals.get("motor_prev")) == "18.5", str(vals.get("motor_prev")))
    c.put(f"/api/form-builder/fields/{fid}", json={"prefill_from": ""})

    print("\n— تصمیم در مرحله: توقف یا برگشت برای مستندسازی —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser, WorkflowAttachment
        from app.models.workflow import (ENTRY_REJECTED, INSTANCE_OPEN,
                                         INSTANCE_STOPPED)
        from app.services.workflow import (WorkflowError, actions_for,
                                           active_workflow, owners_of,
                                           stage_by_number, start_instance,
                                           submit_stage, sync_entries,
                                           take_action)
        wf = active_workflow()
        s2 = next(x for x in wf.stages if x.stage_number == 2)
        mk = AppUser.query.filter_by(username="markaz").one()
        bz = AppUser.query.filter_by(username="bozorg").one()
        kh = AppUser.query.filter_by(username="kahani").one()
        # Stage 2 has two owners; only بزرگمهر may stop it, both may send it
        # back to stage 1 for documents.
        s2.owners = [bz, kh]
        s2.assignee_id = bz.id
        import json as _json
        s2.actions_json = _json.dumps([
            {"id": "stop", "kind": "stop", "label": "نیاز به کشیدن ندارد",
             "user_ids": [bz.id]},
            {"id": "docs", "kind": "return", "label": "مستند تصویری لازم است",
             "target_stage": 1, "needs_docs": True, "user_ids": []}])
        _db.session.commit()

        def fresh():
            i = start_instance({"operation_kind": "کشیدن",
                                "well": "امام رضا 11"}, mk)
            sync_entries(i)
            _db.session.commit()
            submit_stage(i, {"op_jdate": "1405/07/05", "failure": ["هوادهی"],
                             "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                             "fail_aeration_p02": 3, "fail_aeration_p03": 12},
                         mk, stage_number=1)
            return i

        inst = fresh()
        st2 = stage_by_number(inst, 2)
        ids_bz = [a["id"] for a in actions_for(inst, st2, bz)]
        ids_kh = [a["id"] for a in actions_for(inst, st2, kh)]
        check("بزرگمهر هر سه اقدام را دارد", ids_bz == ["forward", "stop", "docs"],
              str(ids_bz))
        check("کاهانی اجازه‌ی توقف ندارد", ids_kh == ["forward", "docs"], str(ids_kh))
        try:
            take_action(inst, 2, kh, "stop", note="نیاز نیست")
            denied = False
        except WorkflowError:
            denied = True
        check("توقف توسط کسی که اجازه ندارد رد می‌شود", denied)
        try:
            take_action(inst, 2, bz, "stop", note="")
            noreason = False
        except WorkflowError:
            noreason = True
        check("توقف بدون دلیل رد می‌شود", noreason)
        take_action(inst, 2, bz, "stop",
                    note="چاه نیاز به کشیدن ندارد و قابل اصلاح در محل است")
        check("بزرگمهر فرایند را متوقف می‌کند",
              inst.status == INSTANCE_STOPPED and bool(inst.outcome_note),
              inst.status)
        check("و مرحله‌های بعد از کارتابل‌ها می‌روند",
              all(e.status != "pending" for e in inst.entries))

        # Send back for documents.
        inst = fresh()
        take_action(inst, 2, kh, "docs", note="عکس تابلو و فیلم خروجی آب لازم است")
        e1 = next(e for e in inst.entries if e.stage_number == 1)
        check("برگشت، مرحله ۱ را دوباره باز می‌کند",
              e1.status == ENTRY_REJECTED and inst.status == INSTANCE_OPEN, e1.status)
        check("و به کارتابل همان کسی می‌رود که آن را پر کرده بود",
              owners_of(inst, stage_by_number(inst, 1)) == [mk.id])
        try:
            submit_stage(inst, {}, mk, stage_number=1)
            blocked = False
        except WorkflowError as exc:
            blocked, why = True, str(exc)
        check("بدون بارگذاری مستند، ثبت دوباره رد می‌شود", blocked,
              why[:60] if blocked else "")
        _db.session.add(WorkflowAttachment(
            instance_id=inst.id, stage_number=1, filename="tablo.jpg",
            stored_name="x.jpg", content_type="image/jpeg", size_bytes=10,
            uploaded_by=mk.id))
        _db.session.commit()
        submit_stage(inst, {}, mk, stage_number=1)
        check("با بارگذاری مستند، مرحله ۱ دوباره ثبت می‌شود",
              e1.status == "submitted", e1.status)
        check("و کار دوباره به مرحله ۲ برمی‌گردد",
              bz.id in owners_of(inst, stage_by_number(inst, 2)))
        s2.actions_json = None
        s2.owners = [bz]
        _db.session.commit()

    print("\n— ارجاع به عقب: برگشت به مرحله‌ی خود او، نه فرم مرحله‌ی بعد —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser
        from app.models.meta import AppMeta
        from app.models.workflow import ENTRY_REJECTED, REFER_CHOOSE
        from app.services.seed import seed_default_return
        from app.services.workflow import (WorkflowError, active_workflow,
                                           actions_for, forward_candidates,
                                           owners_of, referral_choices,
                                           stage_by_number, start_instance,
                                           submit_stage, sync_entries,
                                           take_action)
        wf = active_workflow()
        s2 = next(x for x in wf.stages if x.stage_number == 2)
        mk = AppUser.query.filter_by(username="markaz").one()
        bz = AppUser.query.filter_by(username="bozorg").one()
        AppMeta.set("workflow_default_return_v1", "")
        _db.session.commit()
        seed_default_return()
        check("هر مرحله (جز اولی) یک «برگشت به مرحله‌ی قبل» پیش‌فرض دارد",
              any(a["kind"] == "return" and a["target_stage"] is None
                  for a in s2.actions), str(s2.actions)[:120])
        old_mode = s2.referral_mode
        s2.referral_mode = REFER_CHOOSE
        _db.session.commit()

        inst = start_instance({"operation_kind": "کشیدن",
                               "well": "امام رضا 11"}, mk)
        sync_entries(inst)
        _db.session.commit()
        submit_stage(inst, {"op_jdate": "1405/07/05", "failure": ["هوادهی"],
                            "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                            "fail_aeration_p02": 3, "fail_aeration_p03": 12},
                     mk, stage_number=1)
        st2 = stage_by_number(inst, 2)
        people = [u.id for u in forward_candidates(inst, st2)]
        check("مرکز آبرسانی در فهرست «ارجاع» مرحله ۲ نیست", mk.id not in people,
              str(people))
        check("و فهرست کارتابل هم همین را نشان می‌دهد",
              mk.id not in [u["id"] for u in referral_choices(inst, st2)["users"]])
        try:
            submit_stage(inst, {"review_result": "نیاز به کشیدن ندارد"}, bz,
                         stage_number=2, refer_to=[mk.id])
            refused = False
        except WorkflowError as exc:
            refused, why = True, str(exc)
        check("ارجاع مرحله‌ی بعد به مرکز آبرسانی رد می‌شود", refused,
              why[:70] if refused else "")
        st3 = stage_by_number(inst, 3)
        check("و فرم کارگاه مکانیک به کارتابل مرکز نمی‌رود",
              mk.id not in owners_of(inst, st3))
        back = next(a for a in actions_for(inst, st2, bz) if a["id"] == "back")
        check("«برگشت» مرحله‌های ثبت‌شده‌ی قبل را پیشنهاد می‌کند",
              [c["stage_number"] for c in back["choices"]] == [1],
              str(back["choices"]))
        try:
            take_action(inst, 2, bz, "back", note="علت خرابی ناقص است")
            picked = False
        except WorkflowError:
            picked = True
        check("بدون انتخاب مرحله، برگشت رد می‌شود", picked)
        take_action(inst, 2, bz, "back", note="علت خرابی ناقص است", target_stage=1)
        e1 = next(e for e in inst.entries if e.stage_number == 1)
        check("برگشت، مرحله ۱ خود مرکز را دوباره باز می‌کند",
              e1.status == ENTRY_REJECTED
              and owners_of(inst, stage_by_number(inst, 1)) == [mk.id], e1.status)
        check("و فرایند جلو نمی‌رود (مرحله ۳ هنوز به مرکز نرسیده)",
              mk.id not in owners_of(inst, st3) and inst.current_stage == 1,
              str(inst.current_stage))
        s2.referral_mode = old_mode
        _db.session.commit()

    # One person's powers, from the user editor.
    with app.app_context():
        from app.models import AppUser
        kh = AppUser.query.filter_by(username="kahani").one()
        bz = AppUser.query.filter_by(username="bozorg").one()
        s2 = next(x for x in active_workflow().stages if x.stage_number == 2)
        s2.owners = [bz, kh]
        _db.session.commit()
        kh_id, s2_id, bz_id = kh.id, s2.id, bz.id
    r = c.get(f"/api/workflow/powers/{kh_id}")
    rows = (r.get_json() or {}).get("data", {}).get("rows", [])
    mine = next((x for x in rows if x["key"] == f"{s2_id}:back"), None)
    check("اختیارات هر کاربر در صفحه‌ی کاربران خوانده می‌شود",
          r.status_code == 200 and mine is not None and mine["allowed"],
          str(mine)[:120])
    r = c.put(f"/api/workflow/powers/{kh_id}",
              json={"grants": {f"{s2_id}:back": False}})
    with app.app_context():
        s2 = next(x for x in active_workflow().stages if x.stage_number == 2)
        users = next(a for a in s2.actions if a["id"] == "back")["user_ids"]
    check("گرفتن یک اختیار از کاهانی آن را برای بقیه نگه می‌دارد",
          r.status_code == 200 and users == [bz_id], str(users))
    r = c.put(f"/api/workflow/powers/{bz_id}",
              json={"grants": {f"{s2_id}:back": False}})
    check("گرفتن آخرین دارنده‌ی اختیار رد می‌شود (خالی یعنی همه)",
          r.status_code == 422)
    r = c.put(f"/api/workflow/powers/{kh_id}",
              json={"grants": {f"{s2_id}:back": True}})
    with app.app_context():
        s2 = next(x for x in active_workflow().stages if x.stage_number == 2)
        users = next(a for a in s2.actions if a["id"] == "back")["user_ids"]
        s2.actions_json = None
        s2.owners = [AppUser.query.get(bz_id)]
        _db.session.commit()
    check("و دوباره دادنش هم کار می‌کند", sorted(users) == sorted([bz_id, kh_id]),
          str(users))
    print("\n— تصمیم وابسته به پاسخ: «نیاز به کشیدن ندارد» → توقف —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser
        from app.models.meta import AppMeta
        from app.models.workflow import INSTANCE_STOPPED
        from app.services.seed import seed_review_decision
        from app.services.workflow import (WorkflowError, active_workflow,
                                           forced_decisions, stage_by_number,
                                           start_instance, submit_stage,
                                           sync_entries, take_action)
        wf = active_workflow()
        s2 = next(x for x in wf.stages if x.stage_number == 2)
        mk = AppUser.query.filter_by(username="markaz").one()
        bz = AppUser.query.filter_by(username="bozorg").one()
        AppMeta.set("workflow_review_decision_v1", "")
        _db.session.commit()
        seed_review_decision()
        stop = next((a for a in s2.actions if a["kind"] == "stop"), None)
        check("مرحله‌ی «نتیجه بررسی» توقفِ وابسته به پاسخ دارد",
              stop is not None and stop["when"] == "review_decision=نیاز به کشیدن ندارد",
              str(stop)[:120])
        check("و «برگشت برای اصلاح» کنارش هست",
              any(a["kind"] == "return" for a in s2.actions))

        inst = start_instance({"operation_kind": "کشیدن",
                               "well": "امام رضا 11"}, mk)
        sync_entries(inst)
        _db.session.commit()
        submit_stage(inst, {"op_jdate": "1405/07/05", "failure": ["هوادهی"],
                            "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                            "fail_aeration_p02": 3, "fail_aeration_p03": 12},
                     mk, stage_number=1)
        st2 = stage_by_number(inst, 2)
        no = {"review_decision": "نیاز به کشیدن ندارد"}
        yes = {"review_decision": "نیاز به کشیدن دارد"}
        check("با «ندارد»، تصمیم توقف پیش می‌آید",
              [a["id"] for a in forced_decisions(inst, st2, bz, no)] == ["stop"])
        check("با «دارد»، هیچ تصمیمی تحمیل نمی‌شود",
              forced_decisions(inst, st2, bz, yes) == [])
        try:
            submit_stage(inst, no, bz, stage_number=2)
            blocked = False
        except WorkflowError as exc:
            blocked, why = True, str(exc)
        check("با «ندارد»، «ارسال به مرحله بعد» رد می‌شود", blocked,
              why[:70] if blocked else "")
        try:
            take_action(inst, 2, bz, "stop", note="نیازی نیست", payload=yes)
            refused = False
        except WorkflowError:
            refused = True
        check("با «دارد»، دکمه‌ی توقف پذیرفته نمی‌شود", refused)
        take_action(inst, 2, bz, "stop",
                    note="بررسی نشان داد چاه قابل اصلاح در محل است", payload=no)
        check("با «ندارد» و توضیح، فرایند متوقف می‌شود",
              inst.status == INSTANCE_STOPPED, inst.status)
        check("و پاسخ «ندارد» در پرونده ثبت می‌ماند",
              inst.payload.get("review_decision") == "نیاز به کشیدن ندارد")
        s2_id = s2.id
    r = c.put(f"/api/workflow/stages/{s2_id}", json={"actions": [
        {"id": "stop", "kind": "stop", "label": "x", "when": "review_decision="}]})
    check("شرط بدون پاسخ در فرایندساز رد می‌شود", r.status_code == 422)
    r = c.put(f"/api/workflow/stages/{s2_id}", json={"actions": [
        {"id": "stop", "kind": "stop", "label": "توقف",
         "when": "review_decision=نیاز به کشیدن ندارد"}]})
    check("و شرط درست ذخیره و برگردانده می‌شود",
          r.status_code == 200
          and (r.get_json()["data"]["actions"] or [{}])[0].get("when")
          == "review_decision=نیاز به کشیدن ندارد")
    d = c.get("/api/workflow/definition").get_json()["data"]
    rf = next((f for f in d.get("choice_fields", []) if f["name"] == "review_decision"), None)
    check("پرسش‌های گزینه‌ای با پاسخ‌هایشان به فرایندساز می‌رسند",
          rf is not None and "نیاز به کشیدن ندارد" in rf["options"], str(rf)[:120])
    c.put(f"/api/workflow/stages/{s2_id}", json={"actions": []})
    b = c.get("/api/build").get_json()["data"]
    check("سرور می‌گوید با کدام نسخه اجرا شده", b["stale"] is False, str(b))

    print("\n— فرم‌ساز: انتقال و کپی فیلد بین بخش‌ها —")
    fb = c.get("/api/form-builder?all=1").get_json()["data"]
    sec_a, sec_b = fb["sections"][0], fb["sections"][1]
    r = c.post("/api/form-builder/fields", json={
        "field_name": "pipe_dia_test", "label": "قطر لوله آبده (آزمایشی)",
        "field_type": "select", "section_id": sec_a["id"], "prefill_from": "@self",
        "options": [{"value": "۴ اینچ", "label": "۴ اینچ"},
                    {"value": "۶ اینچ", "label": "۶ اینچ"}]})
    new_id = (r.get_json() or {}).get("data", {}).get("id")
    check("فیلد تازه با «برداشت از سوابق» ساخته می‌شود و آن را نگه می‌دارد",
          r.status_code == 200 and r.get_json()["data"].get("prefill_from") == "@self",
          str((r.get_json() or {}).get("message"))[:80])
    order = [{"id": f["id"], "section_id": sec_b["id"]} for f in sec_b["fields"]]
    order.insert(0, {"id": new_id, "section_id": sec_b["id"]})
    r = c.post("/api/form-builder/reorder", json={"placement": order})
    fb2 = c.get("/api/form-builder?all=1").get_json()["data"]
    moved_to = next(s["code"] for s in fb2["sections"]
                    if any(f["id"] == new_id for f in s["fields"]))
    check("کشیدن فیلد به بخش دیگر ذخیره می‌شود (دیگر برنمی‌گردد)",
          r.status_code == 200 and moved_to == sec_b["code"], moved_to)
    r = c.post(f"/api/form-builder/fields/{new_id}/copy",
               json={"section_id": sec_a["id"]})
    cp = (r.get_json() or {}).get("data") or {}
    check("کپی فیلد در بخش دیگر ساخته می‌شود",
          r.status_code == 200 and cp.get("section_id") == sec_a["id"]
          and cp.get("field_name") == "pipe_dia_test_2", str(cp.get("field_name")))
    check("و گزینه‌ها و «برداشت از سوابق» هم کپی می‌شوند",
          len(cp.get("own_options") or []) == 2 and cp.get("prefill_from") == "@self")

    print("\n— برداشت از سوابق: عقب‌تر از آخرین عملیات —")
    r1 = c.post("/api/records", json={
        "op_jdate": "1404/02/10", "well": "امام رضا 11", "center": "سوران",
        "operation": "نصب", "motor_curr": "18.5", "pump_curr": "233",
        "pipe_dia_test": "۶ اینچ"})
    r2 = c.post("/api/records", json={
        "op_jdate": "1405/07/08", "well": "امام رضا 11", "center": "سوران",
        "operation": "نصب", "motor_curr": "18.5", "pump_curr": "233"})
    check("دو عملیات آزمایشی ثبت شد", r1.status_code == 200 and r2.status_code == 200,
          str((r1.get_json() or {}).get("fields") or (r2.get_json() or {}).get("fields"))[:120])
    prev = c.get("/api/workflow/previous?well=" + "امام رضا 11").get_json()["data"]
    check("مقدار از آخرین عملیاتی که آن را دارد خوانده می‌شود",
          prev["values"].get("pipe_dia_test") == "۶ اینچ", str(prev["values"])[:160])
    with app.app_context():
        from app.extensions import db as _db
        from app.models import FormField, Well
        w = Well.query.filter_by(name="امام رضا 11").first()
        depth_field = FormField.query.get(cp["id"])
        depth_field.prefill_from = "@well:pm_code"
        _db.session.commit()
        want_pm = w.pm_code
    prev = c.get("/api/workflow/previous?well=" + "امام رضا 11").get_json()["data"]
    check("مشخصات ثبت‌شده‌ی خود چاه هم منبع است (کد PM)",
          not want_pm or prev["values"].get("pipe_dia_test_2") == want_pm,
          f"{want_pm} / {prev['values'].get('pipe_dia_test_2')}")

    print("\n— اتصال‌ها: فیلدهای تازه دیده می‌شوند و به فیلد هم وصل می‌شوند —")
    cl = c.get("/api/form-builder/cause-links?field=pipe_dia_test").get_json()
    check("فیلد گزینه‌ای تازه در فهرست «فیلد مبنا» هست",
          any(s["name"] == "pipe_dia_test" for s in cl["data"]["sources"]))
    check("و فیلدها هم برای اتصال پیشنهاد می‌شوند",
          any(f["code"] == "field:pipe_dia_test_2" for f in cl["data"]["field_others"]))
    r = c.put("/api/form-builder/cause-links", json={
        "field": "pipe_dia_test", "links": {"field:pipe_dia_test_2": ["۴ اینچ"]}})
    with app.app_context():
        rule = FormField.query.filter_by(field_name="pipe_dia_test_2").one().visible_when
    check("اتصال یک گزینه به یک فیلد ذخیره می‌شود",
          r.status_code == 200 and rule == "pipe_dia_test=۴ اینچ", str(rule))
    c.put("/api/form-builder/cause-links", json={
        "field": "pipe_dia_test", "links": {"field:pipe_dia_test_2": []}})
    with app.app_context():
        rule = FormField.query.filter_by(field_name="pipe_dia_test_2").one().visible_when
    check("برداشتن همه‌ی گزینه‌ها فیلد را دوباره معمولی می‌کند", rule is None, str(rule))

    print("\n— فرایندساز: فیلدهای قابل ویرایش و قفل در هر مرحله —")
    with app.app_context():
        from app.models import AppUser, FormField, FormSection
        from app.services.workflow import (active_workflow, stage_by_number,
                                           stage_form, start_instance,
                                           submit_stage, sync_entries)
        wf = active_workflow()
        s1 = next(x for x in wf.stages if x.stage_number == 1)
        mk = AppUser.query.filter_by(username="markaz").one()
        sec = FormSection.query.get(sec_b["id"])
        s1_id, sec_code = s1.id, sec.code
        before = [i.to_dict() for i in s1.items]
    items = [{"kind": i["kind"], "id": i["section_id"] or i["field_id"],
              "applies_to": i["applies_to"], "is_optional": i["is_optional"],
              "is_read_only": i["is_read_only"],
              "locked_fields": i.get("locked_fields") or []} for i in before]
    mine = next((i for i in items if i["kind"] == "section" and i["id"] == sec_b["id"]),
                None)
    added = mine is None
    if added:
        mine = {"kind": "section", "id": sec_b["id"], "applies_to": "both",
                "is_optional": True, "is_read_only": False, "locked_fields": []}
        items.append(mine)
    kept_locks = list(mine["locked_fields"])
    mine["locked_fields"] = ["pipe_dia_test", "no_such_field"]
    r = c.put(f"/api/workflow/stages/{s1_id}/items", json={"items": items})
    saved = next((i for i in r.get_json()["data"]["items"]
                  if i["section_id"] == sec_b["id"]), {})
    check("قفل فیلدها برای یک فرم در مرحله ذخیره می‌شود (فقط فیلدهای همان فرم)",
          saved.get("locked_fields") == ["pipe_dia_test"], str(saved.get("locked_fields")))
    with app.app_context():
        inst = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"},
                              AppUser.query.filter_by(username="markaz").one())
        sync_entries(inst)
        _db.session.commit()
        st1 = stage_by_number(inst, 1)
        block = next(b for b in stage_form(inst, st1)["sections"] if b["code"] == sec_code)
        fdict = {f["field_name"]: f for f in block["fields"]}
        check("فیلد قفل‌شده در کارتابل فقط‌خواندنی است",
              fdict["pipe_dia_test"].get("read_only") is True)
        check("و مقدارش از سوابق چاه نشان داده می‌شود",
              fdict["pipe_dia_test"].get("read_only_value") == "۶ اینچ",
              str(fdict["pipe_dia_test"].get("read_only_value")))
        others_open = [n for n, f in fdict.items()
                       if n != "pipe_dia_test" and not f.get("read_only")]
        check("بقیه‌ی فیلدهای همان فرم قابل ویرایش‌اند", bool(others_open) or len(fdict) == 1)
        submit_stage(inst, {"op_jdate": "1405/07/09", "failure": ["هوادهی"],
                            "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                            "fail_aeration_p02": 3, "fail_aeration_p03": 12,
                            "pipe_dia_test": "۴ اینچ"},
                     AppUser.query.filter_by(username="markaz").one(), stage_number=1)
        check("مقدار دست‌کاری‌شده برای فیلد قفل پذیرفته نمی‌شود؛ مقدار سوابق ثبت می‌شود",
              inst.payload.get("pipe_dia_test") == "۶ اینچ",
              str(inst.payload.get("pipe_dia_test")))
    if added:
        items.remove(mine)
    else:
        mine["locked_fields"] = kept_locks
    c.put(f"/api/workflow/stages/{s1_id}/items", json={"items": items})
    for fid in (cp.get("id"), new_id):
        c.delete(f"/api/form-builder/fields/{fid}")

    print("\n— دو فرایند جدا: «فرایند کشیدن» و «فرایند نصب» —")
    with app.app_context():
        from app.models import AppUser, WorkflowDefinition, WorkflowInstance
        from app.services.workflow import (active_workflow, startable_kinds,
                                           start_instance, workflow_for)
        wf = active_workflow()
        s3 = next(x for x in wf.stages if x.stage_number == 3)
        s3.can_start, s3.start_kind = True, "install"
        yq = AppUser.query.filter_by(username="yaghouti").one()
        s3.owners = [yq]
        s3.assignee_id = yq.id
        _db.session.commit()
        source_id = wf.id
        install_titles = [s.title for s in sorted(wf.stages, key=lambda x: x.stage_number)
                          if s.stage_number >= 3 and s.is_active]
    r = c.post(f"/api/workflow/definitions/{source_id}/clone", json={
        "name": "فرایند نصب", "operation_kind": "install",
        "source_kind": "pull", "activate": True})
    new_wf = (r.get_json() or {}).get("data") or {}
    check("«فرایند نصب» از روی فرایند فعلی ساخته و فعال می‌شود",
          r.status_code == 200 and new_wf.get("is_active") is True
          and new_wf.get("operation_kind") == "install",
          str((r.get_json() or {}).get("message"))[:90])
    stages_new = sorted(new_wf.get("stages") or [], key=lambda s: s["stage_number"])
    check("مرحله‌هایش از درِ نصب (مهدی یاقوتی) شروع و از ۱ شماره‌گذاری می‌شوند",
          [s["title"] for s in stages_new] == install_titles
          and stages_new and stages_new[0]["stage_number"] == 1
          and stages_new[0]["can_start"],
          str([(s["stage_number"], s["title"]) for s in stages_new]))
    check("فرم‌ها و متولی‌ها هم کپی شده‌اند",
          stages_new and stages_new[0]["items"]
          and stages_new[0]["owner_names"]
          and "یاقوتی" in "".join(stages_new[0]["owner_names"]),
          str(stages_new[0]["owner_names"] if stages_new else ""))
    with app.app_context():
        src = _db.session.get(WorkflowDefinition, source_id)
        check("فرایند فعلی از این به بعد فقط برای کشیدن است",
              src.operation_kind == "pull" and src.is_active)
        check("و درِ نصب از آن برداشته شده",
              not any(s.can_start and s.start_kind == "install" for s in src.stages))
        check("هر عملیات به فرایند خودش می‌رود",
              workflow_for("pull").id == source_id
              and workflow_for("install").id == new_wf["id"])
        mk = AppUser.query.filter_by(username="markaz").one()
        yq = AppUser.query.filter_by(username="yaghouti").one()
        check("مرکز آبرسانی فقط کشیدن را شروع می‌کند",
              startable_kinds(mk) == ["pull"], str(startable_kinds(mk)))
        check("مهدی یاقوتی فقط نصب را شروع می‌کند",
              startable_kinds(yq) == ["install"], str(startable_kinds(yq)))
        inst = start_instance({"operation_kind": "نصب", "well": "امام رضا 11"}, yq)
        check("نصب در «فرایند نصب» و از مرحله ۱ آن شروع می‌شود",
              inst.workflow_id == new_wf["id"] and inst.entry_stage == 1,
              f"{inst.workflow_id}/{inst.entry_stage}")
        pull = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"}, mk)
        check("کشیدن در فرایند قبلی می‌ماند", pull.workflow_id == source_id)
    r = c.post("/api/workflow/definitions", json={"name": "آزمایشی هر دو"})
    third = r.get_json()["data"]["id"]
    c.post("/api/workflow/stages", json={"workflow_id": third, "title": "یک مرحله"})
    r = c.put(f"/api/workflow/definitions/{third}", json={"is_active": True})
    check("فعال شدن فرایند سوم برای همان عملیات رد می‌شود", r.status_code == 422,
          str(r.get_json().get("message"))[:80])
    c.delete(f"/api/workflow/definitions/{third}")
    # Back to one process for both, for whatever runs after this.
    c.put(f"/api/workflow/definitions/{new_wf['id']}", json={"is_active": False})
    c.put(f"/api/workflow/definitions/{source_id}", json={"operation_kind": None})
    with app.app_context():
        src = _db.session.get(WorkflowDefinition, source_id)
        s3 = next(x for x in src.stages if x.stage_number == 3)
        s3.can_start, s3.start_kind = True, "install"
        _db.session.commit()

    print("\n— فرایند تازه‌ی «خالی» با کاربر تازه: شروع فرایند در کارتابل —")
    r = c.post("/api/users", json={"username": "nasb_new", "password": "pass1234",
                                   "first_name": "کاربر", "last_name": "نصب",
                                   "role": "stage_owner", "must_change_password": False})
    nasb = (r.get_json() or {}).get("data", {}).get("id")
    r = c.post("/api/workflow/definitions", json={"name": "نصب آزمایشی",
                                                  "operation_kind": "install"})
    np_id = r.get_json()["data"]["id"]
    check("فرایند یک‌عملیاتی بدون «مرحله ۰»ِ بی‌متولی ساخته می‌شود",
          not r.get_json()["data"]["stages"], str(r.get_json()["data"]["stages"])[:80])
    c.post("/api/workflow/stages", json={"workflow_id": np_id, "title": "کارگاه نصب"})
    d = c.get(f"/api/workflow/definition?workflow_id={np_id}").get_json()["data"]
    st = d["workflow"]["stages"][0]
    c.put(f"/api/workflow/stages/{st['id']}", json={"title": "کارگاه نصب",
                                                   "owner_ids": [nasb], "applies_to": "both"})
    d = c.get(f"/api/workflow/definition?workflow_id={np_id}").get_json()["data"]
    keys = {x["key"]: x["ok"] for x in d["readiness"]["checks"]}
    check("وضعیت آمادگی می‌گوید چه چیزی کم است (فعال نیست، تداخل دارد)",
          not d["readiness"]["ready"] and keys.get("active") is False
          and keys.get("clash") is False, str(keys))
    check("و مرحله‌ی اول، بی‌تیک هم، مرحله‌ی شروع حساب می‌شود",
          keys.get("door") is True, str(d["readiness"]["checks"])[:160])
    r = c.put(f"/api/workflow/definitions/{np_id}", json={"is_active": True})
    check("فعال کردن با وجود فرایند «هر دو» رد می‌شود و علتش گفته می‌شود",
          r.status_code == 422 and r.get_json().get("clash") is True)
    r = c.put(f"/api/workflow/definitions/{np_id}",
              json={"is_active": True, "resolve_clash": True})
    d = c.get(f"/api/workflow/definition?workflow_id={np_id}").get_json()["data"]
    check("با «فرایند دیگر فقط برای کشیدن باشد» فعال می‌شود",
          r.status_code == 200 and d["readiness"]["ready"],
          str(r.get_json().get("error") or d["readiness"]["checks"])[:160])
    with app.app_context():
        from app.models import AppUser, WorkflowDefinition
        from app.services.workflow import (active_workflow, start_instance,
                                           startable_kinds)
        u = _db.session.get(AppUser, nasb)
        check("کاربر تازه در کارتابل «شروع فرایند — نصب» را دارد",
              startable_kinds(u) == ["install"], str(startable_kinds(u)))
        inst = start_instance({"operation_kind": "نصب", "well": "امام رضا 11"}, u)
        check("و فرایند نصب از مرحله‌ی اولِ فرایند تازه شروع می‌شود",
              inst.workflow_id == np_id and inst.entry_stage == 1,
              f"{inst.workflow_id}/{inst.entry_stage}")
        main = (WorkflowDefinition.query.filter(WorkflowDefinition.id != np_id,
                                                WorkflowDefinition.is_active.is_(True))
                .first())
        check("فرایند اصلی حالا فقط برای کشیدن است", main.operation_kind == "pull")
        main_id = main.id
    r = c.get("/api/workflow/inbox")
    # back to one process for both, for whatever runs after this
    c.put(f"/api/workflow/definitions/{np_id}", json={"is_active": False})
    c.put(f"/api/workflow/definitions/{main_id}", json={"operation_kind": None})
    with app.app_context():
        main = _db.session.get(WorkflowDefinition, main_id)
        s3 = next(x for x in main.stages if x.stage_number == 3)
        s3.can_start, s3.start_kind = True, "install"
        _db.session.commit()

    print("\n— یک فرم یا فیلد، وصل به چند پرسش —")
    fb = c.get("/api/form-builder?all=1").get_json()["data"]
    home = fb["sections"][0]["id"]
    c.post("/api/form-builder/fields", json={
        "field_name": "test_result_ml", "label": "نتیجه تست (آزمایشی)",
        "field_type": "radio", "section_id": home,
        "options": [{"value": "دبی کم", "label": "دبی کم"},
                    {"value": "عادی", "label": "عادی"}]})
    r = c.put("/api/form-builder/cause-links", json={
        "field": "test_result_ml", "links": {"fail_aeration": ["دبی کم"]}})
    with app.app_context():
        from app.models import FormSection
        rule = FormSection.query.filter_by(code="fail_aeration").one().visible_when
    check("وصل کردن فرم به پرسش دوم، اتصال قبلی را پاک نمی‌کند",
          r.status_code == 200 and "failure=" in (rule or "")
          and "test_result_ml=دبی کم" in (rule or ""), str(rule))
    cl = c.get("/api/form-builder/cause-links?field=failure").get_json()["data"]
    check("و فرم هنوز در فهرست «علت خرابی» وصل است",
          any(f["code"] == "fail_aeration" for f in cl["forms"]))
    cl2 = c.get("/api/form-builder/cause-links?field=test_result_ml").get_json()["data"]
    check("و در فهرست پرسش دوم هم وصل نشان داده می‌شود",
          any(f["code"] == "fail_aeration" and f["causes"] == ["دبی کم"]
              for f in cl2["forms"]))
    schema2 = c.get("/api/form-builder").get_json()["data"]
    entry = next((x for x in schema2["conditional"] if x.get("section") == "fail_aeration"),
                 {})
    check("صفحه‌ی فرم هر دو شرط را می‌گیرد", len(entry.get("any") or []) == 2,
          str(entry.get("any")))
    with app.app_context():
        from app.services.records import _hidden_by_condition
        aer = [f.field_name for f in FormSection.query.filter_by(code="fail_aeration")
               .one().fields if f.is_active]
        open_by_test = _hidden_by_condition({"failure": ["سوختن الکتروپمپ"],
                                             "test_result_ml": "دبی کم"})
        open_by_cause = _hidden_by_condition({"failure": ["هوادهی"],
                                              "test_result_ml": "عادی"})
        closed = _hidden_by_condition({"failure": ["سوختن الکتروپمپ"],
                                       "test_result_ml": "عادی"})
    check("با پاسخ پرسش دوم باز می‌شود", not (set(aer) & open_by_test))
    check("با علت خرابی هم هنوز باز می‌شود", not (set(aer) & open_by_cause))
    check("و وقتی هیچ‌کدام برقرار نیست، بسته است", set(aer) <= closed)
    c.put("/api/form-builder/cause-links", json={
        "field": "test_result_ml", "links": {"fail_aeration": []}})
    with app.app_context():
        rule = FormSection.query.filter_by(code="fail_aeration").one().visible_when
    check("برداشتن اتصال پرسش دوم، اتصال علت خرابی را نگه می‌دارد",
          (rule or "").startswith("failure=") and "test_result_ml" in (rule or ""),
          str(rule))

    print("\n— ثبت نهایی: فیلد الزامی‌ای که هیچ مرحله‌ای نپرسیده، مانع نیست —")
    with app.app_context():
        from app.models import AppUser, FormField, WorkflowStageItem
        from app.services.records import ValidationError, create_record
        from app.services.workflow import (WorkflowError, active_workflow,
                                           asked_fields, finalize, start_instance)
        mk = AppUser.query.filter_by(username="markaz").one()
        good = {"op_jdate": "1405/07/10", "well": "امام رضا 11", "center": "سوران",
                "operation": "کشیدن", "operation_kind": "کشیدن",
                "motor_curr": "18.5", "pump_curr": "233", "failure": ["هوادهی"]}
        motor = FormField.query.filter_by(field_name="motor_curr").one()
        # Take «تیپ الکتروموتور فعلی»'s section off every stage, as in a pull
        # whose stages never ask it.
        wf = active_workflow()
        items = [i for s in wf.stages for i in s.items
                 if i.section_id == motor.section_id or i.field_id == motor.id]
        kept = [(i.stage_id, i.section_id, i.field_id, i.sort_order, i.applies_to,
                 i.is_optional, i.is_read_only, i.locked_fields) for i in items]
        for i in items:
            _db.session.delete(i)
        _db.session.commit()
        inst = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"}, mk)
        check("«تیپ الکتروموتور فعلی» در هیچ مرحله‌ای از این اجرا پرسیده نشده",
              "motor_curr" not in asked_fields(inst))
        without = {k: v for k, v in good.items() if k != "motor_curr"}
        try:
            create_record(dict(without))
            plain_blocked = False
        except ValidationError as exc:
            plain_blocked = "motor_curr" in exc.errors
        _db.session.rollback()
        check("ثبت مستقیم رکورد (صفحه‌ی ثبت اطلاعات) هنوز آن را می‌خواهد", plain_blocked)
        inst.set_payload(dict(without))
        try:
            rec = finalize(inst, mk)
            ok_final, why = rec is not None, ""
        except WorkflowError as exc:
            ok_final, why = False, str(exc)
        check("ولی ثبت نهاییِ فرایند به‌خاطر آن متوقف نمی‌شود", ok_final, why[:140])
        for row in kept:
            _db.session.add(WorkflowStageItem(
                stage_id=row[0], section_id=row[1], field_id=row[2], sort_order=row[3],
                applies_to=row[4], is_optional=row[5], is_read_only=row[6],
                locked_fields=row[7]))
        _db.session.commit()

        inst2 = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"}, mk)
        asked = asked_fields(inst2, with_stage=True)
        if "motor_curr" in asked:
            inst2.set_payload({k: v for k, v in good.items() if k != "motor_curr"})
            try:
                finalize(inst2, mk)
                msg = ""
            except WorkflowError as exc:
                msg = str(exc)
            check("اگر فیلدی که مرحله‌ای پرسیده خالی باشد، پیام برچسب و مرحله را می‌گوید",
                  "تیپ الکتروموتور" in msg and "مرحله" in msg and "motor_curr" not in msg,
                  msg[:140])
        _db.session.rollback()

    html = c.get("/inbox").get_data(as_text=True)
    check("نشانی فایل‌های ثابت نسخه دارد (کش مرورگر نسخه‌ی قدیم را نگه ندارد)",
          "inbox.js?v=" in html)

    print("\n— ستون‌های محاسباتی در گزارش‌ساز —")
    # A row with both readings, so the arithmetic is checked against a known
    # answer rather than against whatever the fixture happens to contain.
    _rr = c.post("/api/records", json={
        "op_jdate": "1405/07/03", "well": "امام رضا 11", "center": "سوران",
        "operation": "نصب", "motor_curr": "18.5", "pump_curr": "233",
        "static_level": 120, "dynamic_level": 138})
    check("ردیف آزمایشی برای محاسبه ثبت شد", _rr.status_code == 200,
          str(_rr.get_json().get("fields") or "")[:120])
    with app.app_context():
        from app.reports.builder import (CALCULATIONS, dynamic_fields,
                                         run_builder)
        dyn = dynamic_fields()
        check("فیلدهای فرم‌ساز به گزارش‌ساز می‌رسند", len(dyn) > 60,
              f"{len(dyn)} فیلد")
        check("پارامترهای علت خرابی هم میان آن‌ها هستند",
              any(f["key"].startswith("dyn.fail_") for f in dyn))
        check("توابع محاسباتی تعریف شده‌اند", len(CALCULATIONS) >= 8,
              "، ".join(sorted(CALCULATIONS)))

        out = run_builder({
            "dataset": "records",
            "fields": ["well", "static_level", "dynamic_level"],
            "computed": [
                {"label": "افت سطح", "fn": "diff",
                 "fields": ["dynamic_level", "static_level"], "decimals": 1},
                {"label": "درصد افت", "fn": "pct_change",
                 "fields": ["static_level", "dynamic_level"],
                 "decimals": 1, "suffix": "٪"}],
            "limit": 200})
        heads = [c["label"] for c in out["columns"]]
        check("ستون محاسباتی به جدول اضافه می‌شود",
              heads[-2:] == ["افت سطح", "درصد افت"], "، ".join(heads))
        done = [r for r in out["rows"] if r["calc1"] != "—"]
        check("و برای ردیف‌های دارای مقدار محاسبه می‌شود", bool(done),
              f"{len(done)} ردیف")
        if done:
            r = done[0]
            check("تفاضل درست است",
                  abs(float(r["calc1"])
                      - (float(r["dynamic_level"]) - float(r["static_level"]))) < 0.05,
                  str(r["calc1"]))
        blank = [r for r in out["rows"] if r["static_level"] == "—"]
        check("ردیف بدون مقدار «—» می‌گیرد، نه صفر",
              all(r["calc1"] == "—" for r in blank), f"{len(blank)} ردیف خالی")

        # Division by zero is not an answer.
        out = run_builder({"dataset": "records", "fields": ["well"],
                           "computed": [{"label": "نسبت", "fn": "ratio",
                                         "fields": ["young_wells",
                                                    "young_wells"]}],
                           "limit": 40})
        check("تقسیم بر صفر خطا نمی‌دهد", True)

        try:
            run_builder({"dataset": "records", "fields": ["well"],
                         "computed": [{"fn": "pct_change",
                                       "fields": ["static_level"]}]})
            few = False
        except ValueError as exc:
            few, why = True, str(exc)
        check("محاسبه با فیلد کم رد می‌شود", few, why[:52] if few else "")

        try:
            run_builder({"dataset": "records",
                         "group_by": [dyn[0]["key"]]})
            grouped = False
        except ValueError as exc:
            grouped, gwhy = True, str(exc)
        check("گروه‌بندی روی فیلد فرم‌ساز با پیام روشن رد می‌شود", grouped,
              gwhy[:52] if grouped else "")

        from app.services.exporter import to_xlsx
        out = run_builder({
            "dataset": "records", "fields": ["well", "static_level"],
            "computed": [{"label": "دو برابر", "fn": "product",
                          "fields": ["static_level", "static_level"],
                          "decimals": 0}],
            "limit": 20})
        blob = to_xlsx(out["columns"], out["rows"], "آزمایش")
        check("ستون محاسباتی در خروجی اکسل هست", len(blob) > 3000
              and out["columns"][-1]["label"] == "دو برابر")

    print("\n— تأیید اجباری و دامنه‌ی دید تأییدکننده —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser, WorkflowInstance
        from app.models.workflow import ENTRY_AWAITING, SEE_PICK
        from app.services.workflow import (WorkflowError, active_workflow,
                                           blocked_by, decide_stage,
                                           start_instance, submit_stage,
                                           submitted_summary, sync_entries)
        wf = active_workflow()
        s2 = next(x for x in wf.stages if x.stage_number == 2)
        s3 = next(x for x in wf.stages if x.stage_number == 3)
        boss = AppUser.query.filter_by(role="admin").first()
        # Stage 2 must be signed off by the admin, and it holds the rest up.
        s2.needs_approval = True
        s2.approval_blocks = True
        s2.approver_id = boss.id
        s2.approval_sees = SEE_PICK
        _db.session.commit()

        inst = start_instance({"operation_kind": "کشیدن",
                               "well": "امام رضا 11"}, boss)
        sync_entries(inst)
        _db.session.commit()
        # Stage 1 first, so there is something to choose to share.
        submit_stage(inst, {"failure": ["هوادهی"], "op_jdate": "1405/07/03",
                            "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                            "fail_aeration_p02": 3, "fail_aeration_p03": 12},
                     boss, stage_number=1)
        # Stage 2, sharing nothing but itself.
        submit_stage(inst, {}, boss, stage_number=2,
                     share_stages=[])
        entry2 = next(e for e in inst.entries if e.stage_number == 2)
        check("مرحله برای تأیید منتظر می‌ماند", entry2.status == ENTRY_AWAITING)
        check("و جلوی مرحله‌های بعدی را می‌گیرد",
              blocked_by(inst, 3) is not None)
        try:
            submit_stage(inst, {}, boss, stage_number=3)
            blocked = False
        except WorkflowError as exc:
            blocked, why = True, str(exc)
        check("ثبت مرحله بعد تا تأیید نشدن رد می‌شود", blocked,
              why[:60] if blocked else "")

        seen = [b["stage_number"] for b in submitted_summary(
            inst, except_stage=2, only_stages=entry2.shared_stage_numbers)]
        check("تأییدکننده فقط همان مرحله را می‌بیند", seen == [],
              f"مرحله‌های دیده‌شده: {seen}")

        decide_stage(inst, 2, boss, approved=True, comment="تأیید شد")
        _db.session.commit()
        check("پس از تأیید، راه باز می‌شود", blocked_by(inst, 3) is None)

        # Now the same stage, sharing stage 1 as well.
        inst2 = start_instance({"operation_kind": "کشیدن",
                                "well": "امام رضا 11"}, boss)
        sync_entries(inst2)
        _db.session.commit()
        submit_stage(inst2, {"failure": ["هوادهی"], "op_jdate": "1405/07/03",
                             "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                             "fail_aeration_p02": 3, "fail_aeration_p03": 12},
                     boss, stage_number=1)
        submit_stage(inst2, {}, boss, stage_number=2,
                     share_stages=[1])
        e2 = next(e for e in inst2.entries if e.stage_number == 2)
        seen = sorted(b["stage_number"] for b in submitted_summary(
            inst2, except_stage=2, only_stages=e2.shared_stage_numbers))
        check("و با انتخاب مرحله ۱، همان را هم می‌بیند", seen == [1], str(seen))

        # Turn the block off: the approval still waits, the process carries on.
        s2.approval_blocks = False
        _db.session.commit()
        check("با خاموش‌کردن اجبار، مرحله بعد آزاد می‌شود",
              blocked_by(inst2, 3) is None)
        s2.needs_approval = False
        s2.approval_blocks = False
        _db.session.commit()

    print("\n— فرم هر علت خرابی با انتخاب همان علت باز می‌شود —")
    with app.app_context():
        from app.models import FormSection
        from app.services.records import _hidden_by_condition
        blocks = FormSection.query.filter(
            FormSection.code.like("fail_%"),
            FormSection.is_active.is_(True)).all()
        check("فرم علت‌های خرابی ساخته شده است", len(blocks) == 10,
              f"{len(blocks)} بخش")
        check("هر بخش به علت خودش گره خورده",
              all((b.visible_when or "").startswith("failure=") for b in blocks))
        burn = next(b for b in blocks if b.code == "fail_burn")
        sand = next(b for b in blocks if b.code == "fail_sanding")
        check("همه‌ی پارامترها الزامی‌اند",
              all(f.is_required for b in blocks for f in b.fields))

        # Nothing ticked: none of them may be demanded.
        hidden = _hidden_by_condition({"failure": []})
        check("بدون انتخاب علت، هیچ پارامتری خواسته نمی‌شود",
              all(f.field_name in hidden for b in blocks for f in b.fields))

        # One ticked: only that one opens.
        hidden = _hidden_by_condition({"failure": ["سوختن الکتروپمپ"]})
        check("با «سوختن الکتروپمپ» فقط پارامترهای همان باز می‌شود",
              all(f.field_name not in hidden for f in burn.fields)
              and all(f.field_name in hidden for f in sand.fields))

        # Two ticked: both open — the operator may report several faults.
        hidden = _hidden_by_condition({"failure": ["سوختن الکتروپمپ", "شولات"]})
        check("با دو علت، هر دو فرم باز می‌شود",
              all(f.field_name not in hidden for f in burn.fields)
              and all(f.field_name not in hidden for f in sand.fields))

        # And a comma-joined string reads the same as a list, because that is
        # how a stored answer comes back.
        hidden = _hidden_by_condition({"failure": "سوختن الکتروپمپ، شولات"})
        check("مقدار ذخیره‌شده هم همان‌طور خوانده می‌شود",
              all(f.field_name not in hidden for f in burn.fields))

    rr = c.post("/api/records", json={
        "op_jdate": "1405/07/03", "well": "امام رضا 11", "center": "سوران",
        "operation": "کشیدن", "failure": ["شولات"]})
    check("ثبت رکورد بدون پر کردن پارامترهای علت رد می‌شود",
          rr.status_code == 422,
          "، ".join(sorted((rr.get_json().get("errors") or {}))[:2]))

    print("\n— مرکز از روی چاه خوانده می‌شود و قفل است —")
    with app.app_context():
        from app.models import Well, WorkflowInstance
        from app.services.workflow import (active_workflow, stage_by_number,
                                           stage_form)
        inst = WorkflowInstance.query.get(iid)
        well = inst.well
        centre_field = None
        for number in range(0, 9):
            st = stage_by_number(inst, number)
            if st is None:
                continue
            for block_ in stage_form(inst, st)["sections"]:
                for f in block_["fields"]:
                    if f["field_name"] == "center":
                        centre_field = f
        check("چاه مرکز دارد", well is not None and well.center is not None,
              well.center.label if well and well.center else "—")
        check("مرکز در فرم مرحله نشان داده می‌شود", centre_field is not None)
        if centre_field:
            check("و مقدارش مرکز همان چاه است",
                  centre_field.get("read_only_value") == well.center.label,
                  str(centre_field.get("read_only_value")))
            check("و قابل تغییر نیست", centre_field.get("read_only") is True)
        check("مرکز از همان شروع در فرایند ثبت شده است",
              inst.payload.get("center") == (well.center.label
                                             if well.center else None),
              str(inst.payload.get("center")))

    print("\n— چند متولی برای یک مرحله، و مسیردهی بر اساس مرکز —")
    # The city has eight مراکز آبرسانی. One stage belongs to all of them, and
    # each is shown only the wells of its own مرکز.
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser, LookupCategory, Well
        from app.services.workflow import (active_workflow, owners_of,
                                           start_instance, sync_entries)
        cat = LookupCategory.query.filter_by(code="center").one()
        pair = cat.items[:2]
        people = []
        for item in pair:
            u = AppUser(username=f"c_{item.id}", first_name="مرکز",
                        last_name=item.label, role="stage_owner")
            u.set_password("Aa@123456")
            u.centers = [item]
            _db.session.add(u)
            people.append(u)
        _db.session.flush()
        wf = active_workflow()
        stage = next(x for x in wf.stages if x.stage_number == 1)
        stage.owners = people
        stage.assignee_id = people[0].id
        stage.route_by_center = True
        _db.session.commit()

        check("یک مرحله چند متولی می‌گیرد", len(stage.owner_ids) == 2,
              str([u.full_name for u in stage.all_owners]))

        seen = []
        for item, owner in zip(pair, people):
            well = Well.query.filter_by(center_id=item.id, is_active=True).first()
            if well is None:
                continue
            inst = start_instance({"operation_kind": "کشیدن",
                                   "well": well.name}, owner)
            sync_entries(inst)
            _db.session.commit()
            seen.append(owners_of(inst, stage) == [owner.id])
        check("کار هر چاه فقط به متولی مرکز خودش می‌رسد",
              bool(seen) and all(seen), f"{len(seen)} مرکز آزمایش شد")

        # A well whose مرکز nobody claims stays with the office that opened it
        # rather than fall into a کارتابل nobody reads.
        orphan = Well.query.filter(
            Well.center_id.notin_([i.id for i in pair]),
            Well.is_active.is_(True)).first()
        # a مرکز user picks only their own مرکز's wells
        from app.services.workflow import WorkflowError, well_center_scope
        try:
            start_instance({"operation_kind": "کشیدن", "well": orphan.name}, people[0])
            refused = False
        except WorkflowError:
            refused = True
        _db.session.rollback()
        check("کاربر مرکز فقط چاه‌های مرکز خودش را می‌تواند انتخاب کند",
              refused and well_center_scope(people[0]) == {pair[0].id})
        own_centres = list(people[0].centers)
        people[0].centers = []                 # an office without a مرکز opens it
        _db.session.flush()
        inst = start_instance({"operation_kind": "کشیدن",
                               "well": orphan.name}, people[0])
        sync_entries(inst)
        people[0].centers = own_centres
        _db.session.commit()
        check("چاهی که مرکزش متولی ندارد نزد اداره‌ی شروع‌کننده می‌ماند",
              owners_of(inst, stage) == [people[0].id])

        # Without routing by مرکز: the run is the starting office's alone —
        # the other ادارات on the same door do not see it.
        stage.route_by_center = False
        _db.session.commit()
        check("بدون مسیردهی، کار فقط نزد اداره‌ی شروع‌کننده است",
              owners_of(inst, stage) == [people[0].id])
        inst_b = start_instance({"operation_kind": "کشیدن", "well": orphan.name}, people[1])
        sync_entries(inst_b)
        _db.session.commit()
        check("فرایندِ اداره‌ی دیگر هم فقط نزد خودش است",
              owners_of(inst_b, stage) == [people[1].id])
        from app.services.workflow import stages_of_user as _sou
        check("اداره‌ی دیگر آن را در کارتابل خود نمی‌بیند",
              not _sou(inst, people[1]) and not _sou(inst_b, people[0]))
        boss_ = AppUser.query.filter_by(role="admin").first()
        inst_c = start_instance({"operation_kind": "کشیدن", "well": orphan.name}, boss_)
        sync_entries(inst_c)
        _db.session.commit()
        check("فرایندی که مدیر شروع کند به همه‌ی متولی‌ها می‌رسد",
              sorted(owners_of(inst_c, stage)) == sorted(u.id for u in people))

        # A disabled account must not hold work another owner could do.
        people[0].is_active = False
        stage.route_by_center = True
        _db.session.commit()
        check("کاربر غیرفعال کار را نگه نمی‌دارد",
              people[0].id not in owners_of(inst, stage))
        people[0].is_active = True
        _db.session.commit()


    print("\n— گزارش‌ساز و خروجی گزارش —")
    import json as _json
    import zipfile as _zip
    A = "/api/analytics"
    meta = c.get(A + "/meta").get_json()["data"]
    check("متای گزارش‌ساز: منابع، نمودارها، توابع", len(meta["sources"]) >= 7
          and len(meta["chart_types"]) >= 20 and len(meta["functions"]) > 20)
    rec_fields = c.get(A + "/sources/records/fields").get_json()["data"]["fields"]
    check("فیلدهای منبع رکوردها از فرم‌ساز", any(f["key"] == "center" for f in rec_fields)
          and any(f["key"] == "_date" for f in rec_fields))
    cat = c.get(A + "/catalog").get_json()["data"]["reports"]
    check("الگوهای گزارش‌های ثابت منتشر شده‌اند", len(cat) >= 15, str(len(cat)))
    rb_def = {"source": "records",
              "calcs": [{"key": "c_head", "label": "اختلاف سطح", "formula": "[dynamic_level] - [static_level]"}],
              "groups": [{"field": "center"}],
              "measures": [{"key": "n", "field": "*", "agg": "count", "label": "تعداد"},
                           {"key": "h", "calc": "c_head", "agg": "avg"}],
              "kpis": [{"id": "k1", "title": "کل", "value": {"agg": "count", "field": "*"}, "target": 1}],
              "charts": [{"id": "ch1", "type": "bar", "x": {"field": "center"}, "measures": ["n"]}],
              "tables": [{"id": "t1", "kind": "grouped"}],
              "interactive_filters": [{"id": "f1", "field": "center", "kind": "select"}],
              "drill": {"enabled": True, "path": ["center"]}}
    r = c.post(A + "/reports", json={"name": "گزارش آزمون", "definition": rb_def})
    check("ساخت گزارش", r.status_code == 200)
    rid = r.get_json()["data"]["id"]
    r = c.post(A + "/preview", json={"definition": rb_def}).get_json()
    check("پیش‌نمایش گزارش", r.get("ok") and r["data"]["kpis"][0]["value"] >= 1
          and r["data"]["charts"][0].get("categories") is not None)
    fv = c.post(A + "/formula/validate", json={"source": "records", "formula": "SUM([well]"}).get_json()["data"]
    check("فرمول نادرست رد می‌شود", fv["valid"] is False)
    fv = c.post(A + "/formula/validate", json={"source": "records",
                                               "formula": "ROUND(COUNT() / 2, 1)"}).get_json()["data"]
    check("فرمول تجمیعی معتبر", fv["valid"] and fv["kind"] == "group")
    check("انتشار بدون دسترسی رد می‌شود", c.post(A + f"/reports/{rid}/publish").status_code == 422)
    c.put(A + f"/reports/{rid}/permissions", json={"permissions": [
        {"principal_kind": "role", "principal": "operator", "access": ["view_dashboard", "run", "filter", "export_xlsx"]}]})
    r = c.post(A + f"/reports/{rid}/publish")
    check("انتشار گزارش پس از بررسی", r.status_code == 200
          and r.get_json()["data"]["status"] == "published")
    # an operator: granted dashboard + Excel, not PDF; data gate: operator lacks record.view
    r = c.post("/api/users", json={"username": "rb_op", "password": "pass1234", "first_name": "گزارش",
                                   "last_name": "خوان", "personnel_code": "RB-1", "role": "operator"})
    op_id = r.get_json()["data"]["id"]
    op = app.test_client()
    op.post("/api/login", json={"username": "rb_op", "password": "pass1234"})
    r = op.get(A + f"/reports/{rid}/view")
    check("دسترسی گزارش داده‌ای را که کاربر نمی‌بیند باز نمی‌کند", r.status_code == 403)
    c.put(f"/api/users/{op_id}", json={"permissions": ["record.create", "well.view", "report.view",
                                                       "workflow.act", "record.view"]})
    r = op.get(A + f"/reports/{rid}/view")
    check("کاربر مجاز گزارش را می‌بیند", r.status_code == 200, r.get_data(as_text=True)[:200])
    r = op.post(A + f"/reports/{rid}/run", json={"filters": {"f1": "سوران"}})
    check("اجرای گزارش با فیلتر تعاملی", r.status_code == 200)
    check("خروجی Excel مجاز است",
          op.post(A + f"/reports/{rid}/export/xlsx", json={}).status_code == 200)
    check("خروجی PDF بدون مجوز رد می‌شود",
          op.post(A + f"/reports/{rid}/export/pdf", json={}).status_code == 403)
    check("داده‌ی خام بدون مجوز رد می‌شود",
          op.post(A + f"/reports/{rid}/raw", json={}).status_code == 403)
    check("کاربر عادی به گزارش‌ساز دسترسی ندارد", op.get(A + "/reports").status_code == 403)
    for fmt_ in ("xlsx", "pdf", "docx", "csv", "json", "print"):
        rr = c.post(A + f"/reports/{rid}/export/{fmt_}", json={"include_raw": True})
        check(f"خروجی {fmt_}", rr.status_code == 200 and len(rr.data) > 200)
        if fmt_ == "docx":
            z = _zip.ZipFile(io.BytesIO(rr.data))
            check("Word راست‌به‌چپ و معتبر", "word/document.xml" in z.namelist()
                  and "w:bidi" in z.read("word/document.xml").decode("utf-8"))
    # versions: editing a published report makes a new working version
    r = c.put(A + f"/reports/{rid}", json={"definition": dict(rb_def, charts=[])}).get_json()["data"]
    check("ویرایش گزارش منتشرشده نسخه‌ی تازه می‌سازد",
          r["latest_version"] == 2 and r["published_version"] == 1)
    check("کاربران همچنان نسخه‌ی منتشرشده را می‌بینند",
          len(op.post(A + f"/reports/{rid}/run", json={}).get_json()["data"]["charts"]) == 1)
    snap = op.post(A + f"/reports/{rid}/snapshots", json={"title": "snap"}).get_json()
    check("ذخیره‌ی Snapshot", snap.get("ok"))
    check("بازکردن Snapshot", op.get(A + f"/snapshots/{snap['data']['id']}").status_code == 200)
    deps = c.get(A + "/dependencies?kind=field&ref=center").get_json()["data"]
    check("وابستگی گزارش به فیلد ثبت می‌شود", any(d["id"] == rid for d in deps))
    g = c.post(A + "/groups", json={"name": "گروه آزمون", "member_ids": [op_id]}).get_json()
    check("ساخت گروه کاربری", g.get("ok"))
    ap = c.get(A + f"/access/principal?kind=user&principal={op_id}").get_json()["data"]
    check("دسترسی نهایی کاربر محاسبه می‌شود", "export_xlsx" in (ap["effective"].get(str(rid))
                                                                 or ap["effective"].get(rid) or []))
    check("صفحه‌ی گزارش‌ساز", c.get("/report-builder").status_code == 200)
    check("صفحه‌ی خروجی گزارش", c.get("/reports").status_code == 200)
    sch = c.post(A + f"/reports/{rid}/schedules", json={"frequency": "monthly", "hour": 7}).get_json()
    run_now = c.post(A + f"/schedules/{sch['data']['id']}/run").get_json()
    check("اجرای زمان‌بندی‌شده Snapshot می‌سازد", run_now.get("ok") and "اجرا شد" in run_now["data"]["message"])


    print("\n— فیلد مستند، فیلد محاسباتی، فیلد مشترک و ارجاع برای تأیید —")
    from app.extensions import db as _db2
    from app.models import FormSection as _FS, LookupCategory as _LC, Well as _W
    r = c.post("/api/form-builder/sections", json={"code": "t_rq", "title": "فرم آزمون تأیید",
                                                   "show_on_entry": False})
    check("ساخت فرم فقط-فرایندی", r.status_code == 200, str(r.get_json().get("error")))
    with app.app_context():
        sec_id = _FS.query.filter_by(code="t_rq").first().id
    for payload in (
            {"field_name": "t_a", "label": "عدد الف", "field_type": "number", "is_required": True},
            {"field_name": "t_b", "label": "عدد ب", "field_type": "number", "is_required": True},
            {"field_name": "t_sum", "label": "حاصل", "field_type": "formula",
             "formula": "[t_a] * [t_b] * 0.001", "step": "4"},
            {"field_name": "t_doc", "label": "عکس پلاک", "field_type": "file", "is_required": True,
             "file_accept": "image/*,.pdf", "file_multiple": False},
            {"field_name": "t_mw", "label": "وضعیت چاه (مشترک)", "field_type": "mirror",
             "mirror_of": "ps_well_cementing"}):
        rr = c.post("/api/form-builder/fields", json=dict(payload, section_id=sec_id))
        check(f"ساخت فیلد {payload['field_type']}", rr.status_code == 200,
              str(rr.get_json().get("error")))
    rr = c.post("/api/form-builder/fields", json={"field_name": "t_bad", "label": "بد",
                                                  "field_type": "formula", "section_id": sec_id,
                                                  "formula": "[nope] + 1"})
    check("فرمول با فیلد ناموجود رد می‌شود", rr.status_code == 422)
    rr = c.post("/api/form-builder/fields", json={"field_name": "t_loop", "label": "حلقه",
                                                  "field_type": "mirror", "section_id": sec_id,
                                                  "mirror_of": "t_loop"})
    check("فیلد مشترکِ خودارجاع رد می‌شود", rr.status_code == 422)
    entry_schema = c.get("/api/form-builder?for=entry").get_json()["data"]
    check("فرم فقط-فرایندی در صفحه‌ی ثبت اطلاعات نیست",
          all(sec["code"] != "t_rq" for sec in entry_schema["sections"]))
    check("فرم‌های انتخاب پمپ ساخته شده‌اند",
          len([sec for sec in c.get("/api/form-builder").get_json()["data"]["sections"]
               if sec["code"].startswith("ps_")]) == 9)
    # Two forms on one answer — one of them process-only — both open on the entry page.
    cl = c.get("/api/form-builder/cause-links?field=failure").get_json()["data"]
    check("فرم‌های انتخاب پمپ در «اتصال‌ها» قابل انتخاب‌اند",
          {o["code"] for o in cl["others"] if o.get("process_only")} >= {"ps_select", "ps_collect"})
    two_cause = next(x for x in cl["causes"] if x["forms"])
    before_links = {two_cause["forms"][0]: [two_cause["value"]], "ps_select": []}
    c.put("/api/form-builder/cause-links", json={"field": "failure", "links": {
        two_cause["forms"][0]: [two_cause["value"]], "ps_select": [two_cause["value"]]}})
    cl2 = c.get("/api/form-builder/cause-links?field=failure").get_json()["data"]
    check("دو فرم به یک گزینه وصل می‌شود",
          set(next(x for x in cl2["causes"] if x["value"] == two_cause["value"])["forms"])
          >= {two_cause["forms"][0], "ps_select"})
    es = c.get("/api/form-builder?for=entry").get_json()["data"]
    opened = {r["section"] for r in es["conditional"] if r.get("section")
              and two_cause["value"] in r["value"].split("|")}
    check("در ثبت اطلاعات هر دو فرمِ وصل‌شده باز می‌شوند",
          {two_cause["forms"][0], "ps_select"} <= opened
          and any(sec["code"] == "ps_select" for sec in es["sections"]), str(opened))
    check("فرم فقط-فرایندیِ وصل‌نشده همچنان در ثبت اطلاعات نیست",
          all(sec["code"] != "ps_collect" for sec in es["sections"]))
    c.put("/api/form-builder/cause-links", json={"field": "failure", "links": before_links})

    wfd = c.get("/api/workflow/definition").get_json()["data"]
    st1 = [x for x in wfd["workflow"]["stages"] if x["stage_number"] == 1][0]
    c.put(f"/api/workflow/stages/{st1['id']}", json={
        "owner_ids": [owners["markaz"]], "approval_request_enabled": True,
        "approval_request_user_ids": [owners["bozorg"]]})
    items = [{"kind": i["kind"], "id": i["section_id"] or i["field_id"], "applies_to": i["applies_to"],
              "is_optional": i["is_optional"], "is_read_only": i["is_read_only"]}
             for i in st1["items"]]
    items.append({"kind": "section", "id": sec_id, "applies_to": "both",
                  "approval_user_id": owners["bozorg"], "approval_required": True})
    rr = c.put(f"/api/workflow/stages/{st1['id']}/items", json={"items": items})
    check("تأیید اجباری روی فرم مرحله", rr.status_code == 200
          and any(i.get("approval_user_id") == owners["bozorg"] for i in rr.get_json()["data"]["items"]))
    with app.app_context():
        cat = _LC.query.filter_by(code="failure_reason").first()
        linked = set()
        for sec in _FS.query.filter(_FS.visible_when.isnot(None)).all():
            if (sec.visible_when or "").startswith("failure="):
                linked |= set(sec.visible_when.split("=", 1)[1].split("|"))
        plain_cause = next(i.value for i in cat.items if i.is_active and i.value not in linked)
        well_name = _W.query.filter(_W.is_active.is_(True)).order_by(_W.id.desc()).first().name
    rr = markaz.post("/api/workflow/instances", json={"operation_kind": "کشیدن", "well": well_name})
    check("شروع فرایند آزمون تأیید", rr.status_code == 200, str(rr.get_json().get("error")))
    pid2 = rr.get_json()["data"]["id"]
    det = markaz.get(f"/api/workflow/instances/{pid2}?stage=1").get_json()["data"]
    tsec = [x for x in det["form"]["sections"] if x["code"] == "t_rq"]
    tf = {f["field_name"]: f for f in (tsec[0]["fields"] if tsec else [])}
    check("فرم آزمون در مرحله ۱", bool(tsec), str([x["code"] for x in det["form"]["sections"]][:6]))
    check("فیلد مشترک، فیلد مبدأ را با برچسب خودش نشان می‌دهد",
          tf.get("ps_well_cementing", {}).get("label") == "وضعیت چاه (مشترک)"
          and tf["ps_well_cementing"].get("field_type") == "radio",
          str({k: (v.get("label"), v.get("field_type")) for k, v in tf.items()}))
    check("فیلد محاسباتی و مستند در فرم مرحله", "t_sum" in tf and tf.get("t_doc", {}).get("field_type") == "file")
    check("پنل ارجاع برای تأیید", det["approval_requests"] and det["approval_requests"]["enabled"]
          and det["approval_requests"]["required"][0]["state"] == "not_sent")
    data = {"failure": [plain_cause], "t_a": "20", "t_b": "3", "op_jdate": "1405/07/01",
            "center": "سوران", "ps_well_cementing": "سیمانته"}
    rr = markaz.post(f"/api/workflow/instances/{pid2}/submit", json={"stage_number": 1, "data": data})
    check("بدون مستند الزامی ثبت نمی‌شود", rr.status_code == 422
          and "t_doc" in (rr.get_json().get("fields") or {}), str(rr.get_json())[:160])
    up = markaz.post(f"/api/workflow/instances/{pid2}/attachments",
                     data={"file": (io.BytesIO(b"x" * 10), "note.txt"), "stage_number": "1",
                           "field_name": "t_doc"}, content_type="multipart/form-data")
    check("نوع فایل غیرمجاز برای مستند رد می‌شود", up.status_code == 415)
    big = io.BytesIO(b"%PDF" + b"0" * (70 * 1024 * 1024))
    up = markaz.post(f"/api/workflow/instances/{pid2}/attachments",
                     data={"file": (big, "plaque.pdf"), "stage_number": "1", "field_name": "t_doc"},
                     content_type="multipart/form-data")
    check("مستند بزرگ‌تر از ۶۴ مگابایت هم پذیرفته می‌شود", up.status_code == 200,
          str(up.get_json().get("error") if up.is_json else up.status_code))
    rr = markaz.post(f"/api/workflow/instances/{pid2}/submit", json={"stage_number": 1, "data": data})
    check("بدون تأیید اجباری مرحله نهایی نمی‌شود", rr.status_code == 422
          and "تأیید" in rr.get_json().get("error", ""), str(rr.get_json())[:300])
    blocks = markaz.post(f"/api/workflow/instances/{pid2}/approval-blocks",
                         json={"stage_number": 1, "draft": data}).get_json()["data"]
    check("فهرست فرم‌های قابل ضمیمه", any(b["key"] == "1:t_rq" for b in blocks))
    rule_id = det["approval_requests"]["required"][0]["item_id"]
    rr = markaz.post(f"/api/workflow/instances/{pid2}/approval-requests",
                     json={"stage_number": 1, "keys": ["1:t_rq"], "note": "لطفاً بررسی شود",
                           "draft": data, "rule_item_id": rule_id})
    check("ارسال درخواست تأیید", rr.status_code == 200, str(rr.get_json().get("error")))
    req = rr.get_json()["data"]
    snap_vals = {v["label"]: v["value"] for v in req["snapshot"][0]["values"]}
    check("خلاصه‌ی فرم ضمیمه با مقدار محاسباتی", snap_vals.get("عدد الف") == "20",
          str(snap_vals))
    markaz.post(f"/api/workflow/instances/{pid2}/attachments",
                data={"file": (io.BytesIO(b"doc"), "extra.pdf"), "stage_number": "1",
                      "approval_request_id": str(req["id"])}, content_type="multipart/form-data")
    rr = markaz.post(f"/api/workflow/instances/{pid2}/submit", json={"stage_number": 1, "data": data})
    check("تا پاسخ تأییدکننده، ثبت نهایی ممکن نیست", rr.status_code == 422
          and "انتظار" in rr.get_json().get("error", ""))
    inbox_b = bozorg.get("/api/workflow/inbox").get_json()
    check("درخواست در کارتابل تأییدکننده", any(x["id"] == req["id"] for x in inbox_b["approval_requests"]))
    view = bozorg.get(f"/api/workflow/approval-requests/{req['id']}").get_json()["data"]
    check("تأییدکننده فرم‌ها، توضیح و مستند را می‌بیند",
          view["note"] == "لطفاً بررسی شود" and len(view["attachments"]) == 1 and view["snapshot"])
    check("دیگران درخواست را نمی‌بینند",
          kahani.get(f"/api/workflow/approval-requests/{req['id']}").status_code == 403)
    rr = bozorg.post(f"/api/workflow/approval-requests/{req['id']}/decide", json={"approved": True, "note": "مورد تأیید"})
    check("تأیید درخواست", rr.status_code == 200)
    inbox_m = markaz.get("/api/workflow/inbox").get_json()
    check("پاسخ به کارتابل ارجاع‌دهنده برگشت",
          any(x["id"] == req["id"] and x["status"] == "approved" for x in inbox_m["approval_results"]))
    det = markaz.get(f"/api/workflow/instances/{pid2}?stage=1").get_json()["data"]
    check("پیش‌نویس فرم پس از بازگشت حفظ شده", (det.get("draft") or {}).get("t_a") == "20")
    changed = dict(data, t_a="21")
    rr = markaz.post(f"/api/workflow/instances/{pid2}/submit", json={"stage_number": 1, "data": changed})
    check("تغییر پس از تأیید، تأیید دوباره می‌خواهد", rr.status_code == 422
          and "تغییر" in rr.get_json().get("error", ""), str(rr.get_json().get("error"))[:100])
    rr = markaz.post(f"/api/workflow/instances/{pid2}/submit", json={"stage_number": 1, "data": data})
    check("پس از تأیید، مرحله ثبت می‌شود", rr.status_code == 200, str(rr.get_json().get("error")))
    with app.app_context():
        from app.models import WorkflowInstance as _WI
        pay = _db2.session.get(_WI, pid2).payload
    check("مقدار محاسباتی روی سرور حساب شد", pay.get("t_sum") == 0.06, str(pay.get("t_sum")))
    check("مستند در پاسخ فیلد ثبت شد", isinstance(pay.get("t_doc"), list) and len(pay["t_doc"]) == 1)

    print("\n— تأییدکننده فرم مرحله‌ی بعد را نمی‌گیرد —")
    with app.app_context():
        from app.extensions import db as _db
        from app.models import AppUser as _AU
        from app.models.workflow import REFER_NEXT, REFER_USER
        from app.services.workflow import (_entry_for, _standing_owners, active_workflow,
                                           decide_stage, owners_of, stage_owner_ids,
                                           start_instance, submit_stage, sync_entries)
        wf = active_workflow()
        s1, s2, s3 = [next(x for x in wf.stages if x.stage_number == n) for n in (1, 2, 3)]
        for it in list(s1.items):                 # the compulsory-approval form of the test above
            if it.section is not None and it.section.code == "t_rq":
                _db.session.delete(it)
        boss = _AU.query.filter_by(role="admin").first()
        judge = next(u for u in _AU.query.filter_by(is_active=True).order_by(_AU.id).all()
                     if u.id != boss.id and u.id not in stage_owner_ids(s3))
        # the setting that went wrong: stage 2 approved by X, and «ارجاع به» also X
        s2.needs_approval, s2.approver_id, s2.approval_blocks = True, judge.id, False
        s2.referral_mode, s2.referral_user_id, s2.referral_user_ids = REFER_USER, judge.id, str(judge.id)
        _db.session.commit()
        inst = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"}, boss)
        sync_entries(inst)
        _db.session.commit()
        submit_stage(inst, {"failure": ["هوادهی"], "op_jdate": "1405/07/03",
                            "fail_aeration_p01": ["وضعیت لوله و اتصالات بررسی شده است"],
                            "fail_aeration_p02": 3, "fail_aeration_p03": 12}, boss, stage_number=1)
        submit_stage(inst, {}, boss, stage_number=2)
        decide_stage(inst, 2, judge, approved=True, comment="تأیید")
        standing = [u.id for u in _standing_owners(inst, s3)]
        check("پس از تأیید، مرحله‌ی بعد به تأییدکننده نمی‌رسد", judge.id not in owners_of(inst, s3),
              f"{owners_of(inst, s3)} / judge={judge.id}")
        check("مرحله‌ی بعد نزد متولی خودش است", owners_of(inst, s3) == standing,
              f"{owners_of(inst, s3)} != {standing}")
        # a run the old version already put in the approver's کارتابل is repaired
        e3 = _entry_for(inst, 3)
        e3.referred_to_id, e3.referred_to_ids, e3.referred_by_id = judge.id, str(judge.id), judge.id
        _db.session.commit()
        sync_entries(inst)
        _db.session.commit()
        check("ارجاع اشتباه قبلی به تأییدکننده خودکار اصلاح می‌شود",
              judge.id not in owners_of(inst, s3) and owners_of(inst, s3) == standing)
        s2.needs_approval, s2.approver_id = False, None
        s2.referral_mode, s2.referral_user_id, s2.referral_user_ids = REFER_NEXT, None, None
        _db.session.commit()
        st2_id, st3_id = s2.id, s3.id

        # Stage 4 set aside («بایگانی») by «اقدام مورد نیاز» is still approved
        # first; «ارجاع به» naming the approver and an approval-request approver
        # hands stage 5 to neither of them.
        from app.models.workflow import ENTRY_ARCHIVED, ENTRY_AWAITING
        from app.services.workflow import awaiting_approval, blocked_by, carries_well_install
        s4, s5 = [next(x for x in wf.stages if x.stage_number == n) for n in (4, 5)]
        other = next(u for u in _AU.query.filter_by(is_active=True).order_by(_AU.id).all()
                     if u.id not in (boss.id, judge.id) and u.id not in stage_owner_ids(s5))
        saved4 = (s4.needs_approval, s4.approver_id, s4.approval_blocks, s4.referral_mode,
                  s4.referral_user_id, s4.referral_user_ids, s4.approval_request_user_ids)
        s4.needs_approval, s4.approver_id, s4.approval_blocks = True, judge.id, True
        s4.referral_mode, s4.referral_user_id = REFER_USER, judge.id
        s4.referral_user_ids = f"{judge.id},{other.id}"
        s4.approval_request_user_ids = str(other.id)
        _db.session.commit()
        check("مرحله ۴ فرم «اطلاعات چاه و نصب» را دارد", carries_well_install(s4))
        inst4 = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"}, boss)
        sync_entries(inst4)
        _db.session.commit()
        submit_stage(inst4, {"required_action": "ویدئومتری"}, boss, stage_number=4)
        e4 = _entry_for(inst4, 4)
        check("مرحله‌ی بایگانی‌شده هم برای تأیید می‌رود",
              e4.status == ENTRY_AWAITING and e4.status_after_approval == ENTRY_ARCHIVED,
              f"{e4.status} / {e4.status_after_approval}")
        check("در کارتابل تأییدکننده برای تأیید است",
              [s.stage_number for s in awaiting_approval(inst4, judge)] == [4])
        standing5 = [u.id for u in _standing_owners(inst4, s5)]
        check("مرحله ۵ به تأییدکننده‌ها نمی‌رسد و نزد متولی خودش است",
              owners_of(inst4, s5) == standing5 and other.id not in owners_of(inst4, s5))
        check("تا تأیید نشود مرحله ۵ ثبت نمی‌شود", blocked_by(inst4, 5) is not None)
        decide_stage(inst4, 4, judge, approved=True, comment="تأیید")
        check("پس از تأیید، مرحله ۴ همان «بایگانی» می‌ماند", _entry_for(inst4, 4).status == ENTRY_ARCHIVED)
        check("و مرحله ۵ آزاد و نزد متولی خودش است",
              blocked_by(inst4, 5) is None and owners_of(inst4, s5) == standing5)
        # a run the old version let skip its approval is sent to it now
        s4.needs_approval = False
        _db.session.commit()
        inst5 = start_instance({"operation_kind": "کشیدن", "well": "امام رضا 11"}, boss)
        sync_entries(inst5)
        _db.session.commit()
        submit_stage(inst5, {"required_action": "ویدئومتری"}, boss, stage_number=4)
        s4.needs_approval = True
        _db.session.commit()
        sync_entries(inst5)
        _db.session.commit()
        check("تأییدِ جاافتاده‌ی قبلی خودکار به کارتابل تأییدکننده می‌رود",
              _entry_for(inst5, 4).status == ENTRY_AWAITING
              and _entry_for(inst5, 4).approver_id == judge.id)
        (s4.needs_approval, s4.approver_id, s4.approval_blocks, s4.referral_mode,
         s4.referral_user_id, s4.referral_user_ids, s4.approval_request_user_ids) = saved4
        _db.session.commit()

    print("\n— تأیید گزینه‌ی یک فیلد، و فیلدی که در مرحله‌ی بعد پر می‌شود —")
    c.post("/api/form-builder/sections", json={"code": "t_opt", "title": "فرم تأیید گزینه",
                                               "show_on_entry": False})
    with app.app_context():
        opt_sec = _FS.query.filter_by(code="t_opt").first().id
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_kind", "label": "نوع کار", "field_type": "radio", "section_id": opt_sec,
        "options": [{"value": "عادی", "label": "عادی"}, {"value": "ویژه", "label": "ویژه"}],
        "approval_options": ["ویژه"], "approval_user_id": owners["bozorg"]})
    check("ذخیره‌ی گزینه‌ی نیازمند تأیید", rr.status_code == 200
          and rr.get_json()["data"].get("approval_options") == ["ویژه"], str(rr.get_json())[:200])
    # this test exercises the waiting kind: the rule is marked «الزامی»
    c.put(f"/api/form-builder/fields/{rr.get_json()['data']['id']}", json={"approval_rules": [
        {"options": ["ویژه"], "approvers": [owners["bozorg"]], "required": True}]})
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_txt_ap", "label": "متن", "field_type": "text", "section_id": opt_sec,
        "approval_options": ["x"], "approval_user_id": owners["bozorg"]})
    check("تأیید گزینه برای فیلد غیرانتخابی رد می‌شود", rr.status_code == 422)
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_later2", "label": "دو مرحله", "field_type": "text", "section_id": opt_sec,
        "fill_stage_ids": [st2_id, st3_id]})
    check("دو مرحله از یک فرایند برای پرکردن رد می‌شود", rr.status_code == 422)
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_later", "label": "شماره سریال", "field_type": "text", "section_id": opt_sec,
        "is_required": True, "fill_stage_ids": [st2_id]})
    check("ذخیره‌ی مرحله‌ی پرکردن فیلد", rr.status_code == 200
          and rr.get_json()["data"].get("fill_stage_ids") == [st2_id], str(rr.get_json())[:200])
    fb = c.get("/api/form-builder?all=1").get_json()["data"]
    check("فرم‌ساز فهرست کاربران و مرحله‌ها را دارد",
          fb.get("users") and any(s["id"] == st2_id for s in fb.get("stages") or []))
    # t_opt on stage 1 and on stage 3 — never on stage 2, where t_later is filled
    wfd = c.get("/api/workflow/definition").get_json()["data"]

    def _items(st, extra=None):
        out = [{"kind": i["kind"], "id": i["section_id"] or i["field_id"], "applies_to": i["applies_to"],
                "is_optional": i["is_optional"], "is_read_only": i["is_read_only"],
                "approval_user_id": i.get("approval_user_id"),
                "approval_required": i.get("approval_required"),
                "locked_fields": i.get("locked_fields") or []}
               for i in st["items"]]
        return out + ([extra] if extra else [])
    stages_ = {x["stage_number"]: x for x in wfd["workflow"]["stages"]}
    for n in (1, 3):
        rr = c.put(f"/api/workflow/stages/{stages_[n]['id']}/items", json={"items": _items(
            stages_[n], {"kind": "section", "id": opt_sec, "applies_to": "both",
                         "is_read_only": n == 3})})
        check(f"فرم آزمون روی مرحله {n}", rr.status_code == 200, str(rr.get_json().get("error")))
    rr = markaz.post("/api/workflow/instances", json={"operation_kind": "کشیدن", "well": well_name})
    pid3 = rr.get_json()["data"]["id"]
    det = markaz.get(f"/api/workflow/instances/{pid3}?stage=1").get_json()["data"]
    names1 = {f["field_name"] for b in det["form"]["sections"] for f in b["fields"]}
    check("فیلد مرحله‌ی بعد در مرحله‌ی قبل دیده نمی‌شود", "t_later" not in names1 and "t_kind" in names1)
    check("پنل تأیید پاسخ‌های نیازمند تأیید را زیر نظر دارد",
          "t_kind" in ((det.get("approval_requests") or {}).get("watch") or []))
    data3 = {"failure": [plain_cause], "op_jdate": "1405/07/02", "center": "سوران", "t_kind": "عادی"}
    st = markaz.post(f"/api/workflow/instances/{pid3}/approval-status",
                     json={"stage_number": 1, "draft": data3}).get_json()["data"]
    check("گزینه‌ی عادی تأیید نمی‌خواهد", not [r for r in st["required"] if r.get("kind") == "option"])
    data3["t_kind"] = "ویژه"
    st = markaz.post(f"/api/workflow/instances/{pid3}/approval-status",
                     json={"stage_number": 1, "draft": data3}).get_json()["data"]
    rule = next((r for r in st["required"] if r.get("kind") == "option"), None)
    check("گزینه‌ی «ویژه» تأیید «بزرگمهر» را می‌خواهد", rule is not None
          and rule["approver_id"] == owners["bozorg"] and rule["state"] == "not_sent", str(st)[:300])
    rr = markaz.post(f"/api/workflow/instances/{pid3}/submit", json={"stage_number": 1, "data": data3})
    check("بدون تأیید گزینه، مرحله به بعد نمی‌رود", rr.status_code == 422
          and "ویژه" in rr.get_json().get("error", ""), str(rr.get_json().get("error"))[:200])
    rr = markaz.post(f"/api/workflow/instances/{pid3}/approval-requests", json={
        "stage_number": 1, "keys": [], "note": "گزینه‌ی ویژه", "draft": data3,
        "rule_item_id": rule["item_id"] if rule else None})
    check("ارسال گزینه برای تأیید", rr.status_code == 200, str(rr.get_json().get("error")))
    req3 = rr.get_json()["data"]
    rr = bozorg.post(f"/api/workflow/approval-requests/{req3['id']}/decide",
                     json={"approved": True, "note": "بلامانع"})
    check("تأیید گزینه توسط کاربر تعیین‌شده", rr.status_code == 200)
    rr = markaz.post(f"/api/workflow/instances/{pid3}/submit", json={"stage_number": 1, "data": data3})
    check("پس از تأیید گزینه، مرحله ثبت می‌شود", rr.status_code == 200, str(rr.get_json().get("error")))
    # several approvers on one answer — all must approve — and the approver's choice
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_route", "label": "ارجاع بعدی", "field_type": "radio", "section_id": opt_sec,
        "options": [{"value": "کارگاه مکانیک", "label": "کارگاه مکانیک"},
                    {"value": "تایید و برگشت", "label": "تایید و برگشت"}],
        "fill_stage_ids": [st2_id]})
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_multi", "label": "اقلام", "field_type": "checkbox", "section_id": opt_sec,
        "options": [{"value": "کابل", "label": "کابل"}, {"value": "ترانس", "label": "ترانس"}],
        "approval_options": ["کابل", "ترانس"], "approval_user_ids": [owners["bozorg"], owners["kahani"]],
        "approval_answer_field": "t_route"})
    if rr.status_code == 200:
        c.put(f"/api/form-builder/fields/{rr.get_json()['data']['id']}", json={"approval_rules": [
            {"options": ["کابل", "ترانس"], "approvers": [owners["bozorg"], owners["kahani"]],
             "answer_field": "t_route", "required": True}]})
    check("چند تأییدکننده و پاسخ تأییدکننده ذخیره می‌شود", rr.status_code == 200
          and rr.get_json()["data"]["approval_user_ids"] == [owners["bozorg"], owners["kahani"]]
          and rr.get_json()["data"]["approval_answer_field"] == "t_route", str(rr.get_json())[:200])
    pid4 = markaz.post("/api/workflow/instances", json={"operation_kind": "کشیدن", "well": well_name}
                       ).get_json()["data"]["id"]
    data4 = {"failure": [plain_cause], "op_jdate": "1405/07/02", "center": "سوران",
             "t_kind": "عادی", "t_multi": ["کابل"]}
    st4 = markaz.post(f"/api/workflow/instances/{pid4}/approval-status",
                      json={"stage_number": 1, "draft": data4}).get_json()["data"]
    rules4 = [r for r in st4["required"] if r.get("kind") == "option"]
    check("هر دو تأییدکننده جداگانه لازم‌اند",
          sorted(r["approver_id"] for r in rules4) == sorted([owners["bozorg"], owners["kahani"]]),
          str(rules4)[:300])
    reqs4 = {}
    for rule4 in rules4:
        rr = markaz.post(f"/api/workflow/instances/{pid4}/approval-requests", json={
            "stage_number": 1, "keys": [], "draft": data4, "rule_item_id": rule4["item_id"]})
        reqs4[rule4["approver_id"]] = rr.get_json()["data"]
    check("دو درخواست تأیید ساخته شد", len(reqs4) == 2 and all(r.get("id") for r in reqs4.values()))
    rb = reqs4[owners["bozorg"]]
    check("تأییدکننده گزینه‌های پاسخ را می‌بیند",
          (rb.get("answer") or {}).get("choices") == ["کارگاه مکانیک", "تایید و برگشت"])
    rr = bozorg.post(f"/api/workflow/approval-requests/{rb['id']}/decide", json={"approved": True})
    check("تأیید بدون انتخاب پاسخ رد می‌شود", rr.status_code == 422)
    rr = bozorg.post(f"/api/workflow/approval-requests/{rb['id']}/decide",
                     json={"approved": True, "answer": "کارگاه مکانیک"})
    check("تأیید با پاسخ", rr.status_code == 200
          and rr.get_json()["data"]["answer"]["value"] == "کارگاه مکانیک")
    rr = markaz.post(f"/api/workflow/instances/{pid4}/submit", json={"stage_number": 1, "data": data4})
    check("تا همه‌ی تأییدکننده‌ها تأیید نکنند، مرحله ثبت نمی‌شود", rr.status_code == 422)
    kahani.post(f"/api/workflow/approval-requests/{reqs4[owners['kahani']]['id']}/decide",
                json={"approved": True, "answer": "تایید و برگشت"})
    rr = markaz.post(f"/api/workflow/instances/{pid4}/submit", json={"stage_number": 1, "data": data4})
    check("پس از تأیید همه، مرحله ثبت می‌شود", rr.status_code == 200, str(rr.get_json().get("error")))
    det4 = c.get(f"/api/workflow/instances/{pid4}?stage=2").get_json()["data"]
    route = [f for b in det4["form"]["sections"] for f in b["fields"] if f["field_name"] == "t_route"]
    check("پاسخ تأییدکننده در پرونده ثبت و در مرحله‌ی بعد پیش‌پر است",
          det4["payload"].get("t_route") in ("کارگاه مکانیک", "تایید و برگشت") and bool(route))
    det2 = c.get(f"/api/workflow/instances/{pid3}?stage=2").get_json()["data"]
    fill = [b for b in det2["form"]["sections"] if b["code"] == "fill_t_opt"]
    later = [f for b in fill for f in b["fields"] if f["field_name"] == "t_later"]
    check("فیلد خودکار در کارتابل مرحله‌ی تعیین‌شده برای پر کردن می‌آید",
          bool(later) and not later[0].get("read_only"),
          str([b["code"] for b in det2["form"]["sections"]]))
    check("و فقط فیلدهای همان مرحله، نه کل فرمش",
          {f["field_name"] for b in fill for f in b["fields"]} <= {"t_later", "t_route"})
    rr = c.post(f"/api/workflow/instances/{pid3}/submit", json={"stage_number": 2, "data": {}})
    check("فیلد الزامیِ آن مرحله خالی بماند، ثبت نمی‌شود", rr.status_code == 422
          and "t_later" in (rr.get_json().get("fields") or {}), str(rr.get_json())[:200])
    with app.app_context():
        from app.models import WorkflowInstance as _WI
        from app.services.workflow import stage_form as _sf
        inst3 = _db2.session.get(_WI, pid3)
        e2 = next(e for e in inst3.entries if e.stage_number == 2)
        e2.set_payload({"t_later": "SN-778"})
        e2.status = "submitted"
        _db2.session.commit()
        form3 = _sf(inst3, next(s for s in inst3.workflow.stages if s.stage_number == 3))
        f3 = {f["field_name"]: f for b in form3["sections"] for f in b["fields"]}
    check("در مرحله‌های بعد همان فیلد فقط خواندنی با مقدار ثبت‌شده است",
          f3.get("t_later", {}).get("read_only") and f3["t_later"].get("read_only_value") == "SN-778",
          str(f3.get("t_later"))[:200])
    for n in (1, 3):
        c.put(f"/api/workflow/stages/{stages_[n]['id']}/items", json={"items": _items(stages_[n])})
    c.put(f"/api/form-builder/fields/{[f for s in fb['sections'] for f in s['fields'] if f['field_name'] == 't_later'][0]['id']}",
          json={"fill_stage_ids": []})

    print("\n— فیلد چندعددی، نمودار، چند قاعده‌ی تأیید، پاسخ مانع و پیش‌فرض محاسباتی —")
    from app.analytics.formula import Evaluator as _Ev, parse as _parse
    _ev = _Ev({})
    _q = "QFIT_{}(1.146, 1.05, 0.954, 18.8, 17, 15.3)"
    check("ضریب a خط روند درجه ۲ از مبدأ مثل اکسل",
          round(_ev.row(_parse(_q.format("A")), {}), 3) == 1.932)
    check("ضریب b خط روند درجه ۲ از مبدأ مثل اکسل",
          round(_ev.row(_parse(_q.format("B")), {}), 3) == 14.182)
    check("ROUNDUP مثل اکسل", _ev.row(_parse("ROUNDUP(87.57, 0)"), {}) == 88
          and _ev.row(_parse("ROUNDUP(-1.2, 0)"), {}) == -2)
    check("IFERROR", _ev.row(_parse("IFERROR(1/0, 5)"), {}) == 5)
    check("نتیجه‌ی متنی در فرمول", _ev.row(_parse('IF(0.14 < 0.1, "رس", "سیلت")'), {}) == "سیلت")
    c.post("/api/form-builder/sections", json={"code": "t_r5", "title": "فرم آزمون R5",
                                               "show_on_entry": False})
    with app.app_context():
        r5_sec = _FS.query.filter_by(code="t_r5").first().id
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_r5_phase", "label": "جریان", "field_type": "numbers", "section_id": r5_sec,
        "is_required": True})
    check("فیلد چند مقدار عددی با فاز ۱ تا ۳ ساخته می‌شود", rr.status_code == 200
          and rr.get_json()["data"]["part_labels"] == ["فاز ۱", "فاز ۲", "فاز ۳"], str(rr.get_json())[:200])
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_r5_cause", "label": "علت", "field_type": "checkbox", "section_id": r5_sec,
        "options": [{"value": "برق", "label": "برق"}, {"value": "مکانیک", "label": "مکانیک"}],
        "approval_rules": [{"options": ["برق"], "approvers": [owners["bozorg"]]},
                           {"options": ["مکانیک"], "approvers": [owners["kahani"]]}]})
    check("چند قاعده‌ی تأیید روی یک فیلد", rr.status_code == 200
          and len(rr.get_json()["data"]["approval_rules"]) == 2, str(rr.get_json())[:200])
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_r5_ok", "label": "مسیر آماده است؟", "field_type": "radio", "section_id": r5_sec,
        "options": [{"value": "بله", "label": "بله"}, {"value": "خیر", "label": "خیر"}],
        "block_options": ["خیر"]})
    check("پاسخ مانع ارسال ذخیره می‌شود", rr.status_code == 200
          and rr.get_json()["data"]["block_options"] == ["خیر"])
    c.post("/api/form-builder/fields", json={"field_name": "t_r5_x", "label": "ورودی", "field_type": "number",
                                             "section_id": r5_sec})
    rr = c.post("/api/form-builder/fields", json={"field_name": "t_r5_auto", "label": "خودکار", "field_type": "number",
                                                  "section_id": r5_sec, "default_value": "=[t_r5_x] * 2"})
    check("پیش‌فرض محاسباتی ذخیره می‌شود", rr.status_code == 200)
    rr = c.post("/api/form-builder/fields", json={"field_name": "t_r5_txt", "label": "رده", "field_type": "formula",
                                                  "section_id": r5_sec, "result_type": "text",
                                                  "formula": 'IF([t_r5_x] > 3, "زیاد", "کم")'})
    check("فرمول با نتیجه‌ی متنی", rr.status_code == 200 and rr.get_json()["data"]["result_type"] == "text",
          str(rr.get_json())[:200])
    rr = c.post("/api/form-builder/fields", json={"field_name": "t_r5_chart", "label": "نمودار", "field_type": "chart",
                                                  "section_id": r5_sec, "chart_config": {"series": []}})
    check("نمودار بدون سری رد می‌شود", rr.status_code == 422)
    rr = c.post("/api/form-builder/fields", json={
        "field_name": "t_r5_chart", "label": "نمودار", "field_type": "chart", "section_id": r5_sec, "is_required": True,
        "chart_config": {"series": [{"label": "افت", "x": ["t_r5_x"], "y": ["t_r5_auto"], "trend": "poly2_0"}],
                         "y2_max": 1}})
    check("نمودار با سری ذخیره می‌شود و الزامی نیست", rr.status_code == 200
          and rr.get_json()["data"]["chart_config"]["series"][0]["trend"] == "poly2_0"
          and not rr.get_json()["data"]["is_required"], str(rr.get_json())[:200])
    wfd = c.get("/api/workflow/definition").get_json()["data"]
    stages_ = {x["stage_number"]: x for x in wfd["workflow"]["stages"]}
    rr = c.put(f"/api/workflow/stages/{stages_[1]['id']}/items", json={"items": _items(
        stages_[1], {"kind": "section", "id": r5_sec, "applies_to": "both"})})
    check("فرم آزمون R5 روی مرحله ۱", rr.status_code == 200, str(rr.get_json().get("error")))
    ar = c.get(f"/api/workflow/stages/{stages_[1]['id']}/approval-rules").get_json()["data"]
    mine = [f for f in ar["fields"] if f["field_name"] == "t_r5_cause"]
    check("قاعده‌های تأیید در فرایندساز دیده می‌شوند", mine and len(mine[0]["rules"]) == 2
          and any(f["field_name"] == "t_r5_ok" and f["block_options"] == ["خیر"] for f in ar["fields"]),
          str(ar)[:300])
    pid5 = markaz.post("/api/workflow/instances", json={"operation_kind": "کشیدن", "well": well_name}
                       ).get_json()["data"]["id"]
    data5 = {"failure": [plain_cause], "op_jdate": "1405/07/02", "center": "سوران",
             "t_r5_phase": ["12", "", ""], "t_r5_cause": ["برق"], "t_r5_ok": "خیر", "t_r5_x": "4", "t_r5_auto": ""}
    rr = markaz.post(f"/api/workflow/instances/{pid5}/submit", json={"stage_number": 1, "data": data5})
    check("چندعددی الزامی با خانه‌ی خالی ثبت نمی‌شود", rr.status_code == 422
          and "t_r5_phase" in (rr.get_json().get("fields") or {}), str(rr.get_json())[:200])
    data5["t_r5_phase"] = ["12", "13", "12.5"]
    rr = markaz.post(f"/api/workflow/instances/{pid5}/submit", json={"stage_number": 1, "data": data5})
    check("پاسخ «خیر» مانع ارسال است", rr.status_code == 422 and "خیر" in rr.get_json().get("error", ""),
          str(rr.get_json().get("error"))[:200])
    data5["t_r5_ok"] = "بله"
    st5 = markaz.post(f"/api/workflow/instances/{pid5}/approval-status",
                      json={"stage_number": 1, "draft": data5}).get_json()["data"]
    rules5 = [r for r in st5["required"] if r.get("kind") == "option"]
    check("هر گزینه تأییدکننده‌ی خودش را دارد (برق ← بزرگمهر، نه کاهانی)",
          [r["approver_id"] for r in rules5] == [owners["bozorg"]], str(rules5)[:300])
    rr = markaz.post(f"/api/workflow/instances/{pid5}/approval-requests", json={
        "stage_number": 1, "keys": [], "draft": data5, "rule_item_id": rules5[0]["item_id"] if rules5 else None})
    snap = rr.get_json()["data"].get("snapshot") or []
    seen = {v["label"]: v["value"] for b in snap for v in b.get("values") or []}
    check("تأییدکننده مقدار محاسبه‌شده را هم می‌بیند", seen.get("رده") == "زیاد" and seen.get("جریان", "").startswith("12 / 13"),
          str(seen)[:300])
    bozorg.post(f"/api/workflow/approval-requests/{rr.get_json()['data']['id']}/decide", json={"approved": True})
    rr = markaz.post(f"/api/workflow/instances/{pid5}/submit", json={"stage_number": 1, "data": data5})
    check("پس از تأیید، مرحله ثبت می‌شود", rr.status_code == 200, str(rr.get_json().get("error"))[:200])
    pay5 = c.get(f"/api/workflow/instances/{pid5}").get_json()["data"]["payload"]
    check("پیش‌فرض محاسباتی فیلد خالی را پر می‌کند", pay5.get("t_r5_auto") == 8, str(pay5.get("t_r5_auto")))
    check("فرمول متنی ذخیره می‌شود", pay5.get("t_r5_txt") == "زیاد")
    rep5 = c.post("/api/workflow/stage-report", json={"stage_id": stages_[1]["id"],
                                                      "columns": ["__instance", "t_r5_phase"]}).get_json()["data"]
    row5 = next((r for r in rep5["rows"] if r["__instance"] == pid5), {})
    check("گزارش مرحله سه عدد را خوانا نشان می‌دهد", row5.get("t_r5_phase") == "12 / 13 / 12.5", str(row5))
    with app.app_context():
        from app.analytics.catalogue import form_fields as _ff
        keys5 = {f["key"] for f in _ff()}
    check("گزارش‌ساز برای هر فاز یک ستون عددی دارد",
          {"t_r5_phase", "t_r5_phase__1", "t_r5_phase__3"} <= keys5 and "t_r5_chart" not in keys5)
    # a referral only informs: recorded straight away, the approver is told
    pid6 = markaz.post("/api/workflow/instances", json={"operation_kind": "کشیدن", "well": well_name}
                       ).get_json()["data"]["id"]
    data6 = dict(data5, t_r5_cause=["مکانیک"], t_r5_ok="بله")
    rr = markaz.post(f"/api/workflow/instances/{pid6}/submit", json={"stage_number": 1, "data": data6})
    check("ارجاعِ غیرالزامی مرحله را نگه نمی‌دارد", rr.status_code == 200, str(rr.get_json().get("error"))[:200])
    told = kahani.get("/api/workflow/inbox").get_json()["approval_requests"]
    check("ارجاع پس از ثبت، خودکار برای اطلاع فرستاده شد",
          any(r["instance_id"] == pid6 for r in told), str([r["instance_id"] for r in told]))
    c.put(f"/api/workflow/stages/{stages_[1]['id']}/items", json={"items": _items(stages_[1])})

    print("\n— ترتیب مرحله‌ها، شرط طی‌شدن و شروع خودکار فرایند بعدی —")
    with app.app_context():
        from app.extensions import db as _db6
        from app.models import (WorkflowDefinition as _WD, WorkflowStage as _WS,
                                WorkflowInstance as _WI6, AppUser as _AU6, Well as _W6)
        from app.services import workflow as _wf6
        from app.services.conditions import matches as _m6
        check("قاعده‌ی «هر مقدار» (*)", _m6("180", ["*"]) and not _m6("", ["*"]) and _m6(["a", "b"], ["b"]))
        fa = _WD(code="t_r6a", name="آزمون ترتیب", operation_kind="pull", is_active=False)
        fb6 = _WD(code="t_r6b", name="آزمون فرایند بعد", operation_kind="install", is_active=False)
        _db6.session.add_all([fa, fb6]); _db6.session.flush()
        sa = {n: _WS(workflow_id=fa.id, stage_number=n, title=f"م{n}", applies_to="both") for n in (1, 2, 3, 4)}
        sa[2].waits_for = "1"
        sa[3].waits_for, sa[3].visit_when = "2", "t_dec=بله"
        sa[4].waits_for, sa[4].spawn_workflow_id = "2", fb6.id
        _db6.session.add_all(list(sa.values()) + [_WS(workflow_id=fb6.id, stage_number=1, title="ن۱",
                                                       applies_to="both")])
        _db6.session.flush()
        w6 = _W6.query.filter(_W6.is_active.is_(True)).first()
        inst6 = _WI6(workflow_id=fa.id, operation_kind="pull", well_id=w6.id, current_stage=1,
                     entry_stage=1, status="open")
        inst6.set_payload({"op_jdate": "1405/07/01", "x": 1})
        _db6.session.add(inst6); _db6.session.flush()
        _db6.session.refresh(fa)
        _wf6.sync_entries(inst6)
        check("مرحله‌ی منتظر، پیش از ثبتِ مرحله‌ی قبل باز نیست",
              not _wf6.stage_open(inst6, sa[2]) and _wf6.stage_open(inst6, sa[1]))
        _wf6._entry_for(inst6, 1).status = "submitted"
        check("با ثبت آن باز می‌شود؛ مرحله‌های بعدی هنوز نه",
              _wf6.stage_open(inst6, sa[2]) and not _wf6.stage_open(inst6, sa[3]))
        _wf6._entry_for(inst6, 2).status = "submitted"
        check("دو مرحله‌ی منتظرِ یک مرحله هم‌زمان باز می‌شوند",
              _wf6.stage_open(inst6, sa[3]) and _wf6.stage_open(inst6, sa[4]))
        nums = lambda: {s.stage_number for s in _wf6.applicable_stages(inst6)}
        check("تا پرسشِ شرط جواب نگرفته، مرحله در مسیر است", 3 in nums())
        inst6.set_payload({**inst6.payload, "t_dec": "خیر"})
        check("با پاسخی که در شرط نیست، مرحله طی نمی‌شود", 3 not in nums())
        inst6.set_payload({**inst6.payload, "t_dec": "بله"})
        check("با پاسخ شرط، مرحله دوباره در مسیر است", 3 in nums())
        fb6.is_active = True
        admin6 = _AU6.query.filter_by(username="admin").first()
        sa[4].spawn_when = "t_dec=خیر;x=1"
        check("با شرطِ برقرارنشده، فرایند بعدی شروع نمی‌شود", _wf6.spawn_after(inst6, sa[4], admin6) is None)
        sa[4].spawn_when = "t_dec=بله;x=1"
        check("وقتی همه‌ی شرط‌ها برقرار است شروع می‌شود (شرط‌ها با هم)", _wf6.spawn_ok(inst6, sa[4]))
        child = _wf6.spawn_after(inst6, sa[4], admin6)
        check("پس از ثبت مرحله، فرایند بعدی برای همان چاه شروع می‌شود",
              child is not None and child.parent_id == inst6.id and child.well_id == w6.id
              and child.payload.get("x") == 1 and "op_jdate" not in child.payload)
        check("دوباره شروع نمی‌شود اگر هنوز در جریان است", _wf6.spawn_after(inst6, sa[4], admin6) is None)
        _db6.session.rollback()

    print("\n— کاتالوگ پمپ، فرمول‌های CAT_*، نمودار کاتالوگ و فیلد پنهان در مرحله —")
    with app.app_context():
        import json as _json7
        from app.extensions import db as _db7
        from app.models import (PumpCatalogModel as _PC, FormField as _FF7, FormSection as _FS7,
                                WorkflowDefinition as _WD7, WorkflowStage as _WS7,
                                WorkflowStageItem as _WSI7, WorkflowInstance as _WI7, Well as _W7)
        from app.services import catalogue as _cat
        from app.services import workflow as _wf7
        from app.services.formfields import compute_formulas as _cf7
        check("کاتالوگ گازار بارگذاری شده", _PC.query.count() >= 290, str(_PC.query.count()))
        m = _cat.model("384", 10.0)
        check("مدل 384/10 با موتور 73.5 و جریان 155", m and m["kw"] == 73.5 and m["a"] == 155,
              m and m["full_title"])
        check("CAT_Q در نقطه‌ی کاتالوگ (هد 189 ← 100 m³/h)",
              abs(_cat.call("CAT_Q", ["384", "10", 189]) - 100 / 3.6) < 1e-6)
        check("CAT_H میان دو نقطه خطی خوانده می‌شود",
              abs(_cat.call("CAT_H", ["384", "10", 25]) - 204) < 1e-6)
        check("بیرون از منحنی چیزی برنمی‌گرداند", _cat.call("CAT_Q", ["384", "10", 500]) is None)
        check("مدل a از کاتالوگ", (_cat.call("CAT_TITLE", ["384", "3a"]) or "").startswith("384/3a"))
        sec7 = _FS7(code="t_r7", title="آزمون R7", is_active=True)
        _db7.session.add(sec7); _db7.session.flush()
        for name, ftype, formula, rt in (
                ("t7_t", "text", None, None), ("t7_s", "text", None, None), ("t7_e", "number", None, None),
                ("t7_q", "formula", "CAT_Q([t7_t], [t7_s], 189)", None),
                ("t7_title", "formula", 'CONCAT([t7_s], "a")', "text"),
                ("t7_mean", "formula", "MEAN(2, [t7_e], 4)", None)):
            _db7.session.add(_FF7(section_id=sec7.id, field_name=name, label=name, field_type=ftype,
                                  formula=formula, result_type=rt, is_active=True, is_builtin=False))
        _db7.session.flush()
        got = _cf7({"t7_t": "384", "t7_s": "10"})
        check("فرمول فرم CAT_Q و CONCAT عدد صحیح و MEAN بدون خالی‌ها",
              abs((got.get("t7_q") or 0) - 27.7778) < 1e-3 and got.get("t7_title") == "10a"
              and got.get("t7_mean") == 3, str(got))
        # hidden field on one stage, and the chart handed to the summary
        fh = _FF7(section_id=sec7.id, field_name="t7_hide", label="پنهان", field_type="number",
                  is_active=True, is_builtin=False)
        ch = _FF7(section_id=sec7.id, field_name="t7_chart", label="نمودار", field_type="chart",
                  is_active=True, is_builtin=False, chart_config=_json7.dumps({"series": [
                      {"label": "کاتالوگ", "catalogue": {"type": "t7_t", "stages": "t7_s"}, "x": "head"}]}))
        _db7.session.add_all([fh, ch]); _db7.session.flush()
        wf7 = _WD7(code="t_r7wf", name="آزمون R7", operation_kind="pull", is_active=False)
        _db7.session.add(wf7); _db7.session.flush()
        s1 = _WS7(workflow_id=wf7.id, stage_number=1, title="م۱", applies_to="both")
        _db7.session.add(s1); _db7.session.flush()
        _db7.session.add(_WSI7(stage_id=s1.id, section_id=sec7.id, sort_order=0, applies_to="both",
                               hidden_fields="t7_hide"))
        _db7.session.flush()
        w7 = _W7.query.filter(_W7.is_active.is_(True)).first()
        i7 = _WI7(workflow_id=wf7.id, operation_kind="pull", well_id=w7.id, current_stage=1,
                  entry_stage=1, status="open")
        i7.set_payload({"t7_t": "384", "t7_s": "10"})
        _db7.session.add(i7); _db7.session.flush()
        _db7.session.refresh(wf7); _db7.session.refresh(s1)
        _wf7.sync_entries(i7)
        names = [f["field_name"] for b in _wf7.stage_form(i7, s1)["sections"] for f in b["fields"]]
        check("فیلد «پنهان در این مرحله» نشان داده نمی‌شود", "t7_hide" not in names and "t7_t" in names)
        charts = _wf7.stage_charts(s1, i7.payload)
        check("نمودار مرحله با مقادیرش به خلاصه داده می‌شود",
              len(charts) == 1 and charts[0]["values"] == {"t7_t": "384", "t7_s": "10"})
        _db7.session.rollback()

    print("\n— خروجی PDF بدون کتابخانه‌ی reportlab پیام روشن می‌دهد —")
    import sys as _sys
    _saved = _sys.modules.get("reportlab.platypus")
    _sys.modules["reportlab.platypus"] = None
    try:
        rr = c.post(A + f"/reports/{rid}/export/pdf", json={})
    finally:
        if _saved is not None:
            _sys.modules["reportlab.platypus"] = _saved
        else:
            _sys.modules.pop("reportlab.platypus", None)
    check("نبود reportlab خطای ۵۰۰ نمی‌دهد و راه جایگزین را می‌گوید",
          rr.status_code in (400, 422) and "reportlab" in (rr.get_json() or {}).get("error", "")
          and "چاپ" in rr.get_json()["error"], str(rr.status_code) + str(rr.get_data()[:200]))
    rr = c.post(A + f"/reports/{rid}/export/pdf", json={})
    check("با reportlab، PDF ساخته می‌شود", rr.status_code == 200 and rr.data[:4] == b"%PDF")

    print("\n— فایل‌های WAL جامانده از پایگاه داده‌ی قبلی —")
    import shutil as _sh
    import sqlite3 as _sq
    import tempfile as _tf
    from app.services.bootstrap import set_aside_foreign_wal
    _d = _tf.mkdtemp()
    _a = _sq.connect(_d + "/old.db")
    _a.execute("PRAGMA journal_mode=WAL")
    _a.execute("PRAGMA wal_autocheckpoint=0")
    _a.execute("CREATE TABLE t(x)")
    _a.executemany("INSERT INTO t VALUES (?)", [(str(i) * 20,) for i in range(3000)])
    _a.commit()
    _new = _sq.connect(_d + "/wells.db")
    _new.execute("CREATE TABLE q(y)")
    _new.executemany("INSERT INTO q VALUES (?)", [(i,) for i in range(50)])
    _new.commit()
    _new.close()
    _sh.copy(_d + "/old.db-wal", _d + "/wells.db-wal")     # the old pair left behind
    _sh.copy(_d + "/old.db-shm", _d + "/wells.db-shm")
    _a.close()
    moved = set_aside_foreign_wal(_d + "/wells.db")
    _ok = _sq.connect(_d + "/wells.db")
    check("WAL جامانده‌ی پایگاه داده‌ی دیگر کنار گذاشته می‌شود (حذف نمی‌شود)",
          len(moved) == 2 and all(".stale-" in m for m in moved)
          and _ok.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
          and _ok.execute("SELECT count(*) FROM q").fetchone()[0] == 50, str(moved))
    _ok.close()
    _own = _sq.connect(_d + "/own.db")
    _own.execute("PRAGMA journal_mode=WAL")
    _own.execute("PRAGMA wal_autocheckpoint=0")
    _own.execute("CREATE TABLE z(x)")
    _own.execute("INSERT INTO z VALUES (1)")
    _own.commit()
    check("WAL خودِ پایگاه داده دست نمی‌خورد", set_aside_foreign_wal(_d + "/own.db") == [])
    _own.close()
    _sh.rmtree(_d, ignore_errors=True)

    print("\n— مرتب‌سازی رکوردها و به‌روز بودن گزارش‌ها —")
    for key in ("center_id", "well", "well_pm_code", "op_date", "pump_curr_id", "total_head"):
        for d in ("asc", "desc"):
            rr = c.get(f"/api/records?sort={key}&dir={d}&page_size=5")
            check(f"مرتب‌سازی {key} {d}", rr.status_code == 200 and rr.get_json().get("ok"))
    a = c.get("/api/records?sort=total_head&dir=asc&page_size=500").get_json()["data"]
    vals = [x.get("total_head") for x in a if x.get("total_head") is not None]
    check("مرتب‌سازی عددی صعودی درست است", vals == sorted(vals), str(vals[:5]))
    d_ = c.get("/api/records?sort=total_head&dir=desc&page_size=500").get_json()["data"]
    vals = [x.get("total_head") for x in d_ if x.get("total_head") is not None]
    check("مرتب‌سازی عددی نزولی درست است", vals == sorted(vals, reverse=True))
    from app.services.persian_sort import fa_key
    words = ["گلستان", "مهر", "پارس", "چمران", "بهار", "آبی", "ژاله", "یاس", "كوثر", "زرین", "امید"]
    check("ترتیب الفبای فارسی (پ چ ژ گ ک)", sorted(words, key=fa_key)
          == ["امید", "آبی", "بهار", "پارس", "چمران", "زرین", "ژاله", "كوثر", "گلستان", "مهر", "یاس"])
    check("اعداد داخل متن عددی مقایسه می‌شوند", fa_key("چاه ۲") < fa_key("چاه 10"))
    with app.app_context():
        from app.extensions import db as _db
        got = [r[0] for r in _db.session.execute(_db.text(
            "SELECT v FROM (SELECT 'گل' v UNION ALL SELECT 'پارس' UNION ALL SELECT 'مهر' "
            "UNION ALL SELECT 'بهار') ORDER BY v COLLATE fa"))]
    check("مرتب‌سازی فارسی در پایگاه داده", got == ["بهار", "پارس", "گل", "مهر"], str(got))
    rep_def = {"source": "records", "kpis": [{"id": "k", "title": "جمع هد",
                                              "value": {"agg": "sum", "field": "total_head"}}]}
    before = c.post("/api/analytics/preview", json={"definition": rep_def}).get_json()["data"]["kpis"][0]["value"]
    rec = next(x for x in c.get("/api/records?page_size=200&sort=op_date&dir=asc").get_json()["data"]
               if x.get("motor_curr") and x.get("center"))
    old = rec.get("total_head") or 0
    pr = c.put(f"/api/records/{rec['id']}", json={"total_head": old + 1000})
    check("ویرایش رکورد", pr.status_code == 200, str(pr.get_json())[:300])
    after = c.post("/api/analytics/preview", json={"definition": rep_def}).get_json()["data"]["kpis"][0]["value"]
    check("ویرایش رکورد بلافاصله در گزارش‌ها دیده می‌شود", round((after or 0) - (before or 0)) == 1000,
          f"{before} → {after}")
    c.put(f"/api/records/{rec['id']}", json={"total_head": old or None})

    print("\n— نام برگه‌ی اکسل و نمودارهای فرایند داشبورد —")
    from app.services.exporter import to_xlsx as _to_xlsx
    import io as _io11
    from openpyxl import load_workbook as _lw11
    wb11 = _lw11(_io11.BytesIO(_to_xlsx([{"key": "a", "label": "الف"}], [{"a": 1}],
                                        "سازنده/تعمیرکار: [پمپ]?")))
    check("عنوان گزارش با «/» و «:» خروجی اکسل می‌دهد", wb11.active.title == "سازنده تعمیرکار پمپ",
          wb11.active.title)
    for k11 in ("maker", "executor"):
        r11 = c.get(f"/api/reports/{k11}/export.xlsx")
        check(f"خروجی اکسل گزارش {k11}", r11.status_code == 200, str(r11.status_code))
    d11 = c.get("/api/dashboard").get_json()["data"]
    check("داشبورد نمودارهای فرایند را می‌دهد", isinstance(d11.get("process_charts"), list)
          and all("key" in x and "data" in x for x in d11["process_charts"]),
          str(d11.get("process_charts"))[:200])

    print("\n— R12: تیپ الکتروپمپ، بانک‌های اطلاعاتی، انبار و گزارش متولی —")
    from types import SimpleNamespace as _NS
    from app.services.epump import electropump_label, parse_electropump_label, parse_paren
    check("تیپ الکتروپمپ به شکل 384/10+73.5", electropump_label("384", 10, 73.5) == "384/10+73.5")
    check("تیپ بدون توان موتور 384/10", electropump_label(384, 10.0) == "384/10")
    check("خواندن تیپ با فاصله و ارقام فارسی",
          parse_electropump_label("۳۸۴ / ۱۰ + ۷۳٫۵") == ("384", "10", 73.5))
    check("خواندن تیپ وارونه 73.5+384/10", parse_electropump_label("73.5+384/10") == ("384", "10", 73.5))
    check("خانه‌ی «()» یعنی تیپ خالی", parse_paren("()") == (None, None) and electropump_label("()") is None)
    check("«293 (9)» → تیپ و طبقات", parse_paren("293 (9)") == ("293", "9"))
    _e12 = _Ev({})
    check("LEADNUM × تعداد: نحوه‌ی کشیدن 12 متری × 14 شاخه = 168",
          _e12.row(_parse('LEADNUM("12 متری") * 14'), {}) == 168)
    check("EPUMP در فرمول", _e12.row(_parse("EPUMP(384, 10, 73.5)"), {}) == "384/10+73.5")

    from app.refdata.flowtest import Grid, parse_grid
    def _row(cells, width=12):
        out = [None] * width
        for col, val in cells.items():
            out[col] = val
        return out
    _g = Grid([
        _row({11: "نام چاه", 10: "چاه آزمون 1", 9: "تاریخ آزمایش", 8: "1404/07/12"}),
        _row({11: "تیپ الکتروپمپ", 10: "384/10+73.5", 9: "عمق چاه", 8: 147}),
        _row({11: "عمق نصب", 10: 135, 9: "سطح ایستایی", 8: 122}),
        _row({11: "داده های دبی سنجی"}),
        _row({11: "کارکرد", 5: "آبدهی (l/s)", 4: "سطح پویایی (m)", 3: "فشار (atm)", 0: "آمپرها"}),
        _row({11: "کارکرد 1", 5: 10.8, 4: 129, 3: 1.5, 2: 87, 1: 86, 0: 83}),
        _row({11: "کارکرد 2 (فشار شبکه)", 5: 10.5, 4: 128, 3: 4.3, 2: 117, 1: 117, 0: 113}),
    ], name="sheet1")
    _t = parse_grid(_g)
    check("دبی‌سنجی: مشخصات برگه خوانده می‌شود", bool(_t) and _t["well_name"] == "چاه آزمون 1"
          and _t["test_date"] == "1404/07/12" and _t["well_depth"] == 147
          and _t["install_depth"] == 135 and _t["static_level"] == 122, str(_t)[:300])
    check("دبی‌سنجی: تیپ الکتروپمپ به سه جزء", bool(_t) and (_t["pump_type"], _t["pump_stages"],
                                                          _t["motor_kw"]) == ("384", "10", 73.5))
    _pts = (_t or {}).get("points") or []
    check("دبی‌سنجی: کارکردها و نقطه‌ی فشار شبکه", len(_pts) == 2 and _pts[1]["at_network"]
          and _pts[0]["flow"] == 10.8 and _pts[1]["dynamic_level"] == 128
          and _pts[0].get("amps") == "83/86/87", str(_pts))
    check("دبی‌سنجی: آبدهی و فشار شبکه از نقطه‌ی شبکه", bool(_t) and _t.get("net_flow") == 10.5
          and _t.get("net_pressure") == 4.3)
    check("برگه‌ای که فرم دبی‌سنجی نیست کنار گذاشته می‌شود",
          parse_grid(Grid([["نام چاه", "x"], ["تولید", 5]])) is None)

    r12 = c.get("/api/refdata/summary")
    check("بانک‌های اطلاعاتی: خلاصه", r12.status_code == 200
          and {s_["key"] for s_ in r12.get_json()["data"]["sources"]} >= {"flowtest", "production", "videometry"},
          str(r12.get_json())[:200])
    with app.app_context():
        from app.refdata import bind_path
        _ft_db = bind_path(app, "flowtest")
    check("هر بانک پایگاه داده‌ی جداگانه دارد",
          os.path.dirname(_ft_db) == os.path.join(tmp, "refdata") and os.path.exists(_ft_db), _ft_db)
    _fb12 = c.get("/api/form-builder?all=1").get_json()["data"]
    check("فهرست مقادیر بانک‌ها برای «پر شدن خودکار»",
          any(o.get("value") == "@ref:best.electropump" for o in _fb12.get("prefill_ref", [])))
    check("فهرست تجهیزها و وضعیت‌های انبار در فرم‌ساز از جدول‌های انبار",
          bool(_fb12.get("warehouse", {}).get("equipment_types"))
          and {x["code"] for x in _fb12["warehouse"].get("conditions", [])} >= {"reusable", "scrap"})

    wcat = c.get("/api/warehouse/catalogue").get_json()["data"]
    eq_item = next(i for i in wcat["items"] if i["kind"] == "equipment")
    part_item = next(i for i in wcat["items"] if i["kind"] == "part")
    check("انبار: وضعیت‌ها (قابل استفاده مجدد، تعمیری، نو، اسقاط، مونتاژشده)",
          {x["code"] for x in wcat["conditions"]} >= {"reusable", "repair", "new", "scrap", "assembled"})
    sec12 = c.post("/api/form-builder/sections", json={"code": "t_r12", "title": "انبار آزمون",
                                                       "show_on_entry": False}).get_json()["data"]["id"]
    rf = c.post("/api/form-builder/fields", json={
        "field_name": "t_wh_rows", "label": "تحویل به انبار تجهیزات", "field_type": "wh_lines",
        "section_id": sec12, "wh_config": {"mode": "rows", "warehouse": "equipment", "direction": "in",
                                           "reason": "pull", "conditions": ["reusable", "repair"]}})
    check("فیلد «اقلام انبار» ساخته می‌شود", rf.status_code == 200, str(rf.get_json())[:200])
    pf = c.post("/api/form-builder/fields", json={
        "field_name": "t_wh_parts", "label": "قطعات دمونتاژ", "field_type": "wh_lines",
        "section_id": sec12, "wh_config": {"mode": "parts", "warehouse": "parts", "direction": "in",
                                           "reason": "disassembly", "equipment_type": part_item["category"],
                                           "columns": ["reusable", "scrap"], "equipment_field": "t_eq_code",
                                           "related_action": "جمع آوری تجهیز"}})
    check("فیلد «فرم قطعات» ساخته می‌شود", pf.status_code == 200, str(pf.get_json())[:200])
    pf_id = (pf.get_json() or {}).get("data", {}).get("id")
    up = c.put(f"/api/form-builder/fields/{pf_id}", json={
        "label": "قطعات دمونتاژ", "field_type": "wh_lines",
        "wh_config": {"mode": "parts", "warehouse": "parts", "direction": "in", "reason": "disassembly",
                      "equipment_type": part_item["category"], "columns": ["reusable", "scrap"],
                      "equipment_field": "t_eq_code"}})
    with app.app_context():
        from app.models import FormField as _FF12
        _cfg12 = json.loads(_FF12.query.filter_by(field_name="t_wh_parts").first().wh_config or "{}")
    check("ذخیره‌ی فرم‌ساز تنظیمات «فرم قطعات» را نگه می‌دارد", up.status_code == 200
          and _cfg12.get("mode") == "parts" and _cfg12.get("columns") == ["reusable", "scrap"]
          and _cfg12.get("related_action") == "جمع آوری تجهیز", str(_cfg12))
    bad12 = c.put(f"/api/form-builder/fields/{pf_id}", json={
        "label": "x", "field_type": "wh_lines",
        "wh_config": {"mode": "parts", "warehouse": "parts", "direction": "in", "columns": []}})
    check("«فرم قطعات» بدون ستون شمارش پذیرفته نمی‌شود", bad12.status_code == 422, str(bad12.status_code))

    with app.app_context():
        from app.warehouse.models import WhMovement as _WM, WhPartAction as _WPA
        from app.warehouse.service import post_stage as _post
        from app.extensions import db as _db12
        inst = _NS(id=990001, workflow=None, well=None, well_name_raw="چاه آزمون 1", well_id=None)
        stg = _NS(stage_number=3, title="کشیدن الکتروپمپ")
        data12 = {"t_wh_rows": [{"item_id": eq_item["id"], "condition": "repair", "qty": 1,
                                 "serial": "EM/001"}],
                  "t_wh_parts": [{"item_id": part_item["id"], "reusable": 2, "scrap": 1}],
                  "t_eq_code": "EM/001"}
        n1 = _post(inst, stg, data12)
        _db12.session.commit()
        n2 = _post(inst, stg, data12)
        _db12.session.commit()
        moves = _WM.query.filter_by(instance_id=990001).all()
        acts = _WPA.query.filter_by(instance_id=990001).all()
        check("ثبت مرحله: ورود به انبار تجهیزات و انبار قطعات", n1 == 3 and len(moves) == 3
              and {(m_.warehouse, m_.direction, m_.condition) for m_ in moves}
              == {("equipment", "in", "repair"), ("parts", "in", "reusable"), ("parts", "in", "scrap")},
              f"{n1} {[(m_.warehouse, m_.direction, m_.condition, m_.qty) for m_ in moves]}")
        check("ارسال دوباره‌ی مرحله جایگزین می‌شود، دوبار شمرده نمی‌شود", n2 == 3 and len(moves) == 3)
        check("سابقه‌ی قطعه با کد تجهیز ثبت می‌شود", len(acts) == 2
              and all(a_.equipment_code == "EM/001" and a_.part_action == "collected" for a_ in acts)
              and sorted(a_.qty for a_ in acts) == [1, 2])
    st12 = c.get("/api/warehouse/stock").get_json()["data"]
    check("موجودی انبار از گردش", r12.status_code == 200 and st12, str(st12)[:200])
    man = c.post("/api/warehouse/movements", json={"warehouse": "equipment", "direction": "in",
                                                   "item_id": eq_item["id"], "qty": 1, "condition": "new",
                                                   "reason": "purchase"})
    check("ثبت دستی ورود خرید نو", man.status_code == 200, str(man.get_json())[:200])
    sm12 = c.get("/api/warehouse/summary").get_json()["data"]
    check("گزارش مدیرعامل: خلاصه و قطعات", "parts" in sm12, str(list(sm12))[:200])
    for k12 in ("movements", "stock", "summary", "parts"):
        rx = c.get(f"/api/warehouse/export/{k12}.xlsx")
        check(f"خروجی اکسل انبار: {k12}", rx.status_code == 200 and len(rx.data) > 500, str(rx.status_code))
    mid = (man.get_json() or {}).get("data", {}).get("id")
    check("حذف گردش دستی", c.delete(f"/api/warehouse/movements/{mid}").status_code == 200)

    meta12 = c.get("/api/analytics/meta").get_json()["data"]
    src_keys = {s_["key"] for s_ in meta12.get("sources", [])}
    check("گزارش‌ساز: منابع انبار و بانک‌ها", {"wh_parts", "wh_moves", "ft_tests", "pr_months",
                                               "vm_insp"} <= src_keys, str(sorted(src_keys))[:300])
    for sk in ("wh_parts", "wh_moves", "ft_tests", "pr_months", "vm_insp"):
        flds = c.get(f"/api/analytics/sources/{sk}/fields").get_json()["data"]
        dims = [f_ for f_ in (flds.get("fields") if isinstance(flds, dict) else flds) or []]
        ok12 = bool(dims)
        if ok12:
            prev = c.post("/api/analytics/preview", json={"definition": {
                "source": sk, "kpis": [{"id": "n", "title": "تعداد", "value": {"agg": "count"}}]}})
            ok12 = prev.status_code == 200
        check(f"گزارش‌ساز: منبع {sk} پیش‌نمایش می‌دهد", ok12)

    all_st = c.get("/api/workflow/my-stages").get_json()["data"]["stages"]
    k_st = kahani.get("/api/workflow/my-stages").get_json()["data"]["stages"]
    check("«گزارش مرحله‌ی من»: مدیر همه‌ی مرحله‌ها را می‌بیند", len(all_st) >= len(k_st) > 0,
          f"{len(all_st)} / {len(k_st)}")
    check("«گزارش مرحله‌ی من»: متولی مرحله‌های خودش را می‌بیند", all(x["mine"] for x in k_st))
    own = kahani.post("/api/workflow/stage-report", json={"stage_id": k_st[0]["id"], "all_columns": True})
    check("متولی گزارش مرحله‌ی خودش را می‌گیرد", own.status_code == 200, str(own.status_code))
    ownx = kahani.post("/api/workflow/stage-report/export.xlsx",
                       json={"stage_id": k_st[0]["id"], "all_columns": True})
    check("خروجی اکسل گزارش مرحله‌ی متولی", ownx.status_code == 200, str(ownx.status_code))
    other = [x for x in all_st if x["id"] not in {y["id"] for y in k_st}]
    if other:
        nope = kahani.post("/api/workflow/stage-report", json={"stage_id": other[0]["id"]})
        check("متولی به گزارش مرحله‌ی دیگران دسترسی ندارد", nope.status_code == 403, str(nope.status_code))

    print("\n— R13: نقاط تکرارشونده، طبقات از کاتالوگ، منحنی راندمان، ردیف‌های ثابت انبار —")
    check("EPPART: پمپ و موتور از تیپ الکتروپمپ",
          _e12.row(_parse('EPPART("233/13+24", "pump")'), {}) == "233/13"
          and _e12.row(_parse('EPPART("233/13+24", "motor")'), {}) == 24
          and _e12.row(_parse('EPPART("", "pump")'), {}) is None)
    with app.app_context():
        from app.models import WorkflowStage as _WS13
        _t13 = [st.title for st in _WS13.query.filter_by(stage_number=1).all()]
    check("عنوان مرحله‌ی ۱: «اعلام خرابی مشاهده شده از سمت بهره بردار»",
          "اعلام خرابی مشاهده شده از سمت بهره بردار" in _t13, str(_t13))
    r13 = c.post("/api/form-builder/sections", json={"code": "t_r13_p1", "title": "نقطه ۱ آزمون", "columns": 5,
                                                     "repeat_group": "t_points", "show_on_entry": False})
    sec13 = (r13.get_json() or {}).get("data", {})
    check("بخش با «گروه تکرارشونده» ساخته می‌شود", r13.status_code == 200 and sec13.get("repeat_group") == "t_points"
          and sec13.get("columns") == 5, str(sec13)[:200])
    up13 = c.put(f"/api/form-builder/sections/{sec13.get('id')}", json={"repeat_group": ""})
    check("گروه تکرارشونده برداشته می‌شود", up13.get_json()["data"].get("repeat_group") is None)
    f13 = c.post("/api/form-builder/fields", json={"field_name": "t_r13_type", "label": "تیپ آزمون",
                                                    "field_type": "text", "section_id": sec13.get("id")}).get_json()["data"]
    g13 = c.post("/api/form-builder/fields", json={"field_name": "t_r13_stages", "label": "طبقات آزمون",
                                                    "field_type": "number", "section_id": sec13.get("id"),
                                                    "stages_of": "t_r13_type"})
    check("«طبقات از کاتالوگ» برای تیپِ فیلد دیگر ذخیره می‌شود",
          g13.status_code == 200 and g13.get_json()["data"].get("stages_of") == "t_r13_type", str(g13.get_json())[:200])
    bad13 = c.put(f"/api/form-builder/fields/{g13.get_json()['data']['id']}", json={"stages_of": "no_such_field"})
    check("فیلد تیپ ناموجود پذیرفته نمی‌شود", bad13.status_code == 422, str(bad13.status_code))
    ch13 = c.post("/api/form-builder/fields", json={
        "field_name": "t_r13_chart", "label": "نمودار آزمون", "field_type": "chart", "section_id": sec13.get("id"),
        "chart_config": {"series": [{"label": "راندمان", "x": ["t_r13_stages"], "y": ["t_r13_stages"], "axis": "right",
                                     "scale": 100, "start": [0, 1],
                                     "curve": {"y": "100 * [t_r13_stages] / ([t_r13_stages] + [x])",
                                               "require": "[t_r13_stages] > 0", "note": "داده کم است"}}],
                         "y2_max": 100}})
    with app.app_context():
        from app.models import FormField as _FF13
        _cc13 = json.loads(_FF13.query.filter_by(field_name="t_r13_chart").first().chart_config or "{}")
    _s13 = (_cc13.get("series") or [{}])[0]
    check("منحنی از فرمول، ضریب و نقطه‌ی شروع با نمودار ذخیره می‌شوند", ch13.status_code == 200
          and _s13.get("curve", {}).get("y", "").startswith("100 *") and _s13.get("scale") == 100
          and _s13.get("start") == [0, 1], str(_s13)[:300])
    chb13 = c.put(f"/api/form-builder/fields/{ch13.get_json()['data']['id']}", json={
        "field_type": "chart", "chart_config": {"series": [{"label": "x", "x": ["t_r13_stages"], "y": ["t_r13_stages"],
                                                            "curve": {"y": "[nope] * [x]"}}]}})
    check("فرمول منحنی با فیلد ناموجود رد می‌شود", chb13.status_code == 422, str(chb13.status_code))
    pr13 = c.post("/api/form-builder/fields", json={
        "field_name": "t_r13_rows", "label": "تحویل تجهیز آزمون", "field_type": "wh_lines", "section_id": sec13.get("id"),
        "wh_config": {"mode": "rows", "warehouse": "equipment", "direction": "in", "reason": "pull",
                      "preset": [{"item_code": "EQ-01", "spec": 'EPPART([t_r13_type], "motor")', "qty": 1},
                                 {"item_code": ""}],
                      "preset_lock": True}})
    pid13 = (pr13.get_json() or {}).get("data", {}).get("id")
    c.put(f"/api/form-builder/fields/{pid13}", json={"label": "تحویل تجهیز آزمون", "field_type": "wh_lines",
                                                     "wh_config": {"mode": "rows", "warehouse": "equipment",
                                                                   "direction": "in", "reason": "pull"}})
    with app.app_context():
        _wc13 = json.loads(_FF13.query.filter_by(field_name="t_r13_rows").first().wh_config or "{}")
    check("ردیف‌های ثابت انبار ذخیره و با ذخیره‌ی دوباره حفظ می‌شوند",
          pr13.status_code == 200 and len(_wc13.get("preset") or []) == 1 and _wc13.get("preset_lock") is True,
          str(_wc13)[:300])
    _fb13 = c.get("/api/form-builder?all=1").get_json()["data"]
    check("فهرست کالاهای تجهیز برای ردیف ثابت در فرم‌ساز",
          any(i.get("code") == "EQ-01" for i in _fb13.get("warehouse", {}).get("items", [])))
    _parts13 = [i for i in c.get("/api/warehouse/items").get_json()["data"]["items"] if i["kind"] == "part"][:2]
    bulk13 = c.post("/api/warehouse/items/bulk", json={"items": [
        {"id": _parts13[0]["id"], "name": _parts13[0]["name"] + " آزمون", "is_active": True},
        {"id": None, "name": "قطعه‌ی تازه‌ی آزمون", "code": "T-R13", "kind": "part",
         "category": _parts13[0]["category"], "is_active": True}]})
    _items13 = {i["id"]: i for i in c.get("/api/warehouse/items").get_json()["data"]["items"]}
    check("فهرست قطعات یکجا ذخیره می‌شود (نام تازه و قطعه‌ی تازه)", bulk13.status_code == 200
          and _items13[_parts13[0]["id"]]["name"].endswith("آزمون")
          and _items13[_parts13[0]["id"]]["note"] == _parts13[0]["note"]
          and any(i["code"] == "T-R13" for i in _items13.values()), str(bulk13.get_json())[:200])
    dup13 = c.post("/api/warehouse/items/bulk", json={"items": [{"id": None, "name": "تکراری", "code": "T-R13"}]})
    check("کد تکراری قطعه پذیرفته نمی‌شود", dup13.status_code == 422)

    import zipfile as _zf13
    from app.routes.api_refdata import _unpacked
    _inner13 = io.BytesIO()
    with _zf13.ZipFile(_inner13, "w") as z:
        z.writestr("well_data.json", "[]")
    _outer13 = io.BytesIO()
    with _zf13.ZipFile(_outer13, "w") as z:
        z.writestr("ویدئومتری.xlsx", b"x")
        z.writestr("notes.txt", b"x")
        z.writestr("inner.zip", _inner13.getvalue())
    _m13, _sk13 = _unpacked("videometry", "videometry.zip", _outer13.getvalue())
    check("zip ویدئومتری/روند تولید باز می‌شود (zip در zip، JSON پیش از اکسل)",
          [n for n, _b in _m13] == ["well_data.json", "ویدئومتری.xlsx"] and _sk13 == ["notes.txt"], f"{_m13} {_sk13}")
    check("فایل غیر zip همان‌طور وارد می‌شود", _unpacked("production", "a.xlsx", b"x")[0] == [("a.xlsx", b"x")])

    with app.app_context():
        from app.extensions import db as _db13
        from app.models import Well as _W13
        from app.refdata.models import FlowTest as _FT13, VideoInspection as _VI13
        from app.refdata.profile import well_profile as _wp13
        _w13a, _w13b = _W13.query.order_by(_W13.id).limit(2).all()
        _db13.session.add_all([
            _FT13(well_name=_w13a.name, well_key="k1", main_well_id=_w13a.id, test_date="1404/07/30",
                  test_date_num=14040730, static_level=65.0),
            _VI13(facility_code="T-1", main_well_id=_w13a.id, insp_date="1403/05/12", insp_date_num=14030512,
                  static_level=60.1),
            _FT13(well_name=_w13b.name, well_key="k2", main_well_id=_w13b.id, test_date="1404/01/10",
                  test_date_num=14040110, static_level=88.0)])
        _db13.session.commit()
        _va13, _vb13 = _wp13(_w13a.id)["values"], _wp13(_w13b.id)["values"]
    check("سطح استاتیک: اول ویدئومتری (حتی قدیمی‌تر از دبی‌سنجی)", _va13.get("best.static_level") == 60.1,
          str(_va13.get("best.static_level")))
    check("سطح استاتیک: بدون ویدئومتری، از دبی‌سنجی", _vb13.get("best.static_level") == 88.0,
          str(_vb13.get("best.static_level")))
    from app.services.upgrade_r13 import KEY as _K13
    with app.app_context():
        from app.models import AppMeta as _AM13
        check("تغییرات فرم‌های R13 یک‌بار اجرا و ثبت شد", _AM13.get(_K13) == "done")

    print("\n— R14: تطبیق چاه‌ها، بانک «سوابق سنجش دبی»، اولویت‌ها، آمپر از کاتالوگ —")
    with app.app_context():
        from app.extensions import db as _db14
        from app.models import LookupCategory as _LC14, LookupItem as _LI14, Well as _W14
        from app.refdata.matching import WellIndex as _WI14
        _cat14 = _LC14.query.filter_by(code="center").first()
        _c14 = [_LI14(category_id=_cat14.id, value=v, label=v) for v in ("مرکزالف۱۴", "مرکزب۱۴")]
        _db14.session.add_all(_c14)
        _db14.session.flush()
        _wa = _W14(name="سپاد آزمون 1", pm_code="99/1/1", well_class="9918445", center_id=_c14[0].id)
        _wb = _W14(name="سوران آزمون 8", pm_code="99/2/8", well_class="9915430", center_id=_c14[1].id)
        _wbo = _W14(name="سوران آزمون 8 - قدیم", pm_code="99/2/8", well_class="9915430",
                    center_id=_c14[1].id, is_active=False)
        _db14.session.add_all([_wa, _wb, _wbo])
        _db14.session.commit()
        _ix14 = _WI14()
        _m_copy = _ix14.match(well_class="9918445", name="سوران آزمون 8", center="مرکزب۱۴")
        _m_code = _ix14.match(pm_code="99/2/8", name="سوران آزمون 8 - قدیم", center="مرکزب۱۴")
        _m_cls = _ix14.match(well_class="9915430", name="سوران آزمون 8 - قدیم", center="مرکزب۱۴")
        _m_junk = _ix14.match(well_class="ندارد", name="سوران آزمون 8", center="مرکزب۱۴")
        _m_muni = _ix14.match(name="سوران آزمون 8 (شهرداری)", center="مرکزب۱۴")
        _m_hash = _ix14.match(name="#سوران آزمون 8", center="مرکزب۱۴")
        _m_far = _ix14.match(well_class="9918445", name="چاه بی‌ربط", center="مرکزب۱۴")
        _ids14 = (_wa.id, _wb.id, _wbo.id)
    check("کلاسه‌ی چاه دیگر در فرم کپی‌شده: نام چاه آزمایش‌شده تصمیم می‌گیرد (سوران 8 با کلاسه‌ی سپاد 1)",
          _m_copy == (_ids14[1], "name"), str(_m_copy))
    check("چاه فعال بر نسخه‌ی «قدیم»/ادغام‌شده با همان کد ترجیح دارد",
          _m_code == (_ids14[1], "code") and _m_cls == (_ids14[1], "class"), f"{_m_code} {_m_cls}")
    check("کلاسه‌ی بی‌معنی («ندارد») کنار گذاشته و با نام تطبیق می‌شود", _m_junk == (_ids14[1], "name"), str(_m_junk))
    check("چاه «(شهرداری)» به چاه هم‌نام شرکت وصل نمی‌شود", _m_muni == (None, None), str(_m_muni))
    check("علامت‌های «#» در نام نادیده گرفته می‌شود", _m_hash == (_ids14[1], "name"), str(_m_hash))
    check("کلاسه‌ی اداره‌ی دیگر بدون شباهت نام پذیرفته نمی‌شود", _m_far == (None, None), str(_m_far))

    from openpyxl import Workbook as _WB14
    _wb14 = _WB14()
    _ws14 = _wb14.active
    _ws14.title = "سوابق سنجش دبی"
    _hd14 = ["کد تاسيس", "مرکز استقرار", "نام تاسيس", "تاريخ دبي سنجي", "دليل سنجش دبي", "تیپ الکتروپمپ شناور",
             "توضيحات", "عمق کلي چاه - متر", "عمق نصب الکتروپمپ - متر", "قطر داخلي لوله آبده - ميليمتر",
             "سطح ايستايي - متر", "سطح پويايي - متر", "ميزان آبدهي - Litr/S", "هد - متر",
             "ميزان فشار شبکه- اتمسفر", "شدت جريان 1 با خازن - آمپر", "شدت جريان 2 با خازن - آمپر",
             "شدت جريان 3 با خازن - آمپر", "راندمان  در فشار شبکه", "سطح ايستايي - متر2"]
    _ws14.append(["اطلاعات شناسنامه ای چاه "])
    _ws14.append(["تعداد"] + [0] * (len(_hd14) - 1))
    _ws14.append(_hd14)
    _ws14.append(["99/2/8", "مرکزب۱۴", "سوران آزمون 8 - قديم", "1405/02/01", "نصب پمپ", "384/12 + 92", "پمپ نو",
                  238, 222, 150, 154, 165, 20.8, 256.9, 8, 159, 151, 152, 0.537, 0])
    _ws14.append(["99/2/8", "مرکزب۱۴", "سوران آزمون 8 - قديم", "1405/02/01", "نصب پمپ", None, None,
                  0, 0, 0, 0, 0, 0, None, None, None, None, None, None, 0])
    _ws14.append(["99/1/1", "مرکزالف۱۴", "سپاد آزمون 1", "1403/05/01", "دبي سنجي دوره اي", None, None,
                  0, 200, 0, 0, 120, 9, None, 4, None, None, None, None, 0])
    _x14 = io.BytesIO()
    _wb14.save(_x14)
    imp14 = c.post("/api/refdata/flowrec/import", data={"files": (io.BytesIO(_x14.getvalue()), "سوابق.xlsx")},
                   content_type="multipart/form-data")
    _res14 = ((imp14.get_json() or {}).get("data") or {}).get("results") or [{}]
    check("سوابق سنجش دبی وارد می‌شود (ردیف تکراری همان روز ادغام، همه متصل به چاه با کد تاسیس)",
          imp14.status_code == 200 and _res14[0].get("rows") == 2 and _res14[0].get("merged") == 1
          and _res14[0].get("unmatched") == 0, str(_res14)[:300])
    dup14 = c.post("/api/refdata/flowrec/import", data={"files": (io.BytesIO(_x14.getvalue()), "سوابق.xlsx")},
                   content_type="multipart/form-data").get_json()["data"]["results"][0]
    check("همان فایل دوباره وارد نمی‌شود", dup14.get("duplicate") is True, str(dup14))
    with app.app_context():
        from app.refdata.models import FrRecord as _FR14
        _r14 = _FR14.query.filter_by(facility_code="99/2/8").first()
        _r14b = _FR14.query.filter_by(facility_code="99/1/1").first()
        _r14v = (_r14.main_well_id, _r14.pump_type, _r14.pump_stages, _r14.motor_kw, _r14.efficiency,
                 _r14.amps_with_capacitor, _r14.static_level, _r14.flow, _r14.notes)
        _r14bv = (_r14b.well_depth, _r14b.static_level, _r14b.install_depth, _r14b.main_well_id)
    check("ستون‌ها با برچسب خوانده می‌شوند (تیپ، راندمان ٪، آمپر سه فاز؛ ردیف دوم خالی‌ها را پاک نمی‌کند)",
          _r14v == (_ids14[1], "384", "12", 92.0, 53.7, "159/151/152", 154.0, 20.8, "پمپ نو"), str(_r14v))
    check("صفر به معنی «اندازه‌گیری نشده» خالی می‌ماند", _r14bv == (None, None, 200.0, _ids14[0]), str(_r14bv))

    with app.app_context():
        from app.refdata.models import FlowPoint as _FP14, FlowTest as _FT14, VideoInspection as _VI14
        from app.refdata.profile import well_profile as _wp14
        _t14 = _FT14(well_name="سوران آزمون 8", well_key="k14b", main_well_id=_ids14[1], test_date="1404/01/10",
                     test_date_num=14040110, static_level=150.0, install_depth=210.0, pump_type="345",
                     pump_stages="9", design_flow=18.0, net_flow=15.0)
        _t14.points = [_FP14(point_no=1, flow=18.0, dynamic_level=170.0, pressure=6.0),
                       _FP14(point_no=2, flow=15.0, dynamic_level=168.0, pressure=7.0, at_network=True)]
        _t14a = _FT14(well_name="سپاد آزمون 1", well_key="k14a", main_well_id=_ids14[0], test_date="1404/03/03",
                      test_date_num=14040303, static_level=101.0)
        _t14a.points = [_FP14(point_no=1, flow=9.5, dynamic_level=125.0, pressure=4.0, at_network=True)]
        _db14.session.add_all([_t14, _t14a])
        _db14.session.commit()
        _pb = _wp14(_ids14[1])["values"]
        _pa = _wp14(_ids14[0])["values"]
        _db14.session.add(_VI14(facility_code="99/2/8", main_well_id=_ids14[1], insp_date="1402/11/15",
                                insp_date_num=14021115, static_level=147.0))
        _db14.session.commit()
        _pbv = _wp14(_ids14[1])["values"]
    check("سوابق جدیدتر از آخرین دبی‌سنجی: مقادیر دبی‌سنجی از سوابق (تاریخ، منبع، نقطه‌ی فشار شبکه)",
          _pb.get("ft.test_date") == "1405/02/01" and _pb.get("ft.source") == "سوابق سنجش دبی"
          and _pb.get("ft.q1") == 20.8 and _pb.get("ft.press1_m") == 80.0 and "ft.q2" not in _pb
          and _pb.get("ft.net_flow") == 20.8 and _pb.get("ft.electropump") == "384/12+92", str(_pb)[:400])
    check("نقاط آزمایش قدیمی‌تر مخلوط نمی‌شود و داده‌ی پمپ قبلی کنار می‌رود",
          "ft.dyn2" not in _pb and "ft.design_flow" not in _pb and _pb.get("ft.prev_static") == 150.0
          and _pb.get("ft.prev_install_depth") == 210.0 and _pb.get("ft.last_install_date") == "1405/02/01",
          str({k: _pb.get(k) for k in ("ft.dyn2", "ft.design_flow", "ft.prev_static", "ft.last_install_date")}))
    check("سطح دینامیک و استاتیک از سنجش جدیدتر (بدون ویدئومتری)",
          _pb.get("best.dynamic_level") == 165.0 and _pb.get("best.static_level") == 154.0,
          str((_pb.get("best.dynamic_level"), _pb.get("best.static_level"))))
    check("سطح استاتیک: ویدئومتری بر همه مقدم است", _pbv.get("best.static_level") == 147.0,
          str(_pbv.get("best.static_level")))
    check("دبی‌سنجی جدیدتر از سوابق: آخرین دبی‌سنجی ملاک است",
          _pa.get("ft.test_date") == "1404/03/03" and _pa.get("ft.source") == "بانک دبی‌سنجی"
          and _pa.get("ft.q1") == 9.5 and _pa.get("best.static_level") == 101.0
          and _pa.get("best.dynamic_level") == 125.0 and _pa.get("fr.test_date") == "1403/05/01", str(_pa)[:400])
    check("تیپ الکتروپمپ فعلی از جدیدترین منبع (سوابق)", _pb.get("best.electropump") == "384/12+92",
          str(_pb.get("best.electropump")))

    fl14 = c.get("/api/refdata/flowrecords?q=99/2/8").get_json()["data"]
    check("فهرست سوابق سنجش دبی و جستجو", fl14["total"] == 1 and fl14["rows"][0]["main_well"] == "سوران آزمون 8",
          str(fl14)[:200])
    det14 = c.get(f"/api/refdata/flowrecords/{fl14['rows'][0]['id']}").get_json()["data"]
    check("جزئیات سنجش با برچسب‌های خود کاربرگ",
          any(x["label"] == "سطح ایستایی - متر" and x["value"] == 154 for x in det14["raw_labelled"]),
          str(det14.get("raw_labelled"))[:300])
    check("خروجی اکسل سوابق سنجش دبی", c.get("/api/refdata/flowrec/export.xlsx").status_code == 200)
    _sum14 = {s_["key"]: s_ for s_ in c.get("/api/refdata/summary").get_json()["data"]["sources"]}
    check("کارت «سوابق سنجش دبی» در خلاصه‌ی بانک‌ها", _sum14.get("flowrec", {}).get("rows") == 2
          and _sum14["flowrec"]["file"] == "flowrecords.db", str(_sum14.get("flowrec"))[:200])
    rel14 = c.post("/api/refdata/relink", json={}).get_json()["data"]
    check("اتصال دوباره شامل سوابق سنجش دبی", rel14.get("flowrec", {}).get("linked") == 2, str(rel14))
    lnk14 = c.post(f"/api/refdata/flowrec/{fl14['rows'][0]['id']}/link", json={"well_id": _ids14[0]})
    with app.app_context():
        _ml14 = _FR14.query.filter_by(facility_code="99/2/8").first().match_method
    check("اتصال دستی سوابق به چاه", lnk14.status_code == 200 and _ml14 == "manual", str(lnk14.get_json())[:200])
    c.post(f"/api/refdata/flowrec/{fl14['rows'][0]['id']}/link", json={"well_id": _ids14[1]})
    wl14 = c.get(f"/api/refdata/well?well_id={_ids14[1]}").get_json()["data"]
    check("پرونده‌ی چاه سوابق سنجش دبی را نشان می‌دهد", len(wl14.get("records") or []) == 1
          and any(s_["key"] == "flowrec" for s_ in wl14["profile"]["sources"]), str(wl14.get("records"))[:200])
    with app.app_context():
        from app.analytics.catalogue import get_source as _gs14
        _src14 = _gs14("fr_records")
        _rows14 = _src14._loader(None) if _src14 else []
        from app.models.report import Report as _Rep14
        _tpl14 = _Rep14.query.filter_by(template_key="ref_flowrecords").first()
    check("گزارش‌ساز: منبع «سوابق سنجش دبی» و الگوی گزارش آن",
          _src14 is not None and len(_rows14) == 2 and _tpl14 is not None, f"{_src14} {len(_rows14)} {_tpl14}")

    from app.services import upgrade_r14 as _u14
    with app.app_context():
        from app.models import AppMeta as _AM14, FormField as _FF14, FormSection as _FS14
        check("تغییرات R14 یک‌بار اجرا و ثبت شد", _AM14.get(_u14.KEY) == "done")
        _fbp = _FF14.query.filter_by(field_name="flow_before_pull").first()
        _fbp.prefill_from = None          # an earlier check cleared what the startup set
        _db14.session.commit()
        _u14._last_flow()
        _db14.session.commit()
        _fbp = _FF14.query.filter_by(field_name="flow_before_pull").first()
        check("«آخرین دبی» (دبی قبل از کشیدن) از آخرین دبی بهره‌برداری روند تولید",
              _fbp is not None and _fbp.prefill_from == "@ref:pr.last_flow", str(_fbp and _fbp.prefill_from))
        _amp = _FF14.query.filter_by(field_name="ps_max_amp").first()
        _sec = _amp.section_id
        _db14.session.add_all([_FF14(section_id=_sec, field_name="am_pump", label="تیپ پمپ", field_type="text"),
                               _FF14(section_id=_sec, field_name="no_ste_2", label="طبقات", field_type="text")])
        _amp.field_type, _amp.formula, _amp.prefill_from = "number", None, "@ref:ft.allowed_current"
        _db14.session.commit()
        _done14 = _u14._max_amp()
        _db14.session.commit()
        _amp = _FF14.query.filter_by(field_name="ps_max_amp").first()
        _ampv = (_done14, _amp.field_type, _amp.formula, _amp.prefill_from)
        from app.analytics.formula import Evaluator as _Ev14, parse as _p14
        _cat14v = _Ev14({}).row(_p14('CAT_A("152", "4")'), {})
    check("«حداکثر آمپر مجاز» از کاتالوگ برای تیپ و طبقات پیشنهادی",
          _ampv == (1, "formula", "CAT_A([am_pump], [no_ste_2])", None), str(_ampv))
    check("CAT_A جریان نامی کاتالوگ را برمی‌گرداند", _cat14v == 3.9, str(_cat14v))

    print("\n— R15: کارتابل به تفکیک اداره، جدول ردیفی، جستجو، انبار به تفکیک تیپ —")
    bd15 = c.get("/api/workflow/board?days=3650")
    _bd15 = (bd15.get_json() or {}).get("data") or {}
    check("«وضعیت چاه‌ها» به تفکیک اداره (رفته / مانده)", bd15.status_code == 200
          and isinstance(_bd15.get("offices"), list) and set(_bd15.get("totals", {})) >= {"todo", "done", "waiting"},
          str(_bd15)[:200])
    check("متولی هم «وضعیت چاه‌ها»ی خودش را می‌بیند", kahani.get("/api/workflow/board").status_code == 200)
    _inb15 = (c.get("/api/workflow/inbox").get_json() or {}).get("data") or []
    check("ردیف‌های کارتابل اداره‌ی چاه را دارند", all("well_center" in r for r in _inb15), str(_inb15[:1])[:200])
    with app.app_context():
        from app.models import WorkflowStage as _WS15
        _st15 = _WS15.query.filter(_WS15.stage_number > 0).first()
        _st15_id = _st15.id
    up15 = c.put(f"/api/workflow/stages/{_st15_id}", json={"show_refdata": False})
    with app.app_context():
        _sr15 = _WS15.query.get(_st15_id).show_refdata
    check("فرایندساز: نمایش «اطلاعات بانک‌های اطلاعاتی» در هر مرحله خاموش می‌شود",
          up15.status_code == 200 and _sr15 is False, str(up15.get_json())[:200])
    c.put(f"/api/workflow/stages/{_st15_id}", json={"show_refdata": True})

    g15 = c.post("/api/form-builder/sections", json={"code": "t_r15_grid", "title": "جدول آزمون", "layout": "grid",
                                                     "grid_label": "ردیف {n}", "show_on_entry": False})
    _g15 = (g15.get_json() or {}).get("data") or {}
    check("بخش «جدول ردیفی» ساخته می‌شود", g15.status_code == 200 and _g15.get("layout") == "grid"
          and _g15.get("grid_label") == "ردیف {n}", str(_g15)[:200])
    for n15 in (1, 2):
        c.post("/api/form-builder/fields", json={"field_name": f"t15_q{n15}", "label": f"ردیف {n15} — آبدهی",
                                                 "field_type": "number", "section_id": _g15.get("id")})
    c.post("/api/form-builder/fields", json={"field_name": "t15_note", "label": "یادداشت", "field_type": "text",
                                             "section_id": _g15.get("id")})
    from app.services.gridlayout import grid_cells as _gc15, strip_row as _sr15f
    with app.app_context():
        from app.models import FormSection as _FS15
        _gs15 = _FS15.query.get(_g15.get("id"))
        _cells15 = _gc15(_gs15, [f.field_name for f in _gs15.fields])
        _strip15 = _sr15f("ردیف 2 — آبدهی", _gs15, 2)
    check("ردیف‌ها از شماره‌ی انتهای نام؛ فیلد بی‌شماره بیرون جدول", _cells15 == {"t15_q1": ("t15_q", 1), "t15_q2": ("t15_q", 2)}
          and _strip15 == "آبدهی", f"{_cells15} {_strip15}")
    c.put(f"/api/form-builder/sections/{_g15.get('id')}", json={"layout": ""})
    with app.app_context():
        check("چیدمان دوباره ستونی می‌شود", _FS15.query.get(_g15.get("id")).layout is None)
        from app.models import FormField as _FF15
        _pc15 = _FF15.query.filter_by(field_name="pump_curr").first()
        _mc15 = _FF15.query.filter_by(field_name="motor_curr").first()
    check("تیپ پمپ و موتور دراپ‌دان (با جستجو) است", _pc15 is not None and _pc15.field_type == "select"
          and _mc15.field_type == "select", f"{_pc15 and _pc15.field_type} {_mc15 and _mc15.field_type}")

    from app.services.workflow import chart_inputs as _ci15
    _cin15 = _ci15({"series": [{"x": ["fc_m3min1"], "y": ["fc_eff1"],
                                "curve": {"y": "100 * [fc_b] / ([fc_b] + [fc_a] * [x])",
                                          "require": "AND([fc_b] > 0, [fc_a] >= 0)"}}]})
    check("خلاصه‌ی نمودار: فیلدهای منحنی (a و b) هم فرستاده می‌شوند",
          _cin15 == ["fc_m3min1", "fc_eff1", "fc_b", "fc_a"], str(_cin15))
    from app.warehouse.service import register_type as _rt15, variant_of as _vo15
    check("تیپ از مشخصات: kW، تیپ/طبقه، الکتروپمپ",
          (_vo15("73.5 kW"), _vo15("384 / 10"), _vo15("384/10+73.5"), _vo15("۶۲٫۵")) == ("73.5", "384/10", "384/10+73.5", "62.5"))
    check("تیپ از نام شناسنامه", (_rt15("پمپ شناور6608/15", "پمپ شناور"), _rt15("الکتروموتور شناور30kw", "الکتروموتور شناور"),
                                 _rt15("الکتروموتور شناور247a", "الکتروموتور شناور"),
                                 _rt15("الکتروموتور شناور 9A45", "الکتروموتور شناور"),
                                 _rt15("الکتروموتور شناور24kw9a7a", "الکتروموتور شناور")) == ("6608/15", "30", "24", "45", "24"))
    with app.app_context():
        from app.extensions import db as _db15
        from app.warehouse.models import WhEquipment as _WE15, WhItem as _WI15, WhMovement as _WM15
        # by code: an earlier check renames the first part of the list
        _imp15 = _WI15.query.filter_by(code="PS-01").first()
        _sh15 = _WI15.query.filter_by(code="PS-04").first()
        _rules15 = (_imp15 and _imp15.qty_rule, _imp15 and _imp15.per_type, _sh15 and _sh15.qty_rule)
        _db15.session.add(_WE15(code="MP/T15", name="پمپ شناور384/10", kind="پمپ شناور", type_label="384/10"))
        _db15.session.commit()
        _ids15 = {i.code: i.id for i in _WI15.query.all() if i.code}
    check("پروانه به تعداد طبقات و به تفکیک تیپ؛ شافت یک عدد", _rules15 == ("stages", True, "1"), str(_rules15))

    # a little process of the warehouse forms: join into one electropump, parts by type, equipment out
    sec15 = (c.post("/api/form-builder/sections", json={"code": "t_r15_wh", "title": "انبار آزمون",
                                                       "show_on_entry": False}).get_json() or {}).get("data") or {}
    for fname, ftype in (("t15_pump", "text"), ("t15_stages", "number"), ("t15_pcode", "text")):
        c.post("/api/form-builder/fields", json={"field_name": fname, "label": fname, "field_type": ftype,
                                                 "section_id": sec15.get("id")})
    jn15 = c.post("/api/form-builder/fields", json={
        "field_name": "t15_join", "label": "الکتروپمپ آزمون", "field_type": "wh_lines", "section_id": sec15.get("id"),
        "wh_config": {"mode": "rows", "warehouse": "equipment", "direction": "in", "reason": "assembled",
                      "condition": False, "default_condition": "assembled", "join_into": "EQ-03",
                      "serial_source": "register",
                      "preset": [{"item_code": "EQ-01", "spec": "[t15_stages]", "serial": "[t15_pcode]"}],
                      "preset_lock": True}})
    pt15 = c.post("/api/form-builder/fields", json={
        "field_name": "t15_parts", "label": "قطعات آزمون", "field_type": "wh_lines", "section_id": sec15.get("id"),
        "wh_config": {"mode": "parts", "warehouse": "parts", "direction": "out", "reason": "assembly",
                      "equipment_type": "پمپ شناور", "columns": ["total", "installed_new", "installed_repair"],
                      "list_all": True, "check_stock": True, "stages_field": "t15_stages", "type_field": "t15_pump",
                      "equipment_field": "t15_pcode", "equipment_out": True, "equipment_item": "EQ-02"}})
    with app.app_context():
        _jc15 = _FF15.query.filter_by(field_name="t15_join").first()
        _jcfg15 = json.loads(_jc15.wh_config or "{}")
        _pcfg15 = json.loads(_FF15.query.filter_by(field_name="t15_parts").first().wh_config or "{}")
    check("فرم‌ساز تنظیمات تازه‌ی انبار را نگه می‌دارد (اتصال، پلاک از شناسنامه، بدون وضعیت، کل/نو/تعمیری، موجودی، خروج تجهیز)",
          jn15.status_code == 200 and pt15.status_code == 200 and _jcfg15.get("join_into") == "EQ-03"
          and _jcfg15.get("condition") is False and _jcfg15.get("serial_source") == "register"
          and _jcfg15["preset"][0].get("serial") == "[t15_pcode]" and _pcfg15.get("list_all") is True
          and _pcfg15.get("equipment_out") is True and _pcfg15.get("stages_field") == "t15_stages"
          and _pcfg15.get("columns") == ["total", "installed_new", "installed_repair"], f"{_jcfg15} {_pcfg15}"[:400])
    with app.app_context():
        from app.models import WorkflowInstance as _WIn15
        from app.warehouse.service import post_stage as _ps15, stock as _stk15
        _inst15 = _WIn15.query.first()
        _stage15 = _inst15.workflow.stages[1]
        _db15.session.add(_WM15(jdate="1405/07/01", warehouse="equipment", direction="in", reason="pull",
                                item_id=_ids15["EQ-02"], item_name="پمپ شناور", qty=1, condition="pulled",
                                spec="384/10", variant="384/10", serial="MP/T15"))
        _db15.session.commit()
        _ps15(_inst15, _stage15, {
            "t15_pump": "384", "t15_stages": "10", "t15_pcode": "MP/T15",
            "t15_join": [{"item_id": _ids15["EQ-01"], "spec": "73.5 kW", "serial": "EM/T15"},
                         {"item_id": _ids15["EQ-02"], "spec": "384/10", "serial": "MP/T15"}],
            "t15_parts": [{"item_id": _imp15.id, "item_name": _imp15.name, "total": 10, "installed_new": 6,
                           "installed_repair": 4}]}, None)
        _db15.session.commit()
        _impn15 = _imp15.name
        _mv15 = [(m.item_name, m.direction, m.variant, m.condition, m.qty, m.serial)
                 for m in _WM15.query.filter_by(instance_id=_inst15.id, stage_number=_stage15.stage_number).all()]
        _bal15 = {(r["item_name"], r["variant"], r["condition"]): r["balance"] for r in _stk15()}
    check("دو ردیف الکتروموتور و پمپ یک الکتروپمپ متصل می‌شوند", ("الکتروپمپ کامل (مونتاژشده)", "in", "384/10+73.5",
          "assembled", 1.0, "EM/T15 + MP/T15") in _mv15, str(_mv15))
    check("پروانه به تفکیک تیپ پمپ از انبار قطعات خارج می‌شود (نو و تعمیری)",
          (_impn15, "out", "384", "new", 6.0, "MP/T15") in _mv15 and (_impn15, "out", "384", "reusable", 4.0, "MP/T15") in _mv15,
          str(_mv15))
    check("تجهیز دمونتاژشده با همان تیپ و وضعیت از انبار تجهیزات خارج می‌شود",
          ("پمپ شناور", "out", "384/10", "pulled", 1.0, "MP/T15") in _mv15
          and _bal15.get(("پمپ شناور", "384/10", "pulled")) == 0, f"{_mv15} {_bal15.get(('پمپ شناور', '384/10', 'pulled'))}")
    ss15 = kahani.get("/api/warehouse/stock-summary")
    check("موجودی انبار برای متولی (کاهانی) خوانده می‌شود", ss15.status_code == 200
          and "equipment" in (ss15.get_json() or {}).get("data", {}), str(ss15.status_code))

    # the stock count: export it, change it, import it back (it replaces the opening balance)
    from openpyxl import Workbook as _WB15
    _wbk15 = _WB15()
    _wsk15 = _wbk15.active
    _wsk15.append(["انبار", "کد کالا", "نام کالا", "تیپ", "وضعیت", "تعداد"])
    _wsk15.append(["انبار تجهیزات", "EQ-01", None, "37", "تعمیری", 5])
    _wsk15.append(["انبار قطعات", "PS-01", None, "384", "نو", 40])
    _wsk15.append(["انبار قطعات", None, "قطعه‌ی ناموجود", None, None, 3])
    _xk15 = io.BytesIO()
    _wbk15.save(_xk15)
    op15 = c.post("/api/warehouse/opening/import", data={"file": (io.BytesIO(_xk15.getvalue()), "count.xlsx")},
                  content_type="multipart/form-data")
    _op15 = (op15.get_json() or {}).get("data") or {}
    with app.app_context():
        _bal15 = {(r["item_name"], r["variant"], r["condition"]): r["balance"] for r in _stk15()}
    check("انبارگردانی از اکسل جایگزین موجودی اول دوره می‌شود (به تفکیک تیپ)",
          op15.status_code == 200 and _op15.get("rows") == 2 and _op15.get("unknown_count") == 1
          and _bal15.get(("الکتروموتور شناور", "37", "repair")) == 5, f"{_op15} {_bal15.get(('الکتروموتور شناور', '37', 'repair'))}")
    check("خروجی انبارگردانی و شناسنامه", c.get("/api/warehouse/opening.xlsx").status_code == 200
          and c.get("/api/warehouse/register.xlsx").status_code == 200)
    _wbr15 = _WB15()
    _wbr15.active.append(["کد تجهیز / پلاک", "نوع تجهیز", "تیپ", "شرح"])
    _wbr15.active.append(["EM/R15", "الکتروموتور", "45", "موتور آزمون"])
    _xr15 = io.BytesIO()
    _wbr15.save(_xr15)
    rg15 = c.post("/api/warehouse/register/import", data={"file": (io.BytesIO(_xr15.getvalue()), "reg.xlsx")},
                  content_type="multipart/form-data")
    lst15 = (c.get("/api/warehouse/register?q=EM/R15").get_json() or {}).get("data") or {}
    check("شناسنامه‌ی تجهیزات از اکسل وارد و جستجو می‌شود", rg15.status_code == 200 and lst15.get("total") == 1
          and lst15["rows"][0]["type_label"] == "45" and lst15["rows"][0]["kind"] == "الکتروموتور شناور", str(lst15)[:200])
    with app.app_context():
        from app.analytics.catalogue import get_source as _gsrc15
        from app.models.report import Report as _Rp15
        _src15 = _gsrc15("wh_stock")
        check("گزارش‌ساز: منبع «موجودی انبار به تفکیک تیپ» و الگوی آن",
              _src15 is not None and len(_src15._loader(None)) > 0
              and _Rp15.query.filter_by(template_key="wh_stock_by_type").first() is not None)
        from app.models import AppMeta as _AM15
        from app.services.upgrade_r15 import KEY as _K15
        check("تغییرات R15 یک‌بار اجرا و ثبت شد", _AM15.get(_K15) == "done")

    print("\n— R16: WAL بانک‌ها، آزمایش پمپاژ ← نصب، ورود از اکسل، جایگزینی یا افزودن —")
    import shutil as _sh16
    import sqlite3 as _sq16
    import tempfile as _tf16
    from app.services.bootstrap import set_aside_foreign_wal as _safw16
    _d16 = _tf16.mkdtemp()
    _a16, _b16 = os.path.join(_d16, "other.db"), os.path.join(_d16, "warehouse.db")
    _ca16 = _sq16.connect(_a16)
    _ca16.execute("PRAGMA journal_mode=WAL")
    _ca16.execute("PRAGMA wal_autocheckpoint=0")
    _ca16.execute("CREATE TABLE x (i INTEGER)")
    for _i16 in range(40):
        _ca16.execute("INSERT INTO x VALUES (?)", (_i16,))
    _ca16.commit()
    _cb16 = _sq16.connect(_b16)
    _cb16.execute("CREATE TABLE y (j INTEGER)")
    _cb16.execute("INSERT INTO y VALUES (7)")
    _cb16.commit()
    _cb16.close()
    _sh16.copy(_a16 + "-wal", _b16 + "-wal")
    _sh16.copy(_a16 + "-shm", _b16 + "-shm")
    _moved16 = _safw16(_b16)
    _ca16.close()
    _y16 = _sq16.connect(_b16).execute("SELECT j FROM y").fetchone()
    check("WAL جامانده‌ی دیتابیس دیگر کنار warehouse.db کنار گذاشته می‌شود (لیست قطعات خراب نمی‌شود)",
          len(_moved16) == 2 and not os.path.exists(_b16 + "-wal") and _y16 == (7,), f"{_moved16} {_y16}")
    with app.app_context():
        from app.refdata import BIND_FILES as _BF16
    check("نگهبان WAL برای همه‌ی بانک‌های refdata (از جمله انبار)", "warehouse" in _BF16)

    # a grid read from a test bench workbook
    g16 = (c.post("/api/form-builder/sections", json={"code": "t_r16_grid", "title": "نقاط آزمون", "layout": "grid",
                                                      "grid_label": "نقطه {n}", "grid_import": True,
                                                      "show_on_entry": False}).get_json() or {}).get("data") or {}
    check("بخش جدول ردیفی با «ورود از اکسل»", g16.get("grid_import") is True, str(g16)[:150])
    for n16 in (1, 2, 3):
        c.post("/api/form-builder/fields", json={"field_name": f"t16_h{n16}", "label": f"نقطه {n16} — هد (m)",
                                                 "field_type": "number", "section_id": g16.get("id")})
        c.post("/api/form-builder/fields", json={"field_name": f"t16_q{n16}", "label": f"نقطه {n16} — دبی (l/s)",
                                                 "field_type": "number", "section_id": g16.get("id")})
        c.post("/api/form-builder/fields", json={"field_name": f"t16_p{n16}", "label": f"نقطه {n16} — توان",
                                                 "field_type": "formula", "formula": f"[t16_h{n16}] * [t16_q{n16}]",
                                                 "section_id": g16.get("id")})
    tp16 = c.get("/api/form-builder/sections/t_r16_grid/grid-template.xlsx")
    from openpyxl import Workbook as _WB16, load_workbook as _LW16
    _tw16 = _LW16(io.BytesIO(tp16.data)).active
    check("قالب اکسل جدول: ستون‌های ورودی و ردیف‌ها", tp16.status_code == 200
          and [x.value for x in _tw16[1]] == ["نقطه", "هد (m)", "دبی (l/s)"] and _tw16.max_row == 4,
          str([x.value for x in _tw16[1]]))
    _bw16 = _WB16()
    _bs16 = _bw16.active
    _bs16.append(["آزمایش دستگاه"])
    _bs16.append(["#", "Head (m)", "Q (l/s)"])
    _bs16.append(["نقطه 1", 240, 18])
    _bs16.append(["نقطه 3", 200, "۲۶٫۵"])
    _bx16 = io.BytesIO()
    _bw16.save(_bx16)
    gi16 = c.post("/api/form-builder/sections/t_r16_grid/grid-import",
                  data={"file": (io.BytesIO(_bx16.getvalue()), "bench.xlsx")}, content_type="multipart/form-data")
    _gv16 = ((gi16.get_json() or {}).get("data") or {}).get("values") or {}
    check("ورود از اکسل: نام دیگر ستون‌ها (Head/Q) و شماره‌ی ردیف از ستون اول",
          gi16.status_code == 200 and _gv16 == {"t16_h1": 240.0, "t16_q1": 18.0, "t16_h3": 200.0, "t16_q3": 26.5},
          str(gi16.get_json())[:250])
    check("بخش غیرجدولی ورود از اکسل ندارد",
          c.get("/api/form-builder/sections/t_r15_wh/grid-template.xlsx").status_code == 404)

    # equipment one by one: the pumping test marks it, the install takes only a marked one
    with app.app_context():
        from app.extensions import db as _db16
        from app.warehouse import service as _wh16
        from app.warehouse.models import WhItem as _WI16, WhMovement as _WM16, WhUnitMark as _WU16
        _ep16 = _WI16.query.filter_by(code="EQ-03").first()
        for _s16, _v16 in (("EM/T16 + MP/T16", "384/10+73.5"), ("EM/U16 + MP/U16", "293/12+45.5")):
            _db16.session.add(_WM16(jdate="1405/07/20", date_num=14050720, warehouse="equipment", direction="in",
                                    reason="assembled", item_id=_ep16.id, item_name=_ep16.name, qty=1,
                                    condition="assembled", spec=_v16, variant=_v16, serial=_s16))
        _db16.session.commit()
        _u16 = {u["serial"]: u for u in _wh16.equipment_units(item_id=_ep16.id)}
        check("الکتروپمپ‌های انبار تک‌به‌تک (پلاک، تیپ، آزمایش‌نشده)",
              "EM/T16 + MP/T16" in _u16 and _u16["EM/T16 + MP/T16"]["tested"] is False
              and _u16["EM/T16 + MP/T16"]["variant"] == "384/10+73.5", str(list(_u16))[:200])
        _ep16_id = _ep16.id
    ms16 = (c.post("/api/form-builder/sections", json={"code": "t_r16_wh", "title": "آزمایش و نصب آزمون",
                                                      "show_on_entry": False}).get_json() or {}).get("data") or {}
    c.post("/api/form-builder/fields", json={
        "field_name": "t16_unit", "label": "الکتروپمپ آزمایش‌شده", "field_type": "wh_lines", "section_id": ms16.get("id"),
        "wh_config": {"mode": "mark", "mark": "tested", "item_code": "EQ-03", "warehouse": "equipment"}})
    c.post("/api/form-builder/fields", json={
        "field_name": "t16_out", "label": "الکتروپمپ نصبی", "field_type": "wh_lines", "section_id": ms16.get("id"),
        "wh_config": {"mode": "rows", "warehouse": "equipment", "direction": "out", "reason": "install",
                      "units": True, "require_mark": "tested", "condition": False}})
    with app.app_context():
        from types import SimpleNamespace as _NS16
        _inst16 = _NS16(id=990016, well=None, well_name_raw="چاه آزمون", workflow=None, well_id=None)
        _stg16 = _NS16(stage_number=7, title="آزمایش آزمون")
        _row16 = lambda s: [{"item_id": _ep16_id, "item_name": "الکتروپمپ کامل (مونتاژشده)", "serial": s, "qty": 1}]  # noqa: E731
        _p_untested16 = _wh16.check_stage(_inst16, _stg16, {"t16_out": _row16("EM/T16 + MP/T16")})
        _p_none16 = _wh16.check_stage(_inst16, _stg16, {"t16_unit": []})
        _wh16.post_stage(_inst16, _stg16, {"t16_unit": _row16("EM/T16  +  MP/T16")})
        _db16.session.commit()
        _p_tested16 = _wh16.check_stage(_inst16, _stg16, {"t16_out": _row16("EM/T16 + MP/T16")})
        _p_other16 = _wh16.check_stage(_inst16, _stg16, {"t16_out": _row16("EM/U16 + MP/U16")})
        _p_gone16 = _wh16.check_stage(_inst16, _stg16, {"t16_out": _row16("EM/X + MP/X")})
        _tested16 = {u["serial"]: u["tested"] for u in _wh16.equipment_units(item_id=_ep16_id)}
    check("نصبِ الکتروپمپ آزمایش‌نشده رد می‌شود", len(_p_untested16) == 1 and "آزمایش پمپاژ" in _p_untested16[0],
          str(_p_untested16))
    check("آزمایش پمپاژ بدون انتخاب الکتروپمپ ثبت نمی‌شود", len(_p_none16) == 1, str(_p_none16))
    check("ثبت آزمایش تیک «آزمایش پمپاژ انجام شده» را روی همان پلاک می‌زند (فاصله‌ها یکسان‌سازی)",
          _tested16.get("EM/T16 + MP/T16") is True and _tested16.get("EM/U16 + MP/U16") is False, str(_tested16))
    check("الکتروپمپ آزمایش‌شده برای نصب برداشته می‌شود؛ دیگری نه؛ ناموجود نه",
          _p_tested16 == [] and len(_p_other16) == 1 and "موجود نیست" in (_p_gone16 or [""])[0],
          f"{_p_tested16} {_p_other16} {_p_gone16}")
    un16 = kahani.get("/api/warehouse/units?item_code=EQ-03")
    check("فهرست الکتروپمپ‌ها با وضعیت آزمایش برای فرم‌ها", un16.status_code == 200
          and any(u["serial"] == "EM/T16 + MP/T16" and u["tested"] for u in (un16.get_json() or {})["data"]["units"]))

    # replace or append
    _ow16 = _WB16()
    _ow16.active.append(["انبار", "کد کالا", "نام کالا", "تیپ", "وضعیت", "تعداد", "پلاک / سریال", "آزمایش پمپاژ انجام شده"])
    _ow16.active.append(["انبار تجهیزات", "EQ-03", None, "384/12+92", "مونتاژشده", 1, "EM/O16 + MP/O16", "بله"])
    _ox16 = io.BytesIO()
    _ow16.save(_ox16)
    oa16 = c.post("/api/warehouse/opening/import", data={"file": (io.BytesIO(_ox16.getvalue()), "c.xlsx"), "mode": "append"},
                  content_type="multipart/form-data")
    with app.app_context():
        _n_open16 = _WM16.query.filter_by(reason="opening").count()
        _o_unit16 = {u["serial"]: u["tested"] for u in _wh16.equipment_units()}
    check("انبارگردانی «افزودن»: ردیف‌های قبلی می‌مانند؛ پلاک و آزمایش پمپاژ از فایل",
          oa16.status_code == 200 and (oa16.get_json() or {})["data"]["mode"] == "append" and _n_open16 >= 3
          and _o_unit16.get("EM/O16 + MP/O16") is True, f"{_n_open16} {oa16.get_json()}")
    _rw16 = _WB16()
    _rw16.active.append(["کد تجهیز / پلاک", "نوع تجهیز", "تیپ", "شرح"])
    _rw16.active.append(["EM/R16", "الکتروموتور", "55", "موتور جایگزینی"])
    _rx16 = io.BytesIO()
    _rw16.save(_rx16)
    rr16 = c.post("/api/warehouse/register/import", data={"file": (io.BytesIO(_rx16.getvalue()), "r.xlsx"), "mode": "replace"},
                  content_type="multipart/form-data")
    _rl16 = (c.get("/api/warehouse/register").get_json() or {}).get("data") or {}
    check("شناسنامه «جایگزینی کامل»: فقط ردیف‌های فایل می‌ماند", rr16.status_code == 200 and _rl16.get("total") == 1
          and _rl16["rows"][0]["code"] == "EM/R16", str(_rl16)[:150])
    with app.app_context():
        from app.refdata.models import VideoInspection as _VI16
        _vm_before16 = _VI16.query.count()
    _vw16 = _WB16()
    _vw16.active.append(["هیچ ستونی"])
    _vx16 = io.BytesIO()
    _vw16.save(_vx16)
    vr16 = c.post("/api/refdata/videometry/import", data={"files": (io.BytesIO(_vx16.getvalue()), "v.xlsx"), "mode": "replace"},
                  content_type="multipart/form-data")
    _vres16 = ((vr16.get_json() or {}).get("data") or {}).get("results") or []
    with app.app_context():
        _vm_after16 = _VI16.query.count()
    check("«جایگزینی کامل» با فایل نخواندنی: بانک به حالت قبل برمی‌گردد (پشتیبان در backups)", vr16.status_code == 200
          and any(r.get("restored") for r in _vres16) and _vm_after16 == _vm_before16
          and any(str(r.get("backup") or "").startswith("refdata_videometry_before_replace_") for r in _vres16),
          f"{_vm_before16}→{_vm_after16} {str(_vres16)[:200]}")
    with app.app_context():
        from app.refdata.models import FrRecord as _FR16
        _fr_before16 = _FR16.query.count()
    fr16 = c.post("/api/refdata/flowrec/import", data={"files": (io.BytesIO(_x14.getvalue()), "سوابق.xlsx"), "mode": "replace"},
                  content_type="multipart/form-data")
    with app.app_context():
        _fr_after16 = _FR16.query.count()
    check("«جایگزینی کامل» سوابق سنجش دبی: فقط ردیف‌های فایل می‌ماند (همان فایل دوباره خوانده می‌شود)",
          fr16.status_code == 200 and _fr_after16 == 2, f"{_fr_before16}→{_fr_after16} {str(fr16.get_json())[:200]}")
    with app.app_context():
        from app.models import AppMeta as _AM16
        from app.services.upgrade_r16 import KEY as _K16
        check("تغییرات R16 یک‌بار اجرا و ثبت شد", _AM16.get(_K16) == "done")

    print("\n— ترتیب تب‌ها —")
    check("صفحه اصلی، ثبت اطلاعات است",
          b"page-mode" in c.get("/", follow_redirects=True).data)
    check("داشبورد روی /dashboard است", c.get("/dashboard").status_code == 200)

    print("\n— شبکه —")
    from app.services.network import lan_addresses, port_is_free, primary_lan_ip
    check("تشخیص IP", bool(primary_lan_ip()))
    check("فهرست آدرس‌های شبکه", len(lan_addresses()) >= 1, str(lan_addresses()))
    check("بررسی آزاد بودن پورت", isinstance(port_is_free("0.0.0.0", 59999), bool))

    shutil.rmtree(tmp, ignore_errors=True)
    print("\n" + "=" * 60)
    print(f"  PASSED {len(PASS)}   FAILED {len(FAIL)}")
    if FAIL:
        for name in FAIL:
            print("   ✗ " + name)
    print("=" * 60)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
