"""The warehouse ledger: seeding, posting from the process, stock and reports.

A form field of type «اقلام انبار» (``wh_lines``) is a small table — item,
specification (e.g. «384/10+73.5»), plaque/serial, condition, quantity — and
its settings (``wh_config``) say which warehouse it moves stock in or out of
and why. When a stage is sent, every such answer of that stage is written to
the ledger, replacing what the same stage of the same process wrote before,
so sending a stage again (after a return) never counts the items twice.

Stock is counted by item, *type* and condition: a motor by its kW («73.5»),
a pump by type/stages («384/10»), an electropump by its whole label
(«384/10+73.5»), and a part made for one pump type (impeller, bush, stage)
by that type («384»). The starting balance is the stock count
(«انبارگردانی», reason ``opening``) — a sample until the real count is
imported.
"""
from __future__ import annotations

import io
import json
import logging
import re

from ..extensions import db
from .models import DIRECTIONS, KINDS, REASONS, WAREHOUSES, WhCondition, WhItem, WhMovement

log = logging.getLogger(__name__)

SEED_KEY = "warehouse_seed_v1"

CONDITIONS = [
    ("reusable", "قابل استفاده مجدد", "equipment,part", "#16a34a"),
    ("repair", "تعمیری", "equipment,part", "#d97706"),
    ("new", "نو (خریداری‌شده)", "equipment,part", "#2563eb"),
    ("scrap", "اسقاط", "equipment,part", "#dc2626"),
    ("assembled", "مونتاژشده — آماده نصب", "equipment", "#7c3aed"),
    ("pulled", "کشیده‌شده از چاه (منتظر بررسی)", "equipment", "#64748b"),
]

# A proposal to start with; the real catalogue replaces it («ورود از اکسل»).
ITEMS = [
    ("EQ-01", "الکتروموتور شناور", "equipment", "الکتروموتور", "دستگاه"),
    ("EQ-02", "پمپ شناور", "equipment", "پمپ", "دستگاه"),
    ("EQ-03", "الکتروپمپ کامل (مونتاژشده)", "equipment", "الکتروپمپ", "دستگاه"),
    ("PT-01", "پروانه (ایمپلر)", "part", "پمپ", "عدد"),
    ("PT-02", "دیفیوزر / بدنه طبقه", "part", "پمپ", "عدد"),
    ("PT-03", "شفت پمپ", "part", "پمپ", "عدد"),
    ("PT-04", "بوش و یاتاقان", "part", "پمپ", "عدد"),
    ("PT-05", "کوپلینگ", "part", "پمپ", "عدد"),
    ("PT-06", "سرپمپ (دهانه خروجی)", "part", "پمپ", "عدد"),
    ("PT-07", "صافی / توری مکش", "part", "پمپ", "عدد"),
    ("PT-08", "شیر یک‌طرفه", "part", "پمپ", "عدد"),
    ("PT-09", "مکانیکال سیل", "part", "الکتروموتور", "عدد"),
    ("PT-10", "بلبرینگ", "part", "الکتروموتور", "عدد"),
    ("PT-11", "اورینگ و کاسه‌نمد", "part", "الکتروموتور", "دست"),
    ("PT-12", "سیم‌پیچ الکتروموتور", "part", "الکتروموتور", "دست"),
    ("PT-13", "مایع / روغن الکتروموتور", "part", "الکتروموتور", "لیتر"),
    ("PT-14", "کابل رابط تخت", "part", "کابل", "متر"),
    ("PT-15", "مفصل کابل", "part", "کابل", "عدد"),
    ("PT-16", "گارد کابل", "part", "پمپ", "عدد"),
    ("PT-17", "پیچ و مهره", "part", "اتصالات", "عدد"),
    ("PT-18", "واشر", "part", "اتصالات", "عدد"),
]


def seed_warehouse() -> dict:
    """Condition classes always; the proposed items only into an empty book."""
    added = {"conditions": 0, "items": 0}
    have = {c.code for c in WhCondition.query.all()}
    for n, (code, label, applies, color) in enumerate(CONDITIONS):
        if code not in have:
            db.session.add(WhCondition(code=code, label=label, applies_to=applies,
                                       color=color, sort_order=n))
            added["conditions"] += 1
    if not WhItem.query.count():
        for n, (code, name, kind, cat, unit) in enumerate(ITEMS):
            db.session.add(WhItem(code=code, name=name, kind=kind, category=cat,
                                  unit=unit, sort_order=n, source="پیشنهادی"))
            added["items"] += 1
    if any(added.values()):
        db.session.commit()
    return added


def catalogue() -> dict:
    """What a form needs to draw «اقلام انبار» rows."""
    from .models import WhEquipment
    return {
        "equipment": [{"code": e.code, "name": e.name, "kind": e.kind, "type": e.type_label,
                       "status": e.status, "place": e.last_facility}
                      for e in WhEquipment.query.order_by(WhEquipment.code)],
        "items": [i.to_dict() for i in WhItem.query.filter_by(is_active=True)
                  .order_by(WhItem.kind, WhItem.sort_order, WhItem.id)],
        "conditions": [c.to_dict() for c in WhCondition.query.filter_by(is_active=True)
                       .order_by(WhCondition.sort_order)],
        "warehouses": WAREHOUSES, "kinds": KINDS, "reasons": REASONS,
        "directions": DIRECTIONS,
    }


def parse_config(text) -> dict:
    try:
        cfg = json.loads(text) if isinstance(text, str) and text.strip() else (text or {})
    except ValueError:
        cfg = {}
    cfg = dict(cfg) if isinstance(cfg, dict) else {}
    cfg.setdefault("warehouse", "equipment")
    cfg.setdefault("direction", "in")
    cfg.setdefault("reason", "manual")
    return cfg


def clean_lines(value) -> list:
    """The rows of an answer, as a list of dicts with a positive quantity."""
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError:
            return []
    out = []
    for row in value or []:
        if not isinstance(row, dict):
            continue
        name = str(row.get("item_name") or row.get("item") or "").strip()
        item_id = row.get("item_id")
        if not name and not item_id:
            continue
        try:
            qty = float(str(row.get("qty") or 1).replace("٫", "."))
        except ValueError:
            qty = 1.0
        if qty <= 0:
            continue
        spec = str(row.get("spec") or "").strip() or None
        out.append({"item_id": int(item_id) if str(item_id or "").isdigit() else None,
                    "item_name": name, "qty": qty,
                    "condition": (row.get("condition") or "").strip() or None,
                    "spec": spec,
                    "variant": (str(row.get("variant") or "").strip() or variant_of(spec)),
                    "serial": (str(row.get("serial") or "").strip() or None),
                    "note": (str(row.get("note") or "").strip() or None)})
    return out


_FA = str.maketrans("۰۱۲۳۴۵۶۷۸۹٫", "0123456789.")


def variant_of(spec):
    """The type stock is counted by, out of a specification: «73.5 kW» →
    «73.5», «384 / 10» → «384/10», «384/10+73.5» → «384/10+73.5»."""
    from ..services.epump import electropump_label, fmt_num, parse_electropump_label
    text = str(spec or "").strip().translate(_FA)
    if not text:
        return None
    t, s, m = parse_electropump_label(text)
    if t:
        return electropump_label(t, s, m)
    num = re.search(r"\d+(?:\.\d+)?", text)
    return fmt_num(float(num.group())) if num else text[:60]


def _join_rows(cfg, rows):
    """«الکتروپمپ جوین‌شده»: the motor row and the pump row become one
    electropump — its label «384/10+73.5», its plaque both plaques."""
    from ..services.epump import electropump_label, parse_electropump_label
    target = WhItem.query.filter_by(code=cfg.get("join_into")).first()
    pump_t = pump_s = motor_kw = None
    for r in rows:
        t, s, m = parse_electropump_label(r.get("spec"))
        if t:
            pump_t, pump_s = t, s
            motor_kw = motor_kw or m
        elif r.get("variant"):
            motor_kw = motor_kw or r["variant"]
    spec = electropump_label(pump_t, pump_s, motor_kw) if pump_t else None
    serials = [r["serial"] for r in rows if r.get("serial")]
    return [{"item_id": target.id if target else None,
             "item_name": target.name if target else "الکتروپمپ کامل (مونتاژشده)",
             "qty": 1.0, "condition": cfg.get("default_condition") or "assembled",
             "spec": spec, "variant": variant_of(spec), "serial": " + ".join(serials) or None,
             "note": None}]


PART_COLUMNS = [("total", "تعداد کل"),
                ("installed_new", "نصب — نو"), ("installed_repair", "نصب — تعمیری (کهنه)"),
                ("collected_new", "جمع‌آوری — نو"), ("collected_old", "جمع‌آوری — کهنه"),
                ("reusable", "قابل استفاده مجدد"), ("scrap", "اسقاط")]


def clean_parts(value) -> list:
    """A parts-form answer: rows with at least one count."""
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except ValueError:
            return []
    out = []
    for row in value or []:
        if not isinstance(row, dict) or not (row.get("item_id") or row.get("item_name")):
            continue
        counts = {}
        for key, _label in PART_COLUMNS:
            try:
                n = float(str(row.get(key) or 0).replace("٫", "."))
            except ValueError:
                n = 0
            if n > 0:
                counts[key] = n
        if not counts and not (row.get("note") or "").strip():
            continue
        out.append({"item_id": int(row["item_id"]) if str(row.get("item_id") or "").isdigit() else None,
                    "item_name": str(row.get("item_name") or "").strip(),
                    "code": str(row.get("code") or "").strip() or None,
                    "note": (str(row.get("note") or "").strip() or None), **counts})
    return out


def is_parts_answer(value) -> bool:
    rows = value
    if isinstance(rows, str):
        try:
            rows = json.loads(rows)
        except ValueError:
            return False
    return bool(rows) and isinstance(rows, list) and isinstance(rows[0], dict) and any(
        k in rows[0] for k, _l in PART_COLUMNS)


def parts_text(value) -> str:
    parts = []
    for r in clean_parts(value):
        bits = [f"{lab} {_n(r[k])}" for k, lab in PART_COLUMNS if r.get(k)]
        parts.append(f"{r['item_name']}: " + "، ".join(bits) + (f" ({r['note']})" if r.get("note") else ""))
    return "؛ ".join(parts)


def _n(q):
    return str(int(q)) if float(q).is_integer() else str(q)


def lines_text(value) -> str:
    """One answer as text, for summaries and exports."""
    if is_parts_answer(value):
        return parts_text(value)
    conds = {c.code: c.label for c in WhCondition.query.all()}
    parts = []
    for r in clean_lines(value):
        name = r["item_name"]
        if not name and r["item_id"]:
            item = db.session.get(WhItem, r["item_id"])
            name = item.name if item else "?"
        bits = [name]
        if r["spec"]:
            bits.append(r["spec"])
        if r["serial"]:
            bits.append(f"پلاک {r['serial']}")
        if r["condition"]:
            bits.append(conds.get(r["condition"], r["condition"]))
        q = r["qty"]
        bits.append(f"{int(q) if float(q).is_integer() else q} عدد")
        parts.append(" — ".join(bits))
    return "؛ ".join(parts)


def post_stage(instance, stage, data: dict, user=None) -> int:
    """Write this stage's «اقلام انبار» answers to the ledger (replacing)."""
    from ..models import FormField
    fields = FormField.query.filter_by(field_type="wh_lines", is_active=True).all()
    if not fields:
        return 0
    from ..services.jalali import today_jalali
    jy, jm, jd = today_jalali()
    jdate = f"{jy}/{jm:02d}/{jd:02d}"
    written = 0
    for field in fields:
        if field.field_name not in data:
            continue
        WhMovement.query.filter_by(instance_id=instance.id, stage_number=stage.stage_number,
                                   field_name=field.field_name).delete()
        cfg = parse_config(field.wh_config)
        if cfg.get("mode") == "parts":
            written += _post_parts(instance, stage, field, cfg, data, user, jdate, jy, jm, jd)
            continue
        rows = clean_lines(data.get(field.field_name))
        if cfg.get("join_into") and rows:
            rows = _join_rows(cfg, rows)
        for r in rows:
            item = db.session.get(WhItem, r["item_id"]) if r["item_id"] else None
            db.session.add(WhMovement(
                jdate=jdate, date_num=jy * 10000 + jm * 100 + jd,
                warehouse=cfg["warehouse"], direction=cfg["direction"], reason=cfg["reason"],
                item_id=item.id if item else None,
                item_name=(item.name if item else r["item_name"]) or "?",
                item_kind=item.kind if item else None, unit=item.unit if item else None,
                qty=r["qty"], condition=r["condition"] or cfg.get("default_condition"),
                spec=r["spec"], variant=r.get("variant"), serial=r["serial"], note=r["note"],
                instance_id=instance.id,
                workflow_name=instance.workflow.name if instance.workflow else None,
                stage_number=stage.stage_number, stage_title=stage.title,
                field_name=field.field_name, well_id=instance.well_id,
                well_name=(instance.well.name if instance.well else instance.well_name_raw),
                user_id=getattr(user, "id", None),
                user_name=getattr(user, "full_name", None)))
            written += 1
    return written


def _post_parts(instance, stage, field, cfg, data, user, jdate, jy, jm, jd) -> int:
    """A parts form: counts per part → part actions and parts-warehouse moves.

    نصب (نو / تعمیری) leaves the parts warehouse; what was جمع‌آوری and judged
    قابل استفاده مجدد or اسقاط enters it under that condition. Every count is
    also kept as a part action, the same kind of row as the 98–05 history.

    A part stocked per pump type (impeller, bush, stage) moves under the type
    of the pump on the form (``type_field``). With ``equipment_out`` the
    equipment named on the form (``equipment_field``) leaves the equipment
    warehouse — a motor or pump taken apart is no longer one in stock."""
    from .models import WhPartAction
    WhPartAction.query.filter_by(instance_id=instance.id, stage_number=stage.stage_number,
                                 field_name=field.field_name).delete()
    eq_code = str(data.get(cfg.get("equipment_field") or "", "") or "").strip() or None
    pick = lambda k: (str(data.get(cfg.get(k) or "", "") or "").strip() or None)  # noqa: E731
    well = instance.well.name if instance.well else instance.well_name_raw
    common = dict(jdate=jdate, date_num=jy * 10000 + jm * 100 + jd, year=jy, month=jm,
                  equipment_code=eq_code, equipment_kind=cfg.get("equipment_type"),
                  related_action=cfg.get("related_action"), failure=pick("failure_field"),
                  cause=pick("cause_field"), action_done=pick("action_field"),
                  instance_id=instance.id, stage_number=stage.stage_number,
                  field_name=field.field_name, well_id=instance.well_id, well_name=well,
                  user_name=getattr(user, "full_name", None), source="workflow",
                  facility_name=well)
    moves = 0
    pump_type = _pump_type(data.get(cfg.get("type_field") or "")) if cfg.get("type_field") else None
    if not pump_type and eq_code and "پمپ" in str(cfg.get("equipment_type") or ""):
        # a pump taken apart: its type is the register's (MP/… «6608/15» → 6608)
        from .models import WhEquipment
        reg = WhEquipment.query.filter_by(code=eq_code).first()
        pump_type = _pump_type(reg.type_label) if reg is not None and reg.type_label else None
    if cfg.get("equipment_out") and eq_code:
        moves += _equipment_out(instance, stage, field, cfg, eq_code, user, jdate, jy, jm, jd, well)
    for r in clean_parts(data.get(field.field_name)):
        item = db.session.get(WhItem, r["item_id"]) if r["item_id"] else None
        name = item.name if item else r["item_name"]
        variant = pump_type if (item is not None and item.per_type and pump_type) else None
        if item is not None and item.note and "کد انباری" in item.note:
            code = item.note.replace("کد انباری", "").strip()
        elif item is not None and item.code:
            code = item.code[:-2] if item.code.endswith("-P") else item.code
        else:
            code = r.get("code")
        acts = [("installed", "نو", None, r.get("installed_new")),
                ("installed", "کهنه", None, r.get("installed_repair")),
                ("collected", "نو", None, r.get("collected_new")),
                ("collected", "کهنه", None, r.get("collected_old")),
                ("collected", None, True, r.get("reusable")),
                ("collected", None, False, r.get("scrap"))]
        for pact, state, reuse, q in acts:
            if q:
                db.session.add(WhPartAction(part_action=pact, state=state, reusable=reuse,
                                            part_code=code, part_name=name, qty=q,
                                            part_type=variant, note=r.get("note"), **common))
        # the parts warehouse: a used part put back into a motor or pump
        # comes out of the «قابل استفاده مجدد» stock that disassembly fills
        for key, direction, condition, reason in (
                ("installed_new", "out", "new", "assembly"),
                ("installed_repair", "out", "reusable", "assembly"),
                ("reusable", "in", "reusable", "disassembly"),
                ("scrap", "in", "scrap", "disassembly")):
            if r.get(key):
                db.session.add(WhMovement(
                    jdate=jdate, date_num=jy * 10000 + jm * 100 + jd, warehouse="parts",
                    direction=direction, reason=reason, item_id=item.id if item else None,
                    item_name=name or "?", item_kind="part", unit=item.unit if item else "عدد",
                    qty=r[key], condition=condition, spec=cfg.get("equipment_type"),
                    variant=variant, serial=eq_code, note=r.get("note"), instance_id=instance.id,
                    workflow_name=instance.workflow.name if instance.workflow else None,
                    stage_number=stage.stage_number, stage_title=stage.title,
                    field_name=field.field_name, well_id=instance.well_id, well_name=well,
                    user_id=getattr(user, "id", None), user_name=getattr(user, "full_name", None)))
                moves += 1
    return moves


def _pump_type(value):
    """«384» out of what a pump-type field holds («384», «384|10», «384/10»)."""
    from ..models.catalogue import norm_type
    text = str(value or "").split("|")[0].split("/")[0].strip()
    return norm_type(text) if text else None


def _equipment_out(instance, stage, field, cfg, code, user, jdate, jy, jm, jd, well) -> int:
    """The motor or pump taken apart leaves the equipment warehouse as it came
    in: the same item, type and condition as its last entry under that plaque;
    failing that, the item the form is about, typed from the register."""
    from .models import WhEquipment
    last = (WhMovement.query.filter(WhMovement.warehouse == "equipment",
                                    WhMovement.direction == "in", WhMovement.serial == code)
            .order_by(WhMovement.id.desc()).first())
    item = db.session.get(WhItem, last.item_id) if last is not None and last.item_id else None
    if item is None and cfg.get("equipment_item"):
        item = WhItem.query.filter_by(code=cfg["equipment_item"]).first()
    reg = WhEquipment.query.filter_by(code=code).first()
    variant = (last.variant if last is not None and last.variant
               else variant_of(reg.type_label) if reg is not None and reg.type_label else None)
    db.session.add(WhMovement(
        jdate=jdate, date_num=jy * 10000 + jm * 100 + jd, warehouse="equipment",
        direction="out", reason="disassembly", item_id=item.id if item else None,
        item_name=(item.name if item else cfg.get("equipment_type")) or "تجهیز",
        item_kind="equipment", unit=item.unit if item else "دستگاه", qty=1,
        condition=(last.condition if last is not None else "pulled"),
        spec=(last.spec if last is not None else (reg.type_label if reg else None)),
        variant=variant, serial=code, instance_id=instance.id,
        workflow_name=instance.workflow.name if instance.workflow else None,
        stage_number=stage.stage_number, stage_title=stage.title, field_name=field.field_name,
        well_id=instance.well_id, well_name=well, user_id=getattr(user, "id", None),
        user_name=getattr(user, "full_name", None)))
    return 1


def parts_summary(args) -> dict:
    """Parts installed and collected — by part, equipment kind, year and month."""
    from .models import WhPartAction
    q = WhPartAction.query
    if args.get("equipment_kind"):
        q = q.filter(WhPartAction.equipment_kind == args["equipment_kind"])
    if args.get("source"):
        q = q.filter(WhPartAction.source == args["source"])
    from ..refdata.textnorm import jdate_num
    if args.get("date_from") and jdate_num(args["date_from"]):
        q = q.filter(WhPartAction.date_num >= jdate_num(args["date_from"]))
    if args.get("date_to") and jdate_num(args["date_to"]):
        q = q.filter(WhPartAction.date_num <= jdate_num(args["date_to"]))
    if args.get("equipment_code"):
        q = q.filter(WhPartAction.equipment_code.contains(args["equipment_code"]))
    by_part, by_kind, by_year = {}, {}, {}
    total = {"installed_new": 0, "installed_repair": 0, "installed_unknown": 0,
             "collected": 0, "reusable": 0, "scrap": 0, "rows": 0}
    equipments = set()
    for a in q.all():
        qty = a.qty or 0
        total["rows"] += 1
        if a.equipment_code:
            equipments.add(a.equipment_code)
        key = (a.equipment_kind or "—", a.part_code or "", a.part_name or "?")
        row = by_part.setdefault(key, {"equipment_kind": key[0], "part_code": key[1], "part_name": key[2],
                                       "installed_new": 0, "installed_repair": 0, "installed_unknown": 0,
                                       "collected": 0, "reusable": 0, "scrap": 0})
        kind = by_kind.setdefault(a.equipment_kind or "—", {"equipment_kind": a.equipment_kind or "—",
                                                            "installed": 0, "collected": 0, "reusable": 0, "scrap": 0})
        year = by_year.setdefault(a.year or 0, {"year": a.year, "installed": 0, "collected": 0,
                                                "new": 0, "repair": 0, "reusable": 0, "scrap": 0})
        if a.part_action == "installed":
            k = ("installed_new" if a.state == "نو" else "installed_repair" if a.state == "کهنه"
                 else "installed_unknown")
            row[k] += qty
            total[k] += qty
            kind["installed"] += qty
            year["installed"] += qty
            if a.state == "نو":
                year["new"] += qty
            elif a.state == "کهنه":
                year["repair"] += qty
        else:
            if a.reusable is True:
                row["reusable"] += qty
                total["reusable"] += qty
                kind["reusable"] += qty
                year["reusable"] += qty
            elif a.reusable is False:
                row["scrap"] += qty
                total["scrap"] += qty
                kind["scrap"] += qty
                year["scrap"] += qty
            else:
                row["collected"] += qty
                total["collected"] += qty
                kind["collected"] += qty
                year["collected"] += qty
    parts = sorted(by_part.values(), key=lambda r: -(r["installed_new"] + r["installed_repair"]
                                                      + r["installed_unknown"] + r["collected"]))
    total["equipments"] = len(equipments)
    return {"total": total, "by_part": parts,
            "by_kind": sorted(by_kind.values(), key=lambda r: r["equipment_kind"]),
            "by_year": sorted((v for v in by_year.values() if v["year"]), key=lambda r: r["year"])}


# ── stock and reports ────────────────────────────────────────────────────────
def _filtered(args):
    q = WhMovement.query
    if args.get("warehouse"):
        q = q.filter(WhMovement.warehouse == args["warehouse"])
    if args.get("direction"):
        q = q.filter(WhMovement.direction == args["direction"])
    if args.get("condition"):
        q = q.filter(WhMovement.condition == args["condition"])
    if args.get("reason"):
        q = q.filter(WhMovement.reason == args["reason"])
    if args.get("item_id"):
        q = q.filter(WhMovement.item_id == int(args["item_id"]))
    if args.get("well"):
        q = q.filter(WhMovement.well_name.contains(args["well"]))
    from ..refdata.textnorm import jdate_num
    if args.get("date_from") and jdate_num(args["date_from"]):
        q = q.filter(WhMovement.date_num >= jdate_num(args["date_from"]))
    if args.get("date_to") and jdate_num(args["date_to"]):
        q = q.filter(WhMovement.date_num <= jdate_num(args["date_to"]))
    return q


def movements(args, limit=2000) -> list:
    return [m.to_dict() for m in _filtered(args)
            .order_by(WhMovement.date_num.desc(), WhMovement.id.desc()).limit(limit)]


def stock(args=None) -> list:
    """Balance per warehouse, item, type and condition (in − out)."""
    args = args or {}
    conds = {c.code: c.label for c in WhCondition.query.all()}
    rows = {}
    q = WhMovement.query
    if args.get("warehouse"):
        q = q.filter(WhMovement.warehouse == args["warehouse"])
    for m in q.all():
        key = (m.warehouse, m.item_id or m.item_name, m.variant or "", m.condition or "")
        r = rows.setdefault(key, {"warehouse": m.warehouse,
                                  "warehouse_label": WAREHOUSES.get(m.warehouse, m.warehouse),
                                  "item_id": m.item_id, "item_name": m.item_name,
                                  "variant": m.variant or "",
                                  "unit": m.unit, "condition": m.condition,
                                  "condition_label": conds.get(m.condition, m.condition or "—"),
                                  "in": 0.0, "out": 0.0})
        r["in" if m.direction == "in" else "out"] += m.qty or 0
    out = []
    for r in rows.values():
        r["balance"] = round(r["in"] - r["out"], 3)
        out.append(r)
    out.sort(key=lambda r: (r["warehouse"], str(r["item_name"]), _variant_key(r["variant"]),
                            str(r["condition"])))
    return out


def _variant_key(v):
    """«233/12» before «384/10» before «6608/5»; kW in numeric order."""
    nums = re.findall(r"\d+(?:\.\d+)?", str(v or ""))
    return tuple(float(n) for n in nums) or (float("inf"),)


def stock_summary() -> dict:
    """What a form shows of the stock: the equipment by type (motors, pumps,
    electropumps with how many of each condition) and every part's balance
    by type and condition — {item id: {type: {condition: n}}}."""
    rows = stock()
    equipment, parts = [], {}
    for r in rows:
        if abs(r["balance"]) < 1e-9:
            continue
        if r["warehouse"] == "equipment":
            equipment.append({k: r[k] for k in ("item_id", "item_name", "variant", "condition",
                                                "condition_label", "balance", "unit")})
        elif r["item_id"]:
            parts.setdefault(str(r["item_id"]), {}).setdefault(r["variant"] or "", {})[
                r["condition"] or ""] = r["balance"]
    return {"equipment": equipment, "parts": parts}


# ── the stock count («انبارگردانی»): the balance everything starts from ─────
OPENING_COLUMNS = [("warehouse", "انبار"), ("code", "کد کالا"), ("name", "نام کالا"),
                   ("variant", "تیپ (kW / تیپ پمپ/طبقه / تیپ پمپ)"), ("condition", "وضعیت"),
                   ("qty", "تعداد")]


def opening_xlsx() -> bytes:
    """The count as it stands — also the template to fill and import back."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    conds = {c.code: c.label for c in WhCondition.query.all()}
    wb = Workbook()
    ws = wb.active
    ws.title = "انبارگردانی"
    ws.sheet_view.rightToLeft = True
    ws.append([lab for _k, lab in OPENING_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    items = {i.id: i for i in WhItem.query.all()}
    for r in stock():
        if abs(r["balance"]) < 1e-9:
            continue
        it = items.get(r["item_id"])
        ws.append([WAREHOUSES.get(r["warehouse"], r["warehouse"]), it.code if it else None,
                   r["item_name"], r["variant"] or None, conds.get(r["condition"], r["condition"]),
                   r["balance"]])
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def import_opening(data: bytes, user=None) -> dict:
    """Replace the starting balance with a stock count from Excel (the sample
    count goes with it). Columns: انبار، کد کالا یا نام کالا، تیپ، وضعیت، تعداد."""
    from openpyxl import load_workbook
    from ..refdata.textnorm import norm_label, norm_text, to_float, to_text
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = next((i for i, r in enumerate(rows[:10]) if any(
        isinstance(v, str) and norm_label(v).startswith(("نام کالا", "کد کالا")) for v in r)), None)
    if hdr is None:
        raise ValueError("ستون «نام کالا» یا «کد کالا» در سطرهای اول فایل پیدا نشد.")
    head = [norm_label(v) if isinstance(v, str) else "" for v in rows[hdr]]

    def col(*names):
        return next((c for c, h in enumerate(head) for n in names if h.startswith(n)), None)
    c_wh, c_code, c_name = col("انبار"), col("کد کالا", "کد"), col("نام کالا", "شرح")
    c_var, c_cond, c_qty = col("تیپ"), col("وضعیت"), col("تعداد", "موجودی", "مقدار")
    if c_qty is None:
        raise ValueError("ستون «تعداد» پیدا نشد.")
    conds = {norm_text(c.label): c.code for c in WhCondition.query.all()}
    conds.update({c.code: c.code for c in WhCondition.query.all()})
    by_code = {i.code: i for i in WhItem.query.all() if i.code}
    by_name = {norm_text(i.name): i for i in WhItem.query.all()}
    from ..services.jalali import today_jalali
    jy, jm, jd = today_jalali()
    added, unknown = 0, []
    WhMovement.query.filter_by(reason="opening").delete()
    for r in rows[hdr + 1:]:
        def cell(c):
            return r[c] if c is not None and c < len(r) else None
        qty = to_float(cell(c_qty))
        if not qty or qty <= 0:
            continue
        item = by_code.get(to_text(cell(c_code)) or "") or by_name.get(norm_text(cell(c_name) or ""))
        if item is None:
            unknown.append(to_text(cell(c_name)) or to_text(cell(c_code)) or "?")
            continue
        wh_text = norm_text(cell(c_wh) or "")
        wh = ("equipment" if "تجهیز" in wh_text else "parts" if "قطعه" in wh_text
              else "equipment" if item.kind == "equipment" else "parts")
        cond_text = norm_text(cell(c_cond) or "")
        cond = conds.get(cond_text) or next((code for lab, code in conds.items()
                                             if cond_text and cond_text in lab), None)
        var = to_text(cell(c_var))
        db.session.add(WhMovement(
            jdate=f"{jy}/{jm:02d}/{jd:02d}", date_num=jy * 10000 + jm * 100 + jd, warehouse=wh,
            direction="in", reason="opening", item_id=item.id, item_name=item.name,
            item_kind=item.kind, unit=item.unit, qty=qty, condition=cond or ("new" if wh == "parts" else None),
            spec=var, variant=(variant_of(var) if item.kind == "equipment" else (_pump_type(var) if var else None)),
            note="انبارگردانی", user_id=getattr(user, "id", None),
            user_name=getattr(user, "full_name", None)))
        added += 1
    db.session.commit()
    return {"rows": added, "unknown": unknown[:50], "unknown_count": len(unknown)}


SAMPLE_NOTE = "انبارگردانی نمونه — با فایل انبارگردانی واقعی جایگزین کنید"


def seed_sample_opening() -> int:
    """A sample stock count, so the forms have stock to show and to check
    against until the real count is imported (which replaces it)."""
    if WhMovement.query.filter_by(reason="opening").first():
        return 0
    from ..services.jalali import today_jalali
    jy, jm, jd = today_jalali()
    date = dict(jdate=f"{jy}/{jm:02d}/{jd:02d}", date_num=jy * 10000 + jm * 100 + jd)
    items = {i.code: i for i in WhItem.query.all() if i.code}
    n = 0

    def add(item, wh, variant, cond, qty):
        nonlocal n
        if item is None or qty <= 0:
            return
        db.session.add(WhMovement(warehouse=wh, direction="in", reason="opening", item_id=item.id,
                                  item_name=item.name, item_kind=item.kind, unit=item.unit, qty=qty,
                                  condition=cond, spec=variant, variant=variant, note=SAMPLE_NOTE, **date))
        n += 1
    motors = {"18.5": 2, "30": 3, "37": 4, "45": 3, "55": 3, "62.5": 2, "73.5": 4, "92": 2}
    pumps = {"233/12": 2, "233/16": 1, "293/12": 3, "293/14": 2, "345/9": 2, "384/10": 3,
             "384/12": 2, "6608/15": 2, "6609/12": 1}
    for kw, q in motors.items():
        add(items.get("EQ-01"), "equipment", kw, "repair", q)
    for model, q in pumps.items():
        add(items.get("EQ-02"), "equipment", model, "repair", q)
    for label, q in (("384/10+73.5", 1), ("293/12+45.5", 1)):
        add(items.get("EQ-03"), "equipment", label, "assembled", q)
    types = ["233", "293", "345", "384", "6608", "6609"]
    for k, it in enumerate(WhItem.query.filter_by(kind="part", is_active=True).order_by(WhItem.id)):
        if it.per_type:
            for j, t in enumerate(types):
                add(it, "parts", t, "new", 20 + (k * 7 + j * 5) % 40)
                add(it, "parts", t, "reusable", (k * 3 + j * 4) % 12)
        else:
            add(it, "parts", None, "new", 5 + (k * 11) % 30)
            add(it, "parts", None, "reusable", (k * 5) % 9)
    db.session.commit()
    return n


# ── the equipment register: plaques/serials to pick from ─────────────────────
REGISTER_COLUMNS = [("code", "کد تجهیز / پلاک"), ("kind", "نوع تجهیز"), ("type_label", "تیپ"),
                    ("name", "شرح"), ("maker", "سازنده"), ("property_no", "شماره اموال"),
                    ("status", "وضعیت / محل"), ("last_facility", "آخرین محل"),
                    ("last_date", "آخرین تاریخ")]


def register_type(name, kind):
    """The type written in a register name: «پمپ شناور6608/15» → «6608/15»,
    «الکتروموتور شناور30kw» → «30», «…شناور247a» (24 kW, frame 7A) → «24»,
    «…شناور 9A45» (frame 9A, 45 kW) → «45»."""
    text = str(name or "").translate(_FA).lower().replace(" ", "")
    if "پمپ" in str(kind or "") and "موتور" not in str(kind or ""):
        m = re.search(r"(\d{2,4}[a-z]?)/(\d{1,2}[a-z]?)", text)
        return f"{m.group(1).upper()}/{m.group(2)}" if m else None
    # «9A45» / «10A110»: the frame (9A) first, the kW after it
    m = re.search(r"(?<![\d.])(\d{1,2})a(\d{1,3}(?:\.\d+)?)(?![\d.])", text)
    if m and not re.search(r"\d\s*kw", text):
        return m.group(2)
    m = re.search(r"(\d+(?:\.\d+)?)(\d{1,2}a)$", text)
    if m and float(m.group(1)) < 200:
        return m.group(1)
    m = re.search(r"(\d+(?:\.\d+)?)\s*(kw)?", text)
    return m.group(1) if m else None


def enrich_register() -> int:
    """Fill the type of every register row from its name (the sample register
    the plaque pickers search until the real one is imported)."""
    from .models import WhEquipment
    n = 0
    for e in WhEquipment.query.all():
        if not e.type_label:
            t = register_type(e.name, e.kind)
            if t:
                e.type_label = t
                n += 1
        if not e.status and e.last_facility:
            e.status = e.last_facility
            e.is_sample = True
    db.session.commit()
    return n


def register_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .models import WhEquipment
    wb = Workbook()
    ws = wb.active
    ws.title = "شناسنامه تجهیزات"
    ws.sheet_view.rightToLeft = True
    ws.append([lab for _k, lab in REGISTER_COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for e in WhEquipment.query.order_by(WhEquipment.kind, WhEquipment.code):
        ws.append([getattr(e, k) for k, _l in REGISTER_COLUMNS])
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def import_register(data: bytes) -> dict:
    """Upsert the equipment register from Excel by «کد تجهیز / پلاک»."""
    from openpyxl import load_workbook
    from ..refdata.textnorm import norm_label, norm_text, to_text
    from .models import WhEquipment
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = next((i for i, r in enumerate(rows[:10]) if any(
        isinstance(v, str) and norm_label(v).startswith(("کد تجهیز", "پلاک", "سریال")) for v in r)), None)
    if hdr is None:
        raise ValueError("ستون «کد تجهیز / پلاک» در سطرهای اول فایل پیدا نشد.")
    head = [norm_label(v) if isinstance(v, str) else "" for v in rows[hdr]]
    labels = {k: norm_label(lab) for k, lab in REGISTER_COLUMNS}

    def col(key, *extra):
        names = [labels[key].split(" /")[0]] + list(extra)
        return next((c for c, h in enumerate(head) for n in names if n and h.startswith(n)), None)
    cols = {"code": col("code", "پلاک", "سریال"), "kind": col("kind", "نوع"), "type_label": col("type_label"),
            "name": col("name"), "maker": col("maker"), "property_no": col("property_no"),
            "status": col("status", "محل"), "last_facility": col("last_facility"),
            "last_date": col("last_date")}
    added = updated = 0
    for r in rows[hdr + 1:]:
        def cell(k):
            c = cols.get(k)
            return r[c] if c is not None and c < len(r) else None
        code = to_text(cell("code"))
        if not code:
            continue
        e = WhEquipment.query.filter_by(code=code).first()
        if e is None:
            e = WhEquipment(code=code)
            db.session.add(e)
            added += 1
        else:
            updated += 1
        for k in ("kind", "type_label", "name", "maker", "property_no", "status", "last_facility",
                  "last_date"):
            v = to_text(cell(k))
            if v:
                if k == "kind":
                    v = "الکتروموتور شناور" if "موتور" in norm_text(v) else (
                        "پمپ شناور" if "پمپ" in norm_text(v) else v)
                setattr(e, k, v)
        if not e.type_label:
            e.type_label = register_type(e.name, e.kind)
        e.is_sample = False
    db.session.commit()
    return {"added": added, "updated": updated}


def summary(args) -> dict:
    """The manager's report: what moved, by warehouse, reason and condition."""
    conds = {c.code: c.label for c in WhCondition.query.all()}
    q = _filtered(args).all()
    by_reason, by_cond, by_item = {}, {}, {}
    instances = {"pull": set(), "assembled": set(), "install": set()}
    for m in q:
        key = (m.warehouse, m.direction, m.reason)
        by_reason[key] = by_reason.get(key, 0) + (m.qty or 0)
        ck = (m.warehouse, m.direction, m.condition or "")
        by_cond[ck] = by_cond.get(ck, 0) + (m.qty or 0)
        ik = (m.warehouse, m.item_name, m.condition or "", m.direction)
        by_item[ik] = by_item.get(ik, 0) + (m.qty or 0)
        if m.reason in instances and m.instance_id:
            instances[m.reason].add(m.instance_id)
    return {
        "headline": {
            "electropumps_pulled": len(instances["pull"]),
            "electropumps_assembled": len(instances["assembled"]),
            "electropumps_installed": len(instances["install"]),
            "parts_scrapped": sum(m.qty or 0 for m in q if m.condition == "scrap"),
            "parts_new": sum(m.qty or 0 for m in q if m.condition == "new" and m.direction == "out"),
            "parts_reused": sum(m.qty or 0 for m in q
                                if m.condition == "reusable" and m.direction == "out"),
            "movements": len(q),
        },
        "by_reason": [{"warehouse": WAREHOUSES.get(w, w), "direction": DIRECTIONS.get(d, d),
                       "reason": REASONS.get(r, r), "qty": round(v, 3)}
                      for (w, d, r), v in sorted(by_reason.items())],
        "by_condition": [{"warehouse": WAREHOUSES.get(w, w), "direction": DIRECTIONS.get(d, d),
                          "condition": conds.get(c, c or "—"), "qty": round(v, 3)}
                         for (w, d, c), v in sorted(by_cond.items())],
        "parts": parts_summary(args),
        "by_item": [{"warehouse": WAREHOUSES.get(w, w), "item": i,
                     "condition": conds.get(c, c or "—"), "direction": DIRECTIONS.get(d, d),
                     "qty": round(v, 3)}
                    for (w, i, c, d), v in sorted(by_item.items(), key=lambda kv: (kv[0][0], str(kv[0][1])))],
    }


def add_manual(payload: dict, user=None) -> WhMovement:
    from ..refdata.textnorm import jdate_num, to_jdate
    from ..services.jalali import today_jalali
    wh = payload.get("warehouse")
    if wh not in WAREHOUSES:
        raise ValueError("انبار را انتخاب کنید.")
    direction = payload.get("direction")
    if direction not in DIRECTIONS:
        raise ValueError("ورود یا خروج را مشخص کنید.")
    item = db.session.get(WhItem, int(payload["item_id"])) if str(payload.get("item_id") or "").isdigit() else None
    if item is None:
        raise ValueError("کالا را از فهرست اقلام انتخاب کنید.")
    try:
        qty = float(str(payload.get("qty") or "").replace("٫", "."))
    except ValueError:
        qty = 0
    if qty <= 0:
        raise ValueError("تعداد / مقدار باید بیشتر از صفر باشد.")
    jy, jm, jd = today_jalali()
    jdate = to_jdate(payload.get("jdate")) or f"{jy}/{jm:02d}/{jd:02d}"
    reason = payload.get("reason") if payload.get("reason") in REASONS else "manual"
    spec = (payload.get("spec") or "").strip() or None
    variant = (variant_of(spec) if item.kind == "equipment"
               else (_pump_type(spec) if item.per_type and spec else None))
    m = WhMovement(jdate=jdate, date_num=jdate_num(jdate), warehouse=wh, direction=direction,
                   reason=reason, item_id=item.id, item_name=item.name, item_kind=item.kind,
                   unit=item.unit, qty=qty, condition=payload.get("condition") or None,
                   spec=spec, variant=variant, serial=(payload.get("serial") or None),
                   note=(payload.get("note") or None), well_name=(payload.get("well") or None),
                   user_id=getattr(user, "id", None), user_name=getattr(user, "full_name", None))
    db.session.add(m)
    db.session.commit()
    return m


# ── items to and from Excel ──────────────────────────────────────────────────
ITEM_COLUMNS = [("code", "کد کالا"), ("name", "نام کالا"), ("kind", "نوع (تجهیز/قطعه)"),
                ("category", "گروه"), ("unit", "واحد"), ("is_active", "فعال"),
                ("note", "توضیحات"), ("qty_rule", "تعداد در هر تجهیز (طبقات / عدد)"),
                ("per_type", "موجودی به تفکیک تیپ پمپ")]


def items_xlsx() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "اقلام انبار"
    ws.sheet_view.rightToLeft = True
    ws.append([lab for _k, lab in ITEM_COLUMNS])
    for i in WhItem.query.order_by(WhItem.kind, WhItem.sort_order, WhItem.id):
        ws.append([i.code, i.name, "تجهیز" if i.kind == "equipment" else "قطعه",
                   i.category, i.unit, "بله" if i.is_active else "خیر", i.note,
                   "طبقات" if i.qty_rule == "stages" else i.qty_rule, "بله" if i.per_type else None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def import_items(data: bytes) -> dict:
    """Upsert items from a sheet: «کد کالا»/«نام کالا» columns at least."""
    from openpyxl import load_workbook
    from ..refdata.textnorm import norm_label, norm_text, to_text
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr = next((i for i, r in enumerate(rows[:10]) if any(
        isinstance(v, str) and norm_label(v).startswith(("نام کالا", "نام قطعه", "شرح کالا"))
        for v in r)), None)
    if hdr is None:
        raise ValueError("ستون «نام کالا» در سطرهای اول فایل پیدا نشد.")
    head = [norm_label(v) if isinstance(v, str) else "" for v in rows[hdr]]

    def col(*names):
        return next((c for c, h in enumerate(head) for n in names if h.startswith(n)), None)
    c_code, c_name = col("کد کالا", "کد"), col("نام کالا", "نام قطعه", "شرح کالا")
    c_kind, c_cat, c_unit = col("نوع"), col("گروه", "دسته"), col("واحد")
    c_active, c_note = col("فعال"), col("توضیحات")
    c_rule, c_type = col("تعداد در هر"), col("موجودی به تفکیک")
    added = updated = 0
    top = (db.session.query(db.func.max(WhItem.sort_order)).scalar() or 0) + 1
    for n, r in enumerate(rows[hdr + 1:]):
        def cell(c):
            return r[c] if c is not None and c < len(r) else None
        name = to_text(cell(c_name))
        if not name:
            continue
        code = to_text(cell(c_code))
        item = (WhItem.query.filter_by(code=code).first() if code else None) or \
            WhItem.query.filter_by(name=name).first()
        if item is None:
            item = WhItem(name=name, code=code, sort_order=top + n)
            db.session.add(item)
            added += 1
        else:
            updated += 1
        kind = norm_text(cell(c_kind) or "")
        item.name = name
        item.code = code or item.code
        item.kind = "equipment" if ("تجهیز" in kind or "equip" in kind.lower()) else (
            "part" if kind else item.kind or "part")
        item.category = to_text(cell(c_cat)) or item.category
        item.unit = to_text(cell(c_unit)) or item.unit or "عدد"
        active = norm_text(cell(c_active) or "")
        item.is_active = not active.startswith(("خیر", "0", "no", "غیر"))
        item.note = to_text(cell(c_note)) or item.note
        rule = norm_text(cell(c_rule) or "")
        if c_rule is not None:
            item.qty_rule = ("stages" if "طبق" in rule else
                             (rule if re.fullmatch(r"\d+(\.\d+)?", rule) else None))
        if c_type is not None:
            item.per_type = norm_text(cell(c_type) or "").startswith(("بله", "1", "yes", "دارد"))
        item.source = "ورود از اکسل"
    db.session.commit()
    return {"added": added, "updated": updated}
