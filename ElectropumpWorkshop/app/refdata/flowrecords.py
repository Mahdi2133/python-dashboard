"""سوابق سنجش دبی: the flow-measurement register into flowrecords.db.

Reads the register workbook (sheet «سوابق سنجش دبی»): one row per well and
measurement date — «کد تاسیس» (the main register's PM code), the
electropump, the levels, and one pumping point at network pressure with its
electrics. The header row is found by its «کد تاسیس» and «تاریخ دبی سنجی»
columns, each column by its own label; the statistics rows above it and the
Z-score block to its left (the same labels with a number after them) are
left out. The register writes 0 for a value that was not measured, so a 0
depth, level, flow or pipe is kept as blank. A well measured twice on one
day in the same file keeps the later row's values and the earlier row's
where the later one is blank; a newer file replaces the same measurement.
"""
from __future__ import annotations

import hashlib
import io
import json

from ..extensions import db
from .textnorm import jdate_num, name_key, norm_label, norm_text, to_float, to_jdate, to_text

# key → the column's label in the register (compared normalised, exactly)
HEAD = {
    "facility_code": "کد تاسیس", "center": "مرکز استقرار", "name": "نام تاسیس",
    "address": "آدرس", "power_account": "شماره اشتراک برق", "test_date": "تاریخ دبی سنجی",
    "test_reason": "دلیل سنجش دبی", "pump_label": "تیپ الکتروپمپ شناور",
    "last_rehab_date": "تاریخ آخرین بهسازی", "last_sholat_date": "تاریخ آخرین سابقه شولات",
    "notes": "توضیحات",
    "well_depth": "عمق کلی چاه - متر", "install_depth": "عمق نصب الکتروپمپ - متر",
    "discharge_pipe_mm": "قطر داخلی لوله آبده - میلیمتر",
    "static_level": "سطح ایستایی - متر", "dynamic_level": "سطح پویایی - متر",
    "dynamic_proposed": "سطح پویایی در دبی پیشنهادی", "drawdown": "میزان افت - متر",
    "level_drop": "میزان افت سطح آب چاه - متر", "aquifer_loss": "افت سفره",
    "calc_loss": "افت محاسبه ای", "well_loss": "افت شبکه", "aquifer_coef": "ضریب افت سفره",
    "well_coef": "ضریب افت شبکه", "efficiency_ratio": "راندمان و نرخ افت سفره",
    "flow": "میزان آبدهی - litr/s", "compatible_flow": "دبی سازگار با توان آبدهی",
    "discharge_volume": "حجم تخلیه طی سنجش دبی که همان حجم هدررفت است-مترمکعب",
    "water_column": "ستون آب - متر", "water_change": "تغییر ستون آب- m/lps", "head": "هد - متر",
    "specific_resistance": "مقاومت ویژه", "specific_capacity": "ظرفیت ویژه",
    "pressure_regulated": "آیا فشار تنظیمی است؟", "line_pressure": "فشار خط - بار",
    "net_pressure": "میزان فشار شبکه- اتمسفر",
    "voltage_on": "ولتاژ در حالت روشن - ولت", "voltage_off": "ولتاژ در حالت خاموش - ولت",
    "active_power": "توان اکتیو در فشار شبکه -kw",
    "apparent_power": "توان ظاهری در فشار شبکه - kw", "reactive_power": "توان راکتیو - kvar",
    "cos_phi": "ضریب قدرت / کسینوس فی در فشار شبکه", "capacitor_kvar": "ظرفیت خازن - kvar",
    "capacitor_type": "نوع خازن", "capacitor_pf": "ضریب قدرت خازن",
    "cap_a1": "جریان فاز خازن 1 - آمپر", "cap_a2": "جریان فاز خازن 2 - آمپر",
    "cap_a3": "جریان فاز خازن 3 - آمپر",
    "with_a1": "شدت جریان 1 با خازن - آمپر", "with_a2": "شدت جریان 2 با خازن - آمپر",
    "with_a3": "شدت جریان 3 با خازن - آمپر",
    "without_a1": "شدت جریان 1 بدون خازن - آمپر", "without_a2": "شدت جریان 2 بدون خازن - آمپر",
    "without_a3": "شدت جریان 3 بدون خازن - آمپر",
    "mech_power": "توان مکانیکی در فشار شبکه", "specific_energy": "مصرف ویژه انرژی در فشار شبکه",
    "efficiency": "راندمان در فشار شبکه", "energy_intensity": "شدت انرژی در فشار شبکه",
    "ohm_rs": "مقاومت اهمی سیم پیچی فاز با فاز کیلواهم rs",
    "ohm_rt": "مقاومت اهمی سیم پیچی فاز با فاز کیلواهم rt",
    "ohm_ts": "مقاومت اهمی سیم پیچی فاز با فاز کیلواهم ts",
    "ohm_r": "مقاومت اهمی سیم پیچی فاز با بدنه کیلواهم r",
    "ohm_s": "مقاومت اهمی سیم پیچی فاز با بدنه کیلواهم s",
    "ohm_t": "مقاومت اهمی سیم پیچی فاز با بدنه کیلواهم t",
    "meter_start": "شماره کنتور در ابتدای تخلیه", "meter_end": "شماره کنتور در انتهای تخلیه",
    "meter_status": "وضعیت کنتور",
}
TEXT = {"center", "name", "address", "power_account", "test_reason", "pump_label", "notes",
        "pressure_regulated", "capacitor_type", "meter_status"}
DATES = {"test_date", "last_rehab_date", "last_sholat_date"}
# measured values for which the register's 0 means «not measured»
ZERO_BLANK = {"well_depth", "install_depth", "discharge_pipe_mm", "static_level",
              "dynamic_level", "dynamic_proposed", "flow", "compatible_flow", "water_column",
              "head", "specific_resistance", "specific_capacity", "voltage_on", "voltage_off",
              "active_power", "apparent_power", "reactive_power", "cos_phi", "capacitor_kvar",
              "capacitor_pf", "mech_power", "specific_energy", "efficiency", "energy_intensity",
              "level_drop"}
TRIPLES = {"amps_capacitor": ("cap_a1", "cap_a2", "cap_a3"),
           "amps_with_capacitor": ("with_a1", "with_a2", "with_a3"),
           "amps_without_capacitor": ("without_a1", "without_a2", "without_a3"),
           "ohm_pp": ("ohm_rs", "ohm_rt", "ohm_ts"), "ohm_pg": ("ohm_r", "ohm_s", "ohm_t")}
# the register's label of each key, for the detail view and the raw copy
LABELS = {k: v for k, v in HEAD.items()}


def _header(rows):
    """(row index, {key: column}) of the register's header, or (None, {})."""
    for i, r in enumerate(rows[:20]):
        head = [norm_label(v) if isinstance(v, str) else "" for v in r]
        if norm_label("کد تاسیس") in head and norm_label("تاریخ دبی سنجی") in head:
            cols = {}
            for key, label in HEAD.items():
                nl = norm_label(label)
                hit = next((c for c, h in enumerate(head) if h == nl), None)
                if hit is not None:
                    cols[key] = hit
            return i, cols
    return None, {}


def rows_from_xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr, cols = _header(rows)
        if hdr is None:
            continue
        for r in rows[hdr + 1:]:
            rec = {k: (r[c] if c < len(r) else None) for k, c in cols.items()}
            if rec.get("facility_code") not in (None, "") and to_jdate(rec.get("test_date")):
                yield rec


def _value(key, raw):
    if key in DATES:
        return to_jdate(raw)
    if key in TEXT:
        return to_text(raw)
    f = to_float(raw)
    if f == 0 and key in ZERO_BLANK:
        return None
    return f


def _triple(values):
    """«53/53/53» out of three phase readings (blank when none was read)."""
    parts = [to_float(v) for v in values]
    if not any(p for p in parts):
        return None
    from ..services.epump import fmt_num
    return "/".join(fmt_num(p) if p is not None else "-" for p in parts)


def record_values(rec) -> dict:
    """The row's columns as FrRecord attributes, plus the raw labelled copy."""
    out = {k: _value(k, rec.get(k)) for k in HEAD if k in rec}
    for attr, keys in TRIPLES.items():
        out[attr] = _triple([rec.get(k) for k in keys])
    eff = out.get("efficiency")
    if eff is not None and eff <= 2:            # the register writes 0.537 for 53.7 %
        out["efficiency"] = round(eff * 100, 2)
    from ..services.epump import parse_electropump_label
    out["pump_type"], out["pump_stages"], out["motor_kw"] = parse_electropump_label(out.get("pump_label"))
    out["raw"] = {k: (v if isinstance(v, (int, float, str)) or v is None else str(v))
                  for k, v in rec.items() if v not in (None, "")}
    return out


_SKIP = {"raw", "facility_code", "test_date"} | {k for keys in TRIPLES.values() for k in keys}


def _apply(row, vals, *, merge):
    for k, v in vals.items():
        if k in _SKIP or not hasattr(row, k):
            continue
        if merge and v in (None, ""):
            continue                      # the same day twice: the earlier row fills the blanks
        setattr(row, k, v)
    raw = json.loads(row.raw_json or "{}") if merge else {}
    raw.update({k: v for k, v in vals["raw"].items()
                if not merge or k not in HEAD or _value(k, v) not in (None, "")})
    row.raw_json = json.dumps(raw, ensure_ascii=False)
    row.name_key = name_key(row.name)
    row.test_date_num = jdate_num(row.test_date)


def import_bytes(data: bytes, filename: str, *, force=False, index=None) -> dict:
    from .matching import WellIndex
    from .models import FrRecord, FrSource
    sha1 = hashlib.sha1(data).hexdigest()
    if not force and FrSource.query.filter_by(sha1=sha1, status="ok").first():
        return {"file": filename, "duplicate": True}
    rows = list(rows_from_xlsx(data))
    if not rows:
        return {"file": filename, "rows": 0,
                "error": "ستون‌های «کد تاسیس» و «تاریخ دبی سنجی» در این فایل پیدا نشد."}
    index = index or WellIndex()
    src = FrSource(file_name=filename[:400], sha1=sha1)
    db.session.add(src)
    db.session.flush()
    added = updated = merged = 0
    seen = {}
    for rec in rows:
        code = norm_text(rec["facility_code"])
        date = to_jdate(rec.get("test_date"))
        vals = record_values(rec)
        key = (code, date)
        row = seen.get(key)
        if row is not None:
            _apply(row, vals, merge=True)
            merged += 1
            continue
        row = FrRecord.query.filter_by(facility_code=code, test_date=date).first()
        if row is None:
            row = FrRecord(facility_code=code, test_date=date)
            db.session.add(row)
            added += 1
        else:
            updated += 1
        _apply(row, vals, merge=False)
        row.source_id = src.id
        seen[key] = row
        if row.match_method != "manual":
            wid, how = index.match(pm_code=code, name=row.name, center=row.center)
            row.main_well_id, row.match_method = wid, how
    src.rows = len(seen)
    db.session.commit()
    unmatched = sum(1 for r in seen.values() if not r.main_well_id)
    return {"file": filename, "rows": len(seen), "added": added, "updated": updated,
            "merged": merged, "unmatched": unmatched}


def relink(index=None) -> dict:
    from .matching import WellIndex
    from .models import FrRecord
    index = index or WellIndex()
    linked = 0
    for row in FrRecord.query.all():
        if row.match_method == "manual":
            continue
        wid, how = index.match(pm_code=row.facility_code, name=row.name, center=row.center)
        row.main_well_id, row.match_method = wid, how
        linked += bool(wid)
    db.session.commit()
    return {"linked": linked}
