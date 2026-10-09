"""دبی‌سنجی: reading the flow-test forms into flowtest.db.

The forms come in three generations of one paper template — the 1400-era
layout, the 1401–1404 layout (labels in column V) and the 1404–1405 layout
(labels in column Z, three-phase currents in three cells) — as .xls, .xlsx,
inside zips inside zips, and sometimes a whole well's history in one workbook
with a sheet per test. Nothing here relies on a cell address: every value is
found by the label printed next to it, every test-table column by its header,
every pumping point by its «کارکرد N» label. A sheet without «نام چاه» is not
a flow-test form (helper sheets, pumping-station logs) and is skipped.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import re
import zipfile

from ..extensions import db
from .textnorm import (jdate_num, name_key, norm_label, norm_text, to_float,
                       to_jdate, to_text, year_of)

log = logging.getLogger(__name__)

# key → (label variants, kind, how)
#   kind: num | text | date | year | multi (several cells joined with «/»)
#   how:  left  — the value is just left of the label (RTL forms)
#         below — the value is under the label
#         either — left if it is a number there, else below
KV = {
    "well_name": (["نام چاه"], "text", "left"),
    "test_date": (["تاریخ آزمایش"], "date", "left"),
    "pump_label": (["تیپ الکتروپمپ"], "text", "left"),
    "prev_test_date": (["تاریخ آزمایش قبلی"], "date", "left"),
    "prev_pump_label": (["تیپ الکتروپمپ قبلی", "تیپ قبلی الکتروپمپ"], "text", "left"),
    "drill_year": (["سال حفر"], "year", "left"),
    "well_type": (["نوع چاه"], "text", "left"),
    "last_install_date": (["تاریخ آخرین نصب", "تاریخ نصب الکتروپمپ"], "date", "left"),
    "well_depth": (["عمق چاه"], "num", "left"),
    "install_depth": (["عمق نصب"], "num", "left"),
    "discharge_pipe": (["لوله آبده"], "num", "left"),
    "casing": (["لوله جدار"], "num", "left"),
    "design_flow": (["دبی طراحی", "دبی مجاز"], "num", "left"),
    "impeller_flow": (["دبی پروانه"], "num", "left"),
    "starter": (["سیستم راه انداز", "نوع تابلو راه انداز"], "text", "left"),
    "panel_power": (["توان تابلو"], "num", "left"),
    "voltage": (["ولتاژ روشن/خاموش", "ولتاژ روشن"], "multi", "left"),
    "capacitor_kvar": (["ظرفیت خازن", "ظرفیت نامی خازن"], "text", "left"),
    "pf_capacitor": (["کسینوس فی خازن", "ضریب قدرت با خازن", "ضریب توان باخازن"], "multi", "left"),
    "amps_capacitor": (["آمپر خازن"], "multi", "left"),
    "amps_with_capacitor": (["آمپر با خازن"], "multi", "left"),
    "ohm_pp": (["مقاومت اهمی ف-ف"], "multi", "left"),
    "ohm_pg": (["مقاومت اهمی ف-ب"], "multi", "left"),
    "prev_static": (["سطح ایستایی قبلی"], "num", "left"),
    "well_class": (["کلاسه چاه"], "text", "left"),
    "power_subscription": (["اشتراک برق"], "text", "left"),
    "test_reason": (["دلیل آزمایش"], "text", "left"),
    "network_type": (["نوع شبکه"], "text", "left"),
    "zone": (["پهنه"], "text", "left"),
    "last_sholat": (["آخرین سابقه شولات"], "text", "left"),
    "efficiency": (["راندمان"], "num", "left"),
    "energy_intensity": (["شدت انرژی"], "num", "left"),
    "cos_phi": (["کسینوس ø", "کسینوس"], "num", "left"),
    "allowed_current": (["جریان مجاز"], "num", "left"),
    "static_level": (["سطح ایستایی", "ایستایی اندازه گیری"], "num", "either"),
    "static_video": (["ایستایی ویدئومتری"], "num", "below"),
    "net_flow": (["آبدهی در فشار شبکه"], "num", "left"),
    "net_pressure": (["فشار شبکه"], "num", "left"),
    "water_change": (["تغییرات ستون آب به ازاء دبی"], "num", "left"),
    "line_pressure": (["فشار خط"], "num", "left"),
    "set_pressure": (["فشار تنظیمی"], "text", "left"),
    "pull_reason": (["علت کشیدن پمپ"], "text", "left"),
    "last_rehab_date": (["تاریخ آخرین بهسازی"], "text", "left"),
    "prev_install_depth": (["عمق نصب قبلی"], "num", "left"),
    "meter_status": (["کالیبراسیون/وضعیت کنتور"], "text", "left"),
    "meter_brand": (["برند/سایز کنتور"], "text", "left"),
    "technician": (["تکنسین دبی سنجی", "انجام دهنده"], "text", "left"),
    "computer_code": (["رمز رایانه"], "text", "left"),
    "demand": (["دیماند قرارداد"], "text", "left"),
}
_LABEL_TO_KEY = {}
for _k, (_labels, _kind, _how) in KV.items():
    for _l in _labels:
        _LABEL_TO_KEY[norm_label(_l)] = _k
KNOWN_LABELS = set(_LABEL_TO_KEY)

# test-table headers (prefix match on the normalised header, longest first)
COLUMNS = [
    ("تغییر ستون آب", "water_change"),
    ("ستون آب", "water_column"),
    ("میزان افت", "pipe_loss"),
    ("افت", "pipe_loss"),
    ("هد", "head"),
    ("آبدهی", "flow"),
    ("سطح پویایی", "dynamic_level"),
    ("پویایی", "dynamic_level"),
    ("فشار", "pressure"),
    ("جریان های الکتریکی", "amps"),
    ("آمپرها", "amps"),
    ("آمپر", "amps"),
    ("توان مکانیکی", "mech_power"),
    ("مصرف ویژه انرژی", "specific_energy"),
    ("توان ظاهری", "apparent_power"),
    ("توان اکتیو", "active_power"),
]
_POINT = re.compile(r"^کارکرد\s*(\d)")
_OFFICE_FIX = {"_لشهر": "گلشهر", "لشهر": "گلشهر", "منزل اباد": "منزل آباد"}


# ── a sheet as a grid ────────────────────────────────────────────────────────
class Grid:
    def __init__(self, rows, name=""):
        self.rows = rows
        self.name = name
        self.nr = len(rows)
        self.nc = max((len(r) for r in rows), default=0)
        self.labels = {}                 # normalised text → [(r, c)]
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                if isinstance(v, str) and v.strip():
                    self.labels.setdefault(norm_label(v), []).append((r, c))

    def get(self, r, c):
        if 0 <= r < self.nr and 0 <= c < len(self.rows[r]):
            v = self.rows[r][c]
            if isinstance(v, str):
                v = v.strip()
            return None if v in ("", None) else v
        return None

    def find(self, label):
        return self.labels.get(norm_label(label), [])


def grids_from_bytes(data: bytes, filename: str):
    """Every sheet of an .xls/.xlsx/.xlsm as a Grid."""
    low = filename.lower()
    if low.endswith(".xls"):
        import xlrd
        try:
            wb = xlrd.open_workbook(file_contents=data)
        except Exception as exc:  # noqa: BLE001
            if "encrypted" not in str(exc).lower():
                raise
            wb = xlrd.open_workbook(file_contents=_decrypt(data))
        for sh in wb.sheets():
            rows = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
            yield Grid(rows, sh.name)
    else:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            yield Grid(rows, ws.title)


def _decrypt(data: bytes) -> bytes:
    """Excel's «read-only recommended» encryption uses a published password;
    a workbook locked with a password of its own cannot be read."""
    message = "فایل رمزدار است؛ آن را در اکسل بدون رمز ذخیره کنید و دوباره وارد کنید."
    try:
        import msoffcrypto
        office = msoffcrypto.OfficeFile(io.BytesIO(data))
        office.load_key(password="VelvetSweatshop")
        out = io.BytesIO()
        office.decrypt(out)
        return out.getvalue()
    except (KeyboardInterrupt, SystemExit):
        raise
    except BaseException as exc:  # noqa: BLE001 — missing library, wrong password, broken crypto
        raise ValueError(message) from exc


# ── one sheet → one test ─────────────────────────────────────────────────────
def _left(g, r, c, kind):
    """The value just left of a label (skipping blanks, never past a label)."""
    for cc in range(c - 1, max(-1, c - 5), -1):
        v = g.get(r, cc)
        if v is None:
            continue
        if isinstance(v, str) and norm_label(v) in KNOWN_LABELS:
            return None
        if kind == "multi":
            parts = [v]
            for c2 in range(cc - 1, max(-1, cc - 3), -1):
                w = g.get(r, c2)
                if w is None or (isinstance(w, str) and norm_label(w) in KNOWN_LABELS):
                    break
                parts.insert(0, w)
            return "/".join(to_text(p) for p in parts if to_text(p))
        return v
    return None


def _below(g, r, c):
    for rr in (r + 1, r + 2):
        v = g.get(rr, c)
        if v is not None:
            return v
    return None


def _convert(v, kind):
    if v is None:
        return None
    if kind == "num":
        return to_float(v)
    if kind == "date":
        return to_jdate(v)
    if kind == "year":
        return year_of(v)
    return to_text(v)


def _read_kv(g):
    out = {}
    for key, (labels, kind, how) in KV.items():
        value = None
        for label in labels:
            for (r, c) in g.find(label):
                if how in ("left", "either"):
                    value = _convert(_left(g, r, c, kind), kind)
                if value is None and how in ("below", "either"):
                    value = _convert(_below(g, r, c), kind)
                if value not in (None, ""):
                    break
            if value not in (None, ""):
                break
        if value not in (None, ""):
            out[key] = value
    return out


def _opinions(g):
    found = {}
    for label, cells in g.labels.items():
        if not label.startswith("نظر"):
            continue
        key = ("expert_opinion" if "بهره برداری" in label or "بهرهبرداری" in label
               else "support_opinion" if ("پشتیبانی" in label or "برق" in label) else None)
        if not key or key in found:
            continue
        r, c = cells[0]
        for rr in range(r + 1, min(g.nr, r + 12)):
            v = g.get(rr, c)
            if isinstance(v, str) and len(norm_text(v)) > 2 and norm_label(v) not in KNOWN_LABELS:
                found[key] = norm_text(v)
                break
    return found


def _design_curve(g):
    cells = g.find("داده های طراحی پمپ")
    if not cells:
        return None
    r0, c0 = cells[0]
    for r in range(r0, min(g.nr, r0 + 3)):
        heads = [c for c in range(g.nc) if isinstance(g.get(r, c), str)
                 and norm_label(g.get(r, c)) == "هد"]
        if not heads:
            continue
        ch = heads[0]
        curve = []
        for rr in range(r + 1, min(g.nr, r + 16)):
            h, q = to_float(g.get(rr, ch)), to_float(g.get(rr, ch + 1))
            if h is None or q is None:
                if curve:
                    break
                continue
            curve.append([h, q])
        return curve or None
    return None


def _points(g):
    starts = [rc for rc in g.find("داده های دبی سنجی")]
    if not starts:
        return []
    r0 = min(r for r, _c in starts)
    header = None
    for r in range(r0 + 1, min(g.nr, r0 + 5)):
        if any(isinstance(g.get(r, c), str) and norm_label(g.get(r, c)).startswith("آبدهی")
               for c in range(g.nc)):
            header = r
            break
    if header is None:
        return []
    # the column holding «کارکرد N» bounds the table on the right
    label_col, rows = None, []
    for r in range(header + 1, min(g.nr, header + 12)):
        for c in range(g.nc):
            v = g.get(r, c)
            if isinstance(v, str) and _POINT.match(norm_label(v)):
                if label_col is None or c == label_col:
                    label_col = c
                    rows.append((r, v))
                break
    if label_col is None:
        return []
    cols = []
    for c in range(label_col):
        v = g.get(header, c)
        if not isinstance(v, str):
            continue
        raw = norm_text(v).lower()
        h = norm_label(v)
        if h.startswith("سطح ایستایی") or h.startswith("ایستایی") or h.startswith("شاخص"):
            continue
        for prefix, key in COLUMNS:
            if h.startswith(prefix):
                if key == "flow":
                    key = "flow_m3h" if ("m3" in raw or "m³" in raw) else "flow"
                cols.append((c, key))
                break
    if not cols:
        return []
    spans = {}
    for i, (c, key) in enumerate(cols):
        end = cols[i + 1][0] if i + 1 < len(cols) else label_col
        spans[c] = range(c, end if key == "amps" else c + 1)
    points = []
    for r, label in rows:
        text = norm_text(label)          # with the «(فشار شبکه)» kept
        m = re.search(r"(\d)", text)
        if not m:
            continue
        p = {"point_no": int(m.group(1)), "label": text, "at_network": "فشار" in text}
        for c, key in cols:
            if key == "amps":
                parts = [to_text(g.get(r, cc)) for cc in spans[c]]
                parts = [x for x in parts if x and x not in ("-", "ـ", "&")]
                if parts:
                    p["amps"] = "/".join(parts)
                continue
            f = to_float(g.get(r, c))
            if f is not None and key not in p:
                p[key] = f
        if p.get("flow") is None and p.get("flow_m3h") is not None:
            p["flow"] = round(p["flow_m3h"] / 3.6, 3)
        p.pop("flow_m3h", None)
        if not any(p.get(k) not in (None, 0) for k in ("flow", "head", "dynamic_level", "pressure")):
            continue
        points.append(p)
    # one point per «کارکرد» number (a form sometimes repeats a label)
    seen, unique = set(), []
    for p in points:
        if p["point_no"] in seen:
            continue
        seen.add(p["point_no"])
        unique.append(p)
    return unique


def parse_grid(g, *, fallback_name=None):
    """A test dict, or None for a sheet that is not a flow-test form."""
    # a flow-test form names the well AND has the test table or test date;
    # the production report and the calculation sheets name wells too
    if not g.find("نام چاه") or not (g.find("داده های دبی سنجی") or g.find("تاریخ آزمایش")):
        return None
    kv = _read_kv(g)
    name = kv.get("well_name") or fallback_name
    if not name:
        return None
    test = dict(kv)
    test["well_name"] = norm_text(name)
    test["sheet"] = g.name
    test.update(_opinions(g))
    if not test.get("test_date"):
        test["test_date"] = to_jdate(g.name)
    eff = test.get("efficiency")
    if eff is not None and eff <= 1.5:
        test["efficiency"] = round(eff * 100, 2)
    from ..services.epump import parse_electropump_label
    t, s, m = parse_electropump_label(test.get("pump_label"))
    test["pump_type"], test["pump_stages"], test["motor_kw"] = t, s, m
    test["design_curve"] = _design_curve(g)
    test["points"] = _points(g)
    if test.get("net_flow") is None:
        net = [p for p in test["points"] if p.get("at_network")]
        if net:
            test["net_flow"] = net[0].get("flow")
            test["net_pressure"] = test.get("net_pressure") or net[0].get("pressure")
    return test


# ── files, zips and the import ───────────────────────────────────────────────
def _zip_name(info):
    if info.flag_bits & 0x800:
        return info.filename
    raw = info.filename.encode("cp437")
    for enc in ("utf-8", "cp720"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return info.filename


def iter_files(data: bytes, filename: str, folder: str = "", depth: int = 0):
    """(folder, name, bytes) of every spreadsheet, unpacking zips in zips."""
    low = filename.lower()
    if low.endswith(".zip") and depth < 6:
        z = zipfile.ZipFile(io.BytesIO(data))
        base = os.path.splitext(filename)[0]
        for info in z.infolist():
            if info.is_dir():
                continue
            name = _zip_name(info).replace("\\", "/")
            sub = "/".join(p for p in [folder, base] + name.split("/")[:-1] if p)
            yield from iter_files(z.read(info), name.split("/")[-1], sub, depth + 1)
    elif low.endswith((".xls", ".xlsx", ".xlsm")) and not os.path.basename(low).startswith("~$"):
        yield folder, filename, data
    else:
        yield folder, filename, None            # reported as not a spreadsheet


def office_from(folder: str):
    """«دبي سنجي اداره امام علي 1403» → «امام علی»; a bare «دوستي» folder → «دوستی»."""
    parts = [p for p in folder.split("/") if p]
    for part in reversed(parts):
        text = norm_text(part)
        if "دبی سنجی" not in text:
            continue
        text = re.sub(r"دبی سنجی|اداره|\d{4}", " ", text)
        text = re.sub(r"\s+", " ", text).strip()
        return _OFFICE_FIX.get(text, text) or None
    # no «دبی سنجی …» folder: the first folder under the upload names the office
    for part in parts[1:2]:
        text = re.sub(r"\d{4}", " ", norm_text(part)).strip()
        if text and not re.search(r"\d", text):
            return _OFFICE_FIX.get(text, text)
    return None


COLUMNS_OF_TEST = None


def _test_columns():
    global COLUMNS_OF_TEST
    if COLUMNS_OF_TEST is None:
        from .models import FlowTest
        COLUMNS_OF_TEST = {c.name for c in FlowTest.__table__.columns}
    return COLUMNS_OF_TEST


def import_bytes(data: bytes, filename: str, *, force=False, index=None) -> dict:
    """Import one upload (a spreadsheet or a zip of them)."""
    from .matching import WellIndex
    from .models import FlowPoint, FlowSource, FlowTest
    index = index or WellIndex()
    stats = {"files": 0, "sheets": 0, "tests_added": 0, "tests_updated": 0,
             "duplicates": 0, "skipped_files": [], "errors": [], "unmatched": 0}
    cols = _test_columns()
    for folder, name, blob in iter_files(data, filename):
        if blob is None:
            if not name.lower().endswith((".db", ".tmp")):
                stats["skipped_files"].append(f"{folder}/{name}".strip("/"))
            continue
        stats["files"] += 1
        sha1 = hashlib.sha1(blob).hexdigest()
        if not force and FlowSource.query.filter_by(sha1=sha1, status="ok").first():
            stats["duplicates"] += 1
            continue
        office = office_from(folder)
        src = FlowSource(file_name=name[:400], folder=folder[:400], office=office, sha1=sha1)
        db.session.add(src)
        db.session.flush()
        sheets = 0
        try:
            tests = []
            for g in grids_from_bytes(blob, name):
                sheets += 1
                t = parse_grid(g)
                if t:
                    if not t.get("test_date"):
                        t["test_date"] = to_jdate(name)
                    tests.append(t)
        except Exception as exc:  # noqa: BLE001 — one bad file never stops the rest
            src.status, src.message = "error", str(exc)[:500]
            stats["errors"].append(f"{name}: {str(exc)[:160]}")
            continue
        stats["sheets"] += sheets
        src.sheets = sheets
        src.tests = len(tests)
        for t in tests:
            key = name_key(t["well_name"])
            row = FlowTest.query.filter_by(well_key=key, test_date=t.get("test_date")).first()
            if row is not None and len(row.points) > len(t["points"]) and not force:
                stats["duplicates"] += 1
                continue
            if row is None:
                row = FlowTest(well_key=key)
                db.session.add(row)
                stats["tests_added"] += 1
            else:
                row.points.clear()
                stats["tests_updated"] += 1
            raw = {k: v for k, v in t.items() if k not in ("points", "design_curve")}
            for k, v in t.items():
                if k in cols and k not in ("id", "points", "design_curve", "raw_json"):
                    setattr(row, k, v)
            row.source_id = src.id
            row.office = office
            row.test_date_num = jdate_num(t.get("test_date"))
            row.design_curve = json.dumps(t["design_curve"]) if t.get("design_curve") else None
            row.raw_json = json.dumps(raw, ensure_ascii=False, default=str)
            wid, how = index.match(well_class=t.get("well_class"), name=t["well_name"],
                                   center=office)
            row.main_well_id, row.match_method = wid, how
            if wid is None:
                stats["unmatched"] += 1
            for p in t["points"]:
                row.points.append(FlowPoint(**{k: v for k, v in p.items()
                                               if k in FlowPoint.__table__.columns}))
        db.session.commit()
    return stats


def relink(index=None) -> dict:
    """Re-match every test to the register (after wells were added/renamed)."""
    from .matching import WellIndex
    from .models import FlowTest
    index = index or WellIndex()
    linked = 0
    for row in FlowTest.query.filter(FlowTest.match_method.is_(None) |
                                     (FlowTest.match_method != "manual")).all():
        wid, how = index.match(well_class=row.well_class, name=row.well_name, center=row.office)
        row.main_well_id, row.match_method = wid, how
        linked += bool(wid)
    db.session.commit()
    return {"linked": linked}
