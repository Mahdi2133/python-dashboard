"""ویدئومتری: camera inspections of the wells into videometry.db.

Accepts the inspection workbook («سابقه»: نوع تاسیس، مرکز، کد تاسیس، نام،
آدرس، تاریخ اقدام، عمق چاه، سطح ایستابی، مشبک، ترمیم، پارگی، تغییر جدار،
گرفتگی، توضیحات) and the JSON export of the same data. One row per well and
inspection date; a newer file replaces the same inspection.
"""
from __future__ import annotations

import hashlib
import io
import json
import re

from ..extensions import db
from .textnorm import jdate_num, name_key, norm_label, norm_text, to_float, to_jdate, to_text

HEAD = {
    "kind": ["نوع تاسیس"], "center": ["مرکز استقرار"], "facility_code": ["کد تاسیس"],
    "name": ["نام تاسیس", "نام چاه"], "address": ["آدرس"], "date": ["تاریخ اقدام"],
    "depth": ["عمق چاه"], "swl": ["سطح ایستابی", "سطح ایستآبی"],
    "screen_start": ["عمق شروع مشبک جدار"], "no_screen": ["محدوده عمق جدار فاقد مشبک"],
    "repair": ["محدوده عمق مشاهده ترمیم جدار"], "tear": ["محدوده عمق پارگی جدار"],
    "change": ["محدوده عمق تغییر جدار"], "clog": ["محدوده عمق گرفتگی مشبک"],
    "notes": ["توضیحات"],
}
_RANGE = re.compile(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)(?:\s*:\s*(\S+))?")


def parse_ranges(text):
    """«0 - 125\\n180 - 186.5» / «142 - 218.7 : کم» → [{from, to, sev?}]."""
    if text in (None, ""):
        return []
    if isinstance(text, list):
        return text
    out = []
    for m in _RANGE.finditer(norm_text(str(text).replace("\n", " ; "))):
        item = {"from": float(m.group(1)), "to": float(m.group(2))}
        if m.group(3):
            item["sev"] = m.group(3)
        out.append(item)
    return out


def rows_from_xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = next((i for i, r in enumerate(rows[:10])
                    if any(isinstance(v, str) and norm_label(v) == "کد تاسیس" for v in r)), None)
        if hdr is None:
            continue
        head = [norm_label(v).replace("ستابی", "ستابی") if isinstance(v, str) else "" for v in rows[hdr]]
        cols = {}
        for key, labels in HEAD.items():
            for lab in labels:
                nl = norm_label(lab)
                hit = next((c for c, h in enumerate(head) if h.startswith(nl)), None)
                if hit is not None:
                    cols[key] = hit
                    break
        for r in rows[hdr + 1:]:
            rec = {k: (r[c] if c < len(r) else None) for k, c in cols.items()}
            if rec.get("facility_code"):
                yield rec


def rows_from_json(data: bytes):
    for x in json.loads(data.decode("utf-8-sig")):
        d = x.get("date")
        yield {"kind": x.get("type"), "center": x.get("center"),
               "facility_code": x.get("code"), "name": x.get("name"),
               "address": x.get("address"),
               "date": (d.get("raw") or d.get("fmt")) if isinstance(d, dict) else d,
               "depth": x.get("depth"), "swl": x.get("swl"),
               "screen_start": x.get("screen_start"), "no_screen": x.get("no_screen"),
               "repair": x.get("repair"), "tear": x.get("tear"), "change": x.get("change"),
               "clog": x.get("clog"), "notes": x.get("notes")}


def import_bytes(data: bytes, filename: str, *, force=False, index=None) -> dict:
    from .matching import WellIndex
    from .models import VideoInspection, VideoSource
    sha1 = hashlib.sha1(data).hexdigest()
    if not force and VideoSource.query.filter_by(sha1=sha1, status="ok").first():
        return {"file": filename, "duplicate": True}
    rows = list(rows_from_json(data) if filename.lower().endswith(".json")
                else rows_from_xlsx(data))
    index = index or WellIndex()
    src = VideoSource(file_name=filename[:400], sha1=sha1)
    db.session.add(src)
    db.session.flush()
    added = updated = 0
    for rec in rows:
        code = norm_text(rec["facility_code"])
        date = to_jdate(rec.get("date"))
        row = VideoInspection.query.filter_by(facility_code=code, insp_date=date).first()
        if row is None:
            row = VideoInspection(facility_code=code, insp_date=date)
            db.session.add(row)
            added += 1
        else:
            updated += 1
        row.name = to_text(rec.get("name"))
        row.name_key = name_key(row.name)
        row.center = to_text(rec.get("center"))
        row.address = to_text(rec.get("address"))
        row.insp_date_num = jdate_num(date)
        row.depth = to_float(rec.get("depth"))
        row.static_level = to_float(rec.get("swl"))
        row.screen_start = to_float(rec.get("screen_start"))
        lists = {k: parse_ranges(rec.get(k)) for k in ("no_screen", "repair", "tear", "change", "clog")}
        for k, v in lists.items():
            setattr(row, k, json.dumps(v, ensure_ascii=False) if v else None)
        row.notes = to_text(rec.get("notes"))
        row.defect_count = sum(len(lists[k]) for k in ("repair", "tear", "change", "clog"))
        row.source_id = src.id
        wid, how = index.match(pm_code=code, name=row.name, center=row.center)
        row.main_well_id, row.match_method = wid, how
    src.rows = len(rows)
    db.session.commit()
    return {"file": filename, "rows": len(rows), "added": added, "updated": updated}


def relink(index=None) -> dict:
    from .matching import WellIndex
    from .models import VideoInspection
    index = index or WellIndex()
    linked = 0
    for row in VideoInspection.query.all():
        if row.match_method == "manual":
            continue
        wid, how = index.match(pm_code=row.facility_code, name=row.name, center=row.center)
        row.main_well_id, row.match_method = wid, how
        linked += bool(wid)
    db.session.commit()
    return {"linked": linked}
