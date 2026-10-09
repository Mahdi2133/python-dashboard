"""روند تولید: the yearly production-trend workbooks into production.db.

One workbook per year, one row per well: the well's register columns (code,
center, pump, motor, last rehab…) then twelve months each of production (m³),
running hours, average flow (l/s), pressure and pressure type. Columns are
found by their headers and the month blocks by the group titles above them
(«تولید ماهیانه 1405», «کارکرد ماهیانه …», «دبی متوسط ماهیانه …», …), so a
workbook with a column added or moved still reads correctly. The register
columns of the newest year win; the monthly figures of every year are kept.
"""
from __future__ import annotations

import hashlib
import io
import logging
import re

from ..extensions import db
from .textnorm import name_key, norm_label, norm_text, to_float, to_jdate, to_text

log = logging.getLogger(__name__)

BASE = {
    "name": ["نام چاه"],
    "facility_code": ["کد تاسیس"],
    "center_code": ["مرکز استقرار"],
    "urban_rural": ["شهری/ روستایی", "شهری/روستایی"],
    "zone": ["نام پهنه ای که تولید چاه به آن ارسال می شود", "پهنه"],
    "low_run_reason": ["علت کارکرد کمتر از انتظار"],
    "pump_install_date": ["تاریخ نصب الکتروپمپ"],
    "pump_label": ["تیپ پمپ"],
    "motor_label": ["تیپ الکتروموتور"],
    "last_rehab_failure": ["خرابی مشاهده شده طی آخرین بهسازی"],
    "last_rehab_date": ["تاریخ آخرین بهسازی"],
    "meter_status": ["آخرین وضعیت کنتور در بازه گزارش"],
    "location_status": ["وضعیت تعیین محل چاه"],
    "relocations": ["دفعات صدور برگ جابجایی"],
}
GROUPS = [("نوع فشار ماهیانه", "pressure_type"), ("فشار ماهیانه", "pressure"),
          ("دبی متوسط ماهیانه", "avg_flow"), ("کارکرد ماهیانه", "hours"),
          ("تولید ماهیانه", "production")]


def _year_from(text):
    m = re.search(r"(1[34]\d{2})", norm_text(text or ""))
    return int(m.group(1)) if m else None


def parse_workbook(data: bytes, filename: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header = next((i for i, r in enumerate(rows[:10])
                   if any(norm_label(v) == "کد تاسیس" for v in r if isinstance(v, str))), None)
    if header is None:
        raise ValueError("ستون «کد تاسیس» در سطرهای اول فایل پیدا نشد.")
    head = [norm_label(v) if isinstance(v, str) else "" for v in rows[header]]
    cols = {}
    for key, labels in BASE.items():
        for lab in labels:
            nl = norm_label(lab)
            hit = next((c for c, h in enumerate(head) if h == nl), None)
            if hit is None:
                hit = next((c for c, h in enumerate(head) if h.startswith(nl)), None)
            if hit is not None:
                cols[key] = hit
                break
    groups, year = {}, None
    titles = rows[header - 1] if header else []
    for c, v in enumerate(titles):
        if not isinstance(v, str):
            continue
        t = norm_label(v)
        for prefix, key in GROUPS:
            if t.startswith(norm_label(prefix)) and key not in groups:
                groups[key] = c
                year = year or _year_from(v)
                break
    year = year or _year_from(filename)
    if not year:
        raise ValueError("سال گزارش از عنوان ستون‌ها یا نام فایل خوانده نشد.")
    if "facility_code" not in cols or "name" not in cols:
        raise ValueError("ستون‌های «نام چاه» و «کد تاسیس» لازم است.")
    wells = []
    for r in rows[header + 1:]:
        code = to_text(r[cols["facility_code"]]) if cols["facility_code"] < len(r) else None
        name = to_text(r[cols["name"]]) if cols["name"] < len(r) else None
        if not code or not name:
            continue
        base = {k: (r[c] if c < len(r) else None) for k, c in cols.items()}
        months = {}
        for key, start in groups.items():
            for m in range(12):
                c = start + m
                v = r[c] if c < len(r) else None
                if v in (None, ""):
                    continue
                val = to_text(v) if key == "pressure_type" else to_float(v)
                if val is None:
                    continue
                months.setdefault(m + 1, {})[key] = val
        wells.append({"base": base, "months": months})
    return {"year": year, "wells": wells}


def import_bytes(data: bytes, filename: str, *, force=False, index=None) -> dict:
    from ..services.epump import parse_paren
    from .matching import WellIndex
    from .models import ProdMonth, ProdSource, ProdWell
    sha1 = hashlib.sha1(data).hexdigest()
    if not force and ProdSource.query.filter_by(sha1=sha1, status="ok").first():
        return {"file": filename, "duplicate": True}
    parsed = parse_workbook(data, filename)
    year = parsed["year"]
    index = index or WellIndex()
    src = ProdSource(file_name=filename[:400], year=year, sha1=sha1)
    db.session.add(src)
    existing = {w.facility_code: w for w in ProdWell.query.all()}
    added = months = 0
    for item in parsed["wells"]:
        b = item["base"]
        code = norm_text(b["facility_code"])
        w = existing.get(code)
        if w is None:
            w = ProdWell(facility_code=code, name=norm_text(b["name"]))
            db.session.add(w)
            existing[code] = w
            added += 1
        if not w.snapshot_year or year >= w.snapshot_year:
            w.name = norm_text(b["name"])
            w.name_key = name_key(w.name)
            w.center_code = to_text(b.get("center_code"))
            w.urban_rural = to_text(b.get("urban_rural"))
            w.zone = to_text(b.get("zone"))
            w.low_run_reason = to_text(b.get("low_run_reason"))
            w.pump_install_date = to_jdate(b.get("pump_install_date"))
            w.pump_label = to_text(b.get("pump_label"))
            w.pump_type, w.pump_stages = parse_paren(w.pump_label)
            w.motor_label = to_text(b.get("motor_label"))
            model, kw = parse_paren(w.motor_label)
            w.motor_model, w.motor_kw = model, to_float(kw)
            w.last_rehab_failure = to_text(b.get("last_rehab_failure"))
            w.last_rehab_date = to_jdate(b.get("last_rehab_date"))
            w.meter_status = to_text(b.get("meter_status"))
            w.location_status = to_text(b.get("location_status"))
            rel = to_float(b.get("relocations"))
            w.relocations = int(rel) if rel is not None else None
            w.snapshot_year = year
            wid, how = index.match(pm_code=code, name=w.name)
            w.main_well_id, w.match_method = wid, how
        db.session.flush()
        have = {(m.year, m.month): m for m in w.months if m.year == year}
        for month, vals in item["months"].items():
            row = have.get((year, month))
            if row is None:
                row = ProdMonth(year=year, month=month)
                w.months.append(row)
            for k, v in vals.items():
                setattr(row, k, v)
            months += 1
    src.wells, src.months = len(parsed["wells"]), months
    db.session.commit()
    return {"file": filename, "year": year, "wells": len(parsed["wells"]),
            "wells_added": added, "months": months}


def relink(index=None) -> dict:
    from .matching import WellIndex
    from .models import ProdWell
    index = index or WellIndex()
    linked = 0
    for w in ProdWell.query.all():
        if w.match_method == "manual":
            continue
        wid, how = index.match(pm_code=w.facility_code, name=w.name)
        w.main_well_id, w.match_method = wid, how
        linked += bool(wid)
    db.session.commit()
    return {"linked": linked}
