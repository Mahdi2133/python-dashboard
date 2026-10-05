"""Pump catalogue lookups: a model's curve, and values read off it.

Flows are in l/s throughout, as everywhere else in the forms; the catalogue
stores m³/h, as the manufacturer prints it, and converts (÷ 3.6).

Values between two catalogue points are read off the straight line joining
them; outside the printed range nothing is returned, so no form shows a
number the catalogue does not support.
"""
import io
import logging
import re

from ..extensions import db

log = logging.getLogger(__name__)

SEED_KEY = "pump_catalogue_gazar_v1"
_CACHE = None


def invalidate():
    global _CACHE
    _CACHE = None


def _load():
    global _CACHE
    if _CACHE is None:
        from ..models.catalogue import PumpCatalogModel, norm_stages, norm_type
        cache = {}
        rows = (PumpCatalogModel.query.filter_by(is_active=True)
                .order_by(PumpCatalogModel.sort_order, PumpCatalogModel.id).all())
        for m in rows:
            key = (norm_type(m.pump_type), norm_stages(m.stages))
            if key in cache:
                continue                      # first brand wins
            pts = sorted(((p.q_m3h / 3.6, p.head, p.pump_eff) for p in m.points),
                         key=lambda x: x[0])
            cache[key] = {"id": m.id, "title": m.title, "full_title": m.full_title,
                          "brand": m.brand, "trim": m.trim, "kw": m.motor_kw,
                          "hp": m.motor_hp, "a": m.current_a,
                          "motor_eff": m.motor_eff, "pts": pts}
        _CACHE = cache
    return _CACHE


def model(pump_type, stages):
    from ..models.catalogue import norm_stages, norm_type
    if pump_type in (None, "") or stages in (None, ""):
        return None
    if isinstance(pump_type, list):
        pump_type = pump_type[0] if pump_type else None
    if isinstance(stages, list):
        stages = stages[0] if stages else None
    return _load().get((norm_type(pump_type), norm_stages(stages)))


def compact():
    """Everything a form needs to read curves without asking again."""
    return {f"{t}|{s}": {"title": m["title"], "full_title": m["full_title"],
                         "kw": m["kw"], "a": m["a"],
                         "motor_eff": m["motor_eff"] if m["motor_eff"] is not None else 90.0,
                         "pts": [[round(q, 4), h, e] for q, h, e in m["pts"]]}
            for (t, s), m in _load().items()}


def _interp(x, x1, y1, x2, y2):
    if x2 == x1:
        return y1
    return y1 + (y2 - y1) * (x - x1) / (x2 - x1)


def flow_at_head(m, head):
    """Q (l/s) where the curve reaches ``head``; None outside the curve."""
    if m is None or head is None:
        return None
    pts = m["pts"]
    for (q1, h1, _), (q2, h2, _) in zip(pts, pts[1:]):
        lo, hi = min(h1, h2), max(h1, h2)
        if lo - 1e-9 <= head <= hi + 1e-9:
            return _interp(head, h1, q1, h2, q2)
    return None


def head_at_flow(m, flow):
    if m is None or flow is None:
        return None
    pts = m["pts"]
    for (q1, h1, _), (q2, h2, _) in zip(pts, pts[1:]):
        if q1 - 1e-9 <= flow <= q2 + 1e-9:
            return _interp(flow, q1, h1, q2, h2)
    return None


def eff_at_flow(m, flow):
    if m is None or flow is None:
        return None
    pts = [(q, e) for q, _h, e in m["pts"] if e is not None]
    for (q1, e1), (q2, e2) in zip(pts, pts[1:]):
        if q1 - 1e-9 <= flow <= q2 + 1e-9:
            return _interp(flow, q1, e1, q2, e2)
    return None


def call(name, args):
    """The CAT_* formula functions (see analytics.formula.FUNCTIONS)."""
    def num(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    m = model(args[0] if args else None, args[1] if len(args) > 1 else None)
    if m is None:
        return None
    x = num(args[2]) if len(args) > 2 else None
    if name == "CAT_Q":
        return flow_at_head(m, x)
    if name == "CAT_H":
        return head_at_flow(m, x)
    if name == "CAT_EFF":
        return eff_at_flow(m, x)
    if name == "CAT_KW":
        return m["kw"]
    if name == "CAT_A":
        return m["a"]
    if name == "CAT_MEFF":
        return m["motor_eff"] if m["motor_eff"] is not None else 90.0
    if name == "CAT_TITLE":
        return m["full_title"]
    return None


# ── seed ────────────────────────────────────────────────────────────────────
def seed_catalogue() -> dict:
    from ..models.catalogue import PumpCatalogModel, PumpCatalogPoint
    from ..models.meta import AppMeta
    from .seed_catalogue_data import BRAND, MODELS, MOTOR_EFF
    if AppMeta.get(SEED_KEY):
        return {"pump_catalogue": 0}
    if PumpCatalogModel.query.count():
        AppMeta.set(SEED_KEY, "existing")
        db.session.commit()
        return {"pump_catalogue": 0}
    for n, (t, s, trim, kw, hp, a, kg, lp, lt, page, pts) in enumerate(MODELS):
        m = PumpCatalogModel(brand=BRAND, pump_type=t, stages=s, trim=trim, motor_kw=kw,
                             motor_hp=hp, current_a=a, weight_kg=kg, pump_length_mm=lp,
                             total_length_mm=lt, motor_eff=MOTOR_EFF, sort_order=n,
                             source=f"کاتالوگ گازار — صفحه {page}")
        m.points = [PumpCatalogPoint(q_m3h=q, head=h, pump_eff=e) for q, h, e in pts]
        db.session.add(m)
    AppMeta.set(SEED_KEY, "done")
    db.session.commit()
    invalidate()
    log.info("Pump catalogue seeded: %d models", len(MODELS))
    return {"pump_catalogue": len(MODELS)}


# ── Excel ───────────────────────────────────────────────────────────────────
COLUMNS = [
    ("brand", "برند"), ("pump_type", "تیپ پمپ"), ("stages", "تعداد طبقات"),
    ("trim", "قطر پروانه (مدل a)"), ("motor_kw", "توان الکتروموتور (kW)"),
    ("motor_hp", "توان (HP)"), ("current_a", "جریان (A)"), ("weight_kg", "وزن (kg)"),
    ("pump_length_mm", "طول پمپ (mm)"), ("total_length_mm", "طول مجموعه (mm)"),
    ("motor_eff", "راندمان الکتروموتور (%)"),
    ("q_m3h", "دبی (m³/h)"), ("q_ls", "دبی (l/s)"), ("head", "هد (m)"),
    ("pump_eff", "راندمان پمپ (%)"),
]


def export_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from ..models.catalogue import PumpCatalogModel
    wb = Workbook()
    ws = wb.active
    ws.title = "کاتالوگ پمپ"
    ws.sheet_view.rightToLeft = True
    ws.append([label for _k, label in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for m in PumpCatalogModel.query.order_by(PumpCatalogModel.sort_order, PumpCatalogModel.id):
        for p in m.points:
            row = []
            for key, _label in COLUMNS:
                if key in ("q_m3h", "head", "pump_eff"):
                    row.append(getattr(p, key))
                elif key == "q_ls":
                    row.append(p.q_ls)
                else:
                    row.append(getattr(m, key))
            ws.append(row)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _header_key(text):
    t = re.sub(r"\s+", " ", str(text or "")).strip().lower()
    if not t:
        return None
    if "برند" in t or "brand" in t:
        return "brand"
    if "عنوان" in t:
        return None                                  # «عنوان تیپ پمپ» is derived
    if "تیپ" in t and "پمپ" in t:
        return "pump_type"
    if "طبق" in t or "stage" in t:
        return "stages"
    if "پروانه" in t or "trim" in t:
        return "trim"
    if "راندمان" in t and ("موتور" in t or "motor" in t):
        return "motor_eff"
    if "راندمان" in t and "کل" in t:
        return None
    if "راندمان" in t or "eff" in t:
        return "pump_eff"
    if "hp" in t:
        return "motor_hp"
    if "توان" in t and ("موتور" in t or "kw" in t):
        return "motor_kw"
    if "جریان" in t or "current" in t or t in ("a", "(a)"):
        return "current_a"
    if "وزن" in t or "kg" in t:
        return "weight_kg"
    if "طول" in t and "پمپ" in t:
        return "pump_length_mm"
    if "طول" in t:
        return "total_length_mm"
    if "l/s" in t or "لیتر" in t:
        return "q_ls"
    if "m3" in t or "m³" in t or "مکعب" in t or "مترمکعب" in t or "آبدهی" in t or "دبی" in t:
        return "q_m3h"
    if "هد" in t or "head" in t or "ارتفاع" in t:
        return "head"
    return None


def import_xlsx(data: bytes) -> dict:
    """Upsert every model the sheet lists; models it leaves out stay as they are.

    Accepts this page's own export and the pump-test workbook's «Database»
    sheet (one row per curve point: type, stages, motor, flow, head, …).
    """
    from openpyxl import load_workbook
    from ..models.catalogue import PumpCatalogModel, PumpCatalogPoint, norm_stages, norm_type
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    best = None
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        for n, row in enumerate(rows):
            keys = [_header_key(c) for c in row]
            if "pump_type" in keys and "stages" in keys and "head" in keys:
                best = (ws, n, keys)
                break
            if n > 10:
                break
        if best:
            break
    if not best:
        raise ValueError("ستون‌های «تیپ پمپ»، «تعداد طبقات» و «هد» در فایل پیدا نشد.")
    ws, hdr_row, keys = best
    groups = {}
    order = []
    for n, row in enumerate(ws.iter_rows(values_only=True)):
        if n <= hdr_row:
            continue
        rec = {}
        for k, v in zip(keys, row):
            if k and k not in rec:
                rec[k] = v
        if rec.get("pump_type") in (None, "") or rec.get("stages") in (None, ""):
            continue
        key = (str(rec.get("brand") or "گازار").strip(), norm_type(rec["pump_type"]),
               norm_stages(rec["stages"]))
        if key not in groups:
            groups[key] = {"info": rec, "pts": []}
            order.append(key)

        def num(v):
            try:
                return float(str(v).replace("٫", ".").replace(",", ""))
            except (TypeError, ValueError):
                return None
        q = num(rec.get("q_m3h"))
        if q is None and num(rec.get("q_ls")) is not None:
            q = num(rec.get("q_ls")) * 3.6
        h = num(rec.get("head"))
        if q is None or h is None:
            continue
        e = num(rec.get("pump_eff"))
        groups[key]["pts"].append((q, h, e))
    added = updated = 0
    top = (db.session.query(db.func.max(PumpCatalogModel.sort_order)).scalar() or 0) + 1
    for n, key in enumerate(order):
        g = groups[key]
        if not g["pts"]:
            continue
        brand, t, s = key
        m = PumpCatalogModel.query.filter_by(brand=brand, pump_type=t, stages=s).first()
        if m is None:
            m = PumpCatalogModel(brand=brand, pump_type=t, stages=s, sort_order=top + n,
                                 source="ورود از اکسل")
            db.session.add(m)
            added += 1
        else:
            updated += 1
        info = g["info"]
        for k in ("trim", "motor_kw", "motor_hp", "current_a", "weight_kg",
                  "pump_length_mm", "total_length_mm", "motor_eff"):
            if k in info and info[k] not in (None, ""):
                v = info[k]
                if k != "trim":
                    try:
                        v = float(v)
                    except (TypeError, ValueError):
                        continue
                setattr(m, k, v)
        m.is_active = True
        m.points = [PumpCatalogPoint(q_m3h=q, head=h, pump_eff=e)
                    for q, h, e in sorted(set(g["pts"]))]
    db.session.commit()
    invalidate()
    return {"added": added, "updated": updated, "models": added + updated}
