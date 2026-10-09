"""/api/warehouse — items, conditions, the ledger, stock and the manager's report.

The catalogue is read by every form that has an «اقلام انبار» field; the
ledger and reports need «warehouse.view»; manual movements and item edits
«warehouse.manage».
"""
from __future__ import annotations

import io

from flask import Blueprint, request, send_file

from ..extensions import db
from ..services.audit import record_audit
from ..services.auth import (FORM_READERS, current_user, permission_required,
                             permission_required_any)
from ..warehouse import service as wh
from ..warehouse.models import REASONS, WAREHOUSES, WhCondition, WhItem, WhMovement
from ._helpers import body, fail, ok

bp = Blueprint("api_warehouse", __name__, url_prefix="/api/warehouse")


@bp.get("/catalogue")
@permission_required_any(*FORM_READERS, "warehouse.view")
def catalogue():
    return ok(wh.catalogue())


def _args():
    return {k: request.args.get(k) for k in ("warehouse", "direction", "condition", "reason",
                                              "item_id", "well", "date_from", "date_to")}


@bp.get("/movements")
@permission_required("warehouse.view")
def movements():
    return ok({"rows": wh.movements(_args())})


@bp.post("/movements")
@permission_required("warehouse.manage")
def add_movement():
    try:
        m = wh.add_manual(body(), current_user())
    except ValueError as exc:
        return fail(str(exc), 422)
    record_audit("create", "warehouse", m.id,
                 summary=f"{m.to_dict()['direction_label']} «{m.item_name}» × {m.qty:g} — "
                         f"{WAREHOUSES.get(m.warehouse)}", commit=True)
    return ok(m.to_dict(), message="ثبت شد.")


@bp.delete("/movements/<int:mid>")
@permission_required("warehouse.manage")
def delete_movement(mid):
    m = db.session.get(WhMovement, mid)
    if m is None:
        return fail("ردیف پیدا نشد.", 404)
    if m.instance_id:
        return fail("این ردیف از فرایند ثبت شده است؛ از همان مرحله‌ی فرایند اصلاح کنید.", 409)
    db.session.delete(m)
    db.session.commit()
    return ok({}, message="حذف شد.")


@bp.get("/stock")
@permission_required("warehouse.view")
def stock():
    return ok({"rows": wh.stock(_args())})


@bp.get("/summary")
@permission_required("warehouse.view")
def summary():
    return ok(wh.summary(_args()))


# ── items and conditions ─────────────────────────────────────────────────────
@bp.get("/items")
@permission_required("warehouse.view")
def items():
    return ok({"items": [i.to_dict() for i in WhItem.query.order_by(
        WhItem.kind, WhItem.sort_order, WhItem.id)]})


@bp.post("/items")
@permission_required("warehouse.manage")
def save_item():
    p = body()
    name = (p.get("name") or "").strip()
    if not name:
        return fail("نام کالا را وارد کنید.", 422)
    item = db.session.get(WhItem, int(p["id"])) if p.get("id") else None
    if item is None:
        item = WhItem(sort_order=(db.session.query(db.func.max(WhItem.sort_order)).scalar() or 0) + 1,
                      source="ثبت دستی")
        db.session.add(item)
    code = (p.get("code") or "").strip() or None
    if code and WhItem.query.filter(WhItem.code == code, WhItem.id != (item.id or 0)).first():
        return fail(f"کد «{code}» برای کالای دیگری ثبت شده است.", 422)
    item.name, item.code = name, code
    item.kind = "equipment" if p.get("kind") == "equipment" else "part"
    item.category = (p.get("category") or "").strip() or None
    item.unit = (p.get("unit") or "").strip() or "عدد"
    item.is_active = bool(p.get("is_active", True))
    item.note = (p.get("note") or "").strip() or None
    db.session.commit()
    return ok(item.to_dict(), message="ذخیره شد.")


@bp.post("/items/bulk")
@permission_required("warehouse.manage")
def save_items_bulk():
    """Several items at once — the parts list of one equipment kind as the
    form builder edits it (names, codes, active). Fields not sent are kept."""
    rows = body().get("items") or []
    saved = []
    for p in rows:
        name = (p.get("name") or "").strip()
        item = db.session.get(WhItem, int(p["id"])) if str(p.get("id") or "").isdigit() else None
        if not name:
            if item is None:
                continue
            db.session.rollback()
            return fail(f"نام قطعه‌ی «{item.name}» خالی است.", 422)
        if "code" in p:
            code = (p.get("code") or "").strip() or None
            with db.session.no_autoflush:
                taken = code and WhItem.query.filter(
                    WhItem.code == code, WhItem.id != (item.id if item is not None else 0)).first()
            if taken:
                db.session.rollback()
                return fail(f"کد «{code}» برای کالای دیگری ثبت شده است.", 422)
        if item is None:
            with db.session.no_autoflush:
                last = db.session.query(db.func.max(WhItem.sort_order)).scalar() or 0
            item = WhItem(name=name, sort_order=last + 1, source="فرم‌ساز",
                          kind=p.get("kind") or "part", unit="عدد")
            db.session.add(item)
        if "code" in p:
            item.code = (p.get("code") or "").strip() or None
        item.name = name
        for key in ("category", "unit"):
            if key in p:
                setattr(item, key, (p.get(key) or "").strip() or None)
        if "kind" in p:
            item.kind = "equipment" if p.get("kind") == "equipment" else "part"
        if "is_active" in p:
            item.is_active = bool(p.get("is_active"))
        item.unit = item.unit or "عدد"
        db.session.flush()
        saved.append(item)
    db.session.commit()
    from ..analytics.catalogue import bump_data_version
    bump_data_version()
    return ok({"items": [i.to_dict() for i in saved]}, message=f"{len(saved)} قطعه ذخیره شد.")


@bp.get("/items/export.xlsx")
@permission_required("warehouse.view")
def items_export():
    return send_file(io.BytesIO(wh.items_xlsx()), as_attachment=True,
                     download_name="warehouse_items.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.post("/items/import")
@permission_required("warehouse.manage")
def items_import():
    f = request.files.get("file")
    if f is None:
        return fail("فایلی انتخاب نشده است.", 422)
    data = f.read()
    from ..warehouse.partsbook import import_partsbook, is_partsbook
    try:
        if is_partsbook(data):
            # the workshop's parts workbook: catalogue, equipment and 98–05 history
            res = import_partsbook(data)
            res.update({"added": res["items_added"], "updated": res["items_updated"]})
        else:
            res = wh.import_items(data)
    except ValueError as exc:
        db.session.rollback()
        return fail(str(exc), 422)
    record_audit("import", "warehouse", None,
                 summary=f"ورود اقلام انبار: {res['added']} جدید، {res['updated']} به‌روز", commit=True)
    return ok(res, message="اقلام انبار وارد شد.")


@bp.get("/parts")
@permission_required("warehouse.view")
def parts():
    """Part installed/collected statistics (history 98–05 + parts forms)."""
    args = {k: request.args.get(k) for k in ("equipment_kind", "source", "date_from",
                                              "date_to", "equipment_code")}
    return ok(wh.parts_summary(args))


@bp.get("/part-actions")
@permission_required("warehouse.view")
def part_actions():
    from ..warehouse.models import WhPartAction
    q = WhPartAction.query
    for key in ("equipment_kind", "source", "part_action"):
        if request.args.get(key):
            q = q.filter(getattr(WhPartAction, key) == request.args[key])
    if request.args.get("equipment_code"):
        q = q.filter(WhPartAction.equipment_code.contains(request.args["equipment_code"]))
    if request.args.get("part"):
        q = q.filter(WhPartAction.part_name.contains(request.args["part"]))
    total = q.count()
    page = max(1, request.args.get("page", 1, type=int))
    rows = (q.order_by(WhPartAction.date_num.desc().nullslast(), WhPartAction.id.desc())
            .offset((page - 1) * 100).limit(100).all())
    return ok({"total": total, "page": page, "rows": [{
        "jdate": a.jdate, "equipment_code": a.equipment_code, "equipment_kind": a.equipment_kind,
        "equipment_name": a.equipment_name, "related_action": a.related_action,
        "part_action": "نصب شد" if a.part_action == "installed" else "جمع‌آوری شد",
        "state": a.state, "reusable": ("قابل استفاده مجدد" if a.reusable is True
                                       else "اسقاط" if a.reusable is False else None),
        "part_code": a.part_code, "part_name": a.part_name, "qty": a.qty,
        "failure": a.failure, "cause": a.cause, "source": a.source,
        "well_name": a.well_name, "facility_name": a.facility_name} for a in rows]})


@bp.post("/conditions")
@permission_required("warehouse.manage")
def save_condition():
    p = body()
    code = (p.get("code") or "").strip()
    label = (p.get("label") or "").strip()
    if not code or not label:
        return fail("کد و عنوان وضعیت لازم است.", 422)
    c = WhCondition.query.filter_by(code=code).first() or WhCondition(code=code)
    c.label = label
    c.applies_to = ",".join(x for x in (p.get("applies_to") or ["equipment", "part"])
                            if x in ("equipment", "part"))
    c.color = p.get("color") or c.color
    c.is_active = bool(p.get("is_active", True))
    db.session.add(c)
    db.session.commit()
    return ok(c.to_dict(), message="ذخیره شد.")


# ── exports ──────────────────────────────────────────────────────────────────
def _xlsx(title, head, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


@bp.get("/export/<what>.xlsx")
@permission_required("warehouse.view")
def export(what):
    args = _args()
    conds = {c.code: c.label for c in WhCondition.query.all()}
    if what == "movements":
        rows = wh.movements(args, limit=100000)
        out = _xlsx("گردش انبار", ["تاریخ", "انبار", "ورود/خروج", "علت", "کالا", "مشخصات", "پلاک",
                                    "وضعیت", "مقدار", "واحد", "چاه", "فرایند", "مرحله", "کاربر",
                                    "توضیحات"],
                    [[r["jdate"], r["warehouse_label"], r["direction_label"], r["reason_label"],
                      r["item_name"], r["spec"], r["serial"], conds.get(r["condition"], r["condition"]),
                      r["qty"], r["unit"], r["well_name"], r["workflow_name"], r["stage_title"],
                      r["user_name"], r["note"]] for r in rows])
    elif what == "stock":
        rows = wh.stock(args)
        out = _xlsx("موجودی انبار", ["انبار", "کالا", "وضعیت", "ورود", "خروج", "موجودی", "واحد"],
                    [[r["warehouse_label"], r["item_name"], r["condition_label"], r["in"], r["out"],
                      r["balance"], r["unit"]] for r in rows])
    elif what == "parts":
        from ..warehouse.models import WhPartAction
        q = WhPartAction.query.order_by(WhPartAction.date_num, WhPartAction.id)
        out = _xlsx("سوابق قطعات", ["تاریخ", "منبع", "کد تجهیز", "نوع تجهیز", "نام تجهیز", "اقدام مرتبط",
                                     "اقدام در سطح قطعه", "وضعیت نو/کهنه", "قابل استفاده مجدد؟",
                                     "کد انباری", "شرح قطعه", "تعداد", "خرابی", "علت خرابی",
                                     "اقدام انجام شده", "تاسیس / چاه", "کاربر"],
                    [[a.jdate, "فرایند" if a.source == "workflow" else "سوابق ۹۸–۰۵", a.equipment_code,
                      a.equipment_kind, a.equipment_name, a.related_action,
                      "نصب شد" if a.part_action == "installed" else "جمع‌آوری شد", a.state,
                      ("بلی" if a.reusable is True else "خیر" if a.reusable is False else None),
                      a.part_code, a.part_name, a.qty, a.failure, a.cause, a.action_done,
                      a.well_name or a.facility_name, a.user_name] for a in q])
    elif what == "summary":
        s = wh.summary(args)
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "گزارش مدیرعامل"
        ws.sheet_view.rightToLeft = True
        span = f"{args.get('date_from') or 'ابتدا'} تا {args.get('date_to') or 'امروز'}"
        ws.append([f"گزارش انبار تجهیزات و قطعات کارگاه مکانیک — {span}"])
        ws["A1"].font = Font(bold=True, size=13)
        h = s["headline"]
        for label, key in (("الکتروپمپ کشیده‌شده و تحویل انبار", "electropumps_pulled"),
                           ("الکتروپمپ مونتاژشده", "electropumps_assembled"),
                           ("الکتروپمپ خارج‌شده برای نصب", "electropumps_installed"),
                           ("قطعات و تجهیزات اسقاط", "parts_scrapped"),
                           ("مصرف اقلام نو", "parts_new"),
                           ("مصرف اقلام قابل استفاده مجدد", "parts_reused"),
                           ("تعداد کل ردیف‌های گردش", "movements")):
            ws.append([label, h[key]])
        p = s["parts"]
        ws.append([])
        ws.append(["قطعات (مونتاژ و دمونتاژ)"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for label, key in (("نصب — نو", "installed_new"), ("نصب — کهنه (قابل استفاده مجدد)", "installed_repair"),
                           ("نصب — بی‌وضعیت", "installed_unknown"), ("جمع‌آوری (ارزیابی‌نشده)", "collected"),
                           ("قابل استفاده مجدد", "reusable"), ("اسقاط", "scrap"),
                           ("تعداد تجهیز", "equipments")):
            ws.append([label, p["total"][key]])
        ws.append([])
        ws.append(["تجهیز", "کد انباری", "شرح قطعه", "نصب نو", "نصب کهنه", "نصب بی‌وضعیت",
                   "جمع‌آوری", "قابل استفاده مجدد", "اسقاط"])
        for r in p["by_part"]:
            ws.append([r["equipment_kind"], r["part_code"], r["part_name"], r["installed_new"],
                       r["installed_repair"], r["installed_unknown"], r["collected"],
                       r["reusable"], r["scrap"]])
        ws.append([])
        ws.append(["سال", "نصب", "نصب نو", "نصب کهنه", "جمع‌آوری", "قابل استفاده مجدد", "اسقاط"])
        for r in p["by_year"]:
            ws.append([r["year"], r["installed"], r["new"], r["repair"], r["collected"],
                       r["reusable"], r["scrap"]])
        for title, key, cols in (("بر اساس علت", "by_reason", ("warehouse", "direction", "reason", "qty")),
                                 ("بر اساس وضعیت", "by_condition", ("warehouse", "direction", "condition", "qty")),
                                 ("بر اساس کالا", "by_item", ("warehouse", "item", "condition", "direction", "qty"))):
            ws.append([])
            ws.append([title])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            for r in s[key]:
                ws.append([r[c] for c in cols])
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
    else:
        return fail("خروجی نامعتبر است.", 404)
    return send_file(out, as_attachment=True, download_name=f"warehouse_{what}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
