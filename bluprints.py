################################################################################
# bluprints
# This file was generated automatically.
################################################################################


################################################################################
# FILE: __init__.py
################################################################################




################################################################################
# FILE: __pycache___merged.py
################################################################################

##########################################################################################
# BLUEPRINT : __pycache__
# AUTO GENERATED
##########################################################################################




################################################################################
# FILE: auth_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : auth
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("auth", __name__)

from app.blueprints.auth import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from datetime import datetime

from flask import render_template, redirect, url_for, flash, request, session
from flask_login import login_user, logout_user, login_required, current_user
from flask_wtf import FlaskForm
from wtforms import StringField, PasswordField, BooleanField, SubmitField
from wtforms.validators import DataRequired

from app.extensions import db
from app.blueprints.auth import bp
from app.models.rbac import User, LoginLog

class LoginForm(FlaskForm):
    username = StringField("نام کاربری", validators=[DataRequired()])
    password = PasswordField("گذرواژه", validators=[DataRequired()])
    remember = BooleanField("مرا به خاطر بسپار")
    submit = SubmitField("ورود")


@bp.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard"))

    form = LoginForm()
    if form.validate_on_submit():
        user = db.session.scalar(
            db.select(User).filter_by(username=form.username.data.strip())
        )
        if user is None or not user.check_password(form.password.data):
            flash("نام کاربری یا گذرواژه نادرست است.", "danger")
        elif not user.is_active:
            flash("این حساب غیرفعال است.", "warning")
        else:
            login_user(user, remember=form.remember.data)
            now = datetime.utcnow()
            ip = (request.headers.get("X-Forwarded-For", request.remote_addr) or "").split(",")[0].strip()
            agent = (request.user_agent.string or "")[:255]
            user.last_login_at = now
            user.last_login_ip = ip
            user.last_login_agent = agent
            log = LoginLog(user_id=user.id, login_at=now, ip_address=ip, user_agent=agent)
            db.session.add(log)
            db.session.commit()
            session["login_log_id"] = log.id
            next_page = request.args.get("next")
            return redirect(next_page or url_for("main.dashboard"))

    return render_template("auth/login.html", form=form)


@bp.route("/logout")
@login_required
def logout():
    log_id = session.pop("login_log_id", None)
    if log_id:
        log = db.session.get(LoginLog, log_id)
        if log and log.logout_at is None:
            log.logout_at = datetime.utcnow()
            db.session.commit()
    logout_user()
    flash("از سیستم خارج شدید.", "info")
    return redirect(url_for("auth.login"))





################################################################################
# FILE: documents_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : documents
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("documents", __name__, url_prefix="/documents")

from app.blueprints.documents import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

"""Document/attachment upload, download and delete.

Files are validated against the allowed-extension allowlist, stored under
instance/uploads/ with a random UUID name, and served back with their original
filename. Entity links are polymorphic (entity_type, entity_id).
"""
import os
import uuid

from flask import (current_app, request, redirect, url_for, flash, abort,
                   send_from_directory)
from flask_login import current_user
from werkzeug.utils import secure_filename

from app.extensions import db
from app.blueprints.documents import bp
from app.security import permission_required
from app.models.attachment import Attachment

# where each entity type's page lives, to redirect back after up/delete
_BACK = {"well": ("wells.detail", "well_id")}


def _redirect_back(att_or_type, entity_id=None):
    etype = att_or_type.entity_type if isinstance(att_or_type, Attachment) else att_or_type
    eid = att_or_type.entity_id if isinstance(att_or_type, Attachment) else entity_id
    ep, arg = _BACK.get(etype, ("wells.detail", "well_id"))
    return redirect(url_for(ep, **{arg: eid}) + "#sec-docs")


def _allowed(filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in current_app.config["ALLOWED_UPLOAD_EXT"], ext


@bp.route("/upload", methods=["POST"])
@permission_required("documents", "create")
def upload():
    etype = request.form.get("entity_type", "")
    eid = request.form.get("entity_id", type=int)
    if not etype or not eid:
        abort(400)
    f = request.files.get("file")
    if not f or not f.filename:
        flash("فایلی انتخاب نشده است.", "warning")
        return _redirect_back(etype, eid)
    ok, ext = _allowed(f.filename)
    if not ok:
        flash("نوع فایل مجاز نیست.", "danger")
        return _redirect_back(etype, eid)

    upload_dir = current_app.config["UPLOAD_DIR"]
    os.makedirs(upload_dir, exist_ok=True)
    stored = f"{uuid.uuid4().hex}.{ext}"
    f.save(os.path.join(upload_dir, stored))
    size = os.path.getsize(os.path.join(upload_dir, stored))

    att = Attachment(
        entity_type=etype, entity_id=eid,
        title=(request.form.get("title") or "").strip() or None,
        kind=request.form.get("kind") or "other",
        original_name=secure_filename(f.filename) or f"file.{ext}",
        stored_name=stored, content_type=f.mimetype, size_bytes=size,
        created_by_id=current_user.id)
    db.session.add(att)
    db.session.commit()
    flash("سند بارگذاری شد.", "success")
    return _redirect_back(att)


@bp.route("/<int:att_id>/download")
@permission_required("documents", "view")
def download(att_id):
    att = db.get_or_404(Attachment, att_id)
    return send_from_directory(
        current_app.config["UPLOAD_DIR"], att.stored_name,
        as_attachment=True, download_name=att.original_name)


@bp.route("/<int:att_id>/delete", methods=["POST"])
@permission_required("documents", "delete")
def delete(att_id):
    att = db.get_or_404(Attachment, att_id)
    path = os.path.join(current_app.config["UPLOAD_DIR"], att.stored_name)
    if os.path.exists(path):
        try:
            os.remove(path)
        except OSError:
            pass
    back = _redirect_back(att)
    db.session.delete(att)
    db.session.commit()
    flash("سند حذف شد.", "info")
    return back


def attachments_for(entity_type, entity_id):
    return db.session.scalars(
        db.select(Attachment).filter_by(entity_type=entity_type, entity_id=entity_id)
        .order_by(Attachment.created_at.desc())).all()





################################################################################
# FILE: drilling_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : drilling
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("drilling", __name__)

from app.blueprints.drilling import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, BooleanField, SubmitField
from wtforms.validators import DataRequired, Optional

from app.extensions import db
from app.blueprints.drilling import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.drilling import Drilling
from app.models.audit import RecordHistory
from app.models.constants import REQUEST_TYPES, DRILL_METHODS
from app import workflow


class DrillingForm(FlaskForm):
    request_type = SelectField("نوع پروانه", choices=[("", "—")] + REQUEST_TYPES, validators=[Optional()])
    drill_method = SelectField("روش حفاری", choices=[("", "—")] + DRILL_METHODS, validators=[Optional()])
    executor = StringField("نام مجری", validators=[Optional()])
    contractor = StringField("پیمانکار", validators=[Optional()])
    supervisor = StringField("ناظر", validators=[Optional()])
    credit_source = StringField("محل تأمین اعتبار", validators=[Optional()])
    contract_no = StringField("شماره قرارداد", validators=[Optional()])
    contract_date = JalaliDateField("تاریخ قرارداد", validators=[Optional()])
    start_date = JalaliDateField("تاریخ استقرار دستگاه", validators=[Optional()])
    end_date = JalaliDateField("تاریخ پایان/ترخیص", validators=[Optional()])
    well_depth_permit = FloatField("عمق چاه در پروانه", validators=[Optional()])
    well_depth_actual = FloatField("عمق حفاری واقعی", validators=[Optional()])
    casing_diameter_in = FloatField("قطر لوله جدار (اینچ)", validators=[Optional()])
    casing_total_len = FloatField("طول کلی لوله‌گذاری", validators=[Optional()])
    steel_blank_len = FloatField("لوله فولادی ساده", validators=[Optional()])
    steel_screen_len = FloatField("لوله فولادی مشبک", validators=[Optional()])
    upvc_blank_len = FloatField("لوله UPVC ساده", validators=[Optional()])
    upvc_screen_len = FloatField("لوله UPVC مشبک", validators=[Optional()])
    transition_len = FloatField("قطعه تبدیلی", validators=[Optional()])
    static_level = FloatField("سطح استاتیک", validators=[Optional()])
    max_yield_lps = FloatField("حداکثر آبدهی (l/s)", validators=[Optional()])
    proposed_discharge_lps = FloatField("دبی پیشنهادی (l/s)", validators=[Optional()])
    dynamic_at_proposed = FloatField("سطح دینامیک در دبی پیشنهادی", validators=[Optional()])
    drawdown = FloatField("مقدار افت", validators=[Optional()])
    coeff_a = FloatField("ضریب a", validators=[Optional()])
    coeff_b = FloatField("ضریب b", validators=[Optional()])
    videometry_done = BooleanField("ویدئومتری انجام شده")
    address = StringField("آدرس", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


FIELDS = [
    "request_type", "drill_method", "executor", "contractor", "supervisor",
    "credit_source", "contract_no", "contract_date", "start_date", "end_date",
    "well_depth_permit", "well_depth_actual", "casing_diameter_in", "casing_total_len",
    "steel_blank_len", "steel_screen_len", "upvc_blank_len", "upvc_screen_len",
    "transition_len", "static_level", "max_yield_lps", "proposed_discharge_lps",
    "dynamic_at_proposed", "drawdown", "coeff_a", "coeff_b", "videometry_done",
    "address", "notes",
]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data)


@bp.route("/wells/<int:well_id>/drilling/new", methods=["GET", "POST"])
@permission_required("drilling", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = DrillingForm()
    if form.validate_on_submit():
        rec = Drilling(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("رویداد حفاری ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("drilling.detail", record_id=rec.id))
    return render_template("drilling/form.html", form=form, well=well, title="ثبت حفاری")


@bp.route("/drilling/<int:record_id>")
@permission_required("drilling", "view")
def detail(record_id):
    rec = db.get_or_404(Drilling, record_id)
    history = db.session.scalars(
        db.select(RecordHistory)
        .filter_by(entity_type="drilling", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template(
        "drilling/detail.html",
        rec=rec, well=rec.well, history=history,
        req_labels=dict(REQUEST_TYPES), method_labels=dict(DRILL_METHODS),
        editable=workflow.is_editable(rec),
    )


@bp.route("/drilling/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("drilling", "edit")
def edit(record_id):
    rec = db.get_or_404(Drilling, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد تأیید/ثبت شده و قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("drilling.detail", record_id=rec.id))
    form = DrillingForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("رویداد حفاری به‌روزرسانی شد.", "success")
        return redirect(url_for("drilling.detail", record_id=rec.id))
    return render_template("drilling/form.html", form=form, well=rec.well, title="ویرایش حفاری")


@bp.route("/drilling/<int:record_id>/delete", methods=["POST"])
@permission_required("drilling", "delete")
def delete(record_id):
    rec = db.get_or_404(Drilling, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رویداد حفاری حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


# --- workflow transitions ---
def _transition(record_id, fn, perm_action, success_msg, **kwargs):
    rec = db.get_or_404(Drilling, record_id)
    if not current_user.has_permission("drilling", perm_action):
        abort(403)
    try:
        fn(rec, **kwargs)
        db.session.commit()
        flash(success_msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("drilling.detail", record_id=record_id))


@bp.route("/drilling/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/drilling/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "رکورد تأیید شد.")


@bp.route("/drilling/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    reason = request.form.get("reason", "")
    return _transition(record_id, workflow.reject, "approve", "رکورد برگشت داده شد.", reason=reason)


@bp.route("/drilling/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: exports_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : exports
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("exports", __name__, url_prefix="/exports")

from app.blueprints.exports import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

"""Data export / review section.

Lets an authorized user download the full dataset as Excel for offline review —
in particular to spot duplicate wells (variants of the same physical well such as
«ابوطالب 1 ق» / «ابوطالب 1»), which share a `گروه تکراری` key here.
"""
from datetime import datetime
from io import BytesIO

import pandas as pd
from flask import render_template, send_file

from app.extensions import db
from app.blueprints.exports import bp
from app.security import permission_required
from app.services import data_export

_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(sheets):
    """sheets: list of (sheet_name, DataFrame) -> BytesIO of an .xlsx workbook."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        for name, df in sheets:
            df.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 40)
    buf.seek(0)
    return buf


def _send(buf, prefix):
    stamp = datetime.now().strftime("%Y%m%d")
    return send_file(buf, as_attachment=True, mimetype=_MIME,
                     download_name=f"{prefix}_{stamp}.xlsx")


@bp.route("/")
@permission_required("reports", "view")
def index():
    from app.models.well import Well
    from app.models.production import MonthlyProduction as MP
    df = data_export.wells_df()
    dup_groups = int((df["تعداد در گروه"] > 1).sum())
    dup_group_count = df[df["تعداد در گروه"] > 1]["گروه تکراری"].nunique()
    stats = {
        "wells": int(db.session.scalar(db.select(db.func.count(Well.id))) or 0),
        "dup_rows": dup_groups,
        "dup_groups": int(dup_group_count),
        "production_rows": int(db.session.scalar(db.select(db.func.count(MP.id))) or 0),
    }
    return render_template("exports/index.html", stats=stats)


@bp.route("/wells.xlsx")
@permission_required("reports", "view")
def wells_xlsx():
    return _send(_xlsx([("چاه‌ها", data_export.wells_df())]), "wells")


@bp.route("/duplicates.xlsx")
@permission_required("reports", "view")
def duplicates_xlsx():
    return _send(_xlsx([("تکراری‌های مشکوک", data_export.duplicates_df())]), "duplicates")


@bp.route("/production.xlsx")
@permission_required("reports", "view")
def production_xlsx():
    return _send(_xlsx([("تولید ماهانه", data_export.production_df())]), "production")


@bp.route("/all.xlsx")
@permission_required("reports", "view")
def all_xlsx():
    sheets = [
        ("چاه‌ها", data_export.wells_df()),
        ("تکراری‌های مشکوک", data_export.duplicates_df()),
        ("تولید ماهانه", data_export.production_df()),
    ]
    return _send(_xlsx(sheets), "all_data")





################################################################################
# FILE: finance_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : finance
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("finance", __name__, url_prefix="/finance")

from app.blueprints.finance import routes  # noqa: E402,F401


##########################################################################################
# FILE : routes.py
##########################################################################################

"""ماژول صورت‌وضعیت مالی — فاز ۲: فهرست + فرم + جزئیات."""
from flask import render_template, redirect, url_for, flash
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.finance import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField
from app.models.well import Well
from app.models.finance import FinanceStatement

OP_TYPES = [("", "—"), ("بهسازی", "بهسازی"), ("آزمایش پمپاژ", "آزمایش پمپاژ"),
            ("حفاری", "حفاری"), ("سایر", "سایر")]
KINDS = [("", "—"), ("موقت", "موقت"), ("قطعی", "قطعی"),
         ("وضعیت ۱", "وضعیت ۱"), ("وضعیت ۲", "وضعیت ۲"), ("وضعیت ۳", "وضعیت ۳")]


class StatementForm(FlaskForm):
    op_type = SelectField("نوع عملیات", choices=OP_TYPES, validators=[Optional()])
    op_description = TextAreaField("شرح عملیات (موضوع)", validators=[Optional()])
    contractor = StringField("پیمانکار", validators=[Optional()])
    contract_no = StringField("شماره قرارداد", validators=[Optional()])
    statement_no = StringField("شماره صورت‌وضعیت", validators=[Optional()])
    statement_kind = SelectField("نوع وضعیت", choices=KINDS, validators=[Optional()])
    statement_date = JalaliDateField("تاریخ صورت‌وضعیت", validators=[Optional()])
    coef_region = PersianFloatField("ضریب منطقه", validators=[Optional()])
    coef_overhead = PersianFloatField("ضریب بالاسری", validators=[Optional()])
    coef_contract = PersianFloatField("ضریب پیمان", validators=[Optional()])
    workshop_setup = PersianFloatField("تجهیز کارگاه (ریال)", validators=[Optional()])
    total_before = PersianFloatField("مبلغ کل قبل از ضرایب (ریال)", validators=[Optional()])
    total_after = PersianFloatField("مبلغ کل پس از ضرایب (ریال)", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


FIELDS = ["op_type", "op_description", "contractor", "contract_no", "statement_no", "statement_kind",
          "statement_date", "coef_region", "coef_overhead", "coef_contract",
          "workshop_setup", "total_before", "total_after", "notes"]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data if getattr(form, f).data != "" else None)


@bp.route("/")
@permission_required("finance", "view")
def index():
    statements = db.session.scalars(
        db.select(FinanceStatement).order_by(FinanceStatement.statement_date.desc())
    ).all()
    return render_template("finance/index.html", statements=statements)


@bp.route("/well/<int:well_id>/new", methods=["GET", "POST"])
@permission_required("finance", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = StatementForm()
    if form.validate_on_submit():
        rec = FinanceStatement(well_id=well.id, created_by_id=current_user.id, source="manual")
        _apply(form, rec)
        db.session.add(rec)
        db.session.commit()
        flash("صورت‌وضعیت ثبت شد.", "success")
        return redirect(url_for("finance.detail", record_id=rec.id))
    return render_template("finance/form.html", form=form, well=well, title="ثبت صورت‌وضعیت")


@bp.route("/<int:record_id>")
@permission_required("finance", "view")
def detail(record_id):
    rec = db.get_or_404(FinanceStatement, record_id)
    # تجمیع سهم هر چاه از آیتم‌های این سند (بر اساس تخصیص‌ها + ضریب بالاسری)
    ovh = rec.coef_overhead or 1.0
    agg = {}
    for it in rec.items:
        amt = (it.total_amount or 0) * ovh
        allocs = list(it.allocations)
        tq = sum((a.quantity or 0) for a in allocs)
        if not tq:
            continue
        for a in allocs:
            key = a.well_id if a.well_id else ("name:" + (a.well_name or "?"))
            row = agg.setdefault(key, {
                "well_id": a.well_id,
                "name": (a.well.name if a.well else a.well_name) or "—",
                "qty": 0.0, "cost": 0.0})
            row["qty"] += a.quantity or 0
            row["cost"] += amt * (a.quantity or 0) / tq
    wells_breakdown = sorted(agg.values(), key=lambda r: r["cost"], reverse=True)
    return render_template("finance/detail.html", rec=rec, well=rec.well,
                           wells_breakdown=wells_breakdown)


@bp.route("/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("finance", "edit")
def edit(record_id):
    rec = db.get_or_404(FinanceStatement, record_id)
    form = StatementForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        db.session.commit()
        flash("صورت‌وضعیت به‌روزرسانی شد.", "success")
        return redirect(url_for("finance.detail", record_id=rec.id))
    return render_template("finance/form.html", form=form, well=rec.well, title="ویرایش صورت‌وضعیت")


@bp.route("/<int:record_id>/delete", methods=["POST"])
@permission_required("finance", "delete")
def delete(record_id):
    rec = db.get_or_404(FinanceStatement, record_id)
    db.session.delete(rec)
    db.session.commit()
    flash("صورت‌وضعیت حذف شد.", "info")
    return redirect(url_for("finance.index"))
# ---------- آیتم‌های هزینه (فاز ۳) ----------
from app.models.finance import FinanceItem


class ItemForm(FlaskForm):
    category = StringField("دسته (شرح عملیات کلی)", validators=[Optional()])
    row_no = StringField("شماره ردیف", validators=[Optional()])
    description = TextAreaField("شرح تفصیلی", validators=[Optional()])
    unit = StringField("واحد", validators=[Optional()])
    contract_qty = PersianFloatField("تعداد قرارداد", validators=[Optional()])
    quantity = PersianFloatField("مقدار", validators=[Optional()])
    unit_price = PersianFloatField("مبلغ واحد (ریال)", validators=[Optional()])
    coefficient = PersianFloatField("ضریب", validators=[Optional()])
    total_amount = PersianFloatField("مبلغ کل (ریال)", validators=[Optional()])
    submit = SubmitField("ذخیره")


ITEM_FIELDS = ["category", "row_no", "description", "unit", "contract_qty",
               "quantity", "unit_price", "coefficient", "total_amount"]


@bp.route("/<int:statement_id>/item/new", methods=["GET", "POST"])
@permission_required("finance", "edit")
def item_create(statement_id):
    st = db.get_or_404(FinanceStatement, statement_id)
    form = ItemForm()
    if form.validate_on_submit():
        it = FinanceItem(statement_id=st.id, created_by_id=current_user.id)
        for f in ITEM_FIELDS:
            setattr(it, f, getattr(form, f).data)
        db.session.add(it)
        db.session.commit()
        flash("آیتم هزینه افزوده شد.", "success")
        return redirect(url_for("finance.detail", record_id=st.id))
    return render_template("finance/item_form.html", form=form, st=st, title="افزودن آیتم هزینه")


@bp.route("/item/<int:item_id>/edit", methods=["GET", "POST"])
@permission_required("finance", "edit")
def item_edit(item_id):
    it = db.get_or_404(FinanceItem, item_id)
    form = ItemForm(obj=it)
    if form.validate_on_submit():
        for f in ITEM_FIELDS:
            setattr(it, f, getattr(form, f).data)
        db.session.commit()
        flash("آیتم به‌روزرسانی شد.", "success")
        return redirect(url_for("finance.detail", record_id=it.statement_id))
    return render_template("finance/item_form.html", form=form, st=it.statement,
                           title="ویرایش آیتم هزینه", item_id=it.id)


@bp.route("/item/<int:item_id>/delete", methods=["POST"])
@permission_required("finance", "edit")
def item_delete(item_id):
    it = db.get_or_404(FinanceItem, item_id)
    sid = it.statement_id
    db.session.delete(it)
    db.session.commit()
    flash("آیتم حذف شد.", "info")
    return redirect(url_for("finance.detail", record_id=sid))

# ---------- گزارش‌ها و داشبورد مالی (فاز ۵) ----------
from io import BytesIO
from datetime import datetime
from collections import defaultdict

import pandas as pd
from flask import request, send_file

from app.models.finance import FinanceItemAllocation

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _finance_report(op_type=None, contractor=None, statement_id=None, well_id=None):
    """داده‌ی گزارش مالی را با اعمال فیلترها می‌سازد.

    خروجی: dict شامل سه گروه‌بندی (چاه/سند/نوع عملیات)، KPIها و ردیف‌های جدول.
    هزینه‌ی هر آیتم بین چاه‌های تخصیص‌یافته به نسبت مقدار تقسیم می‌شود.
    """
    q = db.select(FinanceStatement)
    if op_type:
        q = q.where(FinanceStatement.op_type == op_type)
    if contractor:
        q = q.where(FinanceStatement.contractor == contractor)
    if statement_id:
        q = q.where(FinanceStatement.id == statement_id)
    statements = db.session.scalars(q).all()

    by_optype = defaultdict(float)
    by_statement = []                 # (label, total)
    well_cost = defaultdict(float)    # well_name -> cost
    well_meta = {}                    # well_name -> well_id
    well_docs = defaultdict(set)      # well_name -> {statement_id}
    unallocated = 0.0
    grand = 0.0

    for st in statements:
        stmt_total = 0.0
        # ضریب بالاسری در سطح سند اعمال می‌شود تا مبالغ به «مبلغ صورت‌وضعیت» برسند
        ovh = st.coef_overhead or 1.0
        for it in st.items:
            amt = (it.total_amount or 0.0) * ovh
            allocs = list(it.allocations)
            tq = sum((a.quantity or 0) for a in allocs)
            if tq > 0:
                for a in allocs:
                    if well_id and a.well_id != well_id:
                        continue
                    share = amt * (a.quantity or 0) / tq
                    name = a.well.name if a.well else (a.well_name or "—")
                    well_cost[name] += share
                    well_meta[name] = a.well_id
                    well_docs[name].add(st.id)
                    stmt_total += share
            elif st.well_id:            # سند تک‌چاهی (ثبت دستی)
                if well_id and st.well_id != well_id:
                    continue
                name = st.well.name if st.well else "—"
                well_cost[name] += amt
                well_meta[name] = st.well_id
                well_docs[name].add(st.id)
                stmt_total += amt
            else:                        # آیتم بدون تخصیص (مثل ویدئومتری)
                if not well_id:
                    unallocated += amt
                    stmt_total += amt
        # تجهیز کارگاه: هزینه‌ی ثابت سند، به هیچ چاهی تخصیص نمی‌یابد
        if not well_id and st.workshop_setup:
            unallocated += st.workshop_setup
            stmt_total += st.workshop_setup
        if stmt_total:
            by_optype[st.op_type or "نامشخص"] += stmt_total
            label = " ".join(x for x in [st.op_type, st.statement_kind,
                             ("— " + (st.contractor or st.contract_no or f"#{st.id}"))] if x)
            by_statement.append((label, stmt_total))
            grand += stmt_total

    rows = [{
        "well": n, "well_id": well_meta.get(n),
        "cost": round(c), "docs": len(well_docs[n]),
        "matched": well_meta.get(n) is not None,
    } for n, c in well_cost.items()]
    rows.sort(key=lambda r: r["cost"], reverse=True)
    by_statement.sort(key=lambda x: x[1], reverse=True)

    return {
        "rows": rows,
        "by_optype": dict(sorted(by_optype.items(), key=lambda x: x[1], reverse=True)),
        "by_statement": by_statement,
        "top_wells": {r["well"]: r["cost"] for r in rows[:15]},
        "grand": round(grand),
        "unallocated": round(unallocated),
        "well_count": len(rows),
        "doc_count": len(statements),
    }


def _filter_options():
    op_types = [x for x in db.session.scalars(
        db.select(FinanceStatement.op_type).distinct()).all() if x]
    contractors = [x for x in db.session.scalars(
        db.select(FinanceStatement.contractor).distinct()).all() if x]
    statements = db.session.scalars(
        db.select(FinanceStatement).order_by(FinanceStatement.op_type,
                                             FinanceStatement.statement_kind)).all()
    # چاه‌هایی که در مالی هزینه دارند
    wids = set(db.session.scalars(
        db.select(FinanceItemAllocation.well_id).distinct()).all())
    wids |= set(db.session.scalars(
        db.select(FinanceStatement.well_id).distinct()).all())
    wids.discard(None)
    wells = db.session.scalars(
        db.select(Well).where(Well.id.in_(wids)).order_by(Well.name)).all() if wids else []
    return op_types, contractors, statements, wells


def _current_filters():
    def _int(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return None
    return {
        "op_type": request.args.get("op_type") or None,
        "contractor": request.args.get("contractor") or None,
        "statement_id": _int(request.args.get("statement_id")),
        "well_id": _int(request.args.get("well_id")),
    }


@bp.route("/reports")
@permission_required("finance", "view")
def reports():
    f = _current_filters()
    data = _finance_report(**f)
    op_types, contractors, statements, wells = _filter_options()
    cur_qs = {k: v for k, v in f.items() if v}   # برای لینک خروجی اکسل
    return render_template("finance/reports.html", data=data, cur=f, cur_qs=cur_qs,
                           op_types=op_types, contractors=contractors,
                           statements=statements, wells=wells)


@bp.route("/reports/export")
@permission_required("finance", "view")
def reports_export():
    f = _current_filters()
    data = _finance_report(**f)
    df = pd.DataFrame([{
        "چاه": r["well"],
        "هزینه (ریال)": r["cost"],
        "تعداد سند": r["docs"],
        "تطبیق با سامانه": "بله" if r["matched"] else "خیر",
    } for r in data["rows"]])
    by_stmt = pd.DataFrame([{"سند": l, "مبلغ (ریال)": round(v)}
                            for l, v in data["by_statement"]])
    by_op = pd.DataFrame([{"نوع عملیات": k, "مبلغ (ریال)": round(v)}
                          for k, v in data["by_optype"].items()])

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        (df if not df.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="هزینه هر چاه", index=False)
        (by_stmt if not by_stmt.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="به تفکیک سند", index=False)
        (by_op if not by_op.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="به تفکیک عملیات", index=False)
        for name in xl.sheets:
            ws = xl.sheets[name]
            for col in ws.columns:
                w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 45)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return send_file(buf, as_attachment=True, mimetype=_XLSX_MIME,
                     download_name=f"finance_report_{stamp}.xlsx")




################################################################################
# FILE: gis_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : gis
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("gis", __name__, url_prefix="/gis")

from app.blueprints.gis import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, jsonify
from flask_login import login_required

from app.extensions import db
from app.blueprints.gis import bp
from app.security import permission_required
from app.models.well import Well
from app.models.constants import WELL_STATUSES


@bp.route("/map")
@login_required
def map():
    return render_template("gis/map.html")


@bp.route("/analysis")
@permission_required("reports", "view")
def analysis():
    return render_template("gis/analysis.html")


@bp.route("/analysis.json")
@permission_required("reports", "view")
def analysis_json():
    from app.services import spatial
    return jsonify({"points": spatial.metric_points()})


@bp.route("/piezometric.json")
@permission_required("reports", "view")
def piezometric_json():
    from app.services import spatial
    return jsonify(spatial.piezometric_grid() or {})


@bp.route("/interference.json")
@permission_required("reports", "view")
def interference_json():
    from flask import request
    from app.services import spatial
    thr = request.args.get("threshold", 400, type=int)
    return jsonify(spatial.interference(threshold_m=thr))


@bp.route("/suitability.json")
@permission_required("reports", "view")
def suitability_json():
    from app.services import spatial
    return jsonify(spatial.suitability_grid() or {})


@bp.route("/wells.geojson")
@permission_required("wells", "view")
def wells_geojson():
    """Wells that have coordinates, as GeoJSON for Leaflet (themed)."""
    status_labels = dict(WELL_STATUSES)
    from app.models.production import MonthlyProduction as MP
    spec_energy = {}
    for mp in db.session.scalars(db.select(MP).where(MP.energy_kwh.isnot(None))).all():
        se = mp.specific_energy
        if se is not None:
            spec_energy[mp.well_id] = se

    wells = db.session.scalars(
        db.select(Well).where(Well.latitude.isnot(None), Well.longitude.isnot(None))
    ).all()
    features = []
    for w in wells:
        features.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [w.longitude, w.latitude]},
            "properties": {
                "id": w.id, "pm_code": w.display_pm, "name": w.name,
                "status": status_labels.get(w.status, w.status),
                "status_key": w.status or "unknown",
                "zone": w.zone or "", "reservoir": w.destination_reservoir or "",
                "office": w.office.name if w.office else "",
                "elevation": w.ground_elevation,
                "specific_energy": spec_energy.get(w.id),
            },
        })
    return jsonify({"type": "FeatureCollection", "features": features})





################################################################################
# FILE: install_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : install
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("install", __name__)

from app.blueprints.install import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, BooleanField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.install import bp
from app.security import permission_required
from app.utils.forms import (
    JalaliDateField, PersianFloatField as FloatField, PersianIntegerField as IntegerField,
)
from app.models.well import Well
from app.models.pump_asset import PumpInstallation, Supplier
from app.models.audit import RecordHistory
from app.models.constants import EQUIP_CONDITIONS
from app import workflow


def get_or_create_supplier(name, kind=None):
    from app.models.constants import canonical_supplier
    name = canonical_supplier(name)
    if not name:
        return None
    s = db.session.scalar(db.select(Supplier).filter_by(name=name))
    if s is None:
        s = Supplier(name=name, kind=kind)
        db.session.add(s)
        db.session.flush()
    elif kind and not s.kind:
        s.kind = kind
    return s


class InstallForm(FlaskForm):
    install_no = IntegerField("شماره نصب", validators=[Optional()])
    install_date = JalaliDateField("تاریخ نصب", validators=[Optional()])
    pull_date = JalaliDateField("تاریخ کشیدن", validators=[Optional()])
    motor_power_kw = FloatField("توان موتور (kW)", validators=[Optional()])
    motor_condition = SelectField("موتور", choices=[("", "—")] + EQUIP_CONDITIONS, validators=[Optional()])
    pump_type = StringField("تیپ پمپ", validators=[Optional()])
    pump_stages = IntegerField("طبقه", validators=[Optional()])
    pump_condition = SelectField("پمپ", choices=[("", "—")] + EQUIP_CONDITIONS, validators=[Optional()])
    manufacturer_name = StringField("سازنده", validators=[Optional()])
    contractor_name = StringField("پیمانکار", validators=[Optional()])
    install_depth_m = FloatField("عمق نصب", validators=[Optional()])
    well_depth_m = FloatField("عمق چاه", validators=[Optional()])
    static_level = FloatField("سطح استاتیک", validators=[Optional()])
    dynamic_level = FloatField("سطح دینامیک", validators=[Optional()])
    route_loss = FloatField("تلفات مسیر", validators=[Optional()])
    grid_pressure_m = FloatField("فشار شبکه", validators=[Optional()])
    work_shift = StringField("شیفت کاری", validators=[Optional()])
    removal_reason = StringField("علت کشیدن", validators=[Optional()])
    fault_by_operator = TextAreaField("شرح خرابی (بهره‌بردار)", validators=[Optional()])
    fault_by_workshop = TextAreaField("شرح خرابی (کارگاه مکانیک)", validators=[Optional()])
    pm_form_registered = BooleanField("فرم نصب در PM ثبت شده")
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


SIMPLE_FIELDS = [
    "install_no", "install_date", "pull_date", "motor_power_kw", "motor_condition",
    "pump_type", "pump_stages", "pump_condition", "install_depth_m", "well_depth_m",
    "static_level", "dynamic_level", "route_loss", "grid_pressure_m", "work_shift",
    "removal_reason", "fault_by_operator", "fault_by_workshop", "pm_form_registered", "notes",
]


def _apply(form, rec):
    for f in SIMPLE_FIELDS:
        setattr(rec, f, getattr(form, f).data)
    man = get_or_create_supplier(form.manufacturer_name.data, "manufacturer")
    con = get_or_create_supplier(form.contractor_name.data, "contractor")
    rec.manufacturer_id = man.id if man else None
    rec.contractor_id = con.id if con else None


@bp.route("/wells/<int:well_id>/install/new", methods=["GET", "POST"])
@permission_required("install", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = InstallForm()
    if form.validate_on_submit():
        rec = PumpInstallation(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("نصب/کشیدن پمپ ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("install.detail", record_id=rec.id))
    return render_template("install/form.html", form=form, well=well, title="ثبت نصب/کشیدن پمپ")


@bp.route("/install/<int:record_id>")
@permission_required("install", "view")
def detail(record_id):
    rec = db.get_or_404(PumpInstallation, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="pump_installations", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("install/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/install/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("install", "edit")
def edit(record_id):
    rec = db.get_or_404(PumpInstallation, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("install.detail", record_id=rec.id))
    form = InstallForm(obj=rec)
    if request.method == "GET":
        form.manufacturer_name.data = rec.manufacturer.name if rec.manufacturer else ""
        form.contractor_name.data = rec.contractor.name if rec.contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("نصب/کشیدن به‌روزرسانی شد.", "success")
        return redirect(url_for("install.detail", record_id=rec.id))
    return render_template("install/form.html", form=form, well=rec.well, title="ویرایش نصب/کشیدن")


@bp.route("/install/<int:record_id>/delete", methods=["POST"])
@permission_required("install", "delete")
def delete(record_id):
    rec = db.get_or_404(PumpInstallation, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد نصب/کشیدن حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(PumpInstallation, record_id)
    if not current_user.has_permission("install", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("install.detail", record_id=record_id))


@bp.route("/install/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/install/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/install/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/install/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: main_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : main
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("main", __name__)

from app.blueprints.main import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

import re
from datetime import datetime
from io import BytesIO

from flask import render_template, request, flash, redirect, url_for, send_file
from flask_login import login_required, current_user
from flask_wtf import FlaskForm
from wtforms import (StringField, PasswordField, BooleanField, TextAreaField,
                     SelectMultipleField, SubmitField)
from wtforms.validators import DataRequired, Optional, Length
from app.extensions import db
from app.blueprints.main import bp
from app.security import permission_required
from app.models.well import Well
from app.models.org import OrgUnit
from app.models.rbac import User, Role, LoginLog
from app.models.constants import WELL_STATUSES


@bp.route("/")
@login_required
def dashboard():
    well_count = db.session.scalar(db.select(db.func.count(Well.id))) or 0
    org_count = db.session.scalar(db.select(db.func.count(OrgUnit.id))) or 0
    mapped = db.session.scalar(
        db.select(db.func.count(Well.id)).where(Well.latitude.isnot(None))
    ) or 0
    in_circuit = db.session.scalar(
        db.select(db.func.count(Well.id)).where(Well.status == "in_circuit")
    ) or 0

    # Wells grouped by status (for the doughnut + legend).
    status_labels = dict(WELL_STATUSES)
    rows = db.session.execute(
        db.select(Well.status, db.func.count(Well.id)).group_by(Well.status)
    ).all()
    by_status = [
        {"key": s or "unknown", "label": status_labels.get(s, s or "نامشخص"), "count": n}
        for s, n in rows
    ]

    # Urban vs rural split.
    kind_rows = db.session.execute(
        db.select(Well.well_kind, db.func.count(Well.id)).group_by(Well.well_kind)
    ).all()
    kind_counts = {k or "unknown": n for k, n in kind_rows}

    # Coverage of mapped wells (for the progress ring).
    mapped_pct = round((mapped / well_count) * 100) if well_count else 0

    recent_wells = db.session.scalars(
        db.select(Well).order_by(Well.created_at.desc()).limit(6)
    ).all()

    # --- Attention panel (read-only surfacing of existing signals) ---
    from app.models.quality import WaterQuality
    quality_issue_wells = db.session.scalar(
        db.select(db.func.count(db.distinct(WaterQuality.well_id))).where(
            db.or_(
                WaterQuality.is_potable.is_(False),
                WaterQuality.turbidity.isnot(None),
                WaterQuality.sholat.isnot(None),
            )
        )
    ) or 0
    under_rehab = db.session.scalar(
        db.select(db.func.count(Well.id)).where(Well.status == "under_rehab")
    ) or 0

    top_priority = []
    try:
        from app.services import prioritization
        ranked, _ = prioritization.compute(top=5)
        top_priority = ranked
    except Exception:
        top_priority = []

    try:
        from app.services import alerts as alert_svc
        alert_summary = alert_svc.summary(include_predictive=False)  # fast count
    except Exception:
        alert_summary = {"total": 0, "by_severity": {"high": 0, "medium": 0, "low": 0}}

    # --- Operational signals (zone-based dashboard) ---
    from flask import url_for
    from datetime import date, timedelta
    from app.models.pump_asset import PumpInstallation
    from app.models.production import MonthlyProduction as MP
    from app.models.rehab import Rehabilitation
    from app.models.maintenance import MaintenanceRecord

    running_pumps = db.session.scalar(
        db.select(db.func.count(PumpInstallation.id)).where(PumpInstallation.pull_date.is_(None))) or 0
    fail_cutoff = date.today() - timedelta(days=365)
    recent_failures = db.session.scalar(
        db.select(db.func.count(PumpInstallation.id)).where(
            PumpInstallation.pull_date.isnot(None), PumpInstallation.pull_date >= fail_cutoff)) or 0

    e_sum = db.session.execute(
        db.select(db.func.sum(MP.energy_kwh), db.func.sum(MP.production_m3))
        .where(MP.energy_kwh.isnot(None))).first()
    fleet_se = round(e_sum[0] / e_sum[1], 2) if (e_sum and e_sum[0] and e_sum[1]) else None

    tr = db.session.execute(
        db.select(MP.jyear, MP.jmonth, db.func.sum(MP.production_m3))
        .where(MP.production_m3.isnot(None)).group_by(MP.jyear, MP.jmonth)
        .order_by(MP.jyear, MP.jmonth)).all()[-18:]
    trend = {"labels": [f"{y}/{m:02d}" for y, m, _ in tr],
             "series": [round((v or 0) / 1e6, 2) for _, _, v in tr]}

    activity = []
    for i in db.session.scalars(db.select(PumpInstallation).where(
            PumpInstallation.pull_date.isnot(None)).order_by(PumpInstallation.pull_date.desc()).limit(5)).all():
        activity.append({"date": i.pull_date, "well_id": i.well_id, "well": i.well.name,
                         "type": "کشیدن پمپ", "detail": i.removal_reason or "", "icon": "bi-tools",
                         "color": "#2f6bff", "url": url_for("install.detail", record_id=i.id)})
    for r in db.session.scalars(db.select(Rehabilitation).where(
            Rehabilitation.rehab_date.isnot(None)).order_by(Rehabilitation.rehab_date.desc()).limit(5)).all():
        activity.append({"date": r.rehab_date, "well_id": r.well_id, "well": r.well.name,
                         "type": "بهسازی", "detail": r.reason or "", "icon": "bi-arrow-repeat",
                         "color": "#e0463e", "url": url_for("rehab.detail", record_id=r.id)})
    for m in db.session.scalars(db.select(MaintenanceRecord).where(
            MaintenanceRecord.report_date.isnot(None)).order_by(MaintenanceRecord.report_date.desc()).limit(5)).all():
        activity.append({"date": m.report_date, "well_id": m.well_id, "well": m.well.name,
                         "type": "نگهداری", "detail": m.fault_desc or "", "icon": "bi-wrench-adjustable",
                         "color": "#6a7180", "url": url_for("maintenance.detail", record_id=m.id)})
    activity = sorted([a for a in activity if a["date"]], key=lambda a: a["date"], reverse=True)[:7]

    import jdatetime
    today = jdatetime.date.today().strftime("%Y/%m/%d")
    return render_template(
        "main/dashboard.html",
        today=today, well_count=well_count, org_count=org_count,
        mapped=mapped, mapped_pct=mapped_pct, in_circuit=in_circuit,
        by_status=by_status, urban=kind_counts.get("urban", 0),
        rural=kind_counts.get("rural", 0), recent_wells=recent_wells,
        status_labels=status_labels, quality_issue_wells=quality_issue_wells,
        under_rehab=under_rehab, top_priority=top_priority, alert_summary=alert_summary,
        running_pumps=running_pumps, recent_failures=recent_failures,
        fleet_se=fleet_se, trend=trend, activity=activity,
    )


@bp.route("/admin/users")
@permission_required("admin", "view")
def admin_users():
    q = (request.args.get("q") or "").strip()
    query = db.select(User).order_by(User.username)
    users = db.session.scalars(query).unique().all()
    if q:
        ql = q.lower()
        users = [u for u in users if ql in (u.username or "").lower()
                 or ql in (u.full_name or "").lower()
                 or ql in (u.personnel_code or "").lower()]
    return render_template("main/users.html", users=users, q=q)


@bp.route("/admin/economics", methods=["GET", "POST"])
@permission_required("admin", "view")
def admin_economics():
    from flask import request, flash, redirect, url_for
    from flask_login import current_user
    from app.models.economic import EconomicParam
    from app.services import economics
    from app.utils.dates import to_english_digits
    economics.ensure_params()
    rows = db.session.scalars(db.select(EconomicParam).order_by(EconomicParam.id)).all()
    if request.method == "POST":
        if not current_user.has_permission("admin", "edit"):
            flash("مجوز ویرایش ندارید.", "warning")
            return redirect(url_for("main.admin_economics"))
        for p in rows:
            v = to_english_digits(request.form.get(f"v_{p.id}", "")).strip()
            try:
                p.value = float(v)
            except ValueError:
                pass
        db.session.commit()
        flash("پارامترهای اقتصادی به‌روزرسانی شد.", "success")
        return redirect(url_for("main.admin_economics"))
    return render_template("main/economics.html", params=rows,
                           can_edit=current_user.has_permission("admin", "edit"))
# ==================== دسته ۵ — مدیریت کاربران ====================

def _password_errors(pw):
    errs = []
    if len(pw) < 8:
        errs.append("حداقل ۸ کاراکتر")
    if not re.search(r"[A-Z]", pw):
        errs.append("یک حرف بزرگ (A-Z)")
    if not re.search(r"[a-z]", pw):
        errs.append("یک حرف کوچک (a-z)")
    if not re.search(r"\d", pw):
        errs.append("یک عدد")
    if not re.search(r"[^A-Za-z0-9]", pw):
        errs.append("یک کاراکتر ویژه (!@#...)")
    return errs


class UserForm(FlaskForm):
    first_name = StringField("نام", validators=[Optional(), Length(max=60)])
    last_name = StringField("نام خانوادگی", validators=[Optional(), Length(max=60)])
    national_id = StringField("کد ملی", validators=[Optional(), Length(max=10)])
    personnel_code = StringField("کد پرسنلی", validators=[Optional(), Length(max=30)])
    position = StringField("سمت سازمانی", validators=[Optional(), Length(max=120)])
    phone = StringField("شماره تماس", validators=[Optional(), Length(max=20)])
    email = StringField("ایمیل", validators=[Optional(), Length(max=120)])
    username = StringField("نام کاربری", validators=[DataRequired(), Length(max=60)])
    password = PasswordField("رمز عبور", validators=[Optional()])
    is_active = BooleanField("فعال", default=True)
    role_ids = SelectMultipleField("نقش‌ها", coerce=int)
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


def _set_roles(form):
    form.role_ids.choices = [
        (r.id, r.name) for r in db.session.scalars(db.select(Role).order_by(Role.name)).unique()]


@bp.route("/admin/users/new", methods=["GET", "POST"])
@permission_required("admin", "edit")
def user_create():
    form = UserForm()
    _set_roles(form)
    if form.validate_on_submit():
        uname = form.username.data.strip()
        if db.session.scalar(db.select(User).filter_by(username=uname)):
            flash("این نام کاربری قبلاً استفاده شده است.", "danger")
        else:
            pw = form.password.data or ""
            errs = _password_errors(pw)
            if errs:
                flash("رمز عبور ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
            else:
                u = User(username=uname, is_active=bool(form.is_active.data))
                _apply_user(form, u)
                u.set_password(pw)
                db.session.add(u)
                db.session.commit()
                flash("کاربر ایجاد شد.", "success")
                return redirect(url_for("main.admin_users"))
    return render_template("main/user_form.html", form=form, title="ایجاد کاربر", is_new=True)


@bp.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
@permission_required("admin", "edit")
def user_edit(user_id):
    u = db.get_or_404(User, user_id)
    form = UserForm(obj=u)
    _set_roles(form)
    if request.method == "GET":
        form.role_ids.data = [r.id for r in u.roles]
    if form.validate_on_submit():
        uname = form.username.data.strip()
        clash = db.session.scalar(db.select(User).filter(User.username == uname, User.id != u.id))
        if clash:
            flash("این نام کاربری قبلاً استفاده شده است.", "danger")
        else:
            u.username = uname
            u.is_active = bool(form.is_active.data)
            _apply_user(form, u)
            if form.password.data:
                errs = _password_errors(form.password.data)
                if errs:
                    flash("رمز عبور ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
                    return render_template("main/user_form.html", form=form, title="ویرایش کاربر", is_new=False, obj=u)
                u.set_password(form.password.data)
            db.session.commit()
            flash("کاربر به‌روزرسانی شد.", "success")
            return redirect(url_for("main.admin_users"))
    return render_template("main/user_form.html", form=form, title="ویرایش کاربر", is_new=False, obj=u)


def _apply_user(form, u):
    u.first_name = form.first_name.data or None
    u.last_name = form.last_name.data or None
    u.full_name = (" ".join(x for x in [u.first_name, u.last_name] if x)).strip() or None
    u.national_id = form.national_id.data or None
    u.personnel_code = form.personnel_code.data or None
    u.position = form.position.data or None
    u.phone = form.phone.data or None
    u.email = form.email.data or None
    u.notes = form.notes.data or None
    u.roles = db.session.scalars(
        db.select(Role).filter(Role.id.in_(form.role_ids.data or []))).unique().all()


@bp.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@permission_required("admin", "edit")
def user_toggle_active(user_id):
    u = db.get_or_404(User, user_id)
    if u.id == current_user.id:
        flash("نمی‌توانید حساب خودتان را غیرفعال کنید.", "warning")
    else:
        u.is_active = not u.is_active
        db.session.commit()
        flash("وضعیت کاربر تغییر کرد.", "info")
    return redirect(url_for("main.admin_users"))


@bp.route("/admin/users/<int:user_id>/reset-password", methods=["POST"])
@permission_required("admin", "edit")
def user_reset_password(user_id):
    u = db.get_or_404(User, user_id)
    new_pw = (request.form.get("new_password") or "").strip()
    errs = _password_errors(new_pw)
    if errs:
        flash("رمز عبور ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
    else:
        u.set_password(new_pw)
        db.session.commit()
        flash(f"رمز عبور «{u.username}» بازنشانی شد.", "success")
    return redirect(url_for("main.admin_users"))


@bp.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@permission_required("admin", "delete")
def user_delete(user_id):
    u = db.get_or_404(User, user_id)
    if u.id == current_user.id:
        flash("نمی‌توانید حساب خودتان را حذف کنید.", "warning")
    elif u.is_superuser:
        flash("حذف مدیر کل مجاز نیست.", "warning")
    else:
        db.session.delete(u)
        db.session.commit()
        flash("کاربر حذف شد.", "info")
    return redirect(url_for("main.admin_users"))


@bp.route("/admin/users/activity")
@permission_required("admin", "view")
def user_activity():
    logs = db.session.scalars(
        db.select(LoginLog).order_by(LoginLog.login_at.desc()).limit(1000)).all()
    rows = _activity_rows(logs)
    return render_template("main/user_activity.html", rows=rows)


def _fmt_duration(sec):
    if sec is None:
        return "—"
    h, rem = divmod(int(sec), 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}س {m}د"
    if m:
        return f"{m}د {s}ث"
    return f"{s}ث"


def _activity_rows(logs):
    counts = {}
    for lg in logs:
        counts[lg.user_id] = counts.get(lg.user_id, 0) + 1
    out = []
    for lg in logs:
        u = lg.user
        out.append({
            "full_name": (u.full_name if u else None) or (u.username if u else "—"),
            "username": u.username if u else "—",
            "login_at": lg.login_at,
            "logout_at": lg.logout_at,
            "duration": _fmt_duration(lg.duration_seconds),
            "ip": lg.ip_address or "—",
            "agent": lg.user_agent or "—",
            "count": counts.get(lg.user_id, 0),
            "online": lg.logout_at is None,
        })
    return out


@bp.route("/admin/users/activity/export")
@permission_required("admin", "view")
def user_activity_export():
    import pandas as pd
    logs = db.session.scalars(db.select(LoginLog).order_by(LoginLog.login_at.desc())).all()
    rows = _activity_rows(logs)
    df = pd.DataFrame([{
        "نام و نام خانوادگی": r["full_name"],
        "نام کاربری": r["username"],
        "تاریخ/ساعت ورود": r["login_at"].strftime("%Y-%m-%d %H:%M") if r["login_at"] else "",
        "ساعت خروج": r["logout_at"].strftime("%Y-%m-%d %H:%M") if r["logout_at"] else "",
        "مدت حضور": r["duration"],
        "IP": r["ip"],
        "مرورگر/دستگاه": r["agent"],
        "تعداد ورود": r["count"],
        "وضعیت": "آنلاین" if r["online"] else "آفلاین",
    } for r in rows])
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        (df if not df.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="فعالیت کاربران", index=False)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     download_name=f"user_activity_{datetime.now():%Y%m%d}.xlsx")


class ProfileForm(FlaskForm):
    first_name = StringField("نام", validators=[Optional(), Length(max=60)])
    last_name = StringField("نام خانوادگی", validators=[Optional(), Length(max=60)])
    phone = StringField("شماره تماس", validators=[Optional(), Length(max=20)])
    email = StringField("ایمیل", validators=[Optional(), Length(max=120)])
    username = StringField("نام کاربری", validators=[DataRequired(), Length(max=60)])
    current_password = PasswordField("رمز فعلی", validators=[Optional()])
    new_password = PasswordField("رمز جدید", validators=[Optional()])
    submit = SubmitField("ذخیره تغییرات")


@bp.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    form = ProfileForm(obj=current_user)
    if form.validate_on_submit():
        uname = form.username.data.strip()
        clash = db.session.scalar(db.select(User).filter(User.username == uname, User.id != current_user.id))
        if clash:
            flash("این نام کاربری قبلاً استفاده شده است.", "danger")
        else:
            current_user.username = uname
            current_user.first_name = form.first_name.data or None
            current_user.last_name = form.last_name.data or None
            current_user.full_name = (" ".join(x for x in [current_user.first_name, current_user.last_name] if x)).strip() or None
            current_user.phone = form.phone.data or None
            current_user.email = form.email.data or None
            if form.new_password.data:
                if not current_user.check_password(form.current_password.data or ""):
                    flash("رمز فعلی نادرست است.", "danger")
                    return render_template("main/profile.html", form=form)
                errs = _password_errors(form.new_password.data)
                if errs:
                    flash("رمز جدید ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
                    return render_template("main/profile.html", form=form)
                current_user.set_password(form.new_password.data)
            db.session.commit()
            flash("پروفایل به‌روزرسانی شد.", "success")
            return redirect(url_for("main.profile"))
    return render_template("main/profile.html", form=form)




################################################################################
# FILE: maintenance_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : maintenance
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("maintenance", __name__)

from app.blueprints.maintenance import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.maintenance import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.maintenance import MaintenanceRecord
from app.models.pump_asset import PumpInstallation
from app.models.audit import RecordHistory
from app.models.constants import MAINTENANCE_TYPES, MAINTENANCE_CATEGORIES
from app.blueprints.install.routes import get_or_create_supplier
from app.utils.dates import format_jalali as _jd
from app import workflow


class MaintenanceForm(FlaskForm):
    report_date = JalaliDateField("تاریخ گزارش", validators=[Optional()])
    maint_type = SelectField("نوع نگهداری", choices=[("", "—")] + MAINTENANCE_TYPES, validators=[Optional()])
    category = SelectField("دسته", choices=[("", "—")] + MAINTENANCE_CATEGORIES, validators=[Optional()])
    down_from = JalaliDateField("شروع خاموشی", validators=[Optional()])
    down_to = JalaliDateField("پایان خاموشی", validators=[Optional()])
    downtime_hours = FloatField("مدت خاموشی (ساعت)", validators=[Optional()])
    fault_desc = TextAreaField("شرح خرابی", validators=[Optional()])
    action_taken = TextAreaField("اقدام انجام‌شده", validators=[Optional()])
    parts_replaced = StringField("قطعات تعویض‌شده", validators=[Optional()])
    contractor_name = StringField("پیمانکار/مجری", validators=[Optional()])
    cost = FloatField("هزینه", validators=[Optional()])
    pump_installation_id = SelectField("نصب مرتبط", coerce=int, validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_installs(self, well_id):
        installs = db.session.scalars(
            db.select(PumpInstallation).filter_by(well_id=well_id)
            .order_by(PumpInstallation.install_date.desc())).all()
        self.pump_installation_id.choices = [(0, "—")] + [
            (i.id, f"نصب {_jd(i.install_date, '') } — {i.pump_type or ''}".strip())
            for i in installs]


SIMPLE = ["report_date", "maint_type", "category", "down_from", "down_to",
          "downtime_hours", "fault_desc", "action_taken", "parts_replaced", "cost", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data or None)
    rec.pump_installation_id = form.pump_installation_id.data or None
    c = get_or_create_supplier(form.contractor_name.data, "contractor")
    rec.contractor_id = c.id if c else None


@bp.route("/wells/<int:well_id>/maintenance/new", methods=["GET", "POST"])
@permission_required("maintenance", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = MaintenanceForm()
    form.populate_installs(well.id)
    if form.validate_on_submit():
        rec = MaintenanceRecord(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("نگهداری ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("maintenance.detail", record_id=rec.id))
    return render_template("maintenance/form.html", form=form, well=well, title="ثبت نگهداری/تعمیر")


@bp.route("/maintenance/<int:record_id>")
@permission_required("maintenance", "view")
def detail(record_id):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="maintenance_records", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())).all()
    return render_template("maintenance/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           type_labels=dict(MAINTENANCE_TYPES),
                           cat_labels=dict(MAINTENANCE_CATEGORIES))


@bp.route("/maintenance/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("maintenance", "edit")
def edit(record_id):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("maintenance.detail", record_id=rec.id))
    form = MaintenanceForm(obj=rec)
    form.populate_installs(rec.well_id)
    if request.method == "GET":
        form.contractor_name.data = rec.contractor.name if rec.contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("نگهداری به‌روزرسانی شد.", "success")
        return redirect(url_for("maintenance.detail", record_id=rec.id))
    return render_template("maintenance/form.html", form=form, well=rec.well, title="ویرایش نگهداری")


@bp.route("/maintenance/<int:record_id>/delete", methods=["POST"])
@permission_required("maintenance", "delete")
def delete(record_id):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد نگهداری حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    if not current_user.has_permission("maintenance", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("maintenance.detail", record_id=record_id))


@bp.route("/maintenance/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/maintenance/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/maintenance/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/maintenance/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: mechanic_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : mechanic
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("mechanic", __name__, url_prefix="/mechanic")

from app.blueprints.mechanic import routes  # noqa: E402,F401


##########################################################################################
# FILE : routes.py
##########################################################################################

"""ماژول کارگاه مکانیک — فاز ۲: فهرست + فرم ورود/ویرایش/جزئیات."""
from flask import render_template, redirect, url_for, flash
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional, ValidationError

from app.extensions import db
from app.blueprints.mechanic import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField, PersianIntegerField
from app.models.well import Well
from app.models.mechanic import MechanicEvent

OP_TYPES = [
    ("", "—"),
    ("کشیدن", "کشیدن"),
    ("نصب", "نصب"),
    ("جمع‌آوری", "جمع‌آوری"),
    ("نصب جدید", "نصب جدید"),
]
COND = [("", "—"), ("نو", "نو"), ("تعمیری", "تعمیری")]


class MechanicForm(FlaskForm):
    op_date = JalaliDateField("تاریخ عملیات", validators=[Optional()])
    op_type = SelectField("نوع عملیات", choices=OP_TYPES, validators=[Optional()])
    fault_type = SelectField("نوع خرابی", choices=[("", "—"), ("سوختگی", "سوختگی"),
                             ("ایراد مکانیکی", "ایراد مکانیکی"), ("کاهش آبدهی", "کاهش آبدهی"),
                             ("سایر", "سایر")], validators=[Optional()])
    fault_description = TextAreaField("شرح خرابی از نظر بهره‌بردار", validators=[Optional()])
    contractor = StringField("نام پیمانکار", validators=[Optional()])
    pm_form_no = StringField("فرم نصب در PM", validators=[Optional()])

    motor_desc = StringField("مشخصات موتور", validators=[Optional()])
    motor_condition = SelectField("وضعیت موتور", choices=COND, validators=[Optional()])
    pump_desc = StringField("مشخصات پمپ", validators=[Optional()])
    pump_condition = SelectField("وضعیت پمپ", choices=COND, validators=[Optional()])

    tip_change = StringField("تغییر تیپ", validators=[Optional()])
    prev_install_date = JalaliDateField("تاریخ نصب قبلی", validators=[Optional()])
    well_depth = PersianFloatField("عمق چاه (متر)", validators=[Optional()])
    prev_install_depth = PersianFloatField("عمق نصب قبلی (متر)", validators=[Optional()])
    curr_install_depth = PersianFloatField("عمق نصب فعلی (متر)", validators=[Optional()])
    static_level = PersianFloatField("سطح استاتیک (متر)", validators=[Optional()])
    dynamic_level = PersianFloatField("سطح دینامیک (متر)", validators=[Optional()])
    path_loss = PersianFloatField("تلفات مسیر (متر)", validators=[Optional()])
    network_pressure = PersianFloatField("فشار شبکه (متر)", validators=[Optional()])
    total_head = PersianFloatField("هد کلی (متر)", validators=[Optional()])
    design_flow = PersianFloatField("دبی طراحی (l/s)", validators=[Optional()])
    pipe_diameter = PersianFloatField("قطر لوله آبده (اینچ)", validators=[Optional()])

    pt_date = JalaliDateField("تاریخ آزمایش پمپاژ کارگاه", validators=[Optional()])
    pt_pressure = PersianFloatField("فشار آزمایش پمپاژ", validators=[Optional()])
    pt_flow = PersianFloatField("دبی آزمایش پمپاژ (l/s)", validators=[Optional()])

    cable_size = StringField("سایز کابل", validators=[Optional()])
    cable_change = StringField("تغییر سایز/تیپ کابل", validators=[Optional()])
    starter = StringField("راه‌انداز", validators=[Optional()])

    months_worked = PersianIntegerField("تعداد ماه‌های کارکرد", validators=[Optional()])
    mechanic_opinion = TextAreaField("نظر کارگاه مکانیک", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def validate_design_flow(self, field):
        if field.data is not None and field.data < 0:
            raise ValidationError("دبی نمی‌تواند منفی باشد.")


FIELDS = ["op_date", "op_type", "fault_type", "fault_description", "contractor", "pm_form_no",
          "motor_desc", "motor_condition", "pump_desc", "pump_condition",
          "tip_change", "prev_install_date", "well_depth", "prev_install_depth",
          "curr_install_depth", "static_level", "dynamic_level", "path_loss",
          "network_pressure", "total_head", "design_flow", "pipe_diameter",
          "pt_date", "pt_pressure", "pt_flow", "cable_size", "cable_change",
          "starter", "months_worked", "mechanic_opinion", "notes"]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data if getattr(form, f).data not in ("",) else None)


@bp.route("/")
@permission_required("mechanic", "view")
def index():
    events = db.session.scalars(
        db.select(MechanicEvent).join(Well, MechanicEvent.well_id == Well.id)
        .order_by(MechanicEvent.op_date.desc())
    ).all()
    op_types = sorted({e.op_type for e in events if e.op_type})
    return render_template("mechanic/index.html", events=events, op_types=op_types)

@bp.route("/well/<int:well_id>/new", methods=["GET", "POST"])
@permission_required("mechanic", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = MechanicForm()
    if form.validate_on_submit():
        rec = MechanicEvent(well_id=well.id, created_by_id=current_user.id, source="manual")
        _apply(form, rec)
        db.session.add(rec)
        db.session.commit()
        flash("رویداد کارگاه مکانیک ثبت شد.", "success")
        return redirect(url_for("mechanic.detail", record_id=rec.id))
    return render_template("mechanic/form.html", form=form, well=well, title="ثبت رویداد کارگاه مکانیک")


@bp.route("/<int:record_id>")
@permission_required("mechanic", "view")
def detail(record_id):
    rec = db.get_or_404(MechanicEvent, record_id)
    return render_template("mechanic/detail.html", rec=rec, well=rec.well, stages=STAGES)


@bp.route("/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("mechanic", "edit")
def edit(record_id):
    rec = db.get_or_404(MechanicEvent, record_id)
    form = MechanicForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        db.session.commit()
        flash("رویداد به‌روزرسانی شد.", "success")
        return redirect(url_for("mechanic.detail", record_id=rec.id))
    return render_template("mechanic/form.html", form=form, well=rec.well, title="ویرایش رویداد کارگاه مکانیک")


@bp.route("/<int:record_id>/delete", methods=["POST"])
@permission_required("mechanic", "delete")
def delete(record_id):
    rec = db.get_or_404(MechanicEvent, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رویداد حذف شد.", "info")
    return redirect(url_for("mechanic.index"))
# ---------- انبار قطعات (فاز ۳) ----------
from app.models.mechanic_part import MechanicPart
from wtforms import IntegerField


class PartForm(FlaskForm):
    equipment = SelectField("نوع تجهیز", choices=[("موتور", "موتور"), ("پمپ", "پمپ")],
                            validators=[Optional()])
    part_name = StringField("نام قطعه", validators=[Optional()])
    total_count = PersianIntegerField("تعداد کل", validators=[Optional()])
    usable = PersianIntegerField("قابل استفاده", validators=[Optional()])
    scrap = PersianIntegerField("اسقاط", validators=[Optional()])
    new_count = PersianIntegerField("نو", validators=[Optional()])
    repaired = PersianIntegerField("تعمیری", validators=[Optional()])
    inventory_code = StringField("کد انباری قطعه", validators=[Optional()])
    part_type = StringField("نوع قطعه", validators=[Optional()])
    manufacturer = StringField("سازنده قطعه", validators=[Optional()])
    entry_date = JalaliDateField("تاریخ ورود", validators=[Optional()])
    notes = StringField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


@bp.route("/parts")
@permission_required("mechanic", "view")
def parts():
    motor = db.session.scalars(db.select(MechanicPart).filter_by(equipment="موتور")
                               .order_by(MechanicPart.id)).all()
    pump = db.session.scalars(db.select(MechanicPart).filter_by(equipment="پمپ")
                              .order_by(MechanicPart.id)).all()
    return render_template("mechanic/parts.html", motor=motor, pump=pump)


@bp.route("/parts/new", methods=["GET", "POST"])
@permission_required("mechanic", "create")
def part_create():
    form = PartForm()
    if form.validate_on_submit():
        p = MechanicPart(created_by_id=current_user.id)
        for f in ["equipment", "part_name", "total_count", "usable", "scrap",
                  "new_count", "repaired", "inventory_code", "part_type",
                  "manufacturer", "entry_date", "notes"]:
            setattr(p, f, getattr(form, f).data)
        db.session.add(p)
        db.session.commit()
        flash("قطعه ثبت شد.", "success")
        return redirect(url_for("mechanic.parts"))
    return render_template("mechanic/part_form.html", form=form, title="افزودن قطعه", part_id=None)


@bp.route("/parts/<int:part_id>/edit", methods=["GET", "POST"])
@permission_required("mechanic", "edit")
def part_edit(part_id):
    p = db.get_or_404(MechanicPart, part_id)
    form = PartForm(obj=p)
    if form.validate_on_submit():
        for f in ["equipment", "part_name", "total_count", "usable", "scrap",
                  "new_count", "repaired", "inventory_code", "part_type",
                  "manufacturer", "notes"]:
            setattr(p, f, getattr(form, f).data)
        p.updated_by_id = current_user.id
        db.session.commit()
        flash("قطعه به‌روزرسانی شد.", "success")
        return redirect(url_for("mechanic.parts"))
    return render_template("mechanic/part_form.html", form=form, title="ویرایش قطعه", part_id=p.id)


@bp.route("/parts/<int:part_id>/delete", methods=["POST"])
@permission_required("mechanic", "delete")
def part_delete(part_id):
    p = db.get_or_404(MechanicPart, part_id)
    db.session.delete(p)
    db.session.commit()
    flash("قطعه حذف شد.", "info")
    return redirect(url_for("mechanic.parts"))
# ---------- گردش کار مرحله‌ای (فاز ۴) ----------
from app.models.mechanic import MechanicStageLog

# مراحل به ترتیب فلوچارت
STAGES = [
    "اعلام حادثه",
    "تشخیص نوع خرابی",
    "دفتر فنی و مهندسی",
    "ویدئومتری",
    "بهسازی چاه",
    "انتخاب پمپ",
    "پمپاژ آزمایشی",
    "مونتاژ",
    "تست چاله پمپاژ",
    "دمونتاژ",
    "ارجاع به کارگاه نصب",
    "نصب نهایی",
    "پایان‌یافته",
]

FAULT_TYPES = ["سوختگی", "ایراد مکانیکی", "کاهش آبدهی", "سایر"]


@bp.route("/<int:record_id>/stage", methods=["POST"])
@permission_required("mechanic", "edit")
def change_stage(record_id):
    from flask import request
    rec = db.get_or_404(MechanicEvent, record_id)
    new_stage = request.form.get("to_stage", "").strip()
    note = request.form.get("note", "").strip()
    if new_stage and new_stage in STAGES and new_stage != rec.stage:
        log = MechanicStageLog(
            event_id=rec.id, from_stage=rec.stage, to_stage=new_stage,
            note=note or None, changed_by_id=current_user.id)
        rec.stage = new_stage
        db.session.add(log)
        db.session.commit()
        flash(f"مرحله به «{new_stage}» تغییر کرد.", "success")
    else:
        flash("مرحله‌ی معتبری انتخاب نشد.", "warning")
    return redirect(url_for("mechanic.detail", record_id=rec.id))
# ---------- گزارش‌های مدیریتی (فاز ۵) ----------
from sqlalchemy import func
from collections import Counter
import jdatetime

def classify_fault(event):
    """دسته‌بندی خرابی از روی fault_type یا متن fault_description."""
    # اگر نوع خرابی صریح ثبت شده، همان
    if event.fault_type:
        return event.fault_type
    text = (event.fault_description or "").strip()
    if not text:
        return "نامشخص"
    # دسته‌بندی بر اساس کلمات کلیدی
    rules = [
        ("سوختگی", ["سوخت", "سوختن"]),
        ("کاهش آبدهی", ["کاهش دبی", "کاهش آبدهی", "کمبود", "عدم آبده", "عدم ابده", "افت"]),
        ("ایراد مکانیکی", ["صدا", "لرزش", "گیرپاژ", "گیر و پاژ", "گیرو پاژ", "مکانیک"]),
        ("ایراد برقی", ["اهم", "شولات", "شولاتی", "آمپر", "امپر", "برق"]),
        ("هوادهی", ["هوادهی", "هوا"]),
        ("عملیات نصب/کشیدن", ["نصب", "کشیدن", "جمع آوری", "جمع‌آوری", "تجهیز", "جابجایی"]),
        ("تغییر فشار", ["فشار"]),
    ]
    for label, keywords in rules:
        if any(k in text for k in keywords):
            return label
    return "سایر"

@bp.route("/reports")
@permission_required("mechanic", "view")
def reports():
    from flask import request
    # فیلترها
    f_year = request.args.get("year", "").strip()
    f_optype = request.args.get("optype", "").strip()
    f_fault = request.args.get("fault", "").strip()

    q = db.select(MechanicEvent)
    events = db.session.scalars(q).all()

    # تبدیل تاریخ به سال شمسی برای فیلتر و نمودار
    def jyear(d):
        if not d:
            return None
        try:
            return jdatetime.date.fromgregorian(date=d).year
        except Exception:
            return None

    # اعمال فیلترها
    rows = []
    for e in events:
        jy = jyear(e.op_date)
        if f_year and str(jy) != f_year:
            continue
        if f_optype and (e.op_type or "") != f_optype:
            continue
        if f_fault and (e.fault_type or "") != f_fault:
            continue
        rows.append((e, jy))

    # آمار کلی
    total = len(rows)
    by_optype = Counter((e.op_type or "نامشخص") for e, _ in rows)
    by_fault = Counter(classify_fault(e) for e, _ in rows)
    by_year = Counter(jy for _, jy in rows if jy)
    by_contractor = Counter((e.contractor or "نامشخص") for e, _ in rows if e.contractor)
    by_stage = Counter((e.stage or "نامشخص") for e, _ in rows)

    # شاخص‌های بیشتر
    with_pt = sum(1 for e, _ in rows if e.pt_flow is not None)
    by_center = Counter()
    by_month = Counter()
    burnt = 0
    depths = []
    heads = []
    flows = []
    months_worked_list = []
    import jdatetime as _jd
    for e, jy in rows:
        # مرکز از توضیحات یا شرح در دسترس نیست؛ از fault_description رد می‌شویم
        # ماه شمسی برای روند ماهانه
        if e.op_date:
            try:
                jm = _jd.date.fromgregorian(date=e.op_date).month
                by_month[jm] += 1
            except Exception:
                pass
        f = classify_fault(e)
        if f == "سوختگی":
            burnt += 1
        if e.well_depth:
            depths.append(e.well_depth)
        if e.total_head:
            heads.append(e.total_head)
        if e.pt_flow:
            flows.append(e.pt_flow)
        if e.months_worked:
            months_worked_list.append(e.months_worked)

    def avg(lst):
        return round(sum(lst) / len(lst), 1) if lst else 0

    MONTHS_FA = ["", "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
                 "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
    by_month_named = {MONTHS_FA[m]: c for m, c in sorted(by_month.items())}

    stats = {
        "total": total,
        "install": by_optype.get("نصب", 0),
        "pull": by_optype.get("کشیدن", 0),
        "collect": by_optype.get("جمع‌آوری", 0) + by_optype.get("جمع آوری", 0),
        "with_pt": with_pt,
        "burnt": burnt,
        "burnt_pct": round(100 * burnt / total) if total else 0,
        "avg_depth": avg(depths),
        "avg_head": avg(heads),
        "avg_flow": avg(flows),
        "avg_months_worked": avg(months_worked_list),
        "by_optype": dict(by_optype),
        "by_fault": dict(by_fault),
        "by_year": dict(sorted(by_year.items())),
        "by_month": by_month_named,
        "by_contractor": by_contractor.most_common(10),
        "by_stage": dict(by_stage),
    }

    # گزینه‌های فیلتر
    all_years = sorted({jyear(e.op_date) for e in events if jyear(e.op_date)}, reverse=True)
    all_optypes = sorted({e.op_type for e in events if e.op_type})

    return render_template("mechanic/reports.html", stats=stats,
                           all_years=all_years, all_optypes=all_optypes,
                           fault_types=FAULT_TYPES,
                           cur={"year": f_year, "optype": f_optype, "fault": f_fault})
# ---------- گزارش انبار قطعات ----------
@bp.route("/parts/report")
@permission_required("mechanic", "view")
def parts_report():
    parts = db.session.scalars(db.select(MechanicPart).order_by(
        MechanicPart.equipment, MechanicPart.id)).all()

    total_usable = sum(p.usable or 0 for p in parts)
    total_scrap = sum(p.scrap or 0 for p in parts)
    total_new = sum(p.new_count or 0 for p in parts)
    total_repaired = sum(p.repaired or 0 for p in parts)

    # قطعات کم‌موجود: قابل‌استفاده صفر یا کمتر از ۳
    low_stock = [p for p in parts if (p.usable or 0) < 3]
    # قطعات بدون هیچ موجودی
    empty = [p for p in parts if (p.total_count or 0) == 0]

    by_eq = {"موتور": {"usable": 0, "scrap": 0}, "پمپ": {"usable": 0, "scrap": 0}}
    for p in parts:
        if p.equipment in by_eq:
            by_eq[p.equipment]["usable"] += p.usable or 0
            by_eq[p.equipment]["scrap"] += p.scrap or 0

    stats = {
        "total_parts": len(parts),
        "total_usable": total_usable,
        "total_scrap": total_scrap,
        "total_new": total_new,
        "total_repaired": total_repaired,
        "low_stock": low_stock,
        "empty_count": len(empty),
        "by_eq": by_eq,
    }
    parts_data = [{
        "equipment": p.equipment or "نامشخص",
        "name": p.part_name or "—",
        "total": p.total_count or 0,
        "new": p.new_count or 0,
        "repaired": p.repaired or 0,
        "scrap": p.scrap or 0,
        "usable": p.usable or 0,
        "installed": getattr(p, "installed_count", 0) or 0,
    } for p in parts]
    return render_template("mechanic/parts_report.html", stats=stats, parts_data=parts_data)




################################################################################
# FILE: operation_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : operation
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("operation", __name__)

from app.blueprints.operation import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.operation import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.utils.dates import to_english_digits
from app.models.well import Well
from app.models.flow import FlowTest, FlowTestPoint
from app.models.audit import RecordHistory
from app.models.constants import OPERATING_TYPES, FLOW_TEST_REASONS
from app import workflow

POINT_ROWS = 5
POINT_FIELDS = ["operating_type", "discharge_lps", "head_m", "drawdown_m",
                "dynamic_level_m", "pressure_atm", "amperes", "efficiency"]


class FlowTestForm(FlaskForm):
    test_date = JalaliDateField("تاریخ آزمایش", validators=[Optional()])
    test_reason = SelectField("دلیل آزمایش", choices=[("", "—")] + [(r, r) for r in FLOW_TEST_REASONS], validators=[Optional()])
    network_type = StringField("نوع شبکه", validators=[Optional()])
    electropump_type = StringField("تیپ الکتروپمپ", validators=[Optional()])
    electropump_type_prev = StringField("تیپ الکتروپمپ قبلی", validators=[Optional()])
    install_date = JalaliDateField("تاریخ آخرین نصب", validators=[Optional()])
    install_depth = FloatField("عمق نصب", validators=[Optional()])
    well_depth = FloatField("عمق چاه", validators=[Optional()])
    construction_type = StringField("نوع چاه (سیمانته/غیرسیمانته)", validators=[Optional()])
    allowed_q = FloatField("دبی مجاز", validators=[Optional()])
    design_q = FloatField("دبی طراحی", validators=[Optional()])
    license_q = FloatField("دبی پروانه", validators=[Optional()])
    power_subscription = StringField("اشتراک برق", validators=[Optional()])
    static_level = FloatField("سطح ایستایی", validators=[Optional()])
    last_rehab_date = JalaliDateField("تاریخ آخرین بهسازی", validators=[Optional()])
    pull_reason = StringField("علت کشیدن پمپ", validators=[Optional()])
    meter_status = StringField("کالیبراسیون/وضعیت کنتور", validators=[Optional()])
    meter_brand = StringField("برند/سایز کنتور", validators=[Optional()])
    starter_type = StringField("سیستم راه‌انداز", validators=[Optional()])
    capacitor_capacity = FloatField("ظرفیت خازن", validators=[Optional()])
    voltage_on = StringField("ولتاژ روشن", validators=[Optional()])
    voltage_off = StringField("ولتاژ خاموش", validators=[Optional()])
    ohm_ff = StringField("مقاومت اهمی ف-ف", validators=[Optional()])
    ohm_fg = StringField("مقاومت اهمی ف-ب", validators=[Optional()])
    line_pressure = FloatField("فشار خط (bar)", validators=[Optional()])
    regulated_pressure = StringField("فشار تنظیمی", validators=[Optional()])
    discharge_volume_m3 = FloatField("حجم تخلیه (m³)", validators=[Optional()])
    expert_note = TextAreaField("نظر کارشناس", validators=[Optional()])
    submit = SubmitField("ذخیره")


HEADER_FIELDS = [
    "test_date", "test_reason", "network_type", "electropump_type", "electropump_type_prev",
    "install_date", "install_depth", "well_depth", "construction_type", "allowed_q",
    "design_q", "license_q", "power_subscription", "static_level", "last_rehab_date",
    "pull_reason", "meter_status", "meter_brand", "starter_type", "capacitor_capacity",
    "voltage_on", "voltage_off", "ohm_ff", "ohm_fg", "line_pressure",
    "regulated_pressure", "discharge_volume_m3", "expert_note",
]


def _num(s):
    s = to_english_digits((s or "").strip())
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _apply_header(form, rec):
    for f in HEADER_FIELDS:
        setattr(rec, f, getattr(form, f).data)


def _apply_points(rec):
    """Rebuild child points from the manual table in request.form."""
    rec.points.clear()
    for i in range(1, POINT_ROWS + 1):
        vals = {f: request.form.get(f"pt-{i}-{f}", "").strip() for f in POINT_FIELDS}
        if not any(vals.values()):
            continue
        rec.points.append(FlowTestPoint(
            operating_no=i,
            operating_type=vals["operating_type"] or None,
            discharge_lps=_num(vals["discharge_lps"]),
            head_m=_num(vals["head_m"]),
            drawdown_m=_num(vals["drawdown_m"]),
            dynamic_level_m=_num(vals["dynamic_level_m"]),
            pressure_atm=_num(vals["pressure_atm"]),
            amperes=vals["amperes"] or None,
            efficiency=_num(vals["efficiency"]),
        ))


@bp.route("/wells/<int:well_id>/flow/new", methods=["GET", "POST"])
@permission_required("operation", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = FlowTestForm()
    if form.validate_on_submit():
        rec = FlowTest(well_id=well.id, created_by_id=current_user.id)
        _apply_header(form, rec)
        db.session.add(rec)
        _apply_points(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("دبی‌سنجی ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("operation.detail", record_id=rec.id))
    return render_template("operation/form.html", form=form, well=well,
                           title="ثبت دبی‌سنجی", points=[], operating_types=OPERATING_TYPES,
                           point_rows=POINT_ROWS, point_fields=POINT_FIELDS)


@bp.route("/flow/<int:record_id>")
@permission_required("operation", "view")
def detail(record_id):
    rec = db.get_or_404(FlowTest, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="flow_tests", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("operation/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/flow/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("operation", "edit")
def edit(record_id):
    rec = db.get_or_404(FlowTest, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("operation.detail", record_id=rec.id))
    form = FlowTestForm(obj=rec)
    if form.validate_on_submit():
        _apply_header(form, rec)
        _apply_points(rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("دبی‌سنجی به‌روزرسانی شد.", "success")
        return redirect(url_for("operation.detail", record_id=rec.id))
    return render_template("operation/form.html", form=form, well=rec.well,
                           title="ویرایش دبی‌سنجی", points=rec.points,
                           operating_types=OPERATING_TYPES, point_rows=POINT_ROWS,
                           point_fields=POINT_FIELDS)


@bp.route("/flow/<int:record_id>/delete", methods=["POST"])
@permission_required("operation", "delete")
def delete(record_id):
    rec = db.get_or_404(FlowTest, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("دبی‌سنجی حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(FlowTest, record_id)
    if not current_user.has_permission("operation", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("operation.detail", record_id=record_id))


@bp.route("/flow/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/flow/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/flow/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/flow/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: orgs_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : orgs
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("orgs", __name__, url_prefix="/orgs")

from app.blueprints.orgs import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, SubmitField
from wtforms.validators import DataRequired, Optional

from app.extensions import db
from app.blueprints.orgs import bp
from app.security import permission_required
from app.models.org import OrgUnit
from app.models.constants import ORG_UNIT_TYPES


class OrgUnitForm(FlaskForm):
    name = StringField("نام", validators=[DataRequired()])
    code = StringField("کد", validators=[Optional()])
    unit_type = SelectField("نوع", choices=ORG_UNIT_TYPES, validators=[DataRequired()])
    parent_id = SelectField("واحد بالادست", coerce=int, validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_parents(self, exclude_id=None):
        units = db.session.scalars(db.select(OrgUnit).order_by(OrgUnit.name)).all()
        choices = [(0, "— بدون والد —")]
        for u in units:
            if exclude_id and u.id == exclude_id:
                continue
            choices.append((u.id, f"{u.name} ({dict(ORG_UNIT_TYPES).get(u.unit_type, '')})"))
        self.parent_id.choices = choices


@bp.route("/")
@permission_required("orgs", "view")
def list_orgs():
    units = db.session.scalars(
        db.select(OrgUnit).order_by(OrgUnit.unit_type, OrgUnit.name)
    ).all()
    type_labels = dict(ORG_UNIT_TYPES)
    return render_template("orgs/list.html", units=units, type_labels=type_labels)


@bp.route("/new", methods=["GET", "POST"])
@permission_required("orgs", "create")
def create_org():
    form = OrgUnitForm()
    form.populate_parents()
    if form.validate_on_submit():
        unit = OrgUnit(
            name=form.name.data.strip(),
            code=(form.code.data or "").strip() or None,
            unit_type=form.unit_type.data,
            parent_id=form.parent_id.data or None,
        )
        db.session.add(unit)
        db.session.commit()
        flash("واحد سازمانی ایجاد شد.", "success")
        return redirect(url_for("orgs.list_orgs"))
    return render_template("orgs/form.html", form=form, title="واحد سازمانی جدید")


@bp.route("/<int:unit_id>/edit", methods=["GET", "POST"])
@permission_required("orgs", "edit")
def edit_org(unit_id):
    unit = db.get_or_404(OrgUnit, unit_id)
    form = OrgUnitForm(obj=unit)
    form.populate_parents(exclude_id=unit.id)
    if form.validate_on_submit():
        unit.name = form.name.data.strip()
        unit.code = (form.code.data or "").strip() or None
        unit.unit_type = form.unit_type.data
        unit.parent_id = form.parent_id.data or None
        db.session.commit()
        flash("واحد سازمانی به‌روزرسانی شد.", "success")
        return redirect(url_for("orgs.list_orgs"))
    return render_template("orgs/form.html", form=form, title="ویرایش واحد سازمانی")


@bp.route("/<int:unit_id>/delete", methods=["POST"])
@permission_required("orgs", "delete")
def delete_org(unit_id):
    unit = db.get_or_404(OrgUnit, unit_id)
    if unit.children:
        flash("این واحد دارای زیرمجموعه است و قابل حذف نیست.", "danger")
        return redirect(url_for("orgs.list_orgs"))
    db.session.delete(unit)
    db.session.commit()
    flash("واحد سازمانی حذف شد.", "info")
    return redirect(url_for("orgs.list_orgs"))





################################################################################
# FILE: permit_event_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : permit_event
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("permit_event", __name__, url_prefix="/permit-event")

from app.blueprints.permit_event import routes  # noqa: E402,F401


##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, BooleanField, SubmitField
from wtforms.validators import Optional, ValidationError

from app.extensions import db
from app.blueprints.permit_event import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.permit_event import WellPermitEvent
from app.models.audit import RecordHistory
from app import workflow


PERMIT_TYPES = [
    ("", "—"),
    ("بهره‌برداری عادی", "بهره‌برداری عادی"),
    ("حفر", "حفر"),
    ("تغییر محل", "تغییر محل"),
    ("کف‌شکنی", "کف‌شکنی"),
]


class PermitEventForm(FlaskForm):
    permit_type = SelectField("نوع پروانه", choices=PERMIT_TYPES, validators=[Optional()])
    permit_code = StringField("کد آخرین پروانه", validators=[Optional()])
    permit_no = StringField("شماره پروانه", validators=[Optional()])
    permit_date = JalaliDateField("تاریخ صدور", validators=[Optional()])
    expiry_date = JalaliDateField("تاریخ اعتبار/انقضا", validators=[Optional()])
    case_status = StringField("وضعیت پرونده", validators=[Optional()])
    klasse = StringField("کلاسه آب منطقه‌ای", validators=[Optional()])
    request_type = StringField("درخواست جاری (تمدید/صدور/جابجایی)", validators=[Optional()])
    followup_stage = StringField("مرحله پیگیری", validators=[Optional()])
    cost_paid = BooleanField("هزینه پرداخت شده", validators=[Optional()])
    expiry_penalty_rial = FloatField("جریمه انقضا (ریال)", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def validate_expiry_date(self, field):
        if field.data and self.permit_date.data and field.data < self.permit_date.data:
            raise ValidationError("تاریخ اعتبار نمی‌تواند پیش از تاریخ صدور باشد.")

    def validate_expiry_penalty_rial(self, field):
        if field.data is not None and field.data < 0:
            raise ValidationError("جریمه نمی‌تواند منفی باشد.")


SIMPLE = ["permit_type", "permit_code", "permit_no", "permit_date", "expiry_date",
          "case_status", "klasse", "request_type", "followup_stage", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data or None)
    rec.cost_paid = form.cost_paid.data
    rec.expiry_penalty_rial = form.expiry_penalty_rial.data


@bp.route("/wells/<int:well_id>/permit/new", methods=["GET", "POST"])
@permission_required("permit", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = PermitEventForm()
    if form.validate_on_submit():
        rec = WellPermitEvent(well_id=well.id, created_by_id=current_user.id, source="manual")
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("پروانه ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("permit_event.detail", record_id=rec.id))
    return render_template("permit_event/form.html", form=form, well=well, title="ثبت پروانه")


@bp.route("/permit/<int:record_id>")
@permission_required("permit", "view")
def detail(record_id):
    rec = db.get_or_404(WellPermitEvent, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="well_permit_events", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("permit_event/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/permit/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("permit", "edit")
def edit(record_id):
    rec = db.get_or_404(WellPermitEvent, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("permit_event.detail", record_id=rec.id))
    form = PermitEventForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("پروانه به‌روزرسانی شد.", "success")
        return redirect(url_for("permit_event.detail", record_id=rec.id))
    return render_template("permit_event/form.html", form=form, well=rec.well, title="ویرایش پروانه")


@bp.route("/permit/<int:record_id>/delete", methods=["POST"])
@permission_required("permit", "delete")
def delete(record_id):
    rec = db.get_or_404(WellPermitEvent, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد پروانه حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(WellPermitEvent, record_id)
    if not current_user.has_permission("permit", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("permit_event.detail", record_id=record_id))


@bp.route("/permit/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/permit/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/permit/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/permit/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")




################################################################################
# FILE: permit_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : permit
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("permit", __name__, url_prefix="/permits")

from app.blueprints.permit import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

"""وضعیت پروانه چاه‌ها.

صفحه‌ی تحلیل وضعیت پروانه بر پایه‌ی جدول production_trend (ستون‌های well_permit و
permit_expiry) و فرم ورود/ویرایش اطلاعات پروانه‌ی هر چاه (به‌روزرسانی مستقیم همان
جدول). تاریخ اعتبار به‌صورت شمسیِ فشرده (مثل 14050427) است و وضعیت انقضا با مقایسه
با تاریخ امروز محاسبه می‌شود.
"""
from datetime import date
from collections import Counter

import jdatetime
from flask import render_template, request, redirect, url_for, flash
from flask_login import login_required
from sqlalchemy import text

from app.extensions import db
from app.blueprints.permit import bp
from app.models.well import Well
from app.imports.wells_import import match_key


def _digits8(s):
    if not s:
        return None
    d = "".join(ch for ch in str(s) if ch.isdigit())
    return d if len(d) == 8 else None


def _today8():
    t = jdatetime.date.fromgregorian(date=date.today())
    return f"{t.year:04d}{t.month:02d}{t.day:02d}"


def _expiry_state(expiry, today8):
    e = _digits8(expiry)
    if not e:
        return "unknown"
    if e < today8:
        return "expired"
    ey, em = int(e[:4]), int(e[4:6])
    ty, tm = int(today8[:4]), int(today8[4:6])
    months_left = (ey - ty) * 12 + (em - tm)
    return "near" if months_left < 3 else "valid"


STATE_LABEL = {"expired": "منقضی شده", "near": "نزدیک انقضا",
               "valid": "معتبر", "unknown": "نامشخص"}


@bp.route("/")
@login_required
def index():
    today8 = _today8()
    # نگاشت نام چاه به شناسه‌ی چاه سامانه (برای لینک به جزئیات چاه)
    well_by_key = {}
    for w in db.session.scalars(db.select(Well)).all():
        well_by_key.setdefault(match_key(w.name), w.id)
    result = db.session.execute(text(
        "SELECT rowid, well_name, department, well_permit, permit_expiry, "
        "well_type, main_zone FROM production_trend"
    ))
    rows = []
    state_counter = Counter()
    type_counter = Counter()
    for r in result.fetchall():
        rowid, name, dept, permit, expiry, wtype, zone = r
        if not name or not str(name).strip():
            continue
        st = _expiry_state(expiry, today8)
        state_counter[st] += 1
        if permit and str(permit).strip() and str(permit).strip() != "0":
            type_counter[str(permit).strip()] += 1
            wid = well_by_key.get(match_key(name))
        rows.append({
            "well_id": wid, "rowid": rowid, "well_name": name, "office": dept or "—",
            "permit_type": (permit if permit and str(permit) != "0" else "—"),
            "expiry": expiry or "—", "state": st, "state_label": STATE_LABEL[st],
            "well_type": wtype or "—", "zone": zone or "—",
        })
    order = {"expired": 0, "near": 1, "valid": 2, "unknown": 3}
    rows.sort(key=lambda r: (order.get(r["state"], 9), str(r["expiry"])))

    stats = {
        "total": len(rows),
        "expired": state_counter.get("expired", 0),
        "near": state_counter.get("near", 0),
        "valid": state_counter.get("valid", 0),
        "unknown": state_counter.get("unknown", 0),
        "types": type_counter.most_common(),
    }
    return render_template("permit/index.html", rows=rows, stats=stats)


@bp.route("/<int:rowid>/edit", methods=["GET", "POST"])
@login_required
def edit(rowid):
    row = db.session.execute(
        text("SELECT rowid, well_name, department, well_permit, permit_expiry "
             "FROM production_trend WHERE rowid=:i"), {"i": rowid}
    ).fetchone()
    if not row:
        flash("چاه پیدا نشد.", "warning")
        return redirect(url_for("permit.index"))

    if request.method == "POST":
        permit = request.form.get("well_permit", "").strip()
        expiry = request.form.get("permit_expiry", "").strip()
        db.session.execute(
            text("UPDATE production_trend SET well_permit=:p, permit_expiry=:e "
                 "WHERE rowid=:i"),
            {"p": permit or None, "e": expiry or None, "i": rowid},
        )
        db.session.commit()
        flash("وضعیت پروانه به‌روزرسانی شد.", "success")
        return redirect(url_for("permit.index"))

    rec = {"rowid": row[0], "well_name": row[1], "office": row[2],
           "well_permit": row[3] or "", "permit_expiry": row[4] or ""}
    return render_template("permit/form.html", rec=rec)





################################################################################
# FILE: printing_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : printing
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("printing", __name__, url_prefix="/print")

from app.blueprints.printing import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

"""Official print-ready reports (browser Print-to-PDF, RTL/Persian, zero deps)."""
from flask import render_template, url_for

from app.extensions import db
from app.blueprints.printing import bp
from app.security import permission_required
from app.models.well import Well
from app.models.constants import WELL_KINDS, WELL_STATUSES
from app.utils.dates import format_jalali as fj


def _today():
    import jdatetime
    return jdatetime.date.today().strftime("%Y/%m/%d")


@bp.route("/well/<int:well_id>")
@permission_required("wells", "view")
def well_sheet(well_id):
    w = db.get_or_404(Well, well_id)
    from app.models.drilling import Drilling
    from app.models.pump_asset import PumpInstallation
    from app.models.baseline import WellBaseline
    dr = db.session.scalar(db.select(Drilling).filter_by(well_id=w.id)
                           .order_by(Drilling.end_date.desc()))
    inst = db.session.scalar(db.select(PumpInstallation).filter_by(well_id=w.id)
                             .order_by(PumpInstallation.install_date.desc()))
    base = db.session.scalar(db.select(WellBaseline).filter_by(well_id=w.id, is_current=True))
    return render_template(
        "print/well_sheet.html", w=w, tech=w.technical, dr=dr, inst=inst, base=base,
        kind_labels=dict(WELL_KINDS), status_labels=dict(WELL_STATUSES),
        today=_today(), doc_id=w.display_pm, fj=fj,
        back_url=url_for("wells.detail", well_id=w.id))


@bp.route("/pump-test/<int:record_id>")
@permission_required("pump_test", "view")
def pump_test(record_id):
    from app.models.pump_test import PumpTest
    rec = db.get_or_404(PumpTest, record_id)
    return render_template("print/pump_test.html", rec=rec, w=rec.well,
                           today=_today(), doc_id=rec.well.display_pm, fj=fj,
                           back_url=url_for("pump_test.detail", record_id=rec.id))


@bp.route("/flow/<int:record_id>")
@permission_required("operation", "view")
def flow(record_id):
    from app.models.flow import FlowTest
    rec = db.get_or_404(FlowTest, record_id)
    return render_template("print/flow.html", rec=rec, w=rec.well,
                           today=_today(), doc_id=rec.well.display_pm, fj=fj,
                           back_url=url_for("operation.detail", record_id=rec.id))





################################################################################
# FILE: production_trend_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : production_trend
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("production_trend", __name__, url_prefix="/production-trend")

from app.blueprints.production_trend import routes  # noqa: E402,F401



##########################################################################################
# FILE : column_import.py
##########################################################################################

"""افزودن سرستون جدید به جدول production_trend (روش A: ستون واقعی).

کاربر دسته/سال/ماه و فایل اکسل می‌دهد؛ سیستم نام ستون انگلیسی را طبق قرارداد
data_map می‌سازد، ستون را (در صورت نبود) به جدول اضافه می‌کند، و مقادیر را با
تطبیق نام چاه پر می‌کند. نام ستون‌ها اعتبارسنجی می‌شوند تا امن باشند.
"""
import re
import openpyxl
from sqlalchemy import text

from app.extensions import db
from app.imports.wells_import import match_key
from app.blueprints.production_trend.data_map import MONTHS_EN, MONTHS_FA
def _recalc_yearly(category, year):
    """ستون مجموع/میانگین سالانه را از روی ماه‌های موجود همان سال بازمحاسبه می‌کند."""
    cols_now = _existing_columns()
    # ماه‌های موجود این دسته/سال را جمع کن
    month_cols = []
    for mi in range(12):
        mc = make_column_name(category, year, mi)
        if mc in cols_now:
            month_cols.append(mc)
    if not month_cols:
        return

    # نام ستون سالانه
    year_col = make_column_name(category, year, None)
    if not year_col:
        return
    # اگر ستون سالانه نبود بساز
    if year_col not in cols_now:
        sqltype = CATEGORY_SQLTYPE.get(category, "REAL")
        db.session.execute(text(
            f'ALTER TABLE production_trend ADD COLUMN "{year_col}" {sqltype}'))
        db.session.commit()

    # تولید و کارکرد → جمع؛ فشار و دبی → میانگین
    is_sum = category in ("production", "runtime")
    sum_expr = " + ".join(f'COALESCE("{c}",0)' for c in month_cols)
    cnt_expr = " + ".join(f'(CASE WHEN "{c}" IS NOT NULL THEN 1 ELSE 0 END)' for c in month_cols)

    rows = db.session.execute(text(f'SELECT rowid, {sum_expr}, {cnt_expr} FROM production_trend')).fetchall()
    for rid, total, cnt in rows:
        if cnt and cnt > 0:
            val = total if is_sum else round(total / cnt, 2)
        else:
            val = None
        db.session.execute(
            text(f'UPDATE production_trend SET "{year_col}"=:v WHERE rowid=:i'),
            {"v": val, "i": rid})
    db.session.commit()

# دسته‌های پشتیبانی‌شده و الگوی نام ستون
CATEGORIES = {
    "production": "تولید",
    "average_flow": "دبی متوسط",
    "runtime": "کارکرد",
    "well_pressure": "فشار",
}

# نوع داده‌ی هر دسته در SQLite
CATEGORY_SQLTYPE = {
    "production": "INTEGER",
    "average_flow": "REAL",
    "runtime": "INTEGER",
    "well_pressure": "REAL",
}

SAFE_NAME = re.compile(r"^[a-z][a-z0-9_]*$")


def make_column_name(category, year, month_idx):
    """نام ستون انگلیسی طبق قرارداد data_map. month_idx: 0..11 یا None برای سالانه."""
    if month_idx is None:
        # ستون سالانه
        if category == "production":
            return f"production_year_{year}"
        if category == "average_flow":
            return f"average_flow_year_{year}"
        if category == "runtime":
            return f"runtime_year_{year}"
        if category == "well_pressure":
            return f"pressure_year_{year}"
        return None
    m = MONTHS_EN[month_idx]
    if category == "production":
        return f"production_{year}_{m}"
    if category == "average_flow":
        return f"average_flow_{year}_{m}"
    if category == "runtime":
        return f"runtime_{m}_{year}"
    if category == "well_pressure":
        return f"well_pressure_{year}_{m}"
    return None


def _existing_columns():
    rows = db.session.execute(text("PRAGMA table_info(production_trend)")).fetchall()
    return {r[1] for r in rows}


def _find_key_and_value_columns(header):
    """ستون نام چاه و ستون داده را در هدر اکسل تشخیص می‌دهد."""
    name_col = None
    for i, h in enumerate(header):
        if h and any(k in str(h) for k in ("نام چاه", "نام", "چاه")):
            name_col = i
            break
    # ستون داده: اولین ستون عددیِ غیر از ستون نام (ساده: دومین ستون پرشده)
    value_col = None
    for i, h in enumerate(header):
        if i != name_col and h not in (None, ""):
            value_col = i
            break
    return name_col, value_col


def run(path, category, year, month_idx):
    """اکسل را می‌خواند و ستون جدید را می‌سازد/پر می‌کند. آمار برمی‌گرداند."""
    if category not in CATEGORIES:
        raise ValueError("دسته‌ی نامعتبر.")
    col = make_column_name(category, year, month_idx)
    if not col or not SAFE_NAME.match(col):
        raise ValueError(f"نام ستون نامعتبر ساخته شد: {col}")

    sqltype = CATEGORY_SQLTYPE[category]
    stats = {"column": col, "added_column": False, "rows": 0,
             "matched": 0, "unmatched": 0}

    # ۱) افزودن ستون اگر نبود
    if col not in _existing_columns():
        db.session.execute(text(f'ALTER TABLE production_trend ADD COLUMN "{col}" {sqltype}'))
        db.session.commit()
        stats["added_column"] = True

    # ۲) خواندن اکسل
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("فایل اکسل خالی است.")
    header = rows[0]
    name_col, value_col = _find_key_and_value_columns(header)
    if name_col is None or value_col is None:
        raise ValueError("ستون نام چاه یا ستون داده در اکسل پیدا نشد.")

    # ۳) نگاشت نام چاه production_trend → rowid
    pt = db.session.execute(text("SELECT rowid, well_name FROM production_trend")).fetchall()
    by_key = {}
    for rowid, wname in pt:
        by_key.setdefault(match_key(wname), rowid)

    # ۴) پر کردن مقادیر
    for row in rows[1:]:
        if name_col >= len(row):
            continue
        name = row[name_col]
        if not name or not str(name).strip():
            continue
        stats["rows"] += 1
        val = row[value_col] if value_col < len(row) else None
        rid = by_key.get(match_key(name))
        if rid is None:
            stats["unmatched"] += 1
            continue
        db.session.execute(
            text(f'UPDATE production_trend SET "{col}"=:v WHERE rowid=:i'),
            {"v": val, "i": rid},
        )
        stats["matched"] += 1
# محاسبه‌ی خودکار دبی متوسط اگر تولید و کارکرد همان ماه موجود باشند
    if category in ("production", "runtime") and month_idx is not None:
        prod_col = make_column_name("production", year, month_idx)
        run_col = make_column_name("runtime", year, month_idx)
        flow_col = make_column_name("average_flow", year, month_idx)
        cols_now = _existing_columns()
        # فقط اگر هر دو ستون تولید و کارکرد وجود دارند
        if prod_col in cols_now and run_col in cols_now:
            if flow_col not in cols_now:
                db.session.execute(text(
                    f'ALTER TABLE production_trend ADD COLUMN "{flow_col}" REAL'))
                db.session.commit()
            # برای هر چاه دبی را حساب کن: تولید×۱۰۰۰ ÷ (کارکرد×۳۶۰۰)
            rows_calc = db.session.execute(text(
                f'SELECT rowid, "{prod_col}", "{run_col}" FROM production_trend')).fetchall()
            flow_filled = 0
            for rid, prod_v, run_v in rows_calc:
                try:
                    p = float(prod_v) if prod_v not in (None, "") else None
                    h = float(run_v) if run_v not in (None, "") else None
                except (ValueError, TypeError):
                    p, h = None, None
                if p is None or h is None or h <= 0:
                    continue
                flow = round(p * 1000.0 / (h * 3600.0), 2)
                if flow < 0 or flow > 1000:  # اعتبارسنجی: مقدار غیرمنطقی رد شود
                    continue
                db.session.execute(
                    text(f'UPDATE production_trend SET "{flow_col}"=:v WHERE rowid=:i'),
                    {"v": flow, "i": rid})
                flow_filled += 1
            db.session.commit()
            stats["flow_calculated"] = flow_filled
            # به‌روزرسانی ستون مجموع/میانگین سالانه بعد از افزودن ماه جدید
    if month_idx is not None:
        _recalc_yearly(category, year)

    db.session.commit()
    return stats

    db.session.commit()
    return stats


##########################################################################################
# FILE : data_map.py
##########################################################################################

"""نگاشت ستون‌های انگلیسیِ جدول production_trend به کلیدهای فارسیِ اکسلِ داشبورد.

هدف: داشبورد قدیمی (script.js) داده را با کلیدهای فارسی مثل «تولید سال ۱۴۰۴» یا
«دبی متوسط ۱۴۰۴(فروردین)» می‌خواند. این ماژول هر ردیف جدول را به یک dict با همان
کلیدهای فارسی برمی‌گرداند تا کل منطق داشبورد بدون تغییر کار کند.

ستون‌های سری‌زمانی با الگو ساخته می‌شوند؛ ستون‌های ثابت نگاشت دستی دارند.
"""

# ماه‌های شمسی به ترتیب (۱..۱۲) و معادل انگلیسیِ به‌کاررفته در نام ستون‌ها
MONTHS_FA = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
             "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
MONTHS_EN = ["farvardin", "ordibehesht", "khordad", "tir", "mordad", "shahrivar",
             "mehr", "aban", "azar", "dey", "bahman", "esfand"]
YEARS = [1399, 1400, 1401, 1402, 1403, 1404, 1405]

# ---------- نگاشت ستون‌های ثابت: انگلیسی → فارسیِ اکسل ----------
STATIC_MAP = {
    "well_name": "نام چاه",
    "department": "اداره",
    "production_zone_name": "نام پهنه تولید",
    "low_runtime_reason": "علت کارکرد کم",
    "electropump_installation_date": "تاریخ نصب الکتروپمپ",
    "observed_fault_last_rehabilitation": "خرابی مشاهده شده آخرین بهسازی",
    "last_rehabilitation_date": "تاریخ آخرین بهسازی",
    "permit_expiry": "اعتبار پروانه",
    "well_permit": "پروانه چاه",
    "well_type": "نوع چاه",
    "well_location_status": "وضعیت تعیین محل چاه",
    "proposed_permitted_flow_lps": "دبی مجاز پیشنهادی",
    "test_end_date": "تاریخ پایان آزمایش",
    "contractor": "پیمانکار",
    "drilling_year": "سال حفر",
    "max_flow_rate": "حداکثر آبدهی",
    "drawdown_amount": "مقدار افت",
    "static_level": "سطح استاتیک",
    "dynamic_level": "سطح دینامیک",
    "main_zone": "پهنه اصلی",
    "sub_zone": "زیر پهنه",
    "well_status": "وضعیت چاه",
    "full_cycle_count": "تعداد دوره کامل",
    "full_average_months": "میانگین کامل (ماه)",
    "total_average_with_open_installation_months": "میانگین کل با نصب باز (ماه)",
    "contractors": "پیمانکار ها",
}

# نگاشت پسوندهای دوره‌ی نصب/کشیدن (۱..۵ و «باز»)
CYCLE_BASE = {
    "installation": "نصب",
    "motor": "موتور",
    "motor_new_repaired": "موتور نو/تعمیری",
    "pump": "پمپ",
    "manufacturer": "سازنده",
    "pump_new_repaired": "پمپ نو/تعمیری",
    "stage": "طبقه",
    "depth": "عمق",
    "pulling": "کشیدن",
    "interval": None,  # ویژه (interval_1_months → فاصله۱ (ماه))
}


def _build_static_full():
    """نگاشت کامل ستون‌های ثابت شامل ستون‌های دوره‌ای (۱..۵ و open)."""
    m = dict(STATIC_MAP)
    suffixes = [("1", "1"), ("2", "2"), ("3", "3"), ("4", "4"), ("5", "5"),
                ("open", " (باز)")]
    for en_sfx, fa_sfx in suffixes:
        m[f"installation_{en_sfx}"] = f"نصب{fa_sfx}"
        m[f"motor_{en_sfx}"] = f"موتور{fa_sfx}"
        m[f"motor_new_repaired_{en_sfx}"] = f"موتور نو/تعمیری{fa_sfx}"
        m[f"pump_{en_sfx}"] = f"پمپ{fa_sfx}"
        m[f"manufacturer_{en_sfx}"] = f"سازنده{fa_sfx}"
        m[f"pump_new_repaired_{en_sfx}"] = f"پمپ نو/تعمیری{fa_sfx}"
        m[f"stage_{en_sfx}"] = f"طبقه{fa_sfx}"
        m[f"depth_{en_sfx}"] = f"عمق{fa_sfx}"
        if en_sfx != "open":
            m[f"pulling_{en_sfx}"] = f"کشیدن{fa_sfx}"
            m[f"interval_{en_sfx}_months"] = f"فاصله{fa_sfx} (ماه)"
    return m


STATIC_FULL = _build_static_full()


def _ts_map_for_column(col):
    """اگر ستون سری‌زمانی بود، کلید فارسی معادل را برمی‌گرداند، وگرنه None."""
    for y in YEARS:
        ys = str(y)
        # --- تولید ---
        if col == f"production_year_{ys}":
            return f"تولید سال {y}"
        for i, (men, mfa) in enumerate(zip(MONTHS_EN, MONTHS_FA)):
            if col == f"production_{ys}_{men}":
                return f"تولید {y}({mfa})"
        # --- کارکرد (runtime) ---
        if col == f"runtime_year_{ys}":
            return f"کارکرد سال {y}"
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"runtime_{men}_{ys}":
                return f"کارکرد {mfa} {y}"
        # --- دبی متوسط (average_flow) ---
        if col == f"average_flow_year_{ys}":
            return f"دبی متوسط سال {y}"
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"average_flow_{ys}_{men}":
                return f"دبی متوسط {y}({mfa})"
        # --- فشار (well_pressure) ---
        if col == f"pressure_year_{ys}":
            return f"فشار سال {y}"
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"well_pressure_{ys}_{men}":
                return f"فشار چاه {y}({mfa})"
        # --- تاریخ‌های آزمایش ---
        if col == f"test_date_{ys}":
            return f"تاریخ آزمایش {y}"
        if col == f"flow_test_date_{ys}":
            return f"تاریخ آزمایش دبی {y}"
        if col == f"flow_rate_lps_{ys}":
            return f"آبدهی (lit/s) {y}"
        if col == f"pressure_test_date_{ys}":
            return f"تاریخ آزمایش فشار {y}"
        if col == f"pressure_atm_{ys}":
            return f"فشار (atm) {y}"
        # --- پیمانکار سال / دبی سالانه نسبت ---
        if col == f"contractor_{ys}":
            return f"پیمانکار {y}"
        if col == f"average_flow_year_{ys}_vs_1399":
            return f"نسبت دبی متوسط {y} به 1399"
        # --- ترخیص (removal) ماهانه ---
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"removal_{men}_{ys}":
                return f"ترخیص {mfa} {y}"
    # ستون ویژه‌ی سن
    if col.startswith("age_until_") and col.endswith("_months"):
        return "عمر تا 1405/03/01 (ماه)"
    return None


def build_column_map(columns):
    """dict: نام ستون انگلیسی → کلید فارسیِ اکسل، برای همه‌ی ستون‌های موجود."""
    result = {}
    for col in columns:
        if col in STATIC_FULL:
            result[col] = STATIC_FULL[col]
            continue
        fa = _ts_map_for_column(col)
        result[col] = fa if fa else col  # اگر نگاشتی نبود، همان نام انگلیسی
    return result


def rows_to_fa_dicts(columns, rows):
    """ردیف‌های جدول را به لیست dict با کلیدهای فارسی تبدیل می‌کند."""
    cmap = build_column_map(columns)
    fa_keys = [cmap[c] for c in columns]
    out = []
    for row in rows:
        d = {}
        for k, v in zip(fa_keys, row):
            d[k] = v
        out.append(d)
    return out



##########################################################################################
# FILE : routes.py
##########################################################################################

"""تحلیل روند تولید چاه‌ها.

داشبورد کاملِ کلاینت (همان script.js اصلی با همه‌ی تب‌ها و نمودارها) در iframe
نمایش داده می‌شود، اما داده به‌جای فایل اکسل از جدول production_trend خوانده و
با کلیدهای فارسیِ اکسل به‌صورت JSON در اختیار داشبورد گذاشته می‌شود.
"""
from flask import render_template, jsonify, request, redirect, url_for, flash
from flask_login import login_required
from flask_wtf import FlaskForm
from flask_wtf.file import FileField, FileRequired, FileAllowed
from wtforms import SelectField, SubmitField
from wtforms.validators import DataRequired
import os
import tempfile
from sqlalchemy import text

from app.extensions import db
from app.blueprints.production_trend import bp
from app.blueprints.production_trend.data_map import rows_to_fa_dicts


@bp.route("/")
@login_required
def index():
    return render_template("production_trend/index.html")


@bp.route("/api/data")
@login_required
def api_data():
    """کل جدول production_trend را با کلیدهای فارسیِ اکسل برمی‌گرداند."""
    result = db.session.execute(text("SELECT * FROM production_trend"))
    columns = list(result.keys())
    rows = result.fetchall()
    data = rows_to_fa_dicts(columns, [tuple(r) for r in rows])
    # فقط ردیف‌های دارای نام چاه (مطابق فیلتر اصلی داشبورد)
    data = [d for d in data if d.get("نام چاه") and str(d.get("نام چاه")).strip()]
    resp = jsonify(data)
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    resp.headers["Pragma"] = "no-cache"
    resp.headers["Expires"] = "0"
    return resp
CATEGORY_CHOICES = [
    ("production", "تولید"),
    ("runtime", "کارکرد"),
    ("well_pressure", "فشار"),
]

MONTH_CHOICES = [("-1", "سالانه (کل سال)")] + [
    (str(i), m) for i, m in enumerate(
        ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
         "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"])
]

YEAR_CHOICES = [(str(y), str(y)) for y in range(1399, 1411)]


class ColumnImportForm(FlaskForm):
    category = SelectField("دسته‌ی داده", choices=CATEGORY_CHOICES, validators=[DataRequired()])
    year = SelectField("سال", choices=YEAR_CHOICES, validators=[DataRequired()])
    month = SelectField("ماه", choices=MONTH_CHOICES, validators=[DataRequired()])
    excel = FileField("فایل اکسل (ستون اول: نام چاه، ستون دوم: مقدار)",
                      validators=[FileRequired(), FileAllowed(["xlsx", "xls"], "فقط فایل اکسل")])
    submit = SubmitField("افزودن سرستون و ورود داده")


@bp.route("/import-column", methods=["GET", "POST"])
@login_required
def import_column():
    from app.blueprints.production_trend.column_import import run
    form = ColumnImportForm()
    result = None
    if form.validate_on_submit():
        month_idx = int(form.month.data)
        month_arg = None if month_idx == -1 else month_idx
        # ذخیره‌ی موقت فایل آپلودی
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        form.excel.data.save(tmp.name)
        tmp.close()
        try:
            result = run(tmp.name, form.category.data, int(form.year.data), month_arg)
            flash(f"ستون «{result['column']}» ساخته/به‌روزرسانی شد. "
                  f"تطبیق‌خورده: {result['matched']} | بدون تطبیق: {result['unmatched']}", "success")
        except Exception as e:
            flash(f"خطا در ورود داده: {e}", "danger")
        finally:
            try:
                os.unlink(tmp.name)
            except OSError:
                pass  # اگر ویندوز اجازه‌ی حذف نداد، فایل موقت بعداً خودکار پاک می‌شود
        return redirect(url_for("production_trend.import_column"))
    return render_template("production_trend/import_column.html", form=form)




################################################################################
# FILE: pump_select_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : pump_select
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("pump_select", __name__)

from app.blueprints.pump_select import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort, jsonify
from flask_login import current_user, login_required
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.pump_select import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.pump_select import PumpSelection
from app.models.audit import RecordHistory
from app.utils.dates import to_english_digits
from app import workflow


def _num(s):
    s = to_english_digits((s or "").strip())
    try:
        return float(s) if s else None
    except ValueError:
        return None


class PumpSelectionForm(FlaskForm):
    action_needed = StringField("اقدام مورد نیاز", validators=[Optional()])
    jyear = StringField("سال", validators=[Optional()])
    jmonth = StringField("ماه", validators=[Optional()])
    status_done = StringField("وضعیت انجام", validators=[Optional()])
    prev_pump_type = StringField("تیپ پمپ قبلی", validators=[Optional()])
    prev_motor_type = StringField("تیپ موتور قبلی", validators=[Optional()])
    prev_discharge_lps = FloatField("دبی قبلی (l/s)", validators=[Optional()])
    selected_pump_type = StringField("تیپ پمپ پس از بررسی", validators=[Optional()])
    selected_motor_type = StringField("تیپ موتور پس از بررسی", validators=[Optional()])
    target_discharge_lps = FloatField("دبی پس از بررسی (l/s)", validators=[Optional()])
    discharge_increase_lps = FloatField("میزان افزایش دبی", validators=[Optional()])
    selected_head_m = FloatField("هد پمپ انتخابی (m)", validators=[Optional()])
    form_delivery_date = JalaliDateField("تاریخ تحویل فرم", validators=[Optional()])
    pull_date = JalaliDateField("تاریخ کشیدن", validators=[Optional()])
    videometry_date = JalaliDateField("تاریخ ویدئومتری", validators=[Optional()])
    install_date = JalaliDateField("تاریخ نصب", validators=[Optional()])
    verify_flowtest_date = JalaliDateField("تاریخ دبی‌سنجی (صحت‌سنجی)", validators=[Optional()])
    verify_discharge_lps = FloatField("آبدهی دبی‌سنجی", validators=[Optional()])
    verify_head_m = FloatField("هد دبی‌سنجی", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


FIELDS = [
    "action_needed", "jyear", "jmonth", "status_done", "prev_pump_type", "prev_motor_type",
    "prev_discharge_lps", "selected_pump_type", "selected_motor_type", "target_discharge_lps",
    "discharge_increase_lps", "selected_head_m", "form_delivery_date", "pull_date",
    "videometry_date", "install_date", "verify_flowtest_date", "verify_discharge_lps",
    "verify_head_m", "notes",
]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data)


# ---------- catalog-driven selection calculator ----------
@bp.route("/pump-select/catalog.json")
@permission_required("pump_select", "view")
def catalog_json():
    from app.services.catalog import build_library
    return jsonify(build_library())


@bp.route("/pump-select/calculator")
@login_required
def calculator():
    from app.services.catalog import design_params
    well = None
    q, head, src = 10.0, 180, None
    well_id = request.args.get("well", type=int)
    if well_id:
        well = db.session.get(Well, well_id)
        if well:
            q, head, src = design_params(well)
    return render_template("pump_select/calculator.html", well=well, design_q=q,
                           design_head=head, src=src)


@bp.route("/wells/<int:well_id>/pump-select/from-calc", methods=["POST"])
@permission_required("pump_select", "create")
def from_calc(well_id):
    well = db.get_or_404(Well, well_id)
    rec = PumpSelection(
        well_id=well.id, created_by_id=current_user.id,
        action_needed="انتخاب از کاتالوگ",
        selected_pump_type=(request.form.get("model") or "").strip() or None,
        target_discharge_lps=_num(request.form.get("q")),
        selected_head_m=_num(request.form.get("head")),
        notes=(request.form.get("note") or "").strip() or None,
    )
    db.session.add(rec)
    db.session.flush()
    workflow.log_change(rec, "create", detail="از محاسبه‌گر کاتالوگ")
    db.session.commit()
    flash("انتخاب پمپ از محاسبه‌گر ثبت شد (پیش‌نویس).", "success")
    return redirect(url_for("pump_select.detail", record_id=rec.id))


@bp.route("/wells/<int:well_id>/pump-select/new", methods=["GET", "POST"])
@permission_required("pump_select", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = PumpSelectionForm()
    if form.validate_on_submit():
        rec = PumpSelection(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("انتخاب پمپ ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("pump_select.detail", record_id=rec.id))
    return render_template("pump_select/form.html", form=form, well=well, title="ثبت انتخاب پمپ")


@bp.route("/pump-select/<int:record_id>")
@permission_required("pump_select", "view")
def detail(record_id):
    rec = db.get_or_404(PumpSelection, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="pump_selections", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("pump_select/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/pump-select/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("pump_select", "edit")
def edit(record_id):
    rec = db.get_or_404(PumpSelection, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("pump_select.detail", record_id=rec.id))
    form = PumpSelectionForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("انتخاب پمپ به‌روزرسانی شد.", "success")
        return redirect(url_for("pump_select.detail", record_id=rec.id))
    return render_template("pump_select/form.html", form=form, well=rec.well, title="ویرایش انتخاب پمپ")


@bp.route("/pump-select/<int:record_id>/delete", methods=["POST"])
@permission_required("pump_select", "delete")
def delete(record_id):
    rec = db.get_or_404(PumpSelection, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد انتخاب پمپ حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(PumpSelection, record_id)
    if not current_user.has_permission("pump_select", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("pump_select.detail", record_id=record_id))


@bp.route("/pump-select/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/pump-select/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/pump-select/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/pump-select/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: pump_test_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : pump_test
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("pump_test", __name__)

from app.blueprints.pump_test import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.pump_test import bp
from app.security import permission_required
from app.utils.forms import (
    JalaliDateField, PersianFloatField as FloatField, PersianIntegerField as IntegerField,
)
from app.utils.dates import to_english_digits
from app.models.well import Well
from app.models.pump_test import PumpTest, PumpTestStep
from app.models.audit import RecordHistory
from app.models.constants import PUMP_TEST_TYPES
from app import workflow

STEP_ROWS = 8
STEP_FIELDS = ["rpm", "discharge_lps", "observed_drawdown", "calc_drawdown",
               "grid_loss", "aquifer_loss", "efficiency"]


class PumpTestForm(FlaskForm):
    test_date = JalaliDateField("تاریخ آزمایش", validators=[Optional()])
    test_type = SelectField("نوع آزمایش", choices=[("", "—")] + PUMP_TEST_TYPES, validators=[Optional()])
    duration_h = FloatField("مدت شستشو/آزمایش (ساعت)", validators=[Optional()])
    contractor = StringField("پیمانکار", validators=[Optional()])
    consultant = StringField("مشاور", validators=[Optional()])
    employer = StringField("کارفرما", validators=[Optional()])
    contract_no = StringField("شماره قرارداد", validators=[Optional()])
    project_title = StringField("عنوان پروژه", validators=[Optional()])
    static_level = FloatField("سطح استاتیک", validators=[Optional()])
    max_dynamic_level = FloatField("حداکثر سطح دینامیک", validators=[Optional()])
    max_drawdown = FloatField("حداکثر افت چاه", validators=[Optional()])
    max_yield_lps = FloatField("حداکثر آبدهی (l/s)", validators=[Optional()])
    coeff_a = FloatField("ضریب a", validators=[Optional()])
    coeff_b = FloatField("ضریب b", validators=[Optional()])
    proposed_discharge_lps = FloatField("دبی مجاز پیشنهادی", validators=[Optional()])
    proposed_install_depth_m = FloatField("عمق نصب پیشنهادی", validators=[Optional()])
    resulting_drawdown_m = FloatField("میزان افت حاصله", validators=[Optional()])
    motor_type = StringField("نوع موتور", validators=[Optional()])
    motor_power_hp = FloatField("قدرت موتور (HP)", validators=[Optional()])
    gearbox_power_hp = FloatField("قدرت جعبه‌دنده (HP)", validators=[Optional()])
    gearbox_ratio = StringField("تبدیل جعبه‌دنده", validators=[Optional()])
    pump_type = StringField("نوع پمپ", validators=[Optional()])
    pump_stages = IntegerField("تعداد طبقه", validators=[Optional()])
    pump_diameter_in = FloatField("قطر پمپ (اینچ)", validators=[Optional()])
    max_rpm = FloatField("حداکثر دور موتور", validators=[Optional()])
    discharge_pipe_diameter_in = FloatField("قطر لوله آبده (اینچ)", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


HEADER_FIELDS = [
    "test_date", "test_type", "duration_h", "contractor", "consultant", "employer",
    "contract_no", "project_title", "static_level", "max_dynamic_level", "max_drawdown",
    "max_yield_lps", "coeff_a", "coeff_b", "proposed_discharge_lps",
    "proposed_install_depth_m", "resulting_drawdown_m", "motor_type", "motor_power_hp",
    "gearbox_power_hp", "gearbox_ratio", "pump_type", "pump_stages", "pump_diameter_in",
    "max_rpm", "discharge_pipe_diameter_in", "notes",
]


def _num(s):
    s = to_english_digits((s or "").strip())
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _apply_header(form, rec):
    for f in HEADER_FIELDS:
        setattr(rec, f, getattr(form, f).data)


def _apply_steps(rec):
    rec.steps.clear()
    for i in range(1, STEP_ROWS + 1):
        vals = {f: request.form.get(f"st-{i}-{f}", "").strip() for f in STEP_FIELDS}
        if not any(vals.values()):
            continue
        rec.steps.append(PumpTestStep(
            step_no=i,
            rpm=_num(vals["rpm"]),
            discharge_lps=_num(vals["discharge_lps"]),
            observed_drawdown=_num(vals["observed_drawdown"]),
            calc_drawdown=_num(vals["calc_drawdown"]),
            grid_loss=_num(vals["grid_loss"]),
            aquifer_loss=_num(vals["aquifer_loss"]),
            efficiency=_num(vals["efficiency"]),
        ))


@bp.route("/wells/<int:well_id>/pump-test/new", methods=["GET", "POST"])
@permission_required("pump_test", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = PumpTestForm()
    if form.validate_on_submit():
        rec = PumpTest(well_id=well.id, created_by_id=current_user.id)
        _apply_header(form, rec)
        db.session.add(rec)
        _apply_steps(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("آزمایش پمپاژ ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("pump_test.detail", record_id=rec.id))
    return render_template("pump_test/form.html", form=form, well=well,
                           title="ثبت آزمایش پمپاژ", steps=[], step_rows=STEP_ROWS)


@bp.route("/pump-test/<int:record_id>")
@permission_required("pump_test", "view")
def detail(record_id):
    rec = db.get_or_404(PumpTest, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="pump_tests", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("pump_test/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           type_labels=dict(PUMP_TEST_TYPES))


@bp.route("/pump-test/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("pump_test", "edit")
def edit(record_id):
    rec = db.get_or_404(PumpTest, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("pump_test.detail", record_id=rec.id))
    form = PumpTestForm(obj=rec)
    if form.validate_on_submit():
        _apply_header(form, rec)
        _apply_steps(rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("آزمایش پمپاژ به‌روزرسانی شد.", "success")
        return redirect(url_for("pump_test.detail", record_id=rec.id))
    return render_template("pump_test/form.html", form=form, well=rec.well,
                           title="ویرایش آزمایش پمپاژ", steps=rec.steps, step_rows=STEP_ROWS)


@bp.route("/pump-test/<int:record_id>/delete", methods=["POST"])
@permission_required("pump_test", "delete")
def delete(record_id):
    rec = db.get_or_404(PumpTest, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("آزمایش پمپاژ حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(PumpTest, record_id)
    if not current_user.has_permission("pump_test", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("pump_test.detail", record_id=record_id))


@bp.route("/pump-test/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/pump-test/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/pump-test/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/pump-test/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: rehab_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : rehab
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("rehab", __name__)

from app.blueprints.rehab import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.rehab import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.rehab import Rehabilitation
from app.models.audit import RecordHistory
from app.blueprints.install.routes import get_or_create_supplier
from app import workflow


class RehabForm(FlaskForm):
    stage = StringField("مرحله", validators=[Optional()])
    jyear = StringField("سال", validators=[Optional()])
    rehab_date = JalaliDateField("تاریخ بهسازی", validators=[Optional()])
    pumping_end_date = JalaliDateField("تاریخ اتمام پمپاژ", validators=[Optional()])
    rehab_contractor_name = StringField("پیمانکار بهسازی", validators=[Optional()])
    pumping_contractor_name = StringField("پیمانکار پمپاژ", validators=[Optional()])
    pump_type_before = StringField("تیپ پمپ پیش از بهسازی", validators=[Optional()])
    discharge_before_lps = FloatField("دبی قبل (l/s)", validators=[Optional()])
    pump_type_after = StringField("تیپ پمپ پس از بهسازی", validators=[Optional()])
    discharge_after_lps = FloatField("دبی بعد (l/s)", validators=[Optional()])
    reason = StringField("علت بهسازی", validators=[Optional()])
    method = StringField("روش", validators=[Optional()])
    observed_fault = TextAreaField("خرابی مشاهده‌شده", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


SIMPLE = ["stage", "jyear", "rehab_date", "pumping_end_date", "pump_type_before",
          "discharge_before_lps", "pump_type_after", "discharge_after_lps",
          "reason", "method", "observed_fault", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data)
    rc = get_or_create_supplier(form.rehab_contractor_name.data, "contractor")
    pc = get_or_create_supplier(form.pumping_contractor_name.data, "contractor")
    rec.rehab_contractor_id = rc.id if rc else None
    rec.pumping_contractor_id = pc.id if pc else None
    if rec.discharge_before_lps is not None and rec.discharge_after_lps is not None:
        rec.discharge_change_lps = round(rec.discharge_after_lps - rec.discharge_before_lps, 2)


@bp.route("/wells/<int:well_id>/rehab/new", methods=["GET", "POST"])
@permission_required("rehab", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = RehabForm()
    if form.validate_on_submit():
        rec = Rehabilitation(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("بهسازی ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("rehab.detail", record_id=rec.id))
    return render_template("rehab/form.html", form=form, well=well, title="ثبت بهسازی")


@bp.route("/rehab/<int:record_id>")
@permission_required("rehab", "view")
def detail(record_id):
    rec = db.get_or_404(Rehabilitation, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="rehabilitations", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("rehab/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/rehab/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("rehab", "edit")
def edit(record_id):
    rec = db.get_or_404(Rehabilitation, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("rehab.detail", record_id=rec.id))
    form = RehabForm(obj=rec)
    if request.method == "GET":
        form.rehab_contractor_name.data = rec.rehab_contractor.name if rec.rehab_contractor else ""
        form.pumping_contractor_name.data = rec.pumping_contractor.name if rec.pumping_contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("بهسازی به‌روزرسانی شد.", "success")
        return redirect(url_for("rehab.detail", record_id=rec.id))
    return render_template("rehab/form.html", form=form, well=rec.well, title="ویرایش بهسازی")


@bp.route("/rehab/<int:record_id>/delete", methods=["POST"])
@permission_required("rehab", "delete")
def delete(record_id):
    rec = db.get_or_404(Rehabilitation, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد بهسازی حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(Rehabilitation, record_id)
    if not current_user.has_permission("rehab", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("rehab.detail", record_id=record_id))


@bp.route("/rehab/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/rehab/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/rehab/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/rehab/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: relocation_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : relocation
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("relocation", __name__)

from app.blueprints.relocation import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.relocation import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.relocation import RelocationRecord
from app.models.audit import RecordHistory
from app.models.constants import RELOCATION_CANDIDACY, RELOCATION_TYPES
from app import workflow


class RelocationForm(FlaskForm):
    decision_date = JalaliDateField("تاریخ تصمیم", validators=[Optional()])
    candidacy = SelectField("وضعیت نامزدی", choices=[("", "—")] + RELOCATION_CANDIDACY,
                            validators=[Optional()])
    reloc_type = SelectField("نوع جابه‌جایی", choices=[("", "—")] + RELOCATION_TYPES,
                             validators=[Optional()])
    reason = StringField("علت جابه‌جایی", validators=[Optional()])
    location_note = StringField("موقعیت پیشنهادی", validators=[Optional()])
    letter_no = StringField("شماره نامه", validators=[Optional()])
    distance_m = FloatField("فاصله از چاه قبلی (m)", validators=[Optional()])
    new_well_id = SelectField("چاه جانشین", coerce=int, validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_wells(self, exclude_id):
        wells = db.session.scalars(db.select(Well).order_by(Well.name)).all()
        self.new_well_id.choices = [(0, "—")] + [
            (w.id, f"{w.name} ({w.display_pm})") for w in wells if w.id != exclude_id]


SIMPLE = ["decision_date", "candidacy", "reloc_type", "reason",
          "location_note", "letter_no", "distance_m", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data or None)
    rec.new_well_id = form.new_well_id.data or None


@bp.route("/relocation")
@permission_required("relocation", "view")
def list_relocations():
    cand_labels = dict(RELOCATION_CANDIDACY)
    type_labels = dict(RELOCATION_TYPES)
    rows = db.session.scalars(
        db.select(RelocationRecord).order_by(RelocationRecord.candidacy)).all()
    return render_template("relocation/list.html", rows=rows,
                           cand_labels=cand_labels, type_labels=type_labels)


@bp.route("/wells/<int:well_id>/relocation/new", methods=["GET", "POST"])
@permission_required("relocation", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = RelocationForm()
    form.populate_wells(well.id)
    if form.validate_on_submit():
        rec = RelocationRecord(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("جابه‌جایی ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("relocation.detail", record_id=rec.id))
    return render_template("relocation/form.html", form=form, well=well, title="ثبت جابه‌جایی")


@bp.route("/relocation/<int:record_id>")
@permission_required("relocation", "view")
def detail(record_id):
    rec = db.get_or_404(RelocationRecord, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="relocations", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())).all()
    return render_template("relocation/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           cand_labels=dict(RELOCATION_CANDIDACY),
                           type_labels=dict(RELOCATION_TYPES))


@bp.route("/relocation/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("relocation", "edit")
def edit(record_id):
    rec = db.get_or_404(RelocationRecord, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("relocation.detail", record_id=rec.id))
    form = RelocationForm(obj=rec)
    form.populate_wells(rec.well_id)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("جابه‌جایی به‌روزرسانی شد.", "success")
        return redirect(url_for("relocation.detail", record_id=rec.id))
    return render_template("relocation/form.html", form=form, well=rec.well, title="ویرایش جابه‌جایی")


@bp.route("/relocation/<int:record_id>/delete", methods=["POST"])
@permission_required("relocation", "delete")
def delete(record_id):
    rec = db.get_or_404(RelocationRecord, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد جابه‌جایی حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _link_lineage(rec):
    """On approval: wire successor.parent_well = subject well, mark old relocated."""
    if rec.new_well_id and rec.new_well:
        rec.new_well.parent_well_id = rec.well_id
        rec.well.status = "relocated"


def _transition(record_id, fn, perm_action, msg, link=False, **kw):
    rec = db.get_or_404(RelocationRecord, record_id)
    if not current_user.has_permission("relocation", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        if link:
            _link_lineage(rec)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("relocation.detail", record_id=record_id))


@bp.route("/relocation/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/relocation/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve",
                       "تأیید شد و شجره‌نامه‌ی چاه به‌روزرسانی شد.", link=True)


@bp.route("/relocation/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/relocation/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: reports_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : reports
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("reports", __name__, url_prefix="/reports")

from app.blueprints.reports import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from collections import defaultdict

from flask import render_template, request, redirect, url_for, flash
from flask_login import current_user

from app.extensions import db
from app.blueprints.reports import bp
from app.security import permission_required
from app.models.pump_asset import PumpInstallation, Supplier
from app.models.prioritization import PrioritizationCriterion, DEFAULT_CRITERIA
from app.utils.dates import to_english_digits


def _aggregate(installs, key_attr):
    agg = defaultdict(lambda: {"completed": 0, "life_sum": 0.0, "running": 0, "burnt": 0})
    for i in installs:
        sid = getattr(i, key_attr)
        if not sid:
            continue
        a = agg[sid]
        if i.is_running:
            a["running"] += 1
        life = i.useful_life_months
        if life is not None:
            a["completed"] += 1
            a["life_sum"] += life
        if i.removal_reason and "سوخت" in i.removal_reason:
            a["burnt"] += 1
    return agg


def _rows(agg, names, min_completed):
    rows = []
    for sid, a in agg.items():
        if a["completed"] < min_completed:
            continue
        avg = round(a["life_sum"] / a["completed"], 1) if a["completed"] else None
        rows.append({
            "name": names.get(sid, "—"),
            "completed": a["completed"], "running": a["running"],
            "burnt": a["burnt"], "avg_life": avg,
        })
    rows.sort(key=lambda r: (r["avg_life"] is not None, r["avg_life"]), reverse=True)
    return rows


@bp.route("/suppliers")
@permission_required("reports", "view")
def suppliers():
    min_completed = int(request.args.get("min", 5))
    installs = db.session.scalars(db.select(PumpInstallation)).all()
    names = {s.id: s.name for s in db.session.scalars(db.select(Supplier)).all()}
    makers = _rows(_aggregate(installs, "manufacturer_id"), names, min_completed)
    contractors = _rows(_aggregate(installs, "contractor_id"), names, min_completed)
    return render_template("reports/suppliers.html",
                           makers=makers, contractors=contractors,
                           total=len(installs), min_completed=min_completed)


def _ensure_criteria():
    """Seed defaults and add any criteria missing by key (handles upgrades)."""
    existing = {c.key for c in db.session.scalars(db.select(PrioritizationCriterion)).all()}
    added = False
    for key, label, weight, direction in DEFAULT_CRITERIA:
        if key not in existing:
            db.session.add(PrioritizationCriterion(
                key=key, label=label, weight=weight, direction=direction))
            added = True
    if added:
        db.session.commit()


@bp.route("/prioritization", methods=["GET", "POST"])
@permission_required("reports", "view")
def prioritization():
    from app.services import prioritization as svc
    _ensure_criteria()

    if request.method == "POST":
        if not current_user.has_permission("admin", "edit"):
            flash("برای تغییر وزن‌ها مجوز مدیریت لازم است.", "warning")
            return redirect(url_for("reports.prioritization"))
        for c in db.session.scalars(db.select(PrioritizationCriterion)).all():
            w = to_english_digits(request.form.get(f"w_{c.id}", "")).strip()
            try:
                c.weight = float(w)
            except ValueError:
                pass
            c.is_active = request.form.get(f"a_{c.id}") == "on"
        db.session.commit()
        flash("وزن‌ها به‌روزرسانی شد.", "success")
        return redirect(url_for("reports.prioritization"))

    rows, criteria = svc.compute()
    all_criteria = db.session.scalars(
        db.select(PrioritizationCriterion).order_by(PrioritizationCriterion.weight.desc())
    ).all()
    return render_template("reports/prioritization.html",
                           rows=rows[:100], total=len(rows), criteria=all_criteria,
                           can_edit=current_user.has_permission("admin", "edit"))


@bp.route("/energy")
@permission_required("reports", "view")
def energy():
    from app.services import energy_stats
    return render_template("reports/energy.html", e=energy_stats.compute())


@bp.route("/alerts")
@permission_required("reports", "view")
def alerts():
    from app.services import alerts as alert_svc
    items = alert_svc.compute()
    return render_template("reports/alerts.html",
                           alerts=items, summary=alert_svc.summary(items),
                           categories=alert_svc.CATEGORIES)


@bp.route("/analytics")
@permission_required("reports", "view")
def analytics():
    from app.services.analytics import summary
    return render_template("reports/analytics.html", f=summary.fleet())


@bp.route("/zones")
@permission_required("reports", "view")
def zones():
    from app.services import zone_stats
    return render_template("reports/zones.html", z=zone_stats.compute())


@bp.route("/portfolio")
@permission_required("reports", "view")
def portfolio():
    from app.services import optimize, economics
    economics.ensure_params()
    default_budget = economics.get("rehab_cost", 2e9) * 20
    budget = request.args.get("budget", type=float)
    budget_rial = (budget * 1e9) if budget else default_budget
    return render_template("reports/portfolio.html",
                           p=optimize.portfolio(budget_rial),
                           budget_b=round(budget_rial / 1e9, 1),
                           unit_cost_b=round(economics.get("rehab_cost", 2e9) / 1e9, 2))


@bp.route("/dispatch")
@permission_required("reports", "view")
def dispatch():
    from app.services import optimize, zone_stats
    z = zone_stats.compute()
    group = request.args.get("group", "zone")
    key = request.args.get("key")
    demand = request.args.get("demand", type=float)
    result = None
    if key:
        attr = "zone" if group == "zone" else "destination_reservoir"
        result = optimize.field_dispatch(attr, key, demand or 0)
    return render_template("reports/dispatch.html", z=z, group=group,
                           key=key, demand=demand, result=result)





################################################################################
# FILE: videometry_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : videometry
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("videometry", __name__)

from app.blueprints.videometry import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.videometry import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.utils.dates import to_english_digits
from app.models.well import Well
from app.models.videometry import Videometry, VideometryFinding
from app.models.audit import RecordHistory
from app.models.constants import VIDEO_FINDING_TYPES, SEVERITY_LEVELS
from app.blueprints.install.routes import get_or_create_supplier
from app import workflow

FIND_ROWS = 8


class VideometryForm(FlaskForm):
    log_date = JalaliDateField("تاریخ چاه‌نگاری", validators=[Optional()])
    contractor_name = StringField("پیمانکار", validators=[Optional()])
    equipment = StringField("تجهیز/دوربین", validators=[Optional()])
    depth_from_m = FloatField("از عمق (m)", validators=[Optional()])
    depth_to_m = FloatField("تا عمق (m)", validators=[Optional()])
    final_depth_m = FloatField("عمق نهایی چاه (m)", validators=[Optional()])
    water_level_m = FloatField("سطح آب (m)", validators=[Optional()])
    video_file_ref = StringField("مسیر/نام فایل ویدئو", validators=[Optional()])
    summary = TextAreaField("خلاصه‌ی یافته‌ها", validators=[Optional()])
    submit = SubmitField("ذخیره")


HEADER = ["log_date", "equipment", "depth_from_m", "depth_to_m", "final_depth_m",
          "water_level_m", "video_file_ref", "summary"]


def _num(s):
    s = to_english_digits((s or "").strip())
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _apply(form, rec):
    for f in HEADER:
        setattr(rec, f, getattr(form, f).data)
    con = get_or_create_supplier(form.contractor_name.data, "contractor")
    rec.contractor_id = con.id if con else None
    rec.findings.clear()
    for i in range(1, FIND_ROWS + 1):
        depth = request.form.get(f"f-{i}-depth_m", "").strip()
        ftype = request.form.get(f"f-{i}-finding_type", "").strip()
        sev = request.form.get(f"f-{i}-severity", "").strip()
        note = request.form.get(f"f-{i}-note", "").strip()
        if not (depth or ftype or note):
            continue
        rec.findings.append(VideometryFinding(
            depth_m=_num(depth), finding_type=ftype or None,
            severity=sev or None, note=note or None))


@bp.route("/wells/<int:well_id>/videometry/new", methods=["GET", "POST"])
@permission_required("videometry", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = VideometryForm()
    if form.validate_on_submit():
        rec = Videometry(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("چاه‌نگاری ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("videometry.detail", record_id=rec.id))
    return render_template("videometry/form.html", form=form, well=well,
                           title="ثبت چاه‌نگاری", findings=[], find_rows=FIND_ROWS,
                           finding_types=VIDEO_FINDING_TYPES, severities=SEVERITY_LEVELS)


@bp.route("/videometry/<int:record_id>")
@permission_required("videometry", "view")
def detail(record_id):
    rec = db.get_or_404(Videometry, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="videometry_logs", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("videometry/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           type_labels=dict(VIDEO_FINDING_TYPES),
                           sev_labels=dict(SEVERITY_LEVELS))


@bp.route("/videometry/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("videometry", "edit")
def edit(record_id):
    rec = db.get_or_404(Videometry, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("videometry.detail", record_id=rec.id))
    form = VideometryForm(obj=rec)
    if request.method == "GET":
        form.contractor_name.data = rec.contractor.name if rec.contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("چاه‌نگاری به‌روزرسانی شد.", "success")
        return redirect(url_for("videometry.detail", record_id=rec.id))
    return render_template("videometry/form.html", form=form, well=rec.well,
                           title="ویرایش چاه‌نگاری", findings=rec.findings, find_rows=FIND_ROWS,
                           finding_types=VIDEO_FINDING_TYPES, severities=SEVERITY_LEVELS)


@bp.route("/videometry/<int:record_id>/delete", methods=["POST"])
@permission_required("videometry", "delete")
def delete(record_id):
    rec = db.get_or_404(Videometry, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد چاه‌نگاری حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(Videometry, record_id)
    if not current_user.has_permission("videometry", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("videometry.detail", record_id=record_id))


@bp.route("/videometry/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/videometry/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/videometry/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/videometry/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")





################################################################################
# FILE: water_level_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : water_level
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("water_level", __name__)

from app.blueprints.water_level import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import BooleanField, StringField, SubmitField
from wtforms.validators import Optional, DataRequired

from app.extensions import db
from app.blueprints.water_level import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.water_level import WaterLevelLog


class WaterLevelForm(FlaskForm):
    measure_date = JalaliDateField("تاریخ اندازه‌گیری", validators=[DataRequired()])
    static_level = FloatField("تراز آب (عمق تا آب، m)", validators=[DataRequired()])
    is_pumping = BooleanField("در حال پمپاژ")
    notes = StringField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


@bp.route("/wells/<int:well_id>/water-level/new", methods=["GET", "POST"])
@permission_required("water_level", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = WaterLevelForm()
    if form.validate_on_submit():
        rec = WaterLevelLog(well_id=well.id, method="manual",
                            created_by_id=current_user.id)
        form.populate_obj(rec)
        db.session.add(rec)
        db.session.commit()
        flash("تراز آب ثبت شد.", "success")
        return redirect(url_for("wells.detail", well_id=well.id) + "#sec-waterlevel")
    return render_template("water_level/form.html", form=form, well=well, title="ثبت تراز آب")


@bp.route("/water-level/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("water_level", "edit")
def edit(record_id):
    rec = db.get_or_404(WaterLevelLog, record_id)
    form = WaterLevelForm(obj=rec)
    if form.validate_on_submit():
        form.populate_obj(rec)
        rec.updated_by_id = current_user.id
        db.session.commit()
        flash("تراز آب به‌روزرسانی شد.", "success")
        return redirect(url_for("wells.detail", well_id=rec.well_id) + "#sec-waterlevel")
    return render_template("water_level/form.html", form=form, well=rec.well, title="ویرایش تراز آب")


@bp.route("/water-level/<int:record_id>/delete", methods=["POST"])
@permission_required("water_level", "delete")
def delete(record_id):
    rec = db.get_or_404(WaterLevelLog, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد تراز آب حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id) + "#sec-waterlevel")





################################################################################
# FILE: wells_merged.py
################################################################################

##########################################################################################
# BLUEPRINT : wells
# AUTO GENERATED
##########################################################################################


##########################################################################################
# FILE : __init__.py
##########################################################################################

from flask import Blueprint

bp = Blueprint("wells", __name__, url_prefix="/wells")

from app.blueprints.wells import routes  # noqa: E402,F401



##########################################################################################
# FILE : routes.py
##########################################################################################

from flask import render_template, redirect, url_for, flash, request
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, SubmitField
from wtforms.validators import DataRequired, Optional

from app.extensions import db
from app.blueprints.wells import bp
from app.utils.forms import PersianFloatField as FloatField
from app.security import permission_required
from app.models.well import Well
from app.models.org import OrgUnit
from app.models.constants import (
    WELL_KINDS,
    WELL_STATUSES,
    LOCATION_STATUSES,
    ORG_UNIT_TYPES,
    WELL_CONSTRUCTION_TYPES,
)


class WellForm(FlaskForm):
    pm_code = StringField("کد PM", validators=[DataRequired()])
    name = StringField("نام چاه", validators=[DataRequired()])
    well_kind = SelectField("نوع چاه", choices=WELL_KINDS)
    office_id = SelectField("اداره", coerce=int, validators=[Optional()])
    center_id = SelectField("مرکز آبرسانی", coerce=int, validators=[Optional()])
    zone = StringField("پهنه", validators=[Optional()])
    sub_zone = StringField("زیرپهنه", validators=[Optional()])
    utm_x = FloatField("UTM X", validators=[Optional()])
    utm_y = FloatField("UTM Y", validators=[Optional()])
    utm_zone = StringField("Zone", validators=[Optional()])
    latitude = FloatField("عرض جغرافیایی", validators=[Optional()])
    longitude = FloatField("طول جغرافیایی", validators=[Optional()])
    ground_elevation = FloatField("ارتفاع زمین", validators=[Optional()])
    drill_year = StringField("سال حفر", validators=[Optional()])
    location_status = SelectField(
        "وضعیت تعیین محل", choices=[("", "—")] + LOCATION_STATUSES, validators=[Optional()]
    )
    status = SelectField("وضعیت چاه", choices=WELL_STATUSES)
    parent_well_id = SelectField("چاه قبلی (در جابه‌جایی)", coerce=int, validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_choices(self, exclude_well_id=None):
        offices = db.session.scalars(
            db.select(OrgUnit).filter_by(unit_type="office").order_by(OrgUnit.name)
        ).all()
        centers = db.session.scalars(
            db.select(OrgUnit)
            .where(OrgUnit.unit_type.in_(["center", "rural_region"]))
            .order_by(OrgUnit.name)
        ).all()
        self.office_id.choices = [(0, "—")] + [(o.id, o.name) for o in offices]
        self.center_id.choices = [(0, "—")] + [(c.id, c.name) for c in centers]
        wells = db.session.scalars(db.select(Well).order_by(Well.pm_code)).all()
        self.parent_well_id.choices = [(0, "—")] + [
            (w.id, f"{w.pm_code} — {w.name}")
            for w in wells
            if w.id != exclude_well_id
        ]


@bp.route("/")
@permission_required("wells", "view")
def list_wells():
    q = (request.args.get("q") or "").strip()
    query = db.select(Well).order_by(Well.pm_code)
    if q:
        like = f"%{q}%"
        query = db.select(Well).where(
            db.or_(Well.pm_code.ilike(like), Well.name.ilike(like))
        ).order_by(Well.pm_code)
    wells = db.session.scalars(query).all()
    return render_template(
        "wells/list.html",
        wells=wells,
        q=q,
        kind_labels=dict(WELL_KINDS),
        status_labels=dict(WELL_STATUSES),
    )


@bp.route("/<int:well_id>")
@permission_required("wells", "view")
def detail(well_id):
    well = db.get_or_404(Well, well_id)
    from app.services import timeline
    from app.models.baseline import WellBaseline
    from app.models.quality import WaterQuality
    from app.models.production import MonthlyProduction
    baseline = db.session.scalar(
        db.select(WellBaseline).filter_by(well_id=well.id, is_current=True))
    quality = db.session.scalars(
        db.select(WaterQuality).filter_by(well_id=well.id)
        .order_by(WaterQuality.sample_date.desc())).all()

    # Monthly operational time-series (روند تولید) for the trend chart.
    prod_rows = db.session.scalars(
        db.select(MonthlyProduction).filter_by(well_id=well.id)
        .order_by(MonthlyProduction.jyear, MonthlyProduction.jmonth)).all()
    production = {
        "labels": [f"{r.jyear}/{r.jmonth:02d}" for r in prod_rows],
        "discharge": [r.avg_discharge_lps for r in prod_rows],
        "volume": [r.production_m3 for r in prod_rows],
        "hours": [r.run_hours for r in prod_rows],
    }
    energy = next((r for r in prod_rows if r.energy_kwh is not None), None)

    events = timeline.build(well)
    # Position lifecycle events on the production/discharge chart (nearest month).
    chart_events = []
    if prod_rows:
        import jdatetime
        from app.utils.dates import format_jalali
        periods = [r.jyear * 12 + (r.jmonth - 1) for r in prod_rows]
        colors = {
            "drilling": "#c9780b", "pump_test": "#7c4dff", "pump_select": "#11a394",
            "install": "#2f6bff", "operation": "#18a558", "rehab": "#e0463e",
            "videometry": "#0e8fa8", "relocation": "#d6457f", "maintenance": "#6a7180",
        }
        for ev in events:
            d = ev.get("date")
            if not d:
                continue
            jd = jdatetime.date.fromgregorian(date=d)
            p = jd.year * 12 + (jd.month - 1)
            idx = min(range(len(periods)), key=lambda i: abs(periods[i] - p))
            chart_events.append({
                "index": idx,
                "type_label": ev["type_label"],
                "module": ev["module"],
                "color": colors.get(ev["module"], "#6a7180"),
                "icon": ev["icon"],
                "date": format_jalali(d),
                "title": ev.get("title") or "",
                "status": ev.get("status") or "",
                "url": ev["url"],
            })

    # Pump-test step curves (Q vs drawdown / efficiency) — well-performance diagnostic.
    from app.models.pump_test import PumpTest
    from app.models.flow import FlowTest
    from app.utils.dates import format_jalali as _fj
    pump_tests = []
    for t in db.session.scalars(
            db.select(PumpTest).filter_by(well_id=well.id).order_by(PumpTest.test_date)).all():
        pts = sorted(
            [{"q": s.discharge_lps, "dd": s.observed_drawdown,
              "eff": (s.efficiency * 100 if s.efficiency and s.efficiency <= 1 else s.efficiency)}
             for s in t.steps if s.discharge_lps is not None],
            key=lambda x: x["q"])
        if pts:
            pump_tests.append({"date": _fj(t.test_date), "points": pts})

    # Flow-metering operating points (Q vs dynamic level), across tests over time.
    flow_tests = []
    for ft in db.session.scalars(
            db.select(FlowTest).filter_by(well_id=well.id).order_by(FlowTest.test_date)).all():
        pts = [{"q": p.discharge_lps, "dyn": p.dynamic_level_m, "head": p.head_m,
                "eff": (p.efficiency * 100 if p.efficiency and p.efficiency <= 1 else p.efficiency)}
               for p in ft.points if p.discharge_lps is not None]
        if pts:
            flow_tests.append({"date": _fj(ft.test_date), "points": pts})

    # Pump characteristic curve (catalog) + design point overlay for the Q-H chart.
    from app.models.pump_select import PumpSelection
    from app.models.pump_catalog import PumpModel
    from app.models.pump_asset import PumpInstallation
    sels = db.session.scalars(
        db.select(PumpSelection).filter_by(well_id=well.id)
        .order_by(PumpSelection.form_delivery_date)).all()
    dp = next((s for s in reversed(sels)
               if s.target_discharge_lps and s.selected_head_m), None)
    design_point = ({"q": dp.target_discharge_lps, "h": dp.selected_head_m,
                     "pump": dp.selected_pump_type} if dp else None)
    model_name = (dp.selected_pump_type if dp else None) or next(
        (s.selected_pump_type for s in reversed(sels) if s.selected_pump_type), None)
    if not model_name:
        inst = db.session.scalar(
            db.select(PumpInstallation).filter_by(well_id=well.id)
            .order_by(PumpInstallation.install_date.desc()))
        model_name = inst.pump_type if inst else None
    pump_curve = None
    if model_name:
        pm = db.session.scalar(db.select(PumpModel).filter_by(model=model_name))
        if pm and pm.points:
            pts = [{"q": p.flow_lps, "h": p.head_m, "eff": p.efficiency_pct}
                   for p in pm.points if p.flow_lps is not None and p.head_m is not None]
            if pts:
                bep = next((p for p in pm.points if p.is_bep), None)
                bs = next((p for p in pm.points if p.is_beb_start), None)
                be = next((p for p in pm.points if p.is_beb_end), None)
                pump_curve = {
                    "model": pm.model,
                    "points": pts,
                    "bep": ({"q": bep.flow_lps, "h": bep.head_m, "eff": bep.efficiency_pct}
                            if bep and bep.flow_lps is not None else None),
                    "beb": ({"start": bs.flow_lps, "end": be.flow_lps}
                            if bs and be and bs.flow_lps is not None and be.flow_lps is not None else None),
                }

    has_qh = (any(p["head"] is not None for ft in flow_tests for p in ft["points"])
              or design_point is not None or pump_curve is not None)

    from app.blueprints.documents.routes import attachments_for
    attachments = attachments_for("well", well.id)

    from app.services.analytics import summary as analytics_summary
    analytics = analytics_summary.well(well.id)

    # 6.2 pump-replacement ROI (uses energy saving + economic params)
    pump_roi = None
    eo = analytics.get("energy")
    if eo and eo.get("annual_rial_saving") and eo.get("motor_kw"):
        from app.services import economics
        economics.ensure_params()
        invest = (economics.get("pump_price_per_kw") * eo["motor_kw"]
                  + economics.get("pump_install_cost"))
        saving = eo["annual_rial_saving"]
        npv = economics.npv(saving, economics.get("analysis_years"),
                            economics.get("discount_rate")) - invest
        pump_roi = {
            "investment": invest, "annual_saving": saving,
            "payback": economics.payback_years(invest, saving),
            "npv": round(npv), "worth_it": npv > 0,
            "years": int(economics.get("analysis_years")),
        }

    # Groundwater level monitoring (piezometry) — static-level time series.
    from app.models.water_level import WaterLevelLog
    wl_rows = db.session.scalars(
        db.select(WaterLevelLog).filter_by(well_id=well.id)
        .order_by(WaterLevelLog.measure_date)).all()
    water_levels = [
        {"id": w.id, "date": _fj(w.measure_date), "level": w.static_level,
         "method": w.method, "pumping": w.is_pumping}
        for w in wl_rows if w.measure_date and w.static_level is not None]

    # Vertical cross-section (5.6): depths + water levels for the schematic.
    from app.models.drilling import Drilling
    from app.models.pump_asset import PumpInstallation
    from app.models.flow import FlowTestPoint
    tech = well.technical
    dr = db.session.scalar(db.select(Drilling).filter_by(well_id=well.id)
                           .order_by(Drilling.end_date.desc()))
    inst = db.session.scalar(db.select(PumpInstallation).filter_by(well_id=well.id)
                             .order_by(PumpInstallation.install_date.desc()))
    drill_depth = ((dr.well_depth_actual or dr.well_depth_permit) if dr else None) \
        or (tech.drill_depth_m if tech else None) or (inst.well_depth_m if inst else None)
    install_depth = (inst.install_depth_m if inst else None) \
        or (tech.install_depth_m if tech else None)
    static = next((w.static_level for w in reversed(wl_rows)
                   if not w.is_pumping and w.static_level is not None), None)
    dyn = db.session.scalar(
        db.select(FlowTestPoint.dynamic_level_m).join(FlowTest)
        .where(FlowTest.well_id == well.id, FlowTestPoint.dynamic_level_m.isnot(None))
        .order_by(FlowTest.test_date.desc()))
    xsection = None
    if drill_depth or install_depth:
        xsection = {
            "ground_elev": well.ground_elevation,
            "drill_depth": drill_depth, "install_depth": install_depth,
            "static": static, "dynamic": dyn,
            "casing": tech.casing_material if tech else None,
            "construction": well.construction_type,
            "max_depth": max(d for d in [drill_depth, install_depth, static, dyn, 1] if d),
        }

    return render_template(
        "wells/detail.html",
        well=well,
        timeline=events,
        chart_events=chart_events,
        baseline=baseline,
        quality=quality,
        production=production,
        energy=energy,
        pump_tests=pump_tests,
        flow_tests=flow_tests,
        has_qh=has_qh,
        pump_curve=pump_curve,
        design_point=design_point,
        water_levels=water_levels,
        attachments=attachments,
        analytics=analytics,
        pump_roi=pump_roi,
        xsection=xsection,
        kind_labels=dict(WELL_KINDS),
        status_labels=dict(WELL_STATUSES),
        loc_labels=dict(LOCATION_STATUSES),
        constr_labels=dict(WELL_CONSTRUCTION_TYPES),
    )


def _apply_form(form, well):
    well.pm_code = form.pm_code.data.strip()
    well.name = form.name.data.strip()
    well.well_kind = form.well_kind.data
    well.office_id = form.office_id.data or None
    well.center_id = well.office_id   # اداره و مرکز یکی هستند
    well.zone = (form.zone.data or "").strip() or None
    well.sub_zone = (form.sub_zone.data or "").strip() or None
    well.utm_x = form.utm_x.data
    well.utm_y = form.utm_y.data
    well.utm_zone = (form.utm_zone.data or "").strip() or None
    well.latitude = form.latitude.data
    well.longitude = form.longitude.data
    well.ground_elevation = form.ground_elevation.data
    well.drill_year = (form.drill_year.data or "").strip() or None
    well.location_status = form.location_status.data or None
    well.status = form.status.data
    well.parent_well_id = form.parent_well_id.data or None
    well.notes = (form.notes.data or "").strip() or None


@bp.route("/new", methods=["GET", "POST"])
@permission_required("wells", "create")
def create_well():
    form = WellForm()
    form.populate_choices()
    if form.validate_on_submit():
        existing = db.session.scalar(
            db.select(Well).filter_by(pm_code=form.pm_code.data.strip())
        )
        if existing:
            flash("چاهی با این کد PM از قبل وجود دارد.", "danger")
        else:
            well = Well()
            _apply_form(form, well)
            well.created_by_id = current_user.id
            db.session.add(well)
            db.session.commit()
            flash("چاه ثبت شد.", "success")
            return redirect(url_for("wells.detail", well_id=well.id))
    return render_template("wells/form.html", form=form, title="ثبت چاه جدید")


@bp.route("/<int:well_id>/edit", methods=["GET", "POST"])
@permission_required("wells", "edit")
def edit_well(well_id):
    well = db.get_or_404(Well, well_id)
    form = WellForm(obj=well)
    form.populate_choices(exclude_well_id=well.id)
    if form.validate_on_submit():
        clash = db.session.scalar(
            db.select(Well).where(
                Well.pm_code == form.pm_code.data.strip(), Well.id != well.id
            )
        )
        if clash:
            flash("چاه دیگری با این کد PM وجود دارد.", "danger")
        else:
            _apply_form(form, well)
            well.updated_by_id = current_user.id
            db.session.commit()
            flash("چاه به‌روزرسانی شد.", "success")
            return redirect(url_for("wells.detail", well_id=well.id))
    return render_template("wells/form.html", form=form, title="ویرایش چاه")


@bp.route("/<int:well_id>/delete", methods=["POST"])
@permission_required("wells", "delete")
def delete_well(well_id):
    well = db.get_or_404(Well, well_id)
    db.session.delete(well)
    db.session.commit()
    flash("چاه حذف شد.", "info")
    return redirect(url_for("wells.list_wells"))





################################################################################
# FILE: auth\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("auth", __name__)

from app.blueprints.auth import routes  # noqa: E402,F401



################################################################################
# FILE: auth\routes.py
################################################################################

from datetime import datetime

from flask import render_template, redirect, url_for, flash, request, session
from flask_login import login_user, logout_user, login_required, current_user
from flask_wtf import FlaskForm
from wtforms import StringField, PasswordField, BooleanField, SubmitField
from wtforms.validators import DataRequired

from app.extensions import db
from app.blueprints.auth import bp
from app.models.rbac import User, LoginLog

class LoginForm(FlaskForm):
    username = StringField("نام کاربری", validators=[DataRequired()])
    password = PasswordField("گذرواژه", validators=[DataRequired()])
    remember = BooleanField("مرا به خاطر بسپار")
    submit = SubmitField("ورود")


@bp.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("main.dashboard"))

    form = LoginForm()
    if form.validate_on_submit():
        user = db.session.scalar(
            db.select(User).filter_by(username=form.username.data.strip())
        )
        if user is None or not user.check_password(form.password.data):
            flash("نام کاربری یا گذرواژه نادرست است.", "danger")
        elif not user.is_active:
            flash("این حساب غیرفعال است.", "warning")
        else:
            login_user(user, remember=form.remember.data)
            now = datetime.utcnow()
            ip = (request.headers.get("X-Forwarded-For", request.remote_addr) or "").split(",")[0].strip()
            agent = (request.user_agent.string or "")[:255]
            user.last_login_at = now
            user.last_login_ip = ip
            user.last_login_agent = agent
            log = LoginLog(user_id=user.id, login_at=now, ip_address=ip, user_agent=agent)
            db.session.add(log)
            db.session.commit()
            session["login_log_id"] = log.id
            next_page = request.args.get("next")
            return redirect(next_page or url_for("main.dashboard"))

    return render_template("auth/login.html", form=form)


@bp.route("/logout")
@login_required
def logout():
    log_id = session.pop("login_log_id", None)
    if log_id:
        log = db.session.get(LoginLog, log_id)
        if log and log.logout_at is None:
            log.logout_at = datetime.utcnow()
            db.session.commit()
    logout_user()
    flash("از سیستم خارج شدید.", "info")
    return redirect(url_for("auth.login"))



################################################################################
# FILE: documents\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("documents", __name__, url_prefix="/documents")

from app.blueprints.documents import routes  # noqa: E402,F401



################################################################################
# FILE: documents\routes.py
################################################################################

"""Document/attachment upload, download and delete.

Files are validated against the allowed-extension allowlist, stored under
instance/uploads/ with a random UUID name, and served back with their original
filename. Entity links are polymorphic (entity_type, entity_id).
"""
import os
import uuid

from flask import (current_app, request, redirect, url_for, flash, abort,
                   send_from_directory)
from flask_login import current_user
from werkzeug.utils import secure_filename

from app.extensions import db
from app.blueprints.documents import bp
from app.security import permission_required
from app.models.attachment import Attachment

# where each entity type's page lives, to redirect back after up/delete
_BACK = {"well": ("wells.detail", "well_id")}


def _redirect_back(att_or_type, entity_id=None):
    etype = att_or_type.entity_type if isinstance(att_or_type, Attachment) else att_or_type
    eid = att_or_type.entity_id if isinstance(att_or_type, Attachment) else entity_id
    ep, arg = _BACK.get(etype, ("wells.detail", "well_id"))
    return redirect(url_for(ep, **{arg: eid}) + "#sec-docs")


def _allowed(filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext in current_app.config["ALLOWED_UPLOAD_EXT"], ext


@bp.route("/upload", methods=["POST"])
@permission_required("documents", "create")
def upload():
    etype = request.form.get("entity_type", "")
    eid = request.form.get("entity_id", type=int)
    if not etype or not eid:
        abort(400)
    f = request.files.get("file")
    if not f or not f.filename:
        flash("فایلی انتخاب نشده است.", "warning")
        return _redirect_back(etype, eid)
    ok, ext = _allowed(f.filename)
    if not ok:
        flash("نوع فایل مجاز نیست.", "danger")
        return _redirect_back(etype, eid)

    upload_dir = current_app.config["UPLOAD_DIR"]
    os.makedirs(upload_dir, exist_ok=True)
    stored = f"{uuid.uuid4().hex}.{ext}"
    f.save(os.path.join(upload_dir, stored))
    size = os.path.getsize(os.path.join(upload_dir, stored))

    att = Attachment(
        entity_type=etype, entity_id=eid,
        title=(request.form.get("title") or "").strip() or None,
        kind=request.form.get("kind") or "other",
        original_name=secure_filename(f.filename) or f"file.{ext}",
        stored_name=stored, content_type=f.mimetype, size_bytes=size,
        created_by_id=current_user.id)
    db.session.add(att)
    db.session.commit()
    flash("سند بارگذاری شد.", "success")
    return _redirect_back(att)


@bp.route("/<int:att_id>/download")
@permission_required("documents", "view")
def download(att_id):
    att = db.get_or_404(Attachment, att_id)
    return send_from_directory(
        current_app.config["UPLOAD_DIR"], att.stored_name,
        as_attachment=True, download_name=att.original_name)


@bp.route("/<int:att_id>/delete", methods=["POST"])
@permission_required("documents", "delete")
def delete(att_id):
    att = db.get_or_404(Attachment, att_id)
    path = os.path.join(current_app.config["UPLOAD_DIR"], att.stored_name)
    if os.path.exists(path):
        try:
            os.remove(path)
        except OSError:
            pass
    back = _redirect_back(att)
    db.session.delete(att)
    db.session.commit()
    flash("سند حذف شد.", "info")
    return back


def attachments_for(entity_type, entity_id):
    return db.session.scalars(
        db.select(Attachment).filter_by(entity_type=entity_type, entity_id=entity_id)
        .order_by(Attachment.created_at.desc())).all()



################################################################################
# FILE: drilling\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("drilling", __name__)

from app.blueprints.drilling import routes  # noqa: E402,F401



################################################################################
# FILE: drilling\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, BooleanField, SubmitField
from wtforms.validators import DataRequired, Optional

from app.extensions import db
from app.blueprints.drilling import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.drilling import Drilling
from app.models.audit import RecordHistory
from app.models.constants import REQUEST_TYPES, DRILL_METHODS
from app import workflow


class DrillingForm(FlaskForm):
    request_type = SelectField("نوع پروانه", choices=[("", "—")] + REQUEST_TYPES, validators=[Optional()])
    drill_method = SelectField("روش حفاری", choices=[("", "—")] + DRILL_METHODS, validators=[Optional()])
    executor = StringField("نام مجری", validators=[Optional()])
    contractor = StringField("پیمانکار", validators=[Optional()])
    supervisor = StringField("ناظر", validators=[Optional()])
    credit_source = StringField("محل تأمین اعتبار", validators=[Optional()])
    contract_no = StringField("شماره قرارداد", validators=[Optional()])
    contract_date = JalaliDateField("تاریخ قرارداد", validators=[Optional()])
    start_date = JalaliDateField("تاریخ استقرار دستگاه", validators=[Optional()])
    end_date = JalaliDateField("تاریخ پایان/ترخیص", validators=[Optional()])
    well_depth_permit = FloatField("عمق چاه در پروانه", validators=[Optional()])
    well_depth_actual = FloatField("عمق حفاری واقعی", validators=[Optional()])
    casing_diameter_in = FloatField("قطر لوله جدار (اینچ)", validators=[Optional()])
    casing_total_len = FloatField("طول کلی لوله‌گذاری", validators=[Optional()])
    steel_blank_len = FloatField("لوله فولادی ساده", validators=[Optional()])
    steel_screen_len = FloatField("لوله فولادی مشبک", validators=[Optional()])
    upvc_blank_len = FloatField("لوله UPVC ساده", validators=[Optional()])
    upvc_screen_len = FloatField("لوله UPVC مشبک", validators=[Optional()])
    transition_len = FloatField("قطعه تبدیلی", validators=[Optional()])
    static_level = FloatField("سطح استاتیک", validators=[Optional()])
    max_yield_lps = FloatField("حداکثر آبدهی (l/s)", validators=[Optional()])
    proposed_discharge_lps = FloatField("دبی پیشنهادی (l/s)", validators=[Optional()])
    dynamic_at_proposed = FloatField("سطح دینامیک در دبی پیشنهادی", validators=[Optional()])
    drawdown = FloatField("مقدار افت", validators=[Optional()])
    coeff_a = FloatField("ضریب a", validators=[Optional()])
    coeff_b = FloatField("ضریب b", validators=[Optional()])
    videometry_done = BooleanField("ویدئومتری انجام شده")
    address = StringField("آدرس", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


FIELDS = [
    "request_type", "drill_method", "executor", "contractor", "supervisor",
    "credit_source", "contract_no", "contract_date", "start_date", "end_date",
    "well_depth_permit", "well_depth_actual", "casing_diameter_in", "casing_total_len",
    "steel_blank_len", "steel_screen_len", "upvc_blank_len", "upvc_screen_len",
    "transition_len", "static_level", "max_yield_lps", "proposed_discharge_lps",
    "dynamic_at_proposed", "drawdown", "coeff_a", "coeff_b", "videometry_done",
    "address", "notes",
]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data)


@bp.route("/wells/<int:well_id>/drilling/new", methods=["GET", "POST"])
@permission_required("drilling", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = DrillingForm()
    if form.validate_on_submit():
        rec = Drilling(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("رویداد حفاری ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("drilling.detail", record_id=rec.id))
    return render_template("drilling/form.html", form=form, well=well, title="ثبت حفاری")


@bp.route("/drilling/<int:record_id>")
@permission_required("drilling", "view")
def detail(record_id):
    rec = db.get_or_404(Drilling, record_id)
    history = db.session.scalars(
        db.select(RecordHistory)
        .filter_by(entity_type="drilling", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template(
        "drilling/detail.html",
        rec=rec, well=rec.well, history=history,
        req_labels=dict(REQUEST_TYPES), method_labels=dict(DRILL_METHODS),
        editable=workflow.is_editable(rec),
    )


@bp.route("/drilling/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("drilling", "edit")
def edit(record_id):
    rec = db.get_or_404(Drilling, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد تأیید/ثبت شده و قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("drilling.detail", record_id=rec.id))
    form = DrillingForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("رویداد حفاری به‌روزرسانی شد.", "success")
        return redirect(url_for("drilling.detail", record_id=rec.id))
    return render_template("drilling/form.html", form=form, well=rec.well, title="ویرایش حفاری")


@bp.route("/drilling/<int:record_id>/delete", methods=["POST"])
@permission_required("drilling", "delete")
def delete(record_id):
    rec = db.get_or_404(Drilling, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رویداد حفاری حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


# --- workflow transitions ---
def _transition(record_id, fn, perm_action, success_msg, **kwargs):
    rec = db.get_or_404(Drilling, record_id)
    if not current_user.has_permission("drilling", perm_action):
        abort(403)
    try:
        fn(rec, **kwargs)
        db.session.commit()
        flash(success_msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("drilling.detail", record_id=record_id))


@bp.route("/drilling/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/drilling/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "رکورد تأیید شد.")


@bp.route("/drilling/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    reason = request.form.get("reason", "")
    return _transition(record_id, workflow.reject, "approve", "رکورد برگشت داده شد.", reason=reason)


@bp.route("/drilling/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: exports\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("exports", __name__, url_prefix="/exports")

from app.blueprints.exports import routes  # noqa: E402,F401



################################################################################
# FILE: exports\routes.py
################################################################################

"""Data export / review section.

Lets an authorized user download the full dataset as Excel for offline review —
in particular to spot duplicate wells (variants of the same physical well such as
«ابوطالب 1 ق» / «ابوطالب 1»), which share a `گروه تکراری` key here.
"""
from datetime import datetime
from io import BytesIO

import pandas as pd
from flask import render_template, send_file

from app.extensions import db
from app.blueprints.exports import bp
from app.security import permission_required
from app.services import data_export

_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(sheets):
    """sheets: list of (sheet_name, DataFrame) -> BytesIO of an .xlsx workbook."""
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        for name, df in sheets:
            df.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 40)
    buf.seek(0)
    return buf


def _send(buf, prefix):
    stamp = datetime.now().strftime("%Y%m%d")
    return send_file(buf, as_attachment=True, mimetype=_MIME,
                     download_name=f"{prefix}_{stamp}.xlsx")


@bp.route("/")
@permission_required("reports", "view")
def index():
    from app.models.well import Well
    from app.models.production import MonthlyProduction as MP
    df = data_export.wells_df()
    dup_groups = int((df["تعداد در گروه"] > 1).sum())
    dup_group_count = df[df["تعداد در گروه"] > 1]["گروه تکراری"].nunique()
    stats = {
        "wells": int(db.session.scalar(db.select(db.func.count(Well.id))) or 0),
        "dup_rows": dup_groups,
        "dup_groups": int(dup_group_count),
        "production_rows": int(db.session.scalar(db.select(db.func.count(MP.id))) or 0),
    }
    return render_template("exports/index.html", stats=stats)


@bp.route("/wells.xlsx")
@permission_required("reports", "view")
def wells_xlsx():
    return _send(_xlsx([("چاه‌ها", data_export.wells_df())]), "wells")


@bp.route("/duplicates.xlsx")
@permission_required("reports", "view")
def duplicates_xlsx():
    return _send(_xlsx([("تکراری‌های مشکوک", data_export.duplicates_df())]), "duplicates")


@bp.route("/production.xlsx")
@permission_required("reports", "view")
def production_xlsx():
    return _send(_xlsx([("تولید ماهانه", data_export.production_df())]), "production")


@bp.route("/all.xlsx")
@permission_required("reports", "view")
def all_xlsx():
    sheets = [
        ("چاه‌ها", data_export.wells_df()),
        ("تکراری‌های مشکوک", data_export.duplicates_df()),
        ("تولید ماهانه", data_export.production_df()),
    ]
    return _send(_xlsx(sheets), "all_data")



################################################################################
# FILE: finance\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("finance", __name__, url_prefix="/finance")

from app.blueprints.finance import routes  # noqa: E402,F401


################################################################################
# FILE: finance\routes.py
################################################################################

"""ماژول صورت‌وضعیت مالی — فاز ۲: فهرست + فرم + جزئیات."""
from flask import render_template, redirect, url_for, flash
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.finance import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField
from app.models.well import Well
from app.models.finance import FinanceStatement

OP_TYPES = [("", "—"), ("بهسازی", "بهسازی"), ("آزمایش پمپاژ", "آزمایش پمپاژ"),
            ("حفاری", "حفاری"), ("سایر", "سایر")]
KINDS = [("", "—"), ("موقت", "موقت"), ("قطعی", "قطعی"),
         ("وضعیت ۱", "وضعیت ۱"), ("وضعیت ۲", "وضعیت ۲"), ("وضعیت ۳", "وضعیت ۳")]


class StatementForm(FlaskForm):
    op_type = SelectField("نوع عملیات", choices=OP_TYPES, validators=[Optional()])
    op_description = TextAreaField("شرح عملیات (موضوع)", validators=[Optional()])
    contractor = StringField("پیمانکار", validators=[Optional()])
    contract_no = StringField("شماره قرارداد", validators=[Optional()])
    statement_no = StringField("شماره صورت‌وضعیت", validators=[Optional()])
    statement_kind = SelectField("نوع وضعیت", choices=KINDS, validators=[Optional()])
    statement_date = JalaliDateField("تاریخ صورت‌وضعیت", validators=[Optional()])
    coef_region = PersianFloatField("ضریب منطقه", validators=[Optional()])
    coef_overhead = PersianFloatField("ضریب بالاسری", validators=[Optional()])
    coef_contract = PersianFloatField("ضریب پیمان", validators=[Optional()])
    workshop_setup = PersianFloatField("تجهیز کارگاه (ریال)", validators=[Optional()])
    total_before = PersianFloatField("مبلغ کل قبل از ضرایب (ریال)", validators=[Optional()])
    total_after = PersianFloatField("مبلغ کل پس از ضرایب (ریال)", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


FIELDS = ["op_type", "op_description", "contractor", "contract_no", "statement_no", "statement_kind",
          "statement_date", "coef_region", "coef_overhead", "coef_contract",
          "workshop_setup", "total_before", "total_after", "notes"]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data if getattr(form, f).data != "" else None)


@bp.route("/")
@permission_required("finance", "view")
def index():
    statements = db.session.scalars(
        db.select(FinanceStatement).order_by(FinanceStatement.statement_date.desc())
    ).all()
    return render_template("finance/index.html", statements=statements)


@bp.route("/well/<int:well_id>/new", methods=["GET", "POST"])
@permission_required("finance", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = StatementForm()
    if form.validate_on_submit():
        rec = FinanceStatement(well_id=well.id, created_by_id=current_user.id, source="manual")
        _apply(form, rec)
        db.session.add(rec)
        db.session.commit()
        flash("صورت‌وضعیت ثبت شد.", "success")
        return redirect(url_for("finance.detail", record_id=rec.id))
    return render_template("finance/form.html", form=form, well=well, title="ثبت صورت‌وضعیت")


@bp.route("/<int:record_id>")
@permission_required("finance", "view")
def detail(record_id):
    rec = db.get_or_404(FinanceStatement, record_id)
    # تجمیع سهم هر چاه از آیتم‌های این سند (بر اساس تخصیص‌ها + ضریب بالاسری)
    ovh = rec.coef_overhead or 1.0
    agg = {}
    for it in rec.items:
        amt = (it.total_amount or 0) * ovh
        allocs = list(it.allocations)
        tq = sum((a.quantity or 0) for a in allocs)
        if not tq:
            continue
        for a in allocs:
            key = a.well_id if a.well_id else ("name:" + (a.well_name or "?"))
            row = agg.setdefault(key, {
                "well_id": a.well_id,
                "name": (a.well.name if a.well else a.well_name) or "—",
                "qty": 0.0, "cost": 0.0})
            row["qty"] += a.quantity or 0
            row["cost"] += amt * (a.quantity or 0) / tq
    wells_breakdown = sorted(agg.values(), key=lambda r: r["cost"], reverse=True)
    return render_template("finance/detail.html", rec=rec, well=rec.well,
                           wells_breakdown=wells_breakdown)


@bp.route("/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("finance", "edit")
def edit(record_id):
    rec = db.get_or_404(FinanceStatement, record_id)
    form = StatementForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        db.session.commit()
        flash("صورت‌وضعیت به‌روزرسانی شد.", "success")
        return redirect(url_for("finance.detail", record_id=rec.id))
    return render_template("finance/form.html", form=form, well=rec.well, title="ویرایش صورت‌وضعیت")


@bp.route("/<int:record_id>/delete", methods=["POST"])
@permission_required("finance", "delete")
def delete(record_id):
    rec = db.get_or_404(FinanceStatement, record_id)
    db.session.delete(rec)
    db.session.commit()
    flash("صورت‌وضعیت حذف شد.", "info")
    return redirect(url_for("finance.index"))
# ---------- آیتم‌های هزینه (فاز ۳) ----------
from app.models.finance import FinanceItem


class ItemForm(FlaskForm):
    category = StringField("دسته (شرح عملیات کلی)", validators=[Optional()])
    row_no = StringField("شماره ردیف", validators=[Optional()])
    description = TextAreaField("شرح تفصیلی", validators=[Optional()])
    unit = StringField("واحد", validators=[Optional()])
    contract_qty = PersianFloatField("تعداد قرارداد", validators=[Optional()])
    quantity = PersianFloatField("مقدار", validators=[Optional()])
    unit_price = PersianFloatField("مبلغ واحد (ریال)", validators=[Optional()])
    coefficient = PersianFloatField("ضریب", validators=[Optional()])
    total_amount = PersianFloatField("مبلغ کل (ریال)", validators=[Optional()])
    submit = SubmitField("ذخیره")


ITEM_FIELDS = ["category", "row_no", "description", "unit", "contract_qty",
               "quantity", "unit_price", "coefficient", "total_amount"]


@bp.route("/<int:statement_id>/item/new", methods=["GET", "POST"])
@permission_required("finance", "edit")
def item_create(statement_id):
    st = db.get_or_404(FinanceStatement, statement_id)
    form = ItemForm()
    if form.validate_on_submit():
        it = FinanceItem(statement_id=st.id, created_by_id=current_user.id)
        for f in ITEM_FIELDS:
            setattr(it, f, getattr(form, f).data)
        db.session.add(it)
        db.session.commit()
        flash("آیتم هزینه افزوده شد.", "success")
        return redirect(url_for("finance.detail", record_id=st.id))
    return render_template("finance/item_form.html", form=form, st=st, title="افزودن آیتم هزینه")


@bp.route("/item/<int:item_id>/edit", methods=["GET", "POST"])
@permission_required("finance", "edit")
def item_edit(item_id):
    it = db.get_or_404(FinanceItem, item_id)
    form = ItemForm(obj=it)
    if form.validate_on_submit():
        for f in ITEM_FIELDS:
            setattr(it, f, getattr(form, f).data)
        db.session.commit()
        flash("آیتم به‌روزرسانی شد.", "success")
        return redirect(url_for("finance.detail", record_id=it.statement_id))
    return render_template("finance/item_form.html", form=form, st=it.statement,
                           title="ویرایش آیتم هزینه", item_id=it.id)


@bp.route("/item/<int:item_id>/delete", methods=["POST"])
@permission_required("finance", "edit")
def item_delete(item_id):
    it = db.get_or_404(FinanceItem, item_id)
    sid = it.statement_id
    db.session.delete(it)
    db.session.commit()
    flash("آیتم حذف شد.", "info")
    return redirect(url_for("finance.detail", record_id=sid))

# ---------- گزارش‌ها و داشبورد مالی (فاز ۵) ----------
from io import BytesIO
from datetime import datetime
from collections import defaultdict

import pandas as pd
from flask import request, send_file

from app.models.finance import FinanceItemAllocation

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _finance_report(op_type=None, contractor=None, statement_id=None, well_id=None):
    """داده‌ی گزارش مالی را با اعمال فیلترها می‌سازد.

    خروجی: dict شامل سه گروه‌بندی (چاه/سند/نوع عملیات)، KPIها و ردیف‌های جدول.
    هزینه‌ی هر آیتم بین چاه‌های تخصیص‌یافته به نسبت مقدار تقسیم می‌شود.
    """
    q = db.select(FinanceStatement)
    if op_type:
        q = q.where(FinanceStatement.op_type == op_type)
    if contractor:
        q = q.where(FinanceStatement.contractor == contractor)
    if statement_id:
        q = q.where(FinanceStatement.id == statement_id)
    statements = db.session.scalars(q).all()

    by_optype = defaultdict(float)
    by_statement = []                 # (label, total)
    well_cost = defaultdict(float)    # well_name -> cost
    well_meta = {}                    # well_name -> well_id
    well_docs = defaultdict(set)      # well_name -> {statement_id}
    unallocated = 0.0
    grand = 0.0

    for st in statements:
        stmt_total = 0.0
        # ضریب بالاسری در سطح سند اعمال می‌شود تا مبالغ به «مبلغ صورت‌وضعیت» برسند
        ovh = st.coef_overhead or 1.0
        for it in st.items:
            amt = (it.total_amount or 0.0) * ovh
            allocs = list(it.allocations)
            tq = sum((a.quantity or 0) for a in allocs)
            if tq > 0:
                for a in allocs:
                    if well_id and a.well_id != well_id:
                        continue
                    share = amt * (a.quantity or 0) / tq
                    name = a.well.name if a.well else (a.well_name or "—")
                    well_cost[name] += share
                    well_meta[name] = a.well_id
                    well_docs[name].add(st.id)
                    stmt_total += share
            elif st.well_id:            # سند تک‌چاهی (ثبت دستی)
                if well_id and st.well_id != well_id:
                    continue
                name = st.well.name if st.well else "—"
                well_cost[name] += amt
                well_meta[name] = st.well_id
                well_docs[name].add(st.id)
                stmt_total += amt
            else:                        # آیتم بدون تخصیص (مثل ویدئومتری)
                if not well_id:
                    unallocated += amt
                    stmt_total += amt
        # تجهیز کارگاه: هزینه‌ی ثابت سند، به هیچ چاهی تخصیص نمی‌یابد
        if not well_id and st.workshop_setup:
            unallocated += st.workshop_setup
            stmt_total += st.workshop_setup
        if stmt_total:
            by_optype[st.op_type or "نامشخص"] += stmt_total
            label = " ".join(x for x in [st.op_type, st.statement_kind,
                             ("— " + (st.contractor or st.contract_no or f"#{st.id}"))] if x)
            by_statement.append((label, stmt_total))
            grand += stmt_total

    rows = [{
        "well": n, "well_id": well_meta.get(n),
        "cost": round(c), "docs": len(well_docs[n]),
        "matched": well_meta.get(n) is not None,
    } for n, c in well_cost.items()]
    rows.sort(key=lambda r: r["cost"], reverse=True)
    by_statement.sort(key=lambda x: x[1], reverse=True)

    return {
        "rows": rows,
        "by_optype": dict(sorted(by_optype.items(), key=lambda x: x[1], reverse=True)),
        "by_statement": by_statement,
        "top_wells": {r["well"]: r["cost"] for r in rows[:15]},
        "grand": round(grand),
        "unallocated": round(unallocated),
        "well_count": len(rows),
        "doc_count": len(statements),
    }


def _filter_options():
    op_types = [x for x in db.session.scalars(
        db.select(FinanceStatement.op_type).distinct()).all() if x]
    contractors = [x for x in db.session.scalars(
        db.select(FinanceStatement.contractor).distinct()).all() if x]
    statements = db.session.scalars(
        db.select(FinanceStatement).order_by(FinanceStatement.op_type,
                                             FinanceStatement.statement_kind)).all()
    # چاه‌هایی که در مالی هزینه دارند
    wids = set(db.session.scalars(
        db.select(FinanceItemAllocation.well_id).distinct()).all())
    wids |= set(db.session.scalars(
        db.select(FinanceStatement.well_id).distinct()).all())
    wids.discard(None)
    wells = db.session.scalars(
        db.select(Well).where(Well.id.in_(wids)).order_by(Well.name)).all() if wids else []
    return op_types, contractors, statements, wells


def _current_filters():
    def _int(v):
        try:
            return int(v)
        except (TypeError, ValueError):
            return None
    return {
        "op_type": request.args.get("op_type") or None,
        "contractor": request.args.get("contractor") or None,
        "statement_id": _int(request.args.get("statement_id")),
        "well_id": _int(request.args.get("well_id")),
    }


@bp.route("/reports")
@permission_required("finance", "view")
def reports():
    f = _current_filters()
    data = _finance_report(**f)
    op_types, contractors, statements, wells = _filter_options()
    cur_qs = {k: v for k, v in f.items() if v}   # برای لینک خروجی اکسل
    return render_template("finance/reports.html", data=data, cur=f, cur_qs=cur_qs,
                           op_types=op_types, contractors=contractors,
                           statements=statements, wells=wells)


@bp.route("/reports/export")
@permission_required("finance", "view")
def reports_export():
    f = _current_filters()
    data = _finance_report(**f)
    df = pd.DataFrame([{
        "چاه": r["well"],
        "هزینه (ریال)": r["cost"],
        "تعداد سند": r["docs"],
        "تطبیق با سامانه": "بله" if r["matched"] else "خیر",
    } for r in data["rows"]])
    by_stmt = pd.DataFrame([{"سند": l, "مبلغ (ریال)": round(v)}
                            for l, v in data["by_statement"]])
    by_op = pd.DataFrame([{"نوع عملیات": k, "مبلغ (ریال)": round(v)}
                          for k, v in data["by_optype"].items()])

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        (df if not df.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="هزینه هر چاه", index=False)
        (by_stmt if not by_stmt.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="به تفکیک سند", index=False)
        (by_op if not by_op.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="به تفکیک عملیات", index=False)
        for name in xl.sheets:
            ws = xl.sheets[name]
            for col in ws.columns:
                w = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 12), 45)
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    return send_file(buf, as_attachment=True, mimetype=_XLSX_MIME,
                     download_name=f"finance_report_{stamp}.xlsx")


################################################################################
# FILE: gis\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("gis", __name__, url_prefix="/gis")

from app.blueprints.gis import routes  # noqa: E402,F401



################################################################################
# FILE: gis\routes.py
################################################################################

from flask import render_template, jsonify
from flask_login import login_required

from app.extensions import db
from app.blueprints.gis import bp
from app.security import permission_required
from app.models.well import Well
from app.models.constants import WELL_STATUSES


@bp.route("/map")
@login_required
def map():
    return render_template("gis/map.html")


@bp.route("/analysis")
@permission_required("reports", "view")
def analysis():
    return render_template("gis/analysis.html")


@bp.route("/analysis.json")
@permission_required("reports", "view")
def analysis_json():
    from app.services import spatial
    return jsonify({"points": spatial.metric_points()})


@bp.route("/piezometric.json")
@permission_required("reports", "view")
def piezometric_json():
    from app.services import spatial
    return jsonify(spatial.piezometric_grid() or {})


@bp.route("/interference.json")
@permission_required("reports", "view")
def interference_json():
    from flask import request
    from app.services import spatial
    thr = request.args.get("threshold", 400, type=int)
    return jsonify(spatial.interference(threshold_m=thr))


@bp.route("/suitability.json")
@permission_required("reports", "view")
def suitability_json():
    from app.services import spatial
    return jsonify(spatial.suitability_grid() or {})


@bp.route("/wells.geojson")
@permission_required("wells", "view")
def wells_geojson():
    """Wells that have coordinates, as GeoJSON for Leaflet (themed)."""
    status_labels = dict(WELL_STATUSES)
    from app.models.production import MonthlyProduction as MP
    spec_energy = {}
    for mp in db.session.scalars(db.select(MP).where(MP.energy_kwh.isnot(None))).all():
        se = mp.specific_energy
        if se is not None:
            spec_energy[mp.well_id] = se

    wells = db.session.scalars(
        db.select(Well).where(Well.latitude.isnot(None), Well.longitude.isnot(None))
    ).all()
    features = []
    for w in wells:
        features.append({
            "type": "Feature",
            "geometry": {"type": "Point", "coordinates": [w.longitude, w.latitude]},
            "properties": {
                "id": w.id, "pm_code": w.display_pm, "name": w.name,
                "status": status_labels.get(w.status, w.status),
                "status_key": w.status or "unknown",
                "zone": w.zone or "", "reservoir": w.destination_reservoir or "",
                "office": w.office.name if w.office else "",
                "elevation": w.ground_elevation,
                "specific_energy": spec_energy.get(w.id),
            },
        })
    return jsonify({"type": "FeatureCollection", "features": features})



################################################################################
# FILE: install\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("install", __name__)

from app.blueprints.install import routes  # noqa: E402,F401



################################################################################
# FILE: install\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, BooleanField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.install import bp
from app.security import permission_required
from app.utils.forms import (
    JalaliDateField, PersianFloatField as FloatField, PersianIntegerField as IntegerField,
)
from app.models.well import Well
from app.models.pump_asset import PumpInstallation, Supplier
from app.models.audit import RecordHistory
from app.models.constants import EQUIP_CONDITIONS
from app import workflow


def get_or_create_supplier(name, kind=None):
    from app.models.constants import canonical_supplier
    name = canonical_supplier(name)
    if not name:
        return None
    s = db.session.scalar(db.select(Supplier).filter_by(name=name))
    if s is None:
        s = Supplier(name=name, kind=kind)
        db.session.add(s)
        db.session.flush()
    elif kind and not s.kind:
        s.kind = kind
    return s


class InstallForm(FlaskForm):
    install_no = IntegerField("شماره نصب", validators=[Optional()])
    install_date = JalaliDateField("تاریخ نصب", validators=[Optional()])
    pull_date = JalaliDateField("تاریخ کشیدن", validators=[Optional()])
    motor_power_kw = FloatField("توان موتور (kW)", validators=[Optional()])
    motor_condition = SelectField("موتور", choices=[("", "—")] + EQUIP_CONDITIONS, validators=[Optional()])
    pump_type = StringField("تیپ پمپ", validators=[Optional()])
    pump_stages = IntegerField("طبقه", validators=[Optional()])
    pump_condition = SelectField("پمپ", choices=[("", "—")] + EQUIP_CONDITIONS, validators=[Optional()])
    manufacturer_name = StringField("سازنده", validators=[Optional()])
    contractor_name = StringField("پیمانکار", validators=[Optional()])
    install_depth_m = FloatField("عمق نصب", validators=[Optional()])
    well_depth_m = FloatField("عمق چاه", validators=[Optional()])
    static_level = FloatField("سطح استاتیک", validators=[Optional()])
    dynamic_level = FloatField("سطح دینامیک", validators=[Optional()])
    route_loss = FloatField("تلفات مسیر", validators=[Optional()])
    grid_pressure_m = FloatField("فشار شبکه", validators=[Optional()])
    work_shift = StringField("شیفت کاری", validators=[Optional()])
    removal_reason = StringField("علت کشیدن", validators=[Optional()])
    fault_by_operator = TextAreaField("شرح خرابی (بهره‌بردار)", validators=[Optional()])
    fault_by_workshop = TextAreaField("شرح خرابی (کارگاه مکانیک)", validators=[Optional()])
    pm_form_registered = BooleanField("فرم نصب در PM ثبت شده")
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


SIMPLE_FIELDS = [
    "install_no", "install_date", "pull_date", "motor_power_kw", "motor_condition",
    "pump_type", "pump_stages", "pump_condition", "install_depth_m", "well_depth_m",
    "static_level", "dynamic_level", "route_loss", "grid_pressure_m", "work_shift",
    "removal_reason", "fault_by_operator", "fault_by_workshop", "pm_form_registered", "notes",
]


def _apply(form, rec):
    for f in SIMPLE_FIELDS:
        setattr(rec, f, getattr(form, f).data)
    man = get_or_create_supplier(form.manufacturer_name.data, "manufacturer")
    con = get_or_create_supplier(form.contractor_name.data, "contractor")
    rec.manufacturer_id = man.id if man else None
    rec.contractor_id = con.id if con else None


@bp.route("/wells/<int:well_id>/install/new", methods=["GET", "POST"])
@permission_required("install", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = InstallForm()
    if form.validate_on_submit():
        rec = PumpInstallation(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("نصب/کشیدن پمپ ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("install.detail", record_id=rec.id))
    return render_template("install/form.html", form=form, well=well, title="ثبت نصب/کشیدن پمپ")


@bp.route("/install/<int:record_id>")
@permission_required("install", "view")
def detail(record_id):
    rec = db.get_or_404(PumpInstallation, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="pump_installations", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("install/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/install/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("install", "edit")
def edit(record_id):
    rec = db.get_or_404(PumpInstallation, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("install.detail", record_id=rec.id))
    form = InstallForm(obj=rec)
    if request.method == "GET":
        form.manufacturer_name.data = rec.manufacturer.name if rec.manufacturer else ""
        form.contractor_name.data = rec.contractor.name if rec.contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("نصب/کشیدن به‌روزرسانی شد.", "success")
        return redirect(url_for("install.detail", record_id=rec.id))
    return render_template("install/form.html", form=form, well=rec.well, title="ویرایش نصب/کشیدن")


@bp.route("/install/<int:record_id>/delete", methods=["POST"])
@permission_required("install", "delete")
def delete(record_id):
    rec = db.get_or_404(PumpInstallation, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد نصب/کشیدن حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(PumpInstallation, record_id)
    if not current_user.has_permission("install", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("install.detail", record_id=record_id))


@bp.route("/install/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/install/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/install/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/install/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: main\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("main", __name__)

from app.blueprints.main import routes  # noqa: E402,F401



################################################################################
# FILE: main\routes.py
################################################################################

import re
from datetime import datetime
from io import BytesIO

from flask import render_template, request, flash, redirect, url_for, send_file
from flask_login import login_required, current_user
from flask_wtf import FlaskForm
from wtforms import (StringField, PasswordField, BooleanField, TextAreaField,
                     SelectMultipleField, SubmitField)
from wtforms.validators import DataRequired, Optional, Length
from app.extensions import db
from app.blueprints.main import bp
from app.security import permission_required
from app.models.well import Well
from app.models.org import OrgUnit
from app.models.rbac import User, Role, LoginLog
from app.models.constants import WELL_STATUSES


@bp.route("/")
@login_required
def dashboard():
    well_count = db.session.scalar(db.select(db.func.count(Well.id))) or 0
    org_count = db.session.scalar(db.select(db.func.count(OrgUnit.id))) or 0
    mapped = db.session.scalar(
        db.select(db.func.count(Well.id)).where(Well.latitude.isnot(None))
    ) or 0
    in_circuit = db.session.scalar(
        db.select(db.func.count(Well.id)).where(Well.status == "in_circuit")
    ) or 0

    # Wells grouped by status (for the doughnut + legend).
    status_labels = dict(WELL_STATUSES)
    rows = db.session.execute(
        db.select(Well.status, db.func.count(Well.id)).group_by(Well.status)
    ).all()
    by_status = [
        {"key": s or "unknown", "label": status_labels.get(s, s or "نامشخص"), "count": n}
        for s, n in rows
    ]

    # Urban vs rural split.
    kind_rows = db.session.execute(
        db.select(Well.well_kind, db.func.count(Well.id)).group_by(Well.well_kind)
    ).all()
    kind_counts = {k or "unknown": n for k, n in kind_rows}

    # Coverage of mapped wells (for the progress ring).
    mapped_pct = round((mapped / well_count) * 100) if well_count else 0

    recent_wells = db.session.scalars(
        db.select(Well).order_by(Well.created_at.desc()).limit(6)
    ).all()

    # --- Attention panel (read-only surfacing of existing signals) ---
    from app.models.quality import WaterQuality
    quality_issue_wells = db.session.scalar(
        db.select(db.func.count(db.distinct(WaterQuality.well_id))).where(
            db.or_(
                WaterQuality.is_potable.is_(False),
                WaterQuality.turbidity.isnot(None),
                WaterQuality.sholat.isnot(None),
            )
        )
    ) or 0
    under_rehab = db.session.scalar(
        db.select(db.func.count(Well.id)).where(Well.status == "under_rehab")
    ) or 0

    top_priority = []
    try:
        from app.services import prioritization
        ranked, _ = prioritization.compute(top=5)
        top_priority = ranked
    except Exception:
        top_priority = []

    try:
        from app.services import alerts as alert_svc
        alert_summary = alert_svc.summary(include_predictive=False)  # fast count
    except Exception:
        alert_summary = {"total": 0, "by_severity": {"high": 0, "medium": 0, "low": 0}}

    # --- Operational signals (zone-based dashboard) ---
    from flask import url_for
    from datetime import date, timedelta
    from app.models.pump_asset import PumpInstallation
    from app.models.production import MonthlyProduction as MP
    from app.models.rehab import Rehabilitation
    from app.models.maintenance import MaintenanceRecord

    running_pumps = db.session.scalar(
        db.select(db.func.count(PumpInstallation.id)).where(PumpInstallation.pull_date.is_(None))) or 0
    fail_cutoff = date.today() - timedelta(days=365)
    recent_failures = db.session.scalar(
        db.select(db.func.count(PumpInstallation.id)).where(
            PumpInstallation.pull_date.isnot(None), PumpInstallation.pull_date >= fail_cutoff)) or 0

    e_sum = db.session.execute(
        db.select(db.func.sum(MP.energy_kwh), db.func.sum(MP.production_m3))
        .where(MP.energy_kwh.isnot(None))).first()
    fleet_se = round(e_sum[0] / e_sum[1], 2) if (e_sum and e_sum[0] and e_sum[1]) else None

    tr = db.session.execute(
        db.select(MP.jyear, MP.jmonth, db.func.sum(MP.production_m3))
        .where(MP.production_m3.isnot(None)).group_by(MP.jyear, MP.jmonth)
        .order_by(MP.jyear, MP.jmonth)).all()[-18:]
    trend = {"labels": [f"{y}/{m:02d}" for y, m, _ in tr],
             "series": [round((v or 0) / 1e6, 2) for _, _, v in tr]}

    activity = []
    for i in db.session.scalars(db.select(PumpInstallation).where(
            PumpInstallation.pull_date.isnot(None)).order_by(PumpInstallation.pull_date.desc()).limit(5)).all():
        activity.append({"date": i.pull_date, "well_id": i.well_id, "well": i.well.name,
                         "type": "کشیدن پمپ", "detail": i.removal_reason or "", "icon": "bi-tools",
                         "color": "#2f6bff", "url": url_for("install.detail", record_id=i.id)})
    for r in db.session.scalars(db.select(Rehabilitation).where(
            Rehabilitation.rehab_date.isnot(None)).order_by(Rehabilitation.rehab_date.desc()).limit(5)).all():
        activity.append({"date": r.rehab_date, "well_id": r.well_id, "well": r.well.name,
                         "type": "بهسازی", "detail": r.reason or "", "icon": "bi-arrow-repeat",
                         "color": "#e0463e", "url": url_for("rehab.detail", record_id=r.id)})
    for m in db.session.scalars(db.select(MaintenanceRecord).where(
            MaintenanceRecord.report_date.isnot(None)).order_by(MaintenanceRecord.report_date.desc()).limit(5)).all():
        activity.append({"date": m.report_date, "well_id": m.well_id, "well": m.well.name,
                         "type": "نگهداری", "detail": m.fault_desc or "", "icon": "bi-wrench-adjustable",
                         "color": "#6a7180", "url": url_for("maintenance.detail", record_id=m.id)})
    activity = sorted([a for a in activity if a["date"]], key=lambda a: a["date"], reverse=True)[:7]

    import jdatetime
    today = jdatetime.date.today().strftime("%Y/%m/%d")
    return render_template(
        "main/dashboard.html",
        today=today, well_count=well_count, org_count=org_count,
        mapped=mapped, mapped_pct=mapped_pct, in_circuit=in_circuit,
        by_status=by_status, urban=kind_counts.get("urban", 0),
        rural=kind_counts.get("rural", 0), recent_wells=recent_wells,
        status_labels=status_labels, quality_issue_wells=quality_issue_wells,
        under_rehab=under_rehab, top_priority=top_priority, alert_summary=alert_summary,
        running_pumps=running_pumps, recent_failures=recent_failures,
        fleet_se=fleet_se, trend=trend, activity=activity,
    )


@bp.route("/admin/users")
@permission_required("admin", "view")
def admin_users():
    q = (request.args.get("q") or "").strip()
    query = db.select(User).order_by(User.username)
    users = db.session.scalars(query).unique().all()
    if q:
        ql = q.lower()
        users = [u for u in users if ql in (u.username or "").lower()
                 or ql in (u.full_name or "").lower()
                 or ql in (u.personnel_code or "").lower()]
    return render_template("main/users.html", users=users, q=q)


@bp.route("/admin/economics", methods=["GET", "POST"])
@permission_required("admin", "view")
def admin_economics():
    from flask import request, flash, redirect, url_for
    from flask_login import current_user
    from app.models.economic import EconomicParam
    from app.services import economics
    from app.utils.dates import to_english_digits
    economics.ensure_params()
    rows = db.session.scalars(db.select(EconomicParam).order_by(EconomicParam.id)).all()
    if request.method == "POST":
        if not current_user.has_permission("admin", "edit"):
            flash("مجوز ویرایش ندارید.", "warning")
            return redirect(url_for("main.admin_economics"))
        for p in rows:
            v = to_english_digits(request.form.get(f"v_{p.id}", "")).strip()
            try:
                p.value = float(v)
            except ValueError:
                pass
        db.session.commit()
        flash("پارامترهای اقتصادی به‌روزرسانی شد.", "success")
        return redirect(url_for("main.admin_economics"))
    return render_template("main/economics.html", params=rows,
                           can_edit=current_user.has_permission("admin", "edit"))
# ==================== دسته ۵ — مدیریت کاربران ====================

def _password_errors(pw):
    errs = []
    if len(pw) < 8:
        errs.append("حداقل ۸ کاراکتر")
    if not re.search(r"[A-Z]", pw):
        errs.append("یک حرف بزرگ (A-Z)")
    if not re.search(r"[a-z]", pw):
        errs.append("یک حرف کوچک (a-z)")
    if not re.search(r"\d", pw):
        errs.append("یک عدد")
    if not re.search(r"[^A-Za-z0-9]", pw):
        errs.append("یک کاراکتر ویژه (!@#...)")
    return errs


class UserForm(FlaskForm):
    first_name = StringField("نام", validators=[Optional(), Length(max=60)])
    last_name = StringField("نام خانوادگی", validators=[Optional(), Length(max=60)])
    national_id = StringField("کد ملی", validators=[Optional(), Length(max=10)])
    personnel_code = StringField("کد پرسنلی", validators=[Optional(), Length(max=30)])
    position = StringField("سمت سازمانی", validators=[Optional(), Length(max=120)])
    phone = StringField("شماره تماس", validators=[Optional(), Length(max=20)])
    email = StringField("ایمیل", validators=[Optional(), Length(max=120)])
    username = StringField("نام کاربری", validators=[DataRequired(), Length(max=60)])
    password = PasswordField("رمز عبور", validators=[Optional()])
    is_active = BooleanField("فعال", default=True)
    role_ids = SelectMultipleField("نقش‌ها", coerce=int)
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


def _set_roles(form):
    form.role_ids.choices = [
        (r.id, r.name) for r in db.session.scalars(db.select(Role).order_by(Role.name)).unique()]


@bp.route("/admin/users/new", methods=["GET", "POST"])
@permission_required("admin", "edit")
def user_create():
    form = UserForm()
    _set_roles(form)
    if form.validate_on_submit():
        uname = form.username.data.strip()
        if db.session.scalar(db.select(User).filter_by(username=uname)):
            flash("این نام کاربری قبلاً استفاده شده است.", "danger")
        else:
            pw = form.password.data or ""
            errs = _password_errors(pw)
            if errs:
                flash("رمز عبور ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
            else:
                u = User(username=uname, is_active=bool(form.is_active.data))
                _apply_user(form, u)
                u.set_password(pw)
                db.session.add(u)
                db.session.commit()
                flash("کاربر ایجاد شد.", "success")
                return redirect(url_for("main.admin_users"))
    return render_template("main/user_form.html", form=form, title="ایجاد کاربر", is_new=True)


@bp.route("/admin/users/<int:user_id>/edit", methods=["GET", "POST"])
@permission_required("admin", "edit")
def user_edit(user_id):
    u = db.get_or_404(User, user_id)
    form = UserForm(obj=u)
    _set_roles(form)
    if request.method == "GET":
        form.role_ids.data = [r.id for r in u.roles]
    if form.validate_on_submit():
        uname = form.username.data.strip()
        clash = db.session.scalar(db.select(User).filter(User.username == uname, User.id != u.id))
        if clash:
            flash("این نام کاربری قبلاً استفاده شده است.", "danger")
        else:
            u.username = uname
            u.is_active = bool(form.is_active.data)
            _apply_user(form, u)
            if form.password.data:
                errs = _password_errors(form.password.data)
                if errs:
                    flash("رمز عبور ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
                    return render_template("main/user_form.html", form=form, title="ویرایش کاربر", is_new=False, obj=u)
                u.set_password(form.password.data)
            db.session.commit()
            flash("کاربر به‌روزرسانی شد.", "success")
            return redirect(url_for("main.admin_users"))
    return render_template("main/user_form.html", form=form, title="ویرایش کاربر", is_new=False, obj=u)


def _apply_user(form, u):
    u.first_name = form.first_name.data or None
    u.last_name = form.last_name.data or None
    u.full_name = (" ".join(x for x in [u.first_name, u.last_name] if x)).strip() or None
    u.national_id = form.national_id.data or None
    u.personnel_code = form.personnel_code.data or None
    u.position = form.position.data or None
    u.phone = form.phone.data or None
    u.email = form.email.data or None
    u.notes = form.notes.data or None
    u.roles = db.session.scalars(
        db.select(Role).filter(Role.id.in_(form.role_ids.data or []))).unique().all()


@bp.route("/admin/users/<int:user_id>/toggle", methods=["POST"])
@permission_required("admin", "edit")
def user_toggle_active(user_id):
    u = db.get_or_404(User, user_id)
    if u.id == current_user.id:
        flash("نمی‌توانید حساب خودتان را غیرفعال کنید.", "warning")
    else:
        u.is_active = not u.is_active
        db.session.commit()
        flash("وضعیت کاربر تغییر کرد.", "info")
    return redirect(url_for("main.admin_users"))


@bp.route("/admin/users/<int:user_id>/reset-password", methods=["POST"])
@permission_required("admin", "edit")
def user_reset_password(user_id):
    u = db.get_or_404(User, user_id)
    new_pw = (request.form.get("new_password") or "").strip()
    errs = _password_errors(new_pw)
    if errs:
        flash("رمز عبور ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
    else:
        u.set_password(new_pw)
        db.session.commit()
        flash(f"رمز عبور «{u.username}» بازنشانی شد.", "success")
    return redirect(url_for("main.admin_users"))


@bp.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@permission_required("admin", "delete")
def user_delete(user_id):
    u = db.get_or_404(User, user_id)
    if u.id == current_user.id:
        flash("نمی‌توانید حساب خودتان را حذف کنید.", "warning")
    elif u.is_superuser:
        flash("حذف مدیر کل مجاز نیست.", "warning")
    else:
        db.session.delete(u)
        db.session.commit()
        flash("کاربر حذف شد.", "info")
    return redirect(url_for("main.admin_users"))


@bp.route("/admin/users/activity")
@permission_required("admin", "view")
def user_activity():
    logs = db.session.scalars(
        db.select(LoginLog).order_by(LoginLog.login_at.desc()).limit(1000)).all()
    rows = _activity_rows(logs)
    return render_template("main/user_activity.html", rows=rows)


def _fmt_duration(sec):
    if sec is None:
        return "—"
    h, rem = divmod(int(sec), 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h}س {m}د"
    if m:
        return f"{m}د {s}ث"
    return f"{s}ث"


def _activity_rows(logs):
    counts = {}
    for lg in logs:
        counts[lg.user_id] = counts.get(lg.user_id, 0) + 1
    out = []
    for lg in logs:
        u = lg.user
        out.append({
            "full_name": (u.full_name if u else None) or (u.username if u else "—"),
            "username": u.username if u else "—",
            "login_at": lg.login_at,
            "logout_at": lg.logout_at,
            "duration": _fmt_duration(lg.duration_seconds),
            "ip": lg.ip_address or "—",
            "agent": lg.user_agent or "—",
            "count": counts.get(lg.user_id, 0),
            "online": lg.logout_at is None,
        })
    return out


@bp.route("/admin/users/activity/export")
@permission_required("admin", "view")
def user_activity_export():
    import pandas as pd
    logs = db.session.scalars(db.select(LoginLog).order_by(LoginLog.login_at.desc())).all()
    rows = _activity_rows(logs)
    df = pd.DataFrame([{
        "نام و نام خانوادگی": r["full_name"],
        "نام کاربری": r["username"],
        "تاریخ/ساعت ورود": r["login_at"].strftime("%Y-%m-%d %H:%M") if r["login_at"] else "",
        "ساعت خروج": r["logout_at"].strftime("%Y-%m-%d %H:%M") if r["logout_at"] else "",
        "مدت حضور": r["duration"],
        "IP": r["ip"],
        "مرورگر/دستگاه": r["agent"],
        "تعداد ورود": r["count"],
        "وضعیت": "آنلاین" if r["online"] else "آفلاین",
    } for r in rows])
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        (df if not df.empty else pd.DataFrame({"—": []})).to_excel(
            xl, sheet_name="فعالیت کاربران", index=False)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     download_name=f"user_activity_{datetime.now():%Y%m%d}.xlsx")


class ProfileForm(FlaskForm):
    first_name = StringField("نام", validators=[Optional(), Length(max=60)])
    last_name = StringField("نام خانوادگی", validators=[Optional(), Length(max=60)])
    phone = StringField("شماره تماس", validators=[Optional(), Length(max=20)])
    email = StringField("ایمیل", validators=[Optional(), Length(max=120)])
    username = StringField("نام کاربری", validators=[DataRequired(), Length(max=60)])
    current_password = PasswordField("رمز فعلی", validators=[Optional()])
    new_password = PasswordField("رمز جدید", validators=[Optional()])
    submit = SubmitField("ذخیره تغییرات")


@bp.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    form = ProfileForm(obj=current_user)
    if form.validate_on_submit():
        uname = form.username.data.strip()
        clash = db.session.scalar(db.select(User).filter(User.username == uname, User.id != current_user.id))
        if clash:
            flash("این نام کاربری قبلاً استفاده شده است.", "danger")
        else:
            current_user.username = uname
            current_user.first_name = form.first_name.data or None
            current_user.last_name = form.last_name.data or None
            current_user.full_name = (" ".join(x for x in [current_user.first_name, current_user.last_name] if x)).strip() or None
            current_user.phone = form.phone.data or None
            current_user.email = form.email.data or None
            if form.new_password.data:
                if not current_user.check_password(form.current_password.data or ""):
                    flash("رمز فعلی نادرست است.", "danger")
                    return render_template("main/profile.html", form=form)
                errs = _password_errors(form.new_password.data)
                if errs:
                    flash("رمز جدید ضعیف است؛ باید شامل: " + "، ".join(errs), "danger")
                    return render_template("main/profile.html", form=form)
                current_user.set_password(form.new_password.data)
            db.session.commit()
            flash("پروفایل به‌روزرسانی شد.", "success")
            return redirect(url_for("main.profile"))
    return render_template("main/profile.html", form=form)


################################################################################
# FILE: maintenance\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("maintenance", __name__)

from app.blueprints.maintenance import routes  # noqa: E402,F401



################################################################################
# FILE: maintenance\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.maintenance import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.maintenance import MaintenanceRecord
from app.models.pump_asset import PumpInstallation
from app.models.audit import RecordHistory
from app.models.constants import MAINTENANCE_TYPES, MAINTENANCE_CATEGORIES
from app.blueprints.install.routes import get_or_create_supplier
from app.utils.dates import format_jalali as _jd
from app import workflow


class MaintenanceForm(FlaskForm):
    report_date = JalaliDateField("تاریخ گزارش", validators=[Optional()])
    maint_type = SelectField("نوع نگهداری", choices=[("", "—")] + MAINTENANCE_TYPES, validators=[Optional()])
    category = SelectField("دسته", choices=[("", "—")] + MAINTENANCE_CATEGORIES, validators=[Optional()])
    down_from = JalaliDateField("شروع خاموشی", validators=[Optional()])
    down_to = JalaliDateField("پایان خاموشی", validators=[Optional()])
    downtime_hours = FloatField("مدت خاموشی (ساعت)", validators=[Optional()])
    fault_desc = TextAreaField("شرح خرابی", validators=[Optional()])
    action_taken = TextAreaField("اقدام انجام‌شده", validators=[Optional()])
    parts_replaced = StringField("قطعات تعویض‌شده", validators=[Optional()])
    contractor_name = StringField("پیمانکار/مجری", validators=[Optional()])
    cost = FloatField("هزینه", validators=[Optional()])
    pump_installation_id = SelectField("نصب مرتبط", coerce=int, validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_installs(self, well_id):
        installs = db.session.scalars(
            db.select(PumpInstallation).filter_by(well_id=well_id)
            .order_by(PumpInstallation.install_date.desc())).all()
        self.pump_installation_id.choices = [(0, "—")] + [
            (i.id, f"نصب {_jd(i.install_date, '') } — {i.pump_type or ''}".strip())
            for i in installs]


SIMPLE = ["report_date", "maint_type", "category", "down_from", "down_to",
          "downtime_hours", "fault_desc", "action_taken", "parts_replaced", "cost", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data or None)
    rec.pump_installation_id = form.pump_installation_id.data or None
    c = get_or_create_supplier(form.contractor_name.data, "contractor")
    rec.contractor_id = c.id if c else None


@bp.route("/wells/<int:well_id>/maintenance/new", methods=["GET", "POST"])
@permission_required("maintenance", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = MaintenanceForm()
    form.populate_installs(well.id)
    if form.validate_on_submit():
        rec = MaintenanceRecord(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("نگهداری ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("maintenance.detail", record_id=rec.id))
    return render_template("maintenance/form.html", form=form, well=well, title="ثبت نگهداری/تعمیر")


@bp.route("/maintenance/<int:record_id>")
@permission_required("maintenance", "view")
def detail(record_id):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="maintenance_records", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())).all()
    return render_template("maintenance/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           type_labels=dict(MAINTENANCE_TYPES),
                           cat_labels=dict(MAINTENANCE_CATEGORIES))


@bp.route("/maintenance/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("maintenance", "edit")
def edit(record_id):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("maintenance.detail", record_id=rec.id))
    form = MaintenanceForm(obj=rec)
    form.populate_installs(rec.well_id)
    if request.method == "GET":
        form.contractor_name.data = rec.contractor.name if rec.contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("نگهداری به‌روزرسانی شد.", "success")
        return redirect(url_for("maintenance.detail", record_id=rec.id))
    return render_template("maintenance/form.html", form=form, well=rec.well, title="ویرایش نگهداری")


@bp.route("/maintenance/<int:record_id>/delete", methods=["POST"])
@permission_required("maintenance", "delete")
def delete(record_id):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد نگهداری حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(MaintenanceRecord, record_id)
    if not current_user.has_permission("maintenance", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("maintenance.detail", record_id=record_id))


@bp.route("/maintenance/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/maintenance/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/maintenance/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/maintenance/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: mechanic\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("mechanic", __name__, url_prefix="/mechanic")

from app.blueprints.mechanic import routes  # noqa: E402,F401


################################################################################
# FILE: mechanic\routes.py
################################################################################

"""ماژول کارگاه مکانیک — فاز ۲: فهرست + فرم ورود/ویرایش/جزئیات."""
from flask import render_template, redirect, url_for, flash
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional, ValidationError

from app.extensions import db
from app.blueprints.mechanic import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField, PersianIntegerField
from app.models.well import Well
from app.models.mechanic import MechanicEvent

OP_TYPES = [
    ("", "—"),
    ("کشیدن", "کشیدن"),
    ("نصب", "نصب"),
    ("جمع‌آوری", "جمع‌آوری"),
    ("نصب جدید", "نصب جدید"),
]
COND = [("", "—"), ("نو", "نو"), ("تعمیری", "تعمیری")]


class MechanicForm(FlaskForm):
    op_date = JalaliDateField("تاریخ عملیات", validators=[Optional()])
    op_type = SelectField("نوع عملیات", choices=OP_TYPES, validators=[Optional()])
    fault_type = SelectField("نوع خرابی", choices=[("", "—"), ("سوختگی", "سوختگی"),
                             ("ایراد مکانیکی", "ایراد مکانیکی"), ("کاهش آبدهی", "کاهش آبدهی"),
                             ("سایر", "سایر")], validators=[Optional()])
    fault_description = TextAreaField("شرح خرابی از نظر بهره‌بردار", validators=[Optional()])
    contractor = StringField("نام پیمانکار", validators=[Optional()])
    pm_form_no = StringField("فرم نصب در PM", validators=[Optional()])

    motor_desc = StringField("مشخصات موتور", validators=[Optional()])
    motor_condition = SelectField("وضعیت موتور", choices=COND, validators=[Optional()])
    pump_desc = StringField("مشخصات پمپ", validators=[Optional()])
    pump_condition = SelectField("وضعیت پمپ", choices=COND, validators=[Optional()])

    tip_change = StringField("تغییر تیپ", validators=[Optional()])
    prev_install_date = JalaliDateField("تاریخ نصب قبلی", validators=[Optional()])
    well_depth = PersianFloatField("عمق چاه (متر)", validators=[Optional()])
    prev_install_depth = PersianFloatField("عمق نصب قبلی (متر)", validators=[Optional()])
    curr_install_depth = PersianFloatField("عمق نصب فعلی (متر)", validators=[Optional()])
    static_level = PersianFloatField("سطح استاتیک (متر)", validators=[Optional()])
    dynamic_level = PersianFloatField("سطح دینامیک (متر)", validators=[Optional()])
    path_loss = PersianFloatField("تلفات مسیر (متر)", validators=[Optional()])
    network_pressure = PersianFloatField("فشار شبکه (متر)", validators=[Optional()])
    total_head = PersianFloatField("هد کلی (متر)", validators=[Optional()])
    design_flow = PersianFloatField("دبی طراحی (l/s)", validators=[Optional()])
    pipe_diameter = PersianFloatField("قطر لوله آبده (اینچ)", validators=[Optional()])

    pt_date = JalaliDateField("تاریخ آزمایش پمپاژ کارگاه", validators=[Optional()])
    pt_pressure = PersianFloatField("فشار آزمایش پمپاژ", validators=[Optional()])
    pt_flow = PersianFloatField("دبی آزمایش پمپاژ (l/s)", validators=[Optional()])

    cable_size = StringField("سایز کابل", validators=[Optional()])
    cable_change = StringField("تغییر سایز/تیپ کابل", validators=[Optional()])
    starter = StringField("راه‌انداز", validators=[Optional()])

    months_worked = PersianIntegerField("تعداد ماه‌های کارکرد", validators=[Optional()])
    mechanic_opinion = TextAreaField("نظر کارگاه مکانیک", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def validate_design_flow(self, field):
        if field.data is not None and field.data < 0:
            raise ValidationError("دبی نمی‌تواند منفی باشد.")


FIELDS = ["op_date", "op_type", "fault_type", "fault_description", "contractor", "pm_form_no",
          "motor_desc", "motor_condition", "pump_desc", "pump_condition",
          "tip_change", "prev_install_date", "well_depth", "prev_install_depth",
          "curr_install_depth", "static_level", "dynamic_level", "path_loss",
          "network_pressure", "total_head", "design_flow", "pipe_diameter",
          "pt_date", "pt_pressure", "pt_flow", "cable_size", "cable_change",
          "starter", "months_worked", "mechanic_opinion", "notes"]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data if getattr(form, f).data not in ("",) else None)


@bp.route("/")
@permission_required("mechanic", "view")
def index():
    events = db.session.scalars(
        db.select(MechanicEvent).join(Well, MechanicEvent.well_id == Well.id)
        .order_by(MechanicEvent.op_date.desc())
    ).all()
    op_types = sorted({e.op_type for e in events if e.op_type})
    return render_template("mechanic/index.html", events=events, op_types=op_types)

@bp.route("/well/<int:well_id>/new", methods=["GET", "POST"])
@permission_required("mechanic", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = MechanicForm()
    if form.validate_on_submit():
        rec = MechanicEvent(well_id=well.id, created_by_id=current_user.id, source="manual")
        _apply(form, rec)
        db.session.add(rec)
        db.session.commit()
        flash("رویداد کارگاه مکانیک ثبت شد.", "success")
        return redirect(url_for("mechanic.detail", record_id=rec.id))
    return render_template("mechanic/form.html", form=form, well=well, title="ثبت رویداد کارگاه مکانیک")


@bp.route("/<int:record_id>")
@permission_required("mechanic", "view")
def detail(record_id):
    rec = db.get_or_404(MechanicEvent, record_id)
    return render_template("mechanic/detail.html", rec=rec, well=rec.well, stages=STAGES)


@bp.route("/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("mechanic", "edit")
def edit(record_id):
    rec = db.get_or_404(MechanicEvent, record_id)
    form = MechanicForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        db.session.commit()
        flash("رویداد به‌روزرسانی شد.", "success")
        return redirect(url_for("mechanic.detail", record_id=rec.id))
    return render_template("mechanic/form.html", form=form, well=rec.well, title="ویرایش رویداد کارگاه مکانیک")


@bp.route("/<int:record_id>/delete", methods=["POST"])
@permission_required("mechanic", "delete")
def delete(record_id):
    rec = db.get_or_404(MechanicEvent, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رویداد حذف شد.", "info")
    return redirect(url_for("mechanic.index"))
# ---------- انبار قطعات (فاز ۳) ----------
from app.models.mechanic_part import MechanicPart
from wtforms import IntegerField


class PartForm(FlaskForm):
    equipment = SelectField("نوع تجهیز", choices=[("موتور", "موتور"), ("پمپ", "پمپ")],
                            validators=[Optional()])
    part_name = StringField("نام قطعه", validators=[Optional()])
    total_count = PersianIntegerField("تعداد کل", validators=[Optional()])
    usable = PersianIntegerField("قابل استفاده", validators=[Optional()])
    scrap = PersianIntegerField("اسقاط", validators=[Optional()])
    new_count = PersianIntegerField("نو", validators=[Optional()])
    repaired = PersianIntegerField("تعمیری", validators=[Optional()])
    inventory_code = StringField("کد انباری قطعه", validators=[Optional()])
    part_type = StringField("نوع قطعه", validators=[Optional()])
    manufacturer = StringField("سازنده قطعه", validators=[Optional()])
    entry_date = JalaliDateField("تاریخ ورود", validators=[Optional()])
    notes = StringField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


@bp.route("/parts")
@permission_required("mechanic", "view")
def parts():
    motor = db.session.scalars(db.select(MechanicPart).filter_by(equipment="موتور")
                               .order_by(MechanicPart.id)).all()
    pump = db.session.scalars(db.select(MechanicPart).filter_by(equipment="پمپ")
                              .order_by(MechanicPart.id)).all()
    return render_template("mechanic/parts.html", motor=motor, pump=pump)


@bp.route("/parts/new", methods=["GET", "POST"])
@permission_required("mechanic", "create")
def part_create():
    form = PartForm()
    if form.validate_on_submit():
        p = MechanicPart(created_by_id=current_user.id)
        for f in ["equipment", "part_name", "total_count", "usable", "scrap",
                  "new_count", "repaired", "inventory_code", "part_type",
                  "manufacturer", "entry_date", "notes"]:
            setattr(p, f, getattr(form, f).data)
        db.session.add(p)
        db.session.commit()
        flash("قطعه ثبت شد.", "success")
        return redirect(url_for("mechanic.parts"))
    return render_template("mechanic/part_form.html", form=form, title="افزودن قطعه", part_id=None)


@bp.route("/parts/<int:part_id>/edit", methods=["GET", "POST"])
@permission_required("mechanic", "edit")
def part_edit(part_id):
    p = db.get_or_404(MechanicPart, part_id)
    form = PartForm(obj=p)
    if form.validate_on_submit():
        for f in ["equipment", "part_name", "total_count", "usable", "scrap",
                  "new_count", "repaired", "inventory_code", "part_type",
                  "manufacturer", "notes"]:
            setattr(p, f, getattr(form, f).data)
        p.updated_by_id = current_user.id
        db.session.commit()
        flash("قطعه به‌روزرسانی شد.", "success")
        return redirect(url_for("mechanic.parts"))
    return render_template("mechanic/part_form.html", form=form, title="ویرایش قطعه", part_id=p.id)


@bp.route("/parts/<int:part_id>/delete", methods=["POST"])
@permission_required("mechanic", "delete")
def part_delete(part_id):
    p = db.get_or_404(MechanicPart, part_id)
    db.session.delete(p)
    db.session.commit()
    flash("قطعه حذف شد.", "info")
    return redirect(url_for("mechanic.parts"))
# ---------- گردش کار مرحله‌ای (فاز ۴) ----------
from app.models.mechanic import MechanicStageLog

# مراحل به ترتیب فلوچارت
STAGES = [
    "اعلام حادثه",
    "تشخیص نوع خرابی",
    "دفتر فنی و مهندسی",
    "ویدئومتری",
    "بهسازی چاه",
    "انتخاب پمپ",
    "پمپاژ آزمایشی",
    "مونتاژ",
    "تست چاله پمپاژ",
    "دمونتاژ",
    "ارجاع به کارگاه نصب",
    "نصب نهایی",
    "پایان‌یافته",
]

FAULT_TYPES = ["سوختگی", "ایراد مکانیکی", "کاهش آبدهی", "سایر"]


@bp.route("/<int:record_id>/stage", methods=["POST"])
@permission_required("mechanic", "edit")
def change_stage(record_id):
    from flask import request
    rec = db.get_or_404(MechanicEvent, record_id)
    new_stage = request.form.get("to_stage", "").strip()
    note = request.form.get("note", "").strip()
    if new_stage and new_stage in STAGES and new_stage != rec.stage:
        log = MechanicStageLog(
            event_id=rec.id, from_stage=rec.stage, to_stage=new_stage,
            note=note or None, changed_by_id=current_user.id)
        rec.stage = new_stage
        db.session.add(log)
        db.session.commit()
        flash(f"مرحله به «{new_stage}» تغییر کرد.", "success")
    else:
        flash("مرحله‌ی معتبری انتخاب نشد.", "warning")
    return redirect(url_for("mechanic.detail", record_id=rec.id))
# ---------- گزارش‌های مدیریتی (فاز ۵) ----------
from sqlalchemy import func
from collections import Counter
import jdatetime

def classify_fault(event):
    """دسته‌بندی خرابی از روی fault_type یا متن fault_description."""
    # اگر نوع خرابی صریح ثبت شده، همان
    if event.fault_type:
        return event.fault_type
    text = (event.fault_description or "").strip()
    if not text:
        return "نامشخص"
    # دسته‌بندی بر اساس کلمات کلیدی
    rules = [
        ("سوختگی", ["سوخت", "سوختن"]),
        ("کاهش آبدهی", ["کاهش دبی", "کاهش آبدهی", "کمبود", "عدم آبده", "عدم ابده", "افت"]),
        ("ایراد مکانیکی", ["صدا", "لرزش", "گیرپاژ", "گیر و پاژ", "گیرو پاژ", "مکانیک"]),
        ("ایراد برقی", ["اهم", "شولات", "شولاتی", "آمپر", "امپر", "برق"]),
        ("هوادهی", ["هوادهی", "هوا"]),
        ("عملیات نصب/کشیدن", ["نصب", "کشیدن", "جمع آوری", "جمع‌آوری", "تجهیز", "جابجایی"]),
        ("تغییر فشار", ["فشار"]),
    ]
    for label, keywords in rules:
        if any(k in text for k in keywords):
            return label
    return "سایر"

@bp.route("/reports")
@permission_required("mechanic", "view")
def reports():
    from flask import request
    # فیلترها
    f_year = request.args.get("year", "").strip()
    f_optype = request.args.get("optype", "").strip()
    f_fault = request.args.get("fault", "").strip()

    q = db.select(MechanicEvent)
    events = db.session.scalars(q).all()

    # تبدیل تاریخ به سال شمسی برای فیلتر و نمودار
    def jyear(d):
        if not d:
            return None
        try:
            return jdatetime.date.fromgregorian(date=d).year
        except Exception:
            return None

    # اعمال فیلترها
    rows = []
    for e in events:
        jy = jyear(e.op_date)
        if f_year and str(jy) != f_year:
            continue
        if f_optype and (e.op_type or "") != f_optype:
            continue
        if f_fault and (e.fault_type or "") != f_fault:
            continue
        rows.append((e, jy))

    # آمار کلی
    total = len(rows)
    by_optype = Counter((e.op_type or "نامشخص") for e, _ in rows)
    by_fault = Counter(classify_fault(e) for e, _ in rows)
    by_year = Counter(jy for _, jy in rows if jy)
    by_contractor = Counter((e.contractor or "نامشخص") for e, _ in rows if e.contractor)
    by_stage = Counter((e.stage or "نامشخص") for e, _ in rows)

    # شاخص‌های بیشتر
    with_pt = sum(1 for e, _ in rows if e.pt_flow is not None)
    by_center = Counter()
    by_month = Counter()
    burnt = 0
    depths = []
    heads = []
    flows = []
    months_worked_list = []
    import jdatetime as _jd
    for e, jy in rows:
        # مرکز از توضیحات یا شرح در دسترس نیست؛ از fault_description رد می‌شویم
        # ماه شمسی برای روند ماهانه
        if e.op_date:
            try:
                jm = _jd.date.fromgregorian(date=e.op_date).month
                by_month[jm] += 1
            except Exception:
                pass
        f = classify_fault(e)
        if f == "سوختگی":
            burnt += 1
        if e.well_depth:
            depths.append(e.well_depth)
        if e.total_head:
            heads.append(e.total_head)
        if e.pt_flow:
            flows.append(e.pt_flow)
        if e.months_worked:
            months_worked_list.append(e.months_worked)

    def avg(lst):
        return round(sum(lst) / len(lst), 1) if lst else 0

    MONTHS_FA = ["", "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
                 "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
    by_month_named = {MONTHS_FA[m]: c for m, c in sorted(by_month.items())}

    stats = {
        "total": total,
        "install": by_optype.get("نصب", 0),
        "pull": by_optype.get("کشیدن", 0),
        "collect": by_optype.get("جمع‌آوری", 0) + by_optype.get("جمع آوری", 0),
        "with_pt": with_pt,
        "burnt": burnt,
        "burnt_pct": round(100 * burnt / total) if total else 0,
        "avg_depth": avg(depths),
        "avg_head": avg(heads),
        "avg_flow": avg(flows),
        "avg_months_worked": avg(months_worked_list),
        "by_optype": dict(by_optype),
        "by_fault": dict(by_fault),
        "by_year": dict(sorted(by_year.items())),
        "by_month": by_month_named,
        "by_contractor": by_contractor.most_common(10),
        "by_stage": dict(by_stage),
    }

    # گزینه‌های فیلتر
    all_years = sorted({jyear(e.op_date) for e in events if jyear(e.op_date)}, reverse=True)
    all_optypes = sorted({e.op_type for e in events if e.op_type})

    return render_template("mechanic/reports.html", stats=stats,
                           all_years=all_years, all_optypes=all_optypes,
                           fault_types=FAULT_TYPES,
                           cur={"year": f_year, "optype": f_optype, "fault": f_fault})
# ---------- گزارش انبار قطعات ----------
@bp.route("/parts/report")
@permission_required("mechanic", "view")
def parts_report():
    parts = db.session.scalars(db.select(MechanicPart).order_by(
        MechanicPart.equipment, MechanicPart.id)).all()

    total_usable = sum(p.usable or 0 for p in parts)
    total_scrap = sum(p.scrap or 0 for p in parts)
    total_new = sum(p.new_count or 0 for p in parts)
    total_repaired = sum(p.repaired or 0 for p in parts)

    # قطعات کم‌موجود: قابل‌استفاده صفر یا کمتر از ۳
    low_stock = [p for p in parts if (p.usable or 0) < 3]
    # قطعات بدون هیچ موجودی
    empty = [p for p in parts if (p.total_count or 0) == 0]

    by_eq = {"موتور": {"usable": 0, "scrap": 0}, "پمپ": {"usable": 0, "scrap": 0}}
    for p in parts:
        if p.equipment in by_eq:
            by_eq[p.equipment]["usable"] += p.usable or 0
            by_eq[p.equipment]["scrap"] += p.scrap or 0

    stats = {
        "total_parts": len(parts),
        "total_usable": total_usable,
        "total_scrap": total_scrap,
        "total_new": total_new,
        "total_repaired": total_repaired,
        "low_stock": low_stock,
        "empty_count": len(empty),
        "by_eq": by_eq,
    }
    parts_data = [{
        "equipment": p.equipment or "نامشخص",
        "name": p.part_name or "—",
        "total": p.total_count or 0,
        "new": p.new_count or 0,
        "repaired": p.repaired or 0,
        "scrap": p.scrap or 0,
        "usable": p.usable or 0,
        "installed": getattr(p, "installed_count", 0) or 0,
    } for p in parts]
    return render_template("mechanic/parts_report.html", stats=stats, parts_data=parts_data)


################################################################################
# FILE: operation\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("operation", __name__)

from app.blueprints.operation import routes  # noqa: E402,F401



################################################################################
# FILE: operation\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.operation import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.utils.dates import to_english_digits
from app.models.well import Well
from app.models.flow import FlowTest, FlowTestPoint
from app.models.audit import RecordHistory
from app.models.constants import OPERATING_TYPES, FLOW_TEST_REASONS
from app import workflow

POINT_ROWS = 5
POINT_FIELDS = ["operating_type", "discharge_lps", "head_m", "drawdown_m",
                "dynamic_level_m", "pressure_atm", "amperes", "efficiency"]


class FlowTestForm(FlaskForm):
    test_date = JalaliDateField("تاریخ آزمایش", validators=[Optional()])
    test_reason = SelectField("دلیل آزمایش", choices=[("", "—")] + [(r, r) for r in FLOW_TEST_REASONS], validators=[Optional()])
    network_type = StringField("نوع شبکه", validators=[Optional()])
    electropump_type = StringField("تیپ الکتروپمپ", validators=[Optional()])
    electropump_type_prev = StringField("تیپ الکتروپمپ قبلی", validators=[Optional()])
    install_date = JalaliDateField("تاریخ آخرین نصب", validators=[Optional()])
    install_depth = FloatField("عمق نصب", validators=[Optional()])
    well_depth = FloatField("عمق چاه", validators=[Optional()])
    construction_type = StringField("نوع چاه (سیمانته/غیرسیمانته)", validators=[Optional()])
    allowed_q = FloatField("دبی مجاز", validators=[Optional()])
    design_q = FloatField("دبی طراحی", validators=[Optional()])
    license_q = FloatField("دبی پروانه", validators=[Optional()])
    power_subscription = StringField("اشتراک برق", validators=[Optional()])
    static_level = FloatField("سطح ایستایی", validators=[Optional()])
    last_rehab_date = JalaliDateField("تاریخ آخرین بهسازی", validators=[Optional()])
    pull_reason = StringField("علت کشیدن پمپ", validators=[Optional()])
    meter_status = StringField("کالیبراسیون/وضعیت کنتور", validators=[Optional()])
    meter_brand = StringField("برند/سایز کنتور", validators=[Optional()])
    starter_type = StringField("سیستم راه‌انداز", validators=[Optional()])
    capacitor_capacity = FloatField("ظرفیت خازن", validators=[Optional()])
    voltage_on = StringField("ولتاژ روشن", validators=[Optional()])
    voltage_off = StringField("ولتاژ خاموش", validators=[Optional()])
    ohm_ff = StringField("مقاومت اهمی ف-ف", validators=[Optional()])
    ohm_fg = StringField("مقاومت اهمی ف-ب", validators=[Optional()])
    line_pressure = FloatField("فشار خط (bar)", validators=[Optional()])
    regulated_pressure = StringField("فشار تنظیمی", validators=[Optional()])
    discharge_volume_m3 = FloatField("حجم تخلیه (m³)", validators=[Optional()])
    expert_note = TextAreaField("نظر کارشناس", validators=[Optional()])
    submit = SubmitField("ذخیره")


HEADER_FIELDS = [
    "test_date", "test_reason", "network_type", "electropump_type", "electropump_type_prev",
    "install_date", "install_depth", "well_depth", "construction_type", "allowed_q",
    "design_q", "license_q", "power_subscription", "static_level", "last_rehab_date",
    "pull_reason", "meter_status", "meter_brand", "starter_type", "capacitor_capacity",
    "voltage_on", "voltage_off", "ohm_ff", "ohm_fg", "line_pressure",
    "regulated_pressure", "discharge_volume_m3", "expert_note",
]


def _num(s):
    s = to_english_digits((s or "").strip())
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _apply_header(form, rec):
    for f in HEADER_FIELDS:
        setattr(rec, f, getattr(form, f).data)


def _apply_points(rec):
    """Rebuild child points from the manual table in request.form."""
    rec.points.clear()
    for i in range(1, POINT_ROWS + 1):
        vals = {f: request.form.get(f"pt-{i}-{f}", "").strip() for f in POINT_FIELDS}
        if not any(vals.values()):
            continue
        rec.points.append(FlowTestPoint(
            operating_no=i,
            operating_type=vals["operating_type"] or None,
            discharge_lps=_num(vals["discharge_lps"]),
            head_m=_num(vals["head_m"]),
            drawdown_m=_num(vals["drawdown_m"]),
            dynamic_level_m=_num(vals["dynamic_level_m"]),
            pressure_atm=_num(vals["pressure_atm"]),
            amperes=vals["amperes"] or None,
            efficiency=_num(vals["efficiency"]),
        ))


@bp.route("/wells/<int:well_id>/flow/new", methods=["GET", "POST"])
@permission_required("operation", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = FlowTestForm()
    if form.validate_on_submit():
        rec = FlowTest(well_id=well.id, created_by_id=current_user.id)
        _apply_header(form, rec)
        db.session.add(rec)
        _apply_points(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("دبی‌سنجی ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("operation.detail", record_id=rec.id))
    return render_template("operation/form.html", form=form, well=well,
                           title="ثبت دبی‌سنجی", points=[], operating_types=OPERATING_TYPES,
                           point_rows=POINT_ROWS, point_fields=POINT_FIELDS)


@bp.route("/flow/<int:record_id>")
@permission_required("operation", "view")
def detail(record_id):
    rec = db.get_or_404(FlowTest, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="flow_tests", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("operation/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/flow/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("operation", "edit")
def edit(record_id):
    rec = db.get_or_404(FlowTest, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("operation.detail", record_id=rec.id))
    form = FlowTestForm(obj=rec)
    if form.validate_on_submit():
        _apply_header(form, rec)
        _apply_points(rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("دبی‌سنجی به‌روزرسانی شد.", "success")
        return redirect(url_for("operation.detail", record_id=rec.id))
    return render_template("operation/form.html", form=form, well=rec.well,
                           title="ویرایش دبی‌سنجی", points=rec.points,
                           operating_types=OPERATING_TYPES, point_rows=POINT_ROWS,
                           point_fields=POINT_FIELDS)


@bp.route("/flow/<int:record_id>/delete", methods=["POST"])
@permission_required("operation", "delete")
def delete(record_id):
    rec = db.get_or_404(FlowTest, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("دبی‌سنجی حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(FlowTest, record_id)
    if not current_user.has_permission("operation", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("operation.detail", record_id=record_id))


@bp.route("/flow/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/flow/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/flow/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/flow/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: orgs\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("orgs", __name__, url_prefix="/orgs")

from app.blueprints.orgs import routes  # noqa: E402,F401



################################################################################
# FILE: orgs\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, SubmitField
from wtforms.validators import DataRequired, Optional

from app.extensions import db
from app.blueprints.orgs import bp
from app.security import permission_required
from app.models.org import OrgUnit
from app.models.constants import ORG_UNIT_TYPES


class OrgUnitForm(FlaskForm):
    name = StringField("نام", validators=[DataRequired()])
    code = StringField("کد", validators=[Optional()])
    unit_type = SelectField("نوع", choices=ORG_UNIT_TYPES, validators=[DataRequired()])
    parent_id = SelectField("واحد بالادست", coerce=int, validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_parents(self, exclude_id=None):
        units = db.session.scalars(db.select(OrgUnit).order_by(OrgUnit.name)).all()
        choices = [(0, "— بدون والد —")]
        for u in units:
            if exclude_id and u.id == exclude_id:
                continue
            choices.append((u.id, f"{u.name} ({dict(ORG_UNIT_TYPES).get(u.unit_type, '')})"))
        self.parent_id.choices = choices


@bp.route("/")
@permission_required("orgs", "view")
def list_orgs():
    units = db.session.scalars(
        db.select(OrgUnit).order_by(OrgUnit.unit_type, OrgUnit.name)
    ).all()
    type_labels = dict(ORG_UNIT_TYPES)
    return render_template("orgs/list.html", units=units, type_labels=type_labels)


@bp.route("/new", methods=["GET", "POST"])
@permission_required("orgs", "create")
def create_org():
    form = OrgUnitForm()
    form.populate_parents()
    if form.validate_on_submit():
        unit = OrgUnit(
            name=form.name.data.strip(),
            code=(form.code.data or "").strip() or None,
            unit_type=form.unit_type.data,
            parent_id=form.parent_id.data or None,
        )
        db.session.add(unit)
        db.session.commit()
        flash("واحد سازمانی ایجاد شد.", "success")
        return redirect(url_for("orgs.list_orgs"))
    return render_template("orgs/form.html", form=form, title="واحد سازمانی جدید")


@bp.route("/<int:unit_id>/edit", methods=["GET", "POST"])
@permission_required("orgs", "edit")
def edit_org(unit_id):
    unit = db.get_or_404(OrgUnit, unit_id)
    form = OrgUnitForm(obj=unit)
    form.populate_parents(exclude_id=unit.id)
    if form.validate_on_submit():
        unit.name = form.name.data.strip()
        unit.code = (form.code.data or "").strip() or None
        unit.unit_type = form.unit_type.data
        unit.parent_id = form.parent_id.data or None
        db.session.commit()
        flash("واحد سازمانی به‌روزرسانی شد.", "success")
        return redirect(url_for("orgs.list_orgs"))
    return render_template("orgs/form.html", form=form, title="ویرایش واحد سازمانی")


@bp.route("/<int:unit_id>/delete", methods=["POST"])
@permission_required("orgs", "delete")
def delete_org(unit_id):
    unit = db.get_or_404(OrgUnit, unit_id)
    if unit.children:
        flash("این واحد دارای زیرمجموعه است و قابل حذف نیست.", "danger")
        return redirect(url_for("orgs.list_orgs"))
    db.session.delete(unit)
    db.session.commit()
    flash("واحد سازمانی حذف شد.", "info")
    return redirect(url_for("orgs.list_orgs"))



################################################################################
# FILE: permit\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("permit", __name__, url_prefix="/permits")

from app.blueprints.permit import routes  # noqa: E402,F401



################################################################################
# FILE: permit\routes.py
################################################################################

"""وضعیت پروانه چاه‌ها.

صفحه‌ی تحلیل وضعیت پروانه بر پایه‌ی جدول production_trend (ستون‌های well_permit و
permit_expiry) و فرم ورود/ویرایش اطلاعات پروانه‌ی هر چاه (به‌روزرسانی مستقیم همان
جدول). تاریخ اعتبار به‌صورت شمسیِ فشرده (مثل 14050427) است و وضعیت انقضا با مقایسه
با تاریخ امروز محاسبه می‌شود.
"""
from datetime import date
from collections import Counter

import jdatetime
from flask import render_template, request, redirect, url_for, flash
from flask_login import login_required
from sqlalchemy import text

from app.extensions import db
from app.blueprints.permit import bp
from app.models.well import Well
from app.imports.wells_import import match_key


def _digits8(s):
    if not s:
        return None
    d = "".join(ch for ch in str(s) if ch.isdigit())
    return d if len(d) == 8 else None


def _today8():
    t = jdatetime.date.fromgregorian(date=date.today())
    return f"{t.year:04d}{t.month:02d}{t.day:02d}"


def _expiry_state(expiry, today8):
    e = _digits8(expiry)
    if not e:
        return "unknown"
    if e < today8:
        return "expired"
    ey, em = int(e[:4]), int(e[4:6])
    ty, tm = int(today8[:4]), int(today8[4:6])
    months_left = (ey - ty) * 12 + (em - tm)
    return "near" if months_left < 3 else "valid"


STATE_LABEL = {"expired": "منقضی شده", "near": "نزدیک انقضا",
               "valid": "معتبر", "unknown": "نامشخص"}


@bp.route("/")
@login_required
def index():
    today8 = _today8()
    # نگاشت نام چاه به شناسه‌ی چاه سامانه (برای لینک به جزئیات چاه)
    well_by_key = {}
    for w in db.session.scalars(db.select(Well)).all():
        well_by_key.setdefault(match_key(w.name), w.id)
    result = db.session.execute(text(
        "SELECT rowid, well_name, department, well_permit, permit_expiry, "
        "well_type, main_zone FROM production_trend"
    ))
    rows = []
    state_counter = Counter()
    type_counter = Counter()
    for r in result.fetchall():
        rowid, name, dept, permit, expiry, wtype, zone = r
        if not name or not str(name).strip():
            continue
        st = _expiry_state(expiry, today8)
        state_counter[st] += 1
        if permit and str(permit).strip() and str(permit).strip() != "0":
            type_counter[str(permit).strip()] += 1
            wid = well_by_key.get(match_key(name))
        rows.append({
            "well_id": wid, "rowid": rowid, "well_name": name, "office": dept or "—",
            "permit_type": (permit if permit and str(permit) != "0" else "—"),
            "expiry": expiry or "—", "state": st, "state_label": STATE_LABEL[st],
            "well_type": wtype or "—", "zone": zone or "—",
        })
    order = {"expired": 0, "near": 1, "valid": 2, "unknown": 3}
    rows.sort(key=lambda r: (order.get(r["state"], 9), str(r["expiry"])))

    stats = {
        "total": len(rows),
        "expired": state_counter.get("expired", 0),
        "near": state_counter.get("near", 0),
        "valid": state_counter.get("valid", 0),
        "unknown": state_counter.get("unknown", 0),
        "types": type_counter.most_common(),
    }
    return render_template("permit/index.html", rows=rows, stats=stats)


@bp.route("/<int:rowid>/edit", methods=["GET", "POST"])
@login_required
def edit(rowid):
    row = db.session.execute(
        text("SELECT rowid, well_name, department, well_permit, permit_expiry "
             "FROM production_trend WHERE rowid=:i"), {"i": rowid}
    ).fetchone()
    if not row:
        flash("چاه پیدا نشد.", "warning")
        return redirect(url_for("permit.index"))

    if request.method == "POST":
        permit = request.form.get("well_permit", "").strip()
        expiry = request.form.get("permit_expiry", "").strip()
        db.session.execute(
            text("UPDATE production_trend SET well_permit=:p, permit_expiry=:e "
                 "WHERE rowid=:i"),
            {"p": permit or None, "e": expiry or None, "i": rowid},
        )
        db.session.commit()
        flash("وضعیت پروانه به‌روزرسانی شد.", "success")
        return redirect(url_for("permit.index"))

    rec = {"rowid": row[0], "well_name": row[1], "office": row[2],
           "well_permit": row[3] or "", "permit_expiry": row[4] or ""}
    return render_template("permit/form.html", rec=rec)



################################################################################
# FILE: permit_event\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("permit_event", __name__, url_prefix="/permit-event")

from app.blueprints.permit_event import routes  # noqa: E402,F401


################################################################################
# FILE: permit_event\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, BooleanField, SubmitField
from wtforms.validators import Optional, ValidationError

from app.extensions import db
from app.blueprints.permit_event import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.permit_event import WellPermitEvent
from app.models.audit import RecordHistory
from app import workflow


PERMIT_TYPES = [
    ("", "—"),
    ("بهره‌برداری عادی", "بهره‌برداری عادی"),
    ("حفر", "حفر"),
    ("تغییر محل", "تغییر محل"),
    ("کف‌شکنی", "کف‌شکنی"),
]


class PermitEventForm(FlaskForm):
    permit_type = SelectField("نوع پروانه", choices=PERMIT_TYPES, validators=[Optional()])
    permit_code = StringField("کد آخرین پروانه", validators=[Optional()])
    permit_no = StringField("شماره پروانه", validators=[Optional()])
    permit_date = JalaliDateField("تاریخ صدور", validators=[Optional()])
    expiry_date = JalaliDateField("تاریخ اعتبار/انقضا", validators=[Optional()])
    case_status = StringField("وضعیت پرونده", validators=[Optional()])
    klasse = StringField("کلاسه آب منطقه‌ای", validators=[Optional()])
    request_type = StringField("درخواست جاری (تمدید/صدور/جابجایی)", validators=[Optional()])
    followup_stage = StringField("مرحله پیگیری", validators=[Optional()])
    cost_paid = BooleanField("هزینه پرداخت شده", validators=[Optional()])
    expiry_penalty_rial = FloatField("جریمه انقضا (ریال)", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def validate_expiry_date(self, field):
        if field.data and self.permit_date.data and field.data < self.permit_date.data:
            raise ValidationError("تاریخ اعتبار نمی‌تواند پیش از تاریخ صدور باشد.")

    def validate_expiry_penalty_rial(self, field):
        if field.data is not None and field.data < 0:
            raise ValidationError("جریمه نمی‌تواند منفی باشد.")


SIMPLE = ["permit_type", "permit_code", "permit_no", "permit_date", "expiry_date",
          "case_status", "klasse", "request_type", "followup_stage", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data or None)
    rec.cost_paid = form.cost_paid.data
    rec.expiry_penalty_rial = form.expiry_penalty_rial.data


@bp.route("/wells/<int:well_id>/permit/new", methods=["GET", "POST"])
@permission_required("permit", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = PermitEventForm()
    if form.validate_on_submit():
        rec = WellPermitEvent(well_id=well.id, created_by_id=current_user.id, source="manual")
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("پروانه ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("permit_event.detail", record_id=rec.id))
    return render_template("permit_event/form.html", form=form, well=well, title="ثبت پروانه")


@bp.route("/permit/<int:record_id>")
@permission_required("permit", "view")
def detail(record_id):
    rec = db.get_or_404(WellPermitEvent, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="well_permit_events", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("permit_event/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/permit/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("permit", "edit")
def edit(record_id):
    rec = db.get_or_404(WellPermitEvent, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("permit_event.detail", record_id=rec.id))
    form = PermitEventForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("پروانه به‌روزرسانی شد.", "success")
        return redirect(url_for("permit_event.detail", record_id=rec.id))
    return render_template("permit_event/form.html", form=form, well=rec.well, title="ویرایش پروانه")


@bp.route("/permit/<int:record_id>/delete", methods=["POST"])
@permission_required("permit", "delete")
def delete(record_id):
    rec = db.get_or_404(WellPermitEvent, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد پروانه حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(WellPermitEvent, record_id)
    if not current_user.has_permission("permit", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("permit_event.detail", record_id=record_id))


@bp.route("/permit/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/permit/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/permit/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/permit/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")


################################################################################
# FILE: printing\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("printing", __name__, url_prefix="/print")

from app.blueprints.printing import routes  # noqa: E402,F401



################################################################################
# FILE: printing\routes.py
################################################################################

"""Official print-ready reports (browser Print-to-PDF, RTL/Persian, zero deps)."""
from flask import render_template, url_for

from app.extensions import db
from app.blueprints.printing import bp
from app.security import permission_required
from app.models.well import Well
from app.models.constants import WELL_KINDS, WELL_STATUSES
from app.utils.dates import format_jalali as fj


def _today():
    import jdatetime
    return jdatetime.date.today().strftime("%Y/%m/%d")


@bp.route("/well/<int:well_id>")
@permission_required("wells", "view")
def well_sheet(well_id):
    w = db.get_or_404(Well, well_id)
    from app.models.drilling import Drilling
    from app.models.pump_asset import PumpInstallation
    from app.models.baseline import WellBaseline
    dr = db.session.scalar(db.select(Drilling).filter_by(well_id=w.id)
                           .order_by(Drilling.end_date.desc()))
    inst = db.session.scalar(db.select(PumpInstallation).filter_by(well_id=w.id)
                             .order_by(PumpInstallation.install_date.desc()))
    base = db.session.scalar(db.select(WellBaseline).filter_by(well_id=w.id, is_current=True))
    return render_template(
        "print/well_sheet.html", w=w, tech=w.technical, dr=dr, inst=inst, base=base,
        kind_labels=dict(WELL_KINDS), status_labels=dict(WELL_STATUSES),
        today=_today(), doc_id=w.display_pm, fj=fj,
        back_url=url_for("wells.detail", well_id=w.id))


@bp.route("/pump-test/<int:record_id>")
@permission_required("pump_test", "view")
def pump_test(record_id):
    from app.models.pump_test import PumpTest
    rec = db.get_or_404(PumpTest, record_id)
    return render_template("print/pump_test.html", rec=rec, w=rec.well,
                           today=_today(), doc_id=rec.well.display_pm, fj=fj,
                           back_url=url_for("pump_test.detail", record_id=rec.id))


@bp.route("/flow/<int:record_id>")
@permission_required("operation", "view")
def flow(record_id):
    from app.models.flow import FlowTest
    rec = db.get_or_404(FlowTest, record_id)
    return render_template("print/flow.html", rec=rec, w=rec.well,
                           today=_today(), doc_id=rec.well.display_pm, fj=fj,
                           back_url=url_for("operation.detail", record_id=rec.id))



################################################################################
# FILE: production_trend\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("production_trend", __name__, url_prefix="/production-trend")

from app.blueprints.production_trend import routes  # noqa: E402,F401



################################################################################
# FILE: production_trend\column_import.py
################################################################################

"""افزودن سرستون جدید به جدول production_trend (روش A: ستون واقعی).

کاربر دسته/سال/ماه و فایل اکسل می‌دهد؛ سیستم نام ستون انگلیسی را طبق قرارداد
data_map می‌سازد، ستون را (در صورت نبود) به جدول اضافه می‌کند، و مقادیر را با
تطبیق نام چاه پر می‌کند. نام ستون‌ها اعتبارسنجی می‌شوند تا امن باشند.
"""
import re
import openpyxl
from sqlalchemy import text

from app.extensions import db
from app.imports.wells_import import match_key
from app.blueprints.production_trend.data_map import MONTHS_EN, MONTHS_FA
def _recalc_yearly(category, year):
    """ستون مجموع/میانگین سالانه را از روی ماه‌های موجود همان سال بازمحاسبه می‌کند."""
    cols_now = _existing_columns()
    # ماه‌های موجود این دسته/سال را جمع کن
    month_cols = []
    for mi in range(12):
        mc = make_column_name(category, year, mi)
        if mc in cols_now:
            month_cols.append(mc)
    if not month_cols:
        return

    # نام ستون سالانه
    year_col = make_column_name(category, year, None)
    if not year_col:
        return
    # اگر ستون سالانه نبود بساز
    if year_col not in cols_now:
        sqltype = CATEGORY_SQLTYPE.get(category, "REAL")
        db.session.execute(text(
            f'ALTER TABLE production_trend ADD COLUMN "{year_col}" {sqltype}'))
        db.session.commit()

    # تولید و کارکرد → جمع؛ فشار و دبی → میانگین
    is_sum = category in ("production", "runtime")
    sum_expr = " + ".join(f'COALESCE("{c}",0)' for c in month_cols)
    cnt_expr = " + ".join(f'(CASE WHEN "{c}" IS NOT NULL THEN 1 ELSE 0 END)' for c in month_cols)

    rows = db.session.execute(text(f'SELECT rowid, {sum_expr}, {cnt_expr} FROM production_trend')).fetchall()
    for rid, total, cnt in rows:
        if cnt and cnt > 0:
            val = total if is_sum else round(total / cnt, 2)
        else:
            val = None
        db.session.execute(
            text(f'UPDATE production_trend SET "{year_col}"=:v WHERE rowid=:i'),
            {"v": val, "i": rid})
    db.session.commit()

# دسته‌های پشتیبانی‌شده و الگوی نام ستون
CATEGORIES = {
    "production": "تولید",
    "average_flow": "دبی متوسط",
    "runtime": "کارکرد",
    "well_pressure": "فشار",
}

# نوع داده‌ی هر دسته در SQLite
CATEGORY_SQLTYPE = {
    "production": "INTEGER",
    "average_flow": "REAL",
    "runtime": "INTEGER",
    "well_pressure": "REAL",
}

SAFE_NAME = re.compile(r"^[a-z][a-z0-9_]*$")


def make_column_name(category, year, month_idx):
    """نام ستون انگلیسی طبق قرارداد data_map. month_idx: 0..11 یا None برای سالانه."""
    if month_idx is None:
        # ستون سالانه
        if category == "production":
            return f"production_year_{year}"
        if category == "average_flow":
            return f"average_flow_year_{year}"
        if category == "runtime":
            return f"runtime_year_{year}"
        if category == "well_pressure":
            return f"pressure_year_{year}"
        return None
    m = MONTHS_EN[month_idx]
    if category == "production":
        return f"production_{year}_{m}"
    if category == "average_flow":
        return f"average_flow_{year}_{m}"
    if category == "runtime":
        return f"runtime_{m}_{year}"
    if category == "well_pressure":
        return f"well_pressure_{year}_{m}"
    return None


def _existing_columns():
    rows = db.session.execute(text("PRAGMA table_info(production_trend)")).fetchall()
    return {r[1] for r in rows}


def _find_key_and_value_columns(header):
    """ستون نام چاه و ستون داده را در هدر اکسل تشخیص می‌دهد."""
    name_col = None
    for i, h in enumerate(header):
        if h and any(k in str(h) for k in ("نام چاه", "نام", "چاه")):
            name_col = i
            break
    # ستون داده: اولین ستون عددیِ غیر از ستون نام (ساده: دومین ستون پرشده)
    value_col = None
    for i, h in enumerate(header):
        if i != name_col and h not in (None, ""):
            value_col = i
            break
    return name_col, value_col


def run(path, category, year, month_idx):
    """اکسل را می‌خواند و ستون جدید را می‌سازد/پر می‌کند. آمار برمی‌گرداند."""
    if category not in CATEGORIES:
        raise ValueError("دسته‌ی نامعتبر.")
    col = make_column_name(category, year, month_idx)
    if not col or not SAFE_NAME.match(col):
        raise ValueError(f"نام ستون نامعتبر ساخته شد: {col}")

    sqltype = CATEGORY_SQLTYPE[category]
    stats = {"column": col, "added_column": False, "rows": 0,
             "matched": 0, "unmatched": 0}

    # ۱) افزودن ستون اگر نبود
    if col not in _existing_columns():
        db.session.execute(text(f'ALTER TABLE production_trend ADD COLUMN "{col}" {sqltype}'))
        db.session.commit()
        stats["added_column"] = True

    # ۲) خواندن اکسل
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("فایل اکسل خالی است.")
    header = rows[0]
    name_col, value_col = _find_key_and_value_columns(header)
    if name_col is None or value_col is None:
        raise ValueError("ستون نام چاه یا ستون داده در اکسل پیدا نشد.")

    # ۳) نگاشت نام چاه production_trend → rowid
    pt = db.session.execute(text("SELECT rowid, well_name FROM production_trend")).fetchall()
    by_key = {}
    for rowid, wname in pt:
        by_key.setdefault(match_key(wname), rowid)

    # ۴) پر کردن مقادیر
    for row in rows[1:]:
        if name_col >= len(row):
            continue
        name = row[name_col]
        if not name or not str(name).strip():
            continue
        stats["rows"] += 1
        val = row[value_col] if value_col < len(row) else None
        rid = by_key.get(match_key(name))
        if rid is None:
            stats["unmatched"] += 1
            continue
        db.session.execute(
            text(f'UPDATE production_trend SET "{col}"=:v WHERE rowid=:i'),
            {"v": val, "i": rid},
        )
        stats["matched"] += 1
# محاسبه‌ی خودکار دبی متوسط اگر تولید و کارکرد همان ماه موجود باشند
    if category in ("production", "runtime") and month_idx is not None:
        prod_col = make_column_name("production", year, month_idx)
        run_col = make_column_name("runtime", year, month_idx)
        flow_col = make_column_name("average_flow", year, month_idx)
        cols_now = _existing_columns()
        # فقط اگر هر دو ستون تولید و کارکرد وجود دارند
        if prod_col in cols_now and run_col in cols_now:
            if flow_col not in cols_now:
                db.session.execute(text(
                    f'ALTER TABLE production_trend ADD COLUMN "{flow_col}" REAL'))
                db.session.commit()
            # برای هر چاه دبی را حساب کن: تولید×۱۰۰۰ ÷ (کارکرد×۳۶۰۰)
            rows_calc = db.session.execute(text(
                f'SELECT rowid, "{prod_col}", "{run_col}" FROM production_trend')).fetchall()
            flow_filled = 0
            for rid, prod_v, run_v in rows_calc:
                try:
                    p = float(prod_v) if prod_v not in (None, "") else None
                    h = float(run_v) if run_v not in (None, "") else None
                except (ValueError, TypeError):
                    p, h = None, None
                if p is None or h is None or h <= 0:
                    continue
                flow = round(p * 1000.0 / (h * 3600.0), 2)
                if flow < 0 or flow > 1000:  # اعتبارسنجی: مقدار غیرمنطقی رد شود
                    continue
                db.session.execute(
                    text(f'UPDATE production_trend SET "{flow_col}"=:v WHERE rowid=:i'),
                    {"v": flow, "i": rid})
                flow_filled += 1
            db.session.commit()
            stats["flow_calculated"] = flow_filled
            # به‌روزرسانی ستون مجموع/میانگین سالانه بعد از افزودن ماه جدید
    if month_idx is not None:
        _recalc_yearly(category, year)

    db.session.commit()
    return stats

    db.session.commit()
    return stats


################################################################################
# FILE: production_trend\data_map.py
################################################################################

"""نگاشت ستون‌های انگلیسیِ جدول production_trend به کلیدهای فارسیِ اکسلِ داشبورد.

هدف: داشبورد قدیمی (script.js) داده را با کلیدهای فارسی مثل «تولید سال ۱۴۰۴» یا
«دبی متوسط ۱۴۰۴(فروردین)» می‌خواند. این ماژول هر ردیف جدول را به یک dict با همان
کلیدهای فارسی برمی‌گرداند تا کل منطق داشبورد بدون تغییر کار کند.

ستون‌های سری‌زمانی با الگو ساخته می‌شوند؛ ستون‌های ثابت نگاشت دستی دارند.
"""

# ماه‌های شمسی به ترتیب (۱..۱۲) و معادل انگلیسیِ به‌کاررفته در نام ستون‌ها
MONTHS_FA = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
             "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
MONTHS_EN = ["farvardin", "ordibehesht", "khordad", "tir", "mordad", "shahrivar",
             "mehr", "aban", "azar", "dey", "bahman", "esfand"]
YEARS = [1399, 1400, 1401, 1402, 1403, 1404, 1405]

# ---------- نگاشت ستون‌های ثابت: انگلیسی → فارسیِ اکسل ----------
STATIC_MAP = {
    "well_name": "نام چاه",
    "department": "اداره",
    "production_zone_name": "نام پهنه تولید",
    "low_runtime_reason": "علت کارکرد کم",
    "electropump_installation_date": "تاریخ نصب الکتروپمپ",
    "observed_fault_last_rehabilitation": "خرابی مشاهده شده آخرین بهسازی",
    "last_rehabilitation_date": "تاریخ آخرین بهسازی",
    "permit_expiry": "اعتبار پروانه",
    "well_permit": "پروانه چاه",
    "well_type": "نوع چاه",
    "well_location_status": "وضعیت تعیین محل چاه",
    "proposed_permitted_flow_lps": "دبی مجاز پیشنهادی",
    "test_end_date": "تاریخ پایان آزمایش",
    "contractor": "پیمانکار",
    "drilling_year": "سال حفر",
    "max_flow_rate": "حداکثر آبدهی",
    "drawdown_amount": "مقدار افت",
    "static_level": "سطح استاتیک",
    "dynamic_level": "سطح دینامیک",
    "main_zone": "پهنه اصلی",
    "sub_zone": "زیر پهنه",
    "well_status": "وضعیت چاه",
    "full_cycle_count": "تعداد دوره کامل",
    "full_average_months": "میانگین کامل (ماه)",
    "total_average_with_open_installation_months": "میانگین کل با نصب باز (ماه)",
    "contractors": "پیمانکار ها",
}

# نگاشت پسوندهای دوره‌ی نصب/کشیدن (۱..۵ و «باز»)
CYCLE_BASE = {
    "installation": "نصب",
    "motor": "موتور",
    "motor_new_repaired": "موتور نو/تعمیری",
    "pump": "پمپ",
    "manufacturer": "سازنده",
    "pump_new_repaired": "پمپ نو/تعمیری",
    "stage": "طبقه",
    "depth": "عمق",
    "pulling": "کشیدن",
    "interval": None,  # ویژه (interval_1_months → فاصله۱ (ماه))
}


def _build_static_full():
    """نگاشت کامل ستون‌های ثابت شامل ستون‌های دوره‌ای (۱..۵ و open)."""
    m = dict(STATIC_MAP)
    suffixes = [("1", "1"), ("2", "2"), ("3", "3"), ("4", "4"), ("5", "5"),
                ("open", " (باز)")]
    for en_sfx, fa_sfx in suffixes:
        m[f"installation_{en_sfx}"] = f"نصب{fa_sfx}"
        m[f"motor_{en_sfx}"] = f"موتور{fa_sfx}"
        m[f"motor_new_repaired_{en_sfx}"] = f"موتور نو/تعمیری{fa_sfx}"
        m[f"pump_{en_sfx}"] = f"پمپ{fa_sfx}"
        m[f"manufacturer_{en_sfx}"] = f"سازنده{fa_sfx}"
        m[f"pump_new_repaired_{en_sfx}"] = f"پمپ نو/تعمیری{fa_sfx}"
        m[f"stage_{en_sfx}"] = f"طبقه{fa_sfx}"
        m[f"depth_{en_sfx}"] = f"عمق{fa_sfx}"
        if en_sfx != "open":
            m[f"pulling_{en_sfx}"] = f"کشیدن{fa_sfx}"
            m[f"interval_{en_sfx}_months"] = f"فاصله{fa_sfx} (ماه)"
    return m


STATIC_FULL = _build_static_full()


def _ts_map_for_column(col):
    """اگر ستون سری‌زمانی بود، کلید فارسی معادل را برمی‌گرداند، وگرنه None."""
    for y in YEARS:
        ys = str(y)
        # --- تولید ---
        if col == f"production_year_{ys}":
            return f"تولید سال {y}"
        for i, (men, mfa) in enumerate(zip(MONTHS_EN, MONTHS_FA)):
            if col == f"production_{ys}_{men}":
                return f"تولید {y}({mfa})"
        # --- کارکرد (runtime) ---
        if col == f"runtime_year_{ys}":
            return f"کارکرد سال {y}"
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"runtime_{men}_{ys}":
                return f"کارکرد {mfa} {y}"
        # --- دبی متوسط (average_flow) ---
        if col == f"average_flow_year_{ys}":
            return f"دبی متوسط سال {y}"
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"average_flow_{ys}_{men}":
                return f"دبی متوسط {y}({mfa})"
        # --- فشار (well_pressure) ---
        if col == f"pressure_year_{ys}":
            return f"فشار سال {y}"
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"well_pressure_{ys}_{men}":
                return f"فشار چاه {y}({mfa})"
        # --- تاریخ‌های آزمایش ---
        if col == f"test_date_{ys}":
            return f"تاریخ آزمایش {y}"
        if col == f"flow_test_date_{ys}":
            return f"تاریخ آزمایش دبی {y}"
        if col == f"flow_rate_lps_{ys}":
            return f"آبدهی (lit/s) {y}"
        if col == f"pressure_test_date_{ys}":
            return f"تاریخ آزمایش فشار {y}"
        if col == f"pressure_atm_{ys}":
            return f"فشار (atm) {y}"
        # --- پیمانکار سال / دبی سالانه نسبت ---
        if col == f"contractor_{ys}":
            return f"پیمانکار {y}"
        if col == f"average_flow_year_{ys}_vs_1399":
            return f"نسبت دبی متوسط {y} به 1399"
        # --- ترخیص (removal) ماهانه ---
        for men, mfa in zip(MONTHS_EN, MONTHS_FA):
            if col == f"removal_{men}_{ys}":
                return f"ترخیص {mfa} {y}"
    # ستون ویژه‌ی سن
    if col.startswith("age_until_") and col.endswith("_months"):
        return "عمر تا 1405/03/01 (ماه)"
    return None


def build_column_map(columns):
    """dict: نام ستون انگلیسی → کلید فارسیِ اکسل، برای همه‌ی ستون‌های موجود."""
    result = {}
    for col in columns:
        if col in STATIC_FULL:
            result[col] = STATIC_FULL[col]
            continue
        fa = _ts_map_for_column(col)
        result[col] = fa if fa else col  # اگر نگاشتی نبود، همان نام انگلیسی
    return result


def rows_to_fa_dicts(columns, rows):
    """ردیف‌های جدول را به لیست dict با کلیدهای فارسی تبدیل می‌کند."""
    cmap = build_column_map(columns)
    fa_keys = [cmap[c] for c in columns]
    out = []
    for row in rows:
        d = {}
        for k, v in zip(fa_keys, row):
            d[k] = v
        out.append(d)
    return out



################################################################################
# FILE: production_trend\routes.py
################################################################################

"""تحلیل روند تولید چاه‌ها.

داشبورد کاملِ کلاینت (همان script.js اصلی با همه‌ی تب‌ها و نمودارها) در iframe
نمایش داده می‌شود، اما داده به‌جای فایل اکسل از جدول production_trend خوانده و
با کلیدهای فارسیِ اکسل به‌صورت JSON در اختیار داشبورد گذاشته می‌شود.
"""
from flask import render_template, jsonify, request, redirect, url_for, flash
from flask_login import login_required
from flask_wtf import FlaskForm
from flask_wtf.file import FileField, FileRequired, FileAllowed
from wtforms import SelectField, SubmitField
from wtforms.validators import DataRequired
import os
import tempfile
from sqlalchemy import text

from app.extensions import db
from app.blueprints.production_trend import bp
from app.blueprints.production_trend.data_map import rows_to_fa_dicts


@bp.route("/")
@login_required
def index():
    return render_template("production_trend/index.html")


@bp.route("/api/data")
@login_required
def api_data():
    """کل جدول production_trend را با کلیدهای فارسیِ اکسل برمی‌گرداند."""
    result = db.session.execute(text("SELECT * FROM production_trend"))
    columns = list(result.keys())
    rows = result.fetchall()
    data = rows_to_fa_dicts(columns, [tuple(r) for r in rows])
    # فقط ردیف‌های دارای نام چاه (مطابق فیلتر اصلی داشبورد)
    data = [d for d in data if d.get("نام چاه") and str(d.get("نام چاه")).strip()]
    resp = jsonify(data)
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    resp.headers["Pragma"] = "no-cache"
    resp.headers["Expires"] = "0"
    return resp
CATEGORY_CHOICES = [
    ("production", "تولید"),
    ("runtime", "کارکرد"),
    ("well_pressure", "فشار"),
]

MONTH_CHOICES = [("-1", "سالانه (کل سال)")] + [
    (str(i), m) for i, m in enumerate(
        ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
         "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"])
]

YEAR_CHOICES = [(str(y), str(y)) for y in range(1399, 1411)]


class ColumnImportForm(FlaskForm):
    category = SelectField("دسته‌ی داده", choices=CATEGORY_CHOICES, validators=[DataRequired()])
    year = SelectField("سال", choices=YEAR_CHOICES, validators=[DataRequired()])
    month = SelectField("ماه", choices=MONTH_CHOICES, validators=[DataRequired()])
    excel = FileField("فایل اکسل (ستون اول: نام چاه، ستون دوم: مقدار)",
                      validators=[FileRequired(), FileAllowed(["xlsx", "xls"], "فقط فایل اکسل")])
    submit = SubmitField("افزودن سرستون و ورود داده")


@bp.route("/import-column", methods=["GET", "POST"])
@login_required
def import_column():
    from app.blueprints.production_trend.column_import import run
    form = ColumnImportForm()
    result = None
    if form.validate_on_submit():
        month_idx = int(form.month.data)
        month_arg = None if month_idx == -1 else month_idx
        # ذخیره‌ی موقت فایل آپلودی
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        form.excel.data.save(tmp.name)
        tmp.close()
        try:
            result = run(tmp.name, form.category.data, int(form.year.data), month_arg)
            flash(f"ستون «{result['column']}» ساخته/به‌روزرسانی شد. "
                  f"تطبیق‌خورده: {result['matched']} | بدون تطبیق: {result['unmatched']}", "success")
        except Exception as e:
            flash(f"خطا در ورود داده: {e}", "danger")
        finally:
            try:
                os.unlink(tmp.name)
            except OSError:
                pass  # اگر ویندوز اجازه‌ی حذف نداد، فایل موقت بعداً خودکار پاک می‌شود
        return redirect(url_for("production_trend.import_column"))
    return render_template("production_trend/import_column.html", form=form)


################################################################################
# FILE: pump_select\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("pump_select", __name__)

from app.blueprints.pump_select import routes  # noqa: E402,F401



################################################################################
# FILE: pump_select\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort, jsonify
from flask_login import current_user, login_required
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.pump_select import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.pump_select import PumpSelection
from app.models.audit import RecordHistory
from app.utils.dates import to_english_digits
from app import workflow


def _num(s):
    s = to_english_digits((s or "").strip())
    try:
        return float(s) if s else None
    except ValueError:
        return None


class PumpSelectionForm(FlaskForm):
    action_needed = StringField("اقدام مورد نیاز", validators=[Optional()])
    jyear = StringField("سال", validators=[Optional()])
    jmonth = StringField("ماه", validators=[Optional()])
    status_done = StringField("وضعیت انجام", validators=[Optional()])
    prev_pump_type = StringField("تیپ پمپ قبلی", validators=[Optional()])
    prev_motor_type = StringField("تیپ موتور قبلی", validators=[Optional()])
    prev_discharge_lps = FloatField("دبی قبلی (l/s)", validators=[Optional()])
    selected_pump_type = StringField("تیپ پمپ پس از بررسی", validators=[Optional()])
    selected_motor_type = StringField("تیپ موتور پس از بررسی", validators=[Optional()])
    target_discharge_lps = FloatField("دبی پس از بررسی (l/s)", validators=[Optional()])
    discharge_increase_lps = FloatField("میزان افزایش دبی", validators=[Optional()])
    selected_head_m = FloatField("هد پمپ انتخابی (m)", validators=[Optional()])
    form_delivery_date = JalaliDateField("تاریخ تحویل فرم", validators=[Optional()])
    pull_date = JalaliDateField("تاریخ کشیدن", validators=[Optional()])
    videometry_date = JalaliDateField("تاریخ ویدئومتری", validators=[Optional()])
    install_date = JalaliDateField("تاریخ نصب", validators=[Optional()])
    verify_flowtest_date = JalaliDateField("تاریخ دبی‌سنجی (صحت‌سنجی)", validators=[Optional()])
    verify_discharge_lps = FloatField("آبدهی دبی‌سنجی", validators=[Optional()])
    verify_head_m = FloatField("هد دبی‌سنجی", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


FIELDS = [
    "action_needed", "jyear", "jmonth", "status_done", "prev_pump_type", "prev_motor_type",
    "prev_discharge_lps", "selected_pump_type", "selected_motor_type", "target_discharge_lps",
    "discharge_increase_lps", "selected_head_m", "form_delivery_date", "pull_date",
    "videometry_date", "install_date", "verify_flowtest_date", "verify_discharge_lps",
    "verify_head_m", "notes",
]


def _apply(form, rec):
    for f in FIELDS:
        setattr(rec, f, getattr(form, f).data)


# ---------- catalog-driven selection calculator ----------
@bp.route("/pump-select/catalog.json")
@permission_required("pump_select", "view")
def catalog_json():
    from app.services.catalog import build_library
    return jsonify(build_library())


@bp.route("/pump-select/calculator")
@login_required
def calculator():
    from app.services.catalog import design_params
    well = None
    q, head, src = 10.0, 180, None
    well_id = request.args.get("well", type=int)
    if well_id:
        well = db.session.get(Well, well_id)
        if well:
            q, head, src = design_params(well)
    return render_template("pump_select/calculator.html", well=well, design_q=q,
                           design_head=head, src=src)


@bp.route("/wells/<int:well_id>/pump-select/from-calc", methods=["POST"])
@permission_required("pump_select", "create")
def from_calc(well_id):
    well = db.get_or_404(Well, well_id)
    rec = PumpSelection(
        well_id=well.id, created_by_id=current_user.id,
        action_needed="انتخاب از کاتالوگ",
        selected_pump_type=(request.form.get("model") or "").strip() or None,
        target_discharge_lps=_num(request.form.get("q")),
        selected_head_m=_num(request.form.get("head")),
        notes=(request.form.get("note") or "").strip() or None,
    )
    db.session.add(rec)
    db.session.flush()
    workflow.log_change(rec, "create", detail="از محاسبه‌گر کاتالوگ")
    db.session.commit()
    flash("انتخاب پمپ از محاسبه‌گر ثبت شد (پیش‌نویس).", "success")
    return redirect(url_for("pump_select.detail", record_id=rec.id))


@bp.route("/wells/<int:well_id>/pump-select/new", methods=["GET", "POST"])
@permission_required("pump_select", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = PumpSelectionForm()
    if form.validate_on_submit():
        rec = PumpSelection(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("انتخاب پمپ ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("pump_select.detail", record_id=rec.id))
    return render_template("pump_select/form.html", form=form, well=well, title="ثبت انتخاب پمپ")


@bp.route("/pump-select/<int:record_id>")
@permission_required("pump_select", "view")
def detail(record_id):
    rec = db.get_or_404(PumpSelection, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="pump_selections", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("pump_select/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/pump-select/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("pump_select", "edit")
def edit(record_id):
    rec = db.get_or_404(PumpSelection, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("pump_select.detail", record_id=rec.id))
    form = PumpSelectionForm(obj=rec)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("انتخاب پمپ به‌روزرسانی شد.", "success")
        return redirect(url_for("pump_select.detail", record_id=rec.id))
    return render_template("pump_select/form.html", form=form, well=rec.well, title="ویرایش انتخاب پمپ")


@bp.route("/pump-select/<int:record_id>/delete", methods=["POST"])
@permission_required("pump_select", "delete")
def delete(record_id):
    rec = db.get_or_404(PumpSelection, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد انتخاب پمپ حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(PumpSelection, record_id)
    if not current_user.has_permission("pump_select", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("pump_select.detail", record_id=record_id))


@bp.route("/pump-select/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/pump-select/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/pump-select/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/pump-select/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: pump_test\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("pump_test", __name__)

from app.blueprints.pump_test import routes  # noqa: E402,F401



################################################################################
# FILE: pump_test\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.pump_test import bp
from app.security import permission_required
from app.utils.forms import (
    JalaliDateField, PersianFloatField as FloatField, PersianIntegerField as IntegerField,
)
from app.utils.dates import to_english_digits
from app.models.well import Well
from app.models.pump_test import PumpTest, PumpTestStep
from app.models.audit import RecordHistory
from app.models.constants import PUMP_TEST_TYPES
from app import workflow

STEP_ROWS = 8
STEP_FIELDS = ["rpm", "discharge_lps", "observed_drawdown", "calc_drawdown",
               "grid_loss", "aquifer_loss", "efficiency"]


class PumpTestForm(FlaskForm):
    test_date = JalaliDateField("تاریخ آزمایش", validators=[Optional()])
    test_type = SelectField("نوع آزمایش", choices=[("", "—")] + PUMP_TEST_TYPES, validators=[Optional()])
    duration_h = FloatField("مدت شستشو/آزمایش (ساعت)", validators=[Optional()])
    contractor = StringField("پیمانکار", validators=[Optional()])
    consultant = StringField("مشاور", validators=[Optional()])
    employer = StringField("کارفرما", validators=[Optional()])
    contract_no = StringField("شماره قرارداد", validators=[Optional()])
    project_title = StringField("عنوان پروژه", validators=[Optional()])
    static_level = FloatField("سطح استاتیک", validators=[Optional()])
    max_dynamic_level = FloatField("حداکثر سطح دینامیک", validators=[Optional()])
    max_drawdown = FloatField("حداکثر افت چاه", validators=[Optional()])
    max_yield_lps = FloatField("حداکثر آبدهی (l/s)", validators=[Optional()])
    coeff_a = FloatField("ضریب a", validators=[Optional()])
    coeff_b = FloatField("ضریب b", validators=[Optional()])
    proposed_discharge_lps = FloatField("دبی مجاز پیشنهادی", validators=[Optional()])
    proposed_install_depth_m = FloatField("عمق نصب پیشنهادی", validators=[Optional()])
    resulting_drawdown_m = FloatField("میزان افت حاصله", validators=[Optional()])
    motor_type = StringField("نوع موتور", validators=[Optional()])
    motor_power_hp = FloatField("قدرت موتور (HP)", validators=[Optional()])
    gearbox_power_hp = FloatField("قدرت جعبه‌دنده (HP)", validators=[Optional()])
    gearbox_ratio = StringField("تبدیل جعبه‌دنده", validators=[Optional()])
    pump_type = StringField("نوع پمپ", validators=[Optional()])
    pump_stages = IntegerField("تعداد طبقه", validators=[Optional()])
    pump_diameter_in = FloatField("قطر پمپ (اینچ)", validators=[Optional()])
    max_rpm = FloatField("حداکثر دور موتور", validators=[Optional()])
    discharge_pipe_diameter_in = FloatField("قطر لوله آبده (اینچ)", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


HEADER_FIELDS = [
    "test_date", "test_type", "duration_h", "contractor", "consultant", "employer",
    "contract_no", "project_title", "static_level", "max_dynamic_level", "max_drawdown",
    "max_yield_lps", "coeff_a", "coeff_b", "proposed_discharge_lps",
    "proposed_install_depth_m", "resulting_drawdown_m", "motor_type", "motor_power_hp",
    "gearbox_power_hp", "gearbox_ratio", "pump_type", "pump_stages", "pump_diameter_in",
    "max_rpm", "discharge_pipe_diameter_in", "notes",
]


def _num(s):
    s = to_english_digits((s or "").strip())
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _apply_header(form, rec):
    for f in HEADER_FIELDS:
        setattr(rec, f, getattr(form, f).data)


def _apply_steps(rec):
    rec.steps.clear()
    for i in range(1, STEP_ROWS + 1):
        vals = {f: request.form.get(f"st-{i}-{f}", "").strip() for f in STEP_FIELDS}
        if not any(vals.values()):
            continue
        rec.steps.append(PumpTestStep(
            step_no=i,
            rpm=_num(vals["rpm"]),
            discharge_lps=_num(vals["discharge_lps"]),
            observed_drawdown=_num(vals["observed_drawdown"]),
            calc_drawdown=_num(vals["calc_drawdown"]),
            grid_loss=_num(vals["grid_loss"]),
            aquifer_loss=_num(vals["aquifer_loss"]),
            efficiency=_num(vals["efficiency"]),
        ))


@bp.route("/wells/<int:well_id>/pump-test/new", methods=["GET", "POST"])
@permission_required("pump_test", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = PumpTestForm()
    if form.validate_on_submit():
        rec = PumpTest(well_id=well.id, created_by_id=current_user.id)
        _apply_header(form, rec)
        db.session.add(rec)
        _apply_steps(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("آزمایش پمپاژ ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("pump_test.detail", record_id=rec.id))
    return render_template("pump_test/form.html", form=form, well=well,
                           title="ثبت آزمایش پمپاژ", steps=[], step_rows=STEP_ROWS)


@bp.route("/pump-test/<int:record_id>")
@permission_required("pump_test", "view")
def detail(record_id):
    rec = db.get_or_404(PumpTest, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="pump_tests", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("pump_test/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           type_labels=dict(PUMP_TEST_TYPES))


@bp.route("/pump-test/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("pump_test", "edit")
def edit(record_id):
    rec = db.get_or_404(PumpTest, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("pump_test.detail", record_id=rec.id))
    form = PumpTestForm(obj=rec)
    if form.validate_on_submit():
        _apply_header(form, rec)
        _apply_steps(rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("آزمایش پمپاژ به‌روزرسانی شد.", "success")
        return redirect(url_for("pump_test.detail", record_id=rec.id))
    return render_template("pump_test/form.html", form=form, well=rec.well,
                           title="ویرایش آزمایش پمپاژ", steps=rec.steps, step_rows=STEP_ROWS)


@bp.route("/pump-test/<int:record_id>/delete", methods=["POST"])
@permission_required("pump_test", "delete")
def delete(record_id):
    rec = db.get_or_404(PumpTest, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("آزمایش پمپاژ حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(PumpTest, record_id)
    if not current_user.has_permission("pump_test", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("pump_test.detail", record_id=record_id))


@bp.route("/pump-test/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/pump-test/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/pump-test/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/pump-test/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: rehab\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("rehab", __name__)

from app.blueprints.rehab import routes  # noqa: E402,F401



################################################################################
# FILE: rehab\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.rehab import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.rehab import Rehabilitation
from app.models.audit import RecordHistory
from app.blueprints.install.routes import get_or_create_supplier
from app import workflow


class RehabForm(FlaskForm):
    stage = StringField("مرحله", validators=[Optional()])
    jyear = StringField("سال", validators=[Optional()])
    rehab_date = JalaliDateField("تاریخ بهسازی", validators=[Optional()])
    pumping_end_date = JalaliDateField("تاریخ اتمام پمپاژ", validators=[Optional()])
    rehab_contractor_name = StringField("پیمانکار بهسازی", validators=[Optional()])
    pumping_contractor_name = StringField("پیمانکار پمپاژ", validators=[Optional()])
    pump_type_before = StringField("تیپ پمپ پیش از بهسازی", validators=[Optional()])
    discharge_before_lps = FloatField("دبی قبل (l/s)", validators=[Optional()])
    pump_type_after = StringField("تیپ پمپ پس از بهسازی", validators=[Optional()])
    discharge_after_lps = FloatField("دبی بعد (l/s)", validators=[Optional()])
    reason = StringField("علت بهسازی", validators=[Optional()])
    method = StringField("روش", validators=[Optional()])
    observed_fault = TextAreaField("خرابی مشاهده‌شده", validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


SIMPLE = ["stage", "jyear", "rehab_date", "pumping_end_date", "pump_type_before",
          "discharge_before_lps", "pump_type_after", "discharge_after_lps",
          "reason", "method", "observed_fault", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data)
    rc = get_or_create_supplier(form.rehab_contractor_name.data, "contractor")
    pc = get_or_create_supplier(form.pumping_contractor_name.data, "contractor")
    rec.rehab_contractor_id = rc.id if rc else None
    rec.pumping_contractor_id = pc.id if pc else None
    if rec.discharge_before_lps is not None and rec.discharge_after_lps is not None:
        rec.discharge_change_lps = round(rec.discharge_after_lps - rec.discharge_before_lps, 2)


@bp.route("/wells/<int:well_id>/rehab/new", methods=["GET", "POST"])
@permission_required("rehab", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = RehabForm()
    if form.validate_on_submit():
        rec = Rehabilitation(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("بهسازی ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("rehab.detail", record_id=rec.id))
    return render_template("rehab/form.html", form=form, well=well, title="ثبت بهسازی")


@bp.route("/rehab/<int:record_id>")
@permission_required("rehab", "view")
def detail(record_id):
    rec = db.get_or_404(Rehabilitation, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="rehabilitations", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("rehab/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec))


@bp.route("/rehab/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("rehab", "edit")
def edit(record_id):
    rec = db.get_or_404(Rehabilitation, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("rehab.detail", record_id=rec.id))
    form = RehabForm(obj=rec)
    if request.method == "GET":
        form.rehab_contractor_name.data = rec.rehab_contractor.name if rec.rehab_contractor else ""
        form.pumping_contractor_name.data = rec.pumping_contractor.name if rec.pumping_contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("بهسازی به‌روزرسانی شد.", "success")
        return redirect(url_for("rehab.detail", record_id=rec.id))
    return render_template("rehab/form.html", form=form, well=rec.well, title="ویرایش بهسازی")


@bp.route("/rehab/<int:record_id>/delete", methods=["POST"])
@permission_required("rehab", "delete")
def delete(record_id):
    rec = db.get_or_404(Rehabilitation, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد بهسازی حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(Rehabilitation, record_id)
    if not current_user.has_permission("rehab", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("rehab.detail", record_id=record_id))


@bp.route("/rehab/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/rehab/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/rehab/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/rehab/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: relocation\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("relocation", __name__)

from app.blueprints.relocation import routes  # noqa: E402,F401



################################################################################
# FILE: relocation\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SelectField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.relocation import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.relocation import RelocationRecord
from app.models.audit import RecordHistory
from app.models.constants import RELOCATION_CANDIDACY, RELOCATION_TYPES
from app import workflow


class RelocationForm(FlaskForm):
    decision_date = JalaliDateField("تاریخ تصمیم", validators=[Optional()])
    candidacy = SelectField("وضعیت نامزدی", choices=[("", "—")] + RELOCATION_CANDIDACY,
                            validators=[Optional()])
    reloc_type = SelectField("نوع جابه‌جایی", choices=[("", "—")] + RELOCATION_TYPES,
                             validators=[Optional()])
    reason = StringField("علت جابه‌جایی", validators=[Optional()])
    location_note = StringField("موقعیت پیشنهادی", validators=[Optional()])
    letter_no = StringField("شماره نامه", validators=[Optional()])
    distance_m = FloatField("فاصله از چاه قبلی (m)", validators=[Optional()])
    new_well_id = SelectField("چاه جانشین", coerce=int, validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_wells(self, exclude_id):
        wells = db.session.scalars(db.select(Well).order_by(Well.name)).all()
        self.new_well_id.choices = [(0, "—")] + [
            (w.id, f"{w.name} ({w.display_pm})") for w in wells if w.id != exclude_id]


SIMPLE = ["decision_date", "candidacy", "reloc_type", "reason",
          "location_note", "letter_no", "distance_m", "notes"]


def _apply(form, rec):
    for f in SIMPLE:
        setattr(rec, f, getattr(form, f).data or None)
    rec.new_well_id = form.new_well_id.data or None


@bp.route("/relocation")
@permission_required("relocation", "view")
def list_relocations():
    cand_labels = dict(RELOCATION_CANDIDACY)
    type_labels = dict(RELOCATION_TYPES)
    rows = db.session.scalars(
        db.select(RelocationRecord).order_by(RelocationRecord.candidacy)).all()
    return render_template("relocation/list.html", rows=rows,
                           cand_labels=cand_labels, type_labels=type_labels)


@bp.route("/wells/<int:well_id>/relocation/new", methods=["GET", "POST"])
@permission_required("relocation", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = RelocationForm()
    form.populate_wells(well.id)
    if form.validate_on_submit():
        rec = RelocationRecord(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("جابه‌جایی ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("relocation.detail", record_id=rec.id))
    return render_template("relocation/form.html", form=form, well=well, title="ثبت جابه‌جایی")


@bp.route("/relocation/<int:record_id>")
@permission_required("relocation", "view")
def detail(record_id):
    rec = db.get_or_404(RelocationRecord, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="relocations", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())).all()
    return render_template("relocation/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           cand_labels=dict(RELOCATION_CANDIDACY),
                           type_labels=dict(RELOCATION_TYPES))


@bp.route("/relocation/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("relocation", "edit")
def edit(record_id):
    rec = db.get_or_404(RelocationRecord, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("relocation.detail", record_id=rec.id))
    form = RelocationForm(obj=rec)
    form.populate_wells(rec.well_id)
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("جابه‌جایی به‌روزرسانی شد.", "success")
        return redirect(url_for("relocation.detail", record_id=rec.id))
    return render_template("relocation/form.html", form=form, well=rec.well, title="ویرایش جابه‌جایی")


@bp.route("/relocation/<int:record_id>/delete", methods=["POST"])
@permission_required("relocation", "delete")
def delete(record_id):
    rec = db.get_or_404(RelocationRecord, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد جابه‌جایی حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _link_lineage(rec):
    """On approval: wire successor.parent_well = subject well, mark old relocated."""
    if rec.new_well_id and rec.new_well:
        rec.new_well.parent_well_id = rec.well_id
        rec.well.status = "relocated"


def _transition(record_id, fn, perm_action, msg, link=False, **kw):
    rec = db.get_or_404(RelocationRecord, record_id)
    if not current_user.has_permission("relocation", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        if link:
            _link_lineage(rec)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("relocation.detail", record_id=record_id))


@bp.route("/relocation/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/relocation/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve",
                       "تأیید شد و شجره‌نامه‌ی چاه به‌روزرسانی شد.", link=True)


@bp.route("/relocation/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/relocation/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: reports\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("reports", __name__, url_prefix="/reports")

from app.blueprints.reports import routes  # noqa: E402,F401



################################################################################
# FILE: reports\routes.py
################################################################################

from collections import defaultdict

from flask import render_template, request, redirect, url_for, flash
from flask_login import current_user

from app.extensions import db
from app.blueprints.reports import bp
from app.security import permission_required
from app.models.pump_asset import PumpInstallation, Supplier
from app.models.prioritization import PrioritizationCriterion, DEFAULT_CRITERIA
from app.utils.dates import to_english_digits


def _aggregate(installs, key_attr):
    agg = defaultdict(lambda: {"completed": 0, "life_sum": 0.0, "running": 0, "burnt": 0})
    for i in installs:
        sid = getattr(i, key_attr)
        if not sid:
            continue
        a = agg[sid]
        if i.is_running:
            a["running"] += 1
        life = i.useful_life_months
        if life is not None:
            a["completed"] += 1
            a["life_sum"] += life
        if i.removal_reason and "سوخت" in i.removal_reason:
            a["burnt"] += 1
    return agg


def _rows(agg, names, min_completed):
    rows = []
    for sid, a in agg.items():
        if a["completed"] < min_completed:
            continue
        avg = round(a["life_sum"] / a["completed"], 1) if a["completed"] else None
        rows.append({
            "name": names.get(sid, "—"),
            "completed": a["completed"], "running": a["running"],
            "burnt": a["burnt"], "avg_life": avg,
        })
    rows.sort(key=lambda r: (r["avg_life"] is not None, r["avg_life"]), reverse=True)
    return rows


@bp.route("/suppliers")
@permission_required("reports", "view")
def suppliers():
    min_completed = int(request.args.get("min", 5))
    installs = db.session.scalars(db.select(PumpInstallation)).all()
    names = {s.id: s.name for s in db.session.scalars(db.select(Supplier)).all()}
    makers = _rows(_aggregate(installs, "manufacturer_id"), names, min_completed)
    contractors = _rows(_aggregate(installs, "contractor_id"), names, min_completed)
    return render_template("reports/suppliers.html",
                           makers=makers, contractors=contractors,
                           total=len(installs), min_completed=min_completed)


def _ensure_criteria():
    """Seed defaults and add any criteria missing by key (handles upgrades)."""
    existing = {c.key for c in db.session.scalars(db.select(PrioritizationCriterion)).all()}
    added = False
    for key, label, weight, direction in DEFAULT_CRITERIA:
        if key not in existing:
            db.session.add(PrioritizationCriterion(
                key=key, label=label, weight=weight, direction=direction))
            added = True
    if added:
        db.session.commit()


@bp.route("/prioritization", methods=["GET", "POST"])
@permission_required("reports", "view")
def prioritization():
    from app.services import prioritization as svc
    _ensure_criteria()

    if request.method == "POST":
        if not current_user.has_permission("admin", "edit"):
            flash("برای تغییر وزن‌ها مجوز مدیریت لازم است.", "warning")
            return redirect(url_for("reports.prioritization"))
        for c in db.session.scalars(db.select(PrioritizationCriterion)).all():
            w = to_english_digits(request.form.get(f"w_{c.id}", "")).strip()
            try:
                c.weight = float(w)
            except ValueError:
                pass
            c.is_active = request.form.get(f"a_{c.id}") == "on"
        db.session.commit()
        flash("وزن‌ها به‌روزرسانی شد.", "success")
        return redirect(url_for("reports.prioritization"))

    rows, criteria = svc.compute()
    all_criteria = db.session.scalars(
        db.select(PrioritizationCriterion).order_by(PrioritizationCriterion.weight.desc())
    ).all()
    return render_template("reports/prioritization.html",
                           rows=rows[:100], total=len(rows), criteria=all_criteria,
                           can_edit=current_user.has_permission("admin", "edit"))


@bp.route("/energy")
@permission_required("reports", "view")
def energy():
    from app.services import energy_stats
    return render_template("reports/energy.html", e=energy_stats.compute())


@bp.route("/alerts")
@permission_required("reports", "view")
def alerts():
    from app.services import alerts as alert_svc
    items = alert_svc.compute()
    return render_template("reports/alerts.html",
                           alerts=items, summary=alert_svc.summary(items),
                           categories=alert_svc.CATEGORIES)


@bp.route("/analytics")
@permission_required("reports", "view")
def analytics():
    from app.services.analytics import summary
    return render_template("reports/analytics.html", f=summary.fleet())


@bp.route("/zones")
@permission_required("reports", "view")
def zones():
    from app.services import zone_stats
    return render_template("reports/zones.html", z=zone_stats.compute())


@bp.route("/portfolio")
@permission_required("reports", "view")
def portfolio():
    from app.services import optimize, economics
    economics.ensure_params()
    default_budget = economics.get("rehab_cost", 2e9) * 20
    budget = request.args.get("budget", type=float)
    budget_rial = (budget * 1e9) if budget else default_budget
    return render_template("reports/portfolio.html",
                           p=optimize.portfolio(budget_rial),
                           budget_b=round(budget_rial / 1e9, 1),
                           unit_cost_b=round(economics.get("rehab_cost", 2e9) / 1e9, 2))


@bp.route("/dispatch")
@permission_required("reports", "view")
def dispatch():
    from app.services import optimize, zone_stats
    z = zone_stats.compute()
    group = request.args.get("group", "zone")
    key = request.args.get("key")
    demand = request.args.get("demand", type=float)
    result = None
    if key:
        attr = "zone" if group == "zone" else "destination_reservoir"
        result = optimize.field_dispatch(attr, key, demand or 0)
    return render_template("reports/dispatch.html", z=z, group=group,
                           key=key, demand=demand, result=result)



################################################################################
# FILE: videometry\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("videometry", __name__)

from app.blueprints.videometry import routes  # noqa: E402,F401



################################################################################
# FILE: videometry\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request, abort
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, TextAreaField, SubmitField
from wtforms.validators import Optional

from app.extensions import db
from app.blueprints.videometry import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.utils.dates import to_english_digits
from app.models.well import Well
from app.models.videometry import Videometry, VideometryFinding
from app.models.audit import RecordHistory
from app.models.constants import VIDEO_FINDING_TYPES, SEVERITY_LEVELS
from app.blueprints.install.routes import get_or_create_supplier
from app import workflow

FIND_ROWS = 8


class VideometryForm(FlaskForm):
    log_date = JalaliDateField("تاریخ چاه‌نگاری", validators=[Optional()])
    contractor_name = StringField("پیمانکار", validators=[Optional()])
    equipment = StringField("تجهیز/دوربین", validators=[Optional()])
    depth_from_m = FloatField("از عمق (m)", validators=[Optional()])
    depth_to_m = FloatField("تا عمق (m)", validators=[Optional()])
    final_depth_m = FloatField("عمق نهایی چاه (m)", validators=[Optional()])
    water_level_m = FloatField("سطح آب (m)", validators=[Optional()])
    video_file_ref = StringField("مسیر/نام فایل ویدئو", validators=[Optional()])
    summary = TextAreaField("خلاصه‌ی یافته‌ها", validators=[Optional()])
    submit = SubmitField("ذخیره")


HEADER = ["log_date", "equipment", "depth_from_m", "depth_to_m", "final_depth_m",
          "water_level_m", "video_file_ref", "summary"]


def _num(s):
    s = to_english_digits((s or "").strip())
    try:
        return float(s) if s else None
    except ValueError:
        return None


def _apply(form, rec):
    for f in HEADER:
        setattr(rec, f, getattr(form, f).data)
    con = get_or_create_supplier(form.contractor_name.data, "contractor")
    rec.contractor_id = con.id if con else None
    rec.findings.clear()
    for i in range(1, FIND_ROWS + 1):
        depth = request.form.get(f"f-{i}-depth_m", "").strip()
        ftype = request.form.get(f"f-{i}-finding_type", "").strip()
        sev = request.form.get(f"f-{i}-severity", "").strip()
        note = request.form.get(f"f-{i}-note", "").strip()
        if not (depth or ftype or note):
            continue
        rec.findings.append(VideometryFinding(
            depth_m=_num(depth), finding_type=ftype or None,
            severity=sev or None, note=note or None))


@bp.route("/wells/<int:well_id>/videometry/new", methods=["GET", "POST"])
@permission_required("videometry", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = VideometryForm()
    if form.validate_on_submit():
        rec = Videometry(well_id=well.id, created_by_id=current_user.id)
        _apply(form, rec)
        db.session.add(rec)
        db.session.flush()
        workflow.log_change(rec, "create")
        db.session.commit()
        flash("چاه‌نگاری ثبت شد (پیش‌نویس).", "success")
        return redirect(url_for("videometry.detail", record_id=rec.id))
    return render_template("videometry/form.html", form=form, well=well,
                           title="ثبت چاه‌نگاری", findings=[], find_rows=FIND_ROWS,
                           finding_types=VIDEO_FINDING_TYPES, severities=SEVERITY_LEVELS)


@bp.route("/videometry/<int:record_id>")
@permission_required("videometry", "view")
def detail(record_id):
    rec = db.get_or_404(Videometry, record_id)
    history = db.session.scalars(
        db.select(RecordHistory).filter_by(entity_type="videometry_logs", entity_id=rec.id)
        .order_by(RecordHistory.timestamp.desc())
    ).all()
    return render_template("videometry/detail.html", rec=rec, well=rec.well,
                           history=history, editable=workflow.is_editable(rec),
                           type_labels=dict(VIDEO_FINDING_TYPES),
                           sev_labels=dict(SEVERITY_LEVELS))


@bp.route("/videometry/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("videometry", "edit")
def edit(record_id):
    rec = db.get_or_404(Videometry, record_id)
    if not workflow.is_editable(rec):
        flash("این رکورد قفل است؛ ابتدا باید به پیش‌نویس بازگردانده شود.", "warning")
        return redirect(url_for("videometry.detail", record_id=rec.id))
    form = VideometryForm(obj=rec)
    if request.method == "GET":
        form.contractor_name.data = rec.contractor.name if rec.contractor else ""
    if form.validate_on_submit():
        _apply(form, rec)
        rec.updated_by_id = current_user.id
        workflow.log_change(rec, "update")
        db.session.commit()
        flash("چاه‌نگاری به‌روزرسانی شد.", "success")
        return redirect(url_for("videometry.detail", record_id=rec.id))
    return render_template("videometry/form.html", form=form, well=rec.well,
                           title="ویرایش چاه‌نگاری", findings=rec.findings, find_rows=FIND_ROWS,
                           finding_types=VIDEO_FINDING_TYPES, severities=SEVERITY_LEVELS)


@bp.route("/videometry/<int:record_id>/delete", methods=["POST"])
@permission_required("videometry", "delete")
def delete(record_id):
    rec = db.get_or_404(Videometry, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد چاه‌نگاری حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id))


def _transition(record_id, fn, perm_action, msg, **kw):
    rec = db.get_or_404(Videometry, record_id)
    if not current_user.has_permission("videometry", perm_action):
        abort(403)
    try:
        fn(rec, **kw)
        db.session.commit()
        flash(msg, "success")
    except workflow.WorkflowError as e:
        db.session.rollback()
        flash(str(e), "danger")
    return redirect(url_for("videometry.detail", record_id=record_id))


@bp.route("/videometry/<int:record_id>/submit", methods=["POST"])
def submit(record_id):
    return _transition(record_id, workflow.submit, "edit", "برای تأیید ثبت شد.")


@bp.route("/videometry/<int:record_id>/approve", methods=["POST"])
def approve(record_id):
    return _transition(record_id, workflow.approve, "approve", "تأیید شد.")


@bp.route("/videometry/<int:record_id>/reject", methods=["POST"])
def reject(record_id):
    return _transition(record_id, workflow.reject, "approve", "برگشت داده شد.",
                       reason=request.form.get("reason", ""))


@bp.route("/videometry/<int:record_id>/revert", methods=["POST"])
def revert(record_id):
    return _transition(record_id, workflow.revert_to_draft, "approve", "به پیش‌نویس بازگردانده شد.")



################################################################################
# FILE: water_level\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("water_level", __name__)

from app.blueprints.water_level import routes  # noqa: E402,F401



################################################################################
# FILE: water_level\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import BooleanField, StringField, SubmitField
from wtforms.validators import Optional, DataRequired

from app.extensions import db
from app.blueprints.water_level import bp
from app.security import permission_required
from app.utils.forms import JalaliDateField, PersianFloatField as FloatField
from app.models.well import Well
from app.models.water_level import WaterLevelLog


class WaterLevelForm(FlaskForm):
    measure_date = JalaliDateField("تاریخ اندازه‌گیری", validators=[DataRequired()])
    static_level = FloatField("تراز آب (عمق تا آب، m)", validators=[DataRequired()])
    is_pumping = BooleanField("در حال پمپاژ")
    notes = StringField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")


@bp.route("/wells/<int:well_id>/water-level/new", methods=["GET", "POST"])
@permission_required("water_level", "create")
def create(well_id):
    well = db.get_or_404(Well, well_id)
    form = WaterLevelForm()
    if form.validate_on_submit():
        rec = WaterLevelLog(well_id=well.id, method="manual",
                            created_by_id=current_user.id)
        form.populate_obj(rec)
        db.session.add(rec)
        db.session.commit()
        flash("تراز آب ثبت شد.", "success")
        return redirect(url_for("wells.detail", well_id=well.id) + "#sec-waterlevel")
    return render_template("water_level/form.html", form=form, well=well, title="ثبت تراز آب")


@bp.route("/water-level/<int:record_id>/edit", methods=["GET", "POST"])
@permission_required("water_level", "edit")
def edit(record_id):
    rec = db.get_or_404(WaterLevelLog, record_id)
    form = WaterLevelForm(obj=rec)
    if form.validate_on_submit():
        form.populate_obj(rec)
        rec.updated_by_id = current_user.id
        db.session.commit()
        flash("تراز آب به‌روزرسانی شد.", "success")
        return redirect(url_for("wells.detail", well_id=rec.well_id) + "#sec-waterlevel")
    return render_template("water_level/form.html", form=form, well=rec.well, title="ویرایش تراز آب")


@bp.route("/water-level/<int:record_id>/delete", methods=["POST"])
@permission_required("water_level", "delete")
def delete(record_id):
    rec = db.get_or_404(WaterLevelLog, record_id)
    well_id = rec.well_id
    db.session.delete(rec)
    db.session.commit()
    flash("رکورد تراز آب حذف شد.", "info")
    return redirect(url_for("wells.detail", well_id=well_id) + "#sec-waterlevel")



################################################################################
# FILE: wells\__init__.py
################################################################################

from flask import Blueprint

bp = Blueprint("wells", __name__, url_prefix="/wells")

from app.blueprints.wells import routes  # noqa: E402,F401



################################################################################
# FILE: wells\routes.py
################################################################################

from flask import render_template, redirect, url_for, flash, request
from flask_login import current_user
from flask_wtf import FlaskForm
from wtforms import StringField, SelectField, TextAreaField, SubmitField
from wtforms.validators import DataRequired, Optional

from app.extensions import db
from app.blueprints.wells import bp
from app.utils.forms import PersianFloatField as FloatField
from app.security import permission_required
from app.models.well import Well
from app.models.org import OrgUnit
from app.models.constants import (
    WELL_KINDS,
    WELL_STATUSES,
    LOCATION_STATUSES,
    ORG_UNIT_TYPES,
    WELL_CONSTRUCTION_TYPES,
)


class WellForm(FlaskForm):
    pm_code = StringField("کد PM", validators=[DataRequired()])
    name = StringField("نام چاه", validators=[DataRequired()])
    well_kind = SelectField("نوع چاه", choices=WELL_KINDS)
    office_id = SelectField("اداره", coerce=int, validators=[Optional()])
    center_id = SelectField("مرکز آبرسانی", coerce=int, validators=[Optional()])
    zone = StringField("پهنه", validators=[Optional()])
    sub_zone = StringField("زیرپهنه", validators=[Optional()])
    utm_x = FloatField("UTM X", validators=[Optional()])
    utm_y = FloatField("UTM Y", validators=[Optional()])
    utm_zone = StringField("Zone", validators=[Optional()])
    latitude = FloatField("عرض جغرافیایی", validators=[Optional()])
    longitude = FloatField("طول جغرافیایی", validators=[Optional()])
    ground_elevation = FloatField("ارتفاع زمین", validators=[Optional()])
    drill_year = StringField("سال حفر", validators=[Optional()])
    location_status = SelectField(
        "وضعیت تعیین محل", choices=[("", "—")] + LOCATION_STATUSES, validators=[Optional()]
    )
    status = SelectField("وضعیت چاه", choices=WELL_STATUSES)
    parent_well_id = SelectField("چاه قبلی (در جابه‌جایی)", coerce=int, validators=[Optional()])
    notes = TextAreaField("توضیحات", validators=[Optional()])
    submit = SubmitField("ذخیره")

    def populate_choices(self, exclude_well_id=None):
        offices = db.session.scalars(
            db.select(OrgUnit).filter_by(unit_type="office").order_by(OrgUnit.name)
        ).all()
        centers = db.session.scalars(
            db.select(OrgUnit)
            .where(OrgUnit.unit_type.in_(["center", "rural_region"]))
            .order_by(OrgUnit.name)
        ).all()
        self.office_id.choices = [(0, "—")] + [(o.id, o.name) for o in offices]
        self.center_id.choices = [(0, "—")] + [(c.id, c.name) for c in centers]
        wells = db.session.scalars(db.select(Well).order_by(Well.pm_code)).all()
        self.parent_well_id.choices = [(0, "—")] + [
            (w.id, f"{w.pm_code} — {w.name}")
            for w in wells
            if w.id != exclude_well_id
        ]


@bp.route("/")
@permission_required("wells", "view")
def list_wells():
    q = (request.args.get("q") or "").strip()
    query = db.select(Well).order_by(Well.pm_code)
    if q:
        like = f"%{q}%"
        query = db.select(Well).where(
            db.or_(Well.pm_code.ilike(like), Well.name.ilike(like))
        ).order_by(Well.pm_code)
    wells = db.session.scalars(query).all()
    return render_template(
        "wells/list.html",
        wells=wells,
        q=q,
        kind_labels=dict(WELL_KINDS),
        status_labels=dict(WELL_STATUSES),
    )


@bp.route("/<int:well_id>")
@permission_required("wells", "view")
def detail(well_id):
    well = db.get_or_404(Well, well_id)
    from app.services import timeline
    from app.models.baseline import WellBaseline
    from app.models.quality import WaterQuality
    from app.models.production import MonthlyProduction
    baseline = db.session.scalar(
        db.select(WellBaseline).filter_by(well_id=well.id, is_current=True))
    quality = db.session.scalars(
        db.select(WaterQuality).filter_by(well_id=well.id)
        .order_by(WaterQuality.sample_date.desc())).all()

    # Monthly operational time-series (روند تولید) for the trend chart.
    prod_rows = db.session.scalars(
        db.select(MonthlyProduction).filter_by(well_id=well.id)
        .order_by(MonthlyProduction.jyear, MonthlyProduction.jmonth)).all()
    production = {
        "labels": [f"{r.jyear}/{r.jmonth:02d}" for r in prod_rows],
        "discharge": [r.avg_discharge_lps for r in prod_rows],
        "volume": [r.production_m3 for r in prod_rows],
        "hours": [r.run_hours for r in prod_rows],
    }
    energy = next((r for r in prod_rows if r.energy_kwh is not None), None)

    events = timeline.build(well)
    # Position lifecycle events on the production/discharge chart (nearest month).
    chart_events = []
    if prod_rows:
        import jdatetime
        from app.utils.dates import format_jalali
        periods = [r.jyear * 12 + (r.jmonth - 1) for r in prod_rows]
        colors = {
            "drilling": "#c9780b", "pump_test": "#7c4dff", "pump_select": "#11a394",
            "install": "#2f6bff", "operation": "#18a558", "rehab": "#e0463e",
            "videometry": "#0e8fa8", "relocation": "#d6457f", "maintenance": "#6a7180",
        }
        for ev in events:
            d = ev.get("date")
            if not d:
                continue
            jd = jdatetime.date.fromgregorian(date=d)
            p = jd.year * 12 + (jd.month - 1)
            idx = min(range(len(periods)), key=lambda i: abs(periods[i] - p))
            chart_events.append({
                "index": idx,
                "type_label": ev["type_label"],
                "module": ev["module"],
                "color": colors.get(ev["module"], "#6a7180"),
                "icon": ev["icon"],
                "date": format_jalali(d),
                "title": ev.get("title") or "",
                "status": ev.get("status") or "",
                "url": ev["url"],
            })

    # Pump-test step curves (Q vs drawdown / efficiency) — well-performance diagnostic.
    from app.models.pump_test import PumpTest
    from app.models.flow import FlowTest
    from app.utils.dates import format_jalali as _fj
    pump_tests = []
    for t in db.session.scalars(
            db.select(PumpTest).filter_by(well_id=well.id).order_by(PumpTest.test_date)).all():
        pts = sorted(
            [{"q": s.discharge_lps, "dd": s.observed_drawdown,
              "eff": (s.efficiency * 100 if s.efficiency and s.efficiency <= 1 else s.efficiency)}
             for s in t.steps if s.discharge_lps is not None],
            key=lambda x: x["q"])
        if pts:
            pump_tests.append({"date": _fj(t.test_date), "points": pts})

    # Flow-metering operating points (Q vs dynamic level), across tests over time.
    flow_tests = []
    for ft in db.session.scalars(
            db.select(FlowTest).filter_by(well_id=well.id).order_by(FlowTest.test_date)).all():
        pts = [{"q": p.discharge_lps, "dyn": p.dynamic_level_m, "head": p.head_m,
                "eff": (p.efficiency * 100 if p.efficiency and p.efficiency <= 1 else p.efficiency)}
               for p in ft.points if p.discharge_lps is not None]
        if pts:
            flow_tests.append({"date": _fj(ft.test_date), "points": pts})

    # Pump characteristic curve (catalog) + design point overlay for the Q-H chart.
    from app.models.pump_select import PumpSelection
    from app.models.pump_catalog import PumpModel
    from app.models.pump_asset import PumpInstallation
    sels = db.session.scalars(
        db.select(PumpSelection).filter_by(well_id=well.id)
        .order_by(PumpSelection.form_delivery_date)).all()
    dp = next((s for s in reversed(sels)
               if s.target_discharge_lps and s.selected_head_m), None)
    design_point = ({"q": dp.target_discharge_lps, "h": dp.selected_head_m,
                     "pump": dp.selected_pump_type} if dp else None)
    model_name = (dp.selected_pump_type if dp else None) or next(
        (s.selected_pump_type for s in reversed(sels) if s.selected_pump_type), None)
    if not model_name:
        inst = db.session.scalar(
            db.select(PumpInstallation).filter_by(well_id=well.id)
            .order_by(PumpInstallation.install_date.desc()))
        model_name = inst.pump_type if inst else None
    pump_curve = None
    if model_name:
        pm = db.session.scalar(db.select(PumpModel).filter_by(model=model_name))
        if pm and pm.points:
            pts = [{"q": p.flow_lps, "h": p.head_m, "eff": p.efficiency_pct}
                   for p in pm.points if p.flow_lps is not None and p.head_m is not None]
            if pts:
                bep = next((p for p in pm.points if p.is_bep), None)
                bs = next((p for p in pm.points if p.is_beb_start), None)
                be = next((p for p in pm.points if p.is_beb_end), None)
                pump_curve = {
                    "model": pm.model,
                    "points": pts,
                    "bep": ({"q": bep.flow_lps, "h": bep.head_m, "eff": bep.efficiency_pct}
                            if bep and bep.flow_lps is not None else None),
                    "beb": ({"start": bs.flow_lps, "end": be.flow_lps}
                            if bs and be and bs.flow_lps is not None and be.flow_lps is not None else None),
                }

    has_qh = (any(p["head"] is not None for ft in flow_tests for p in ft["points"])
              or design_point is not None or pump_curve is not None)

    from app.blueprints.documents.routes import attachments_for
    attachments = attachments_for("well", well.id)

    from app.services.analytics import summary as analytics_summary
    analytics = analytics_summary.well(well.id)

    # 6.2 pump-replacement ROI (uses energy saving + economic params)
    pump_roi = None
    eo = analytics.get("energy")
    if eo and eo.get("annual_rial_saving") and eo.get("motor_kw"):
        from app.services import economics
        economics.ensure_params()
        invest = (economics.get("pump_price_per_kw") * eo["motor_kw"]
                  + economics.get("pump_install_cost"))
        saving = eo["annual_rial_saving"]
        npv = economics.npv(saving, economics.get("analysis_years"),
                            economics.get("discount_rate")) - invest
        pump_roi = {
            "investment": invest, "annual_saving": saving,
            "payback": economics.payback_years(invest, saving),
            "npv": round(npv), "worth_it": npv > 0,
            "years": int(economics.get("analysis_years")),
        }

    # Groundwater level monitoring (piezometry) — static-level time series.
    from app.models.water_level import WaterLevelLog
    wl_rows = db.session.scalars(
        db.select(WaterLevelLog).filter_by(well_id=well.id)
        .order_by(WaterLevelLog.measure_date)).all()
    water_levels = [
        {"id": w.id, "date": _fj(w.measure_date), "level": w.static_level,
         "method": w.method, "pumping": w.is_pumping}
        for w in wl_rows if w.measure_date and w.static_level is not None]

    # Vertical cross-section (5.6): depths + water levels for the schematic.
    from app.models.drilling import Drilling
    from app.models.pump_asset import PumpInstallation
    from app.models.flow import FlowTestPoint
    tech = well.technical
    dr = db.session.scalar(db.select(Drilling).filter_by(well_id=well.id)
                           .order_by(Drilling.end_date.desc()))
    inst = db.session.scalar(db.select(PumpInstallation).filter_by(well_id=well.id)
                             .order_by(PumpInstallation.install_date.desc()))
    drill_depth = ((dr.well_depth_actual or dr.well_depth_permit) if dr else None) \
        or (tech.drill_depth_m if tech else None) or (inst.well_depth_m if inst else None)
    install_depth = (inst.install_depth_m if inst else None) \
        or (tech.install_depth_m if tech else None)
    static = next((w.static_level for w in reversed(wl_rows)
                   if not w.is_pumping and w.static_level is not None), None)
    dyn = db.session.scalar(
        db.select(FlowTestPoint.dynamic_level_m).join(FlowTest)
        .where(FlowTest.well_id == well.id, FlowTestPoint.dynamic_level_m.isnot(None))
        .order_by(FlowTest.test_date.desc()))
    xsection = None
    if drill_depth or install_depth:
        xsection = {
            "ground_elev": well.ground_elevation,
            "drill_depth": drill_depth, "install_depth": install_depth,
            "static": static, "dynamic": dyn,
            "casing": tech.casing_material if tech else None,
            "construction": well.construction_type,
            "max_depth": max(d for d in [drill_depth, install_depth, static, dyn, 1] if d),
        }

    return render_template(
        "wells/detail.html",
        well=well,
        timeline=events,
        chart_events=chart_events,
        baseline=baseline,
        quality=quality,
        production=production,
        energy=energy,
        pump_tests=pump_tests,
        flow_tests=flow_tests,
        has_qh=has_qh,
        pump_curve=pump_curve,
        design_point=design_point,
        water_levels=water_levels,
        attachments=attachments,
        analytics=analytics,
        pump_roi=pump_roi,
        xsection=xsection,
        kind_labels=dict(WELL_KINDS),
        status_labels=dict(WELL_STATUSES),
        loc_labels=dict(LOCATION_STATUSES),
        constr_labels=dict(WELL_CONSTRUCTION_TYPES),
    )


def _apply_form(form, well):
    well.pm_code = form.pm_code.data.strip()
    well.name = form.name.data.strip()
    well.well_kind = form.well_kind.data
    well.office_id = form.office_id.data or None
    well.center_id = well.office_id   # اداره و مرکز یکی هستند
    well.zone = (form.zone.data or "").strip() or None
    well.sub_zone = (form.sub_zone.data or "").strip() or None
    well.utm_x = form.utm_x.data
    well.utm_y = form.utm_y.data
    well.utm_zone = (form.utm_zone.data or "").strip() or None
    well.latitude = form.latitude.data
    well.longitude = form.longitude.data
    well.ground_elevation = form.ground_elevation.data
    well.drill_year = (form.drill_year.data or "").strip() or None
    well.location_status = form.location_status.data or None
    well.status = form.status.data
    well.parent_well_id = form.parent_well_id.data or None
    well.notes = (form.notes.data or "").strip() or None


@bp.route("/new", methods=["GET", "POST"])
@permission_required("wells", "create")
def create_well():
    form = WellForm()
    form.populate_choices()
    if form.validate_on_submit():
        existing = db.session.scalar(
            db.select(Well).filter_by(pm_code=form.pm_code.data.strip())
        )
        if existing:
            flash("چاهی با این کد PM از قبل وجود دارد.", "danger")
        else:
            well = Well()
            _apply_form(form, well)
            well.created_by_id = current_user.id
            db.session.add(well)
            db.session.commit()
            flash("چاه ثبت شد.", "success")
            return redirect(url_for("wells.detail", well_id=well.id))
    return render_template("wells/form.html", form=form, title="ثبت چاه جدید")


@bp.route("/<int:well_id>/edit", methods=["GET", "POST"])
@permission_required("wells", "edit")
def edit_well(well_id):
    well = db.get_or_404(Well, well_id)
    form = WellForm(obj=well)
    form.populate_choices(exclude_well_id=well.id)
    if form.validate_on_submit():
        clash = db.session.scalar(
            db.select(Well).where(
                Well.pm_code == form.pm_code.data.strip(), Well.id != well.id
            )
        )
        if clash:
            flash("چاه دیگری با این کد PM وجود دارد.", "danger")
        else:
            _apply_form(form, well)
            well.updated_by_id = current_user.id
            db.session.commit()
            flash("چاه به‌روزرسانی شد.", "success")
            return redirect(url_for("wells.detail", well_id=well.id))
    return render_template("wells/form.html", form=form, title="ویرایش چاه")


@bp.route("/<int:well_id>/delete", methods=["POST"])
@permission_required("wells", "delete")
def delete_well(well_id):
    well = db.get_or_404(Well, well_id)
    db.session.delete(well)
    db.session.commit()
    flash("چاه حذف شد.", "info")
    return redirect(url_for("wells.list_wells"))


