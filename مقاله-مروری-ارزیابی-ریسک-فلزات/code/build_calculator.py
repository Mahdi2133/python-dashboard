# -*- coding: utf-8 -*-
"""Excel calculator for the integrated aggregate-exposure equation (live formulas), prefilled with the Honghu Lake example."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
import ma_data as D
from ma_engine import combine_groups

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Integrated_HRA_Calculator.xlsx'
HEAD = PatternFill('solid', fgColor='DCE6F2')
INPUT = PatternFill('solid', fgColor='FFF7D6')
BOLD = Font(bold=True)
thin = Side(style='thin', color='B0B0B0')
BOX = Border(bottom=thin)
wb = Workbook()


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = BOLD; cell.fill = HEAD; cell.alignment = Alignment(wrap_text=True, vertical='top')


def widths(ws, w):
    for i, x in enumerate(w):
        ws.column_dimensions[chr(65 + i)].width = x


# ------------------------------------------------------------------ README
ws = wb.active; ws.title = 'README'
lines = [
    ('Integrated aggregate-exposure calculator for metals (water + diet + air)', True),
    ('Implements Equations 1-9 of the revised manuscript (Section 4). Yellow cells are inputs; all other numbers are live formulas.', False),
    ('Prefilled with the Honghu Lake illustration (Zhang et al. 2017; lake water + fish; no air). Replace with your own co-located data.', False),
    ('', False),
    ('Sheets', True),
    ('Inputs      - population exposure parameters (BW, water intake, EF, ED, lifetime, inhalation rate, route-to-route factor).', False),
    ('TRV         - toxicity reference values (verified Oct 2026). Check the live agency entry before publication.', False),
    ('Water_Air   - concentration in treated tap water (µg/L) and ambient PM10 (ng/m3) for each metal.', False),
    ('Food        - one row per metal x food group: concentration as eaten (µg/kg wet wt), intake (kg/d), species fraction S, relative bioavailability RBA.', False),
    ('Results     - route doses, HQs, aggregate HI, source shares, MOE, ILCR and context-specific water value per metal.', False),
    ('', False),
    ('Rules', True),
    ('1. All media must describe the same population, place and period.', False),
    ('2. Food concentrations must be wet weight as eaten: C_ww = C_dw x (1 - moisture).', False),
    ('3. For As enter inorganic As (or total As with S = iAs/tAs); for Cr choose CrVI or CrIII (or run both); for fish Hg use MeHg.', False),
    ('4. Pb and iAs: read the MOE columns (no threshold). MOE <= 1 = concern; for Pb SBP/CKD, MOE >= 10 reassuring.', False),
    ('5. Non-detects: LOD/2 only if < 15% ND; otherwise estimate the mean by Kaplan-Meier/ROS before entering.', False),
    ('6. This sheet is deterministic. For P95 and Pr(HI>1) use the Python Monte Carlo script (integrated_hra.py).', False),
    ('', False),
    ('راهنمای فارسی', True),
    ('سلول‌های زرد ورودی هستند. داده‌های آب، غذا و هوا باید از یک جمعیت، یک مکان و یک بازه زمانی باشند.', False),
    ('غلظت غذا باید بر پایه وزن تر (همان‌طور که خورده می‌شود) باشد. برای آرسنیک فقط شکل معدنی و برای کروم شکل شش‌ظرفیتی/سه‌ظرفیتی را جدا وارد کنید.', False),
    ('برای سرب و آرسنیک معدنی به ستون MOE نگاه کنید؛ HI برای فلزات آستانه‌دار است. ستون سهم منابع (f) سهم واقعی آب/غذا/هوا را نشان می‌دهد.', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = Font(bold=b, size=13 if i == 1 else 11)
ws.column_dimensions['A'].width = 140

# ------------------------------------------------------------------ Inputs
ws = wb.create_sheet('Inputs')
header(ws, 1, ['Parameter', 'Symbol', 'Value', 'Unit', 'Note / source'])
inputs = [
    ('Body weight', 'BW', 61.6, 'kg', 'Population-specific (Honghu: Zhang 2017). WHO default 60; EPA 80'),
    ('Drinking-water intake', 'IRw', 2.0, 'L/d', 'WHO default 2 L/d'),
    ('Exposure frequency', 'EF', 365, 'd/y', 'EPA residential 350'),
    ('Exposure duration', 'ED', 30, 'y', 'EPA residential 26'),
    ('Averaging time, threshold effects', 'AT_nc', '=C5*365', 'd', 'ED x 365'),
    ('Lifetime', 'LT', 70, 'y', 'EPA 70 y'),
    ('Averaging time, cancer', 'AT_c', '=C7*365', 'd', 'LT x 365'),
    ('Fraction of day exposed to ambient air', 'ET', 1, '-', '24 h / 24 h residential'),
    ('Inhalation rate (MOE only)', 'InhR', 20, 'm3/d', 'Adult default'),
    ('Route-to-route factor F_inh/F_oral (MOE only)', 'RR', 2, '-', 'ECHA R.8 default'),
    ('Allocation floor', 'P_min', 0.2, '-', 'WHO/EPA'),
    ('Allocation ceiling', 'P_max', 0.8, '-', 'WHO/EPA'),
]
for i, row in enumerate(inputs, 2):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        if j == 3 and not str(v).startswith('='):
            c.fill = INPUT
widths(ws, [42, 10, 10, 8, 60])
# named-like absolute references
REF = {k: "Inputs!$C$%d" % (i + 2) for i, (_, k, *_r) in enumerate(inputs)}

# ------------------------------------------------------------------ TRV
ws = wb.create_sheet('TRV')
header(ws, 1, ['Metal key', 'RfD water (µg/kg/d)', 'RfD food (µg/kg/d)', 'RfC (µg/m3)', 'Oral CSF (per µg/kg/d)',
               'IUR (per µg/m3)', 'BMDL 1 (µg/kg/d)', 'BMDL 1 end-point', 'BMDL 2 (µg/kg/d)', 'BMDL 2 end-point',
               'Target organ', 'Source'])
trv = [
    ('iAs', 0.06, 0.06, 0.015, 0.032, 4.3e-3, 0.06, 'skin cancer BMDL05 (EFSA 2024)', None, None, 'skin/vascular',
     'IRIS 2025 RfD 6e-5 mg/kg-d, CSF 32 per mg/kg-d (linear only <0.2 µg/kg/d); OEHHA REL; IRIS IUR'),
    ('Cd', 0.5, 1.0, 0.01, None, 1.8e-3, None, None, None, None, 'kidney', 'IRIS (water/food); ATSDR inhalation MRL; IRIS IUR'),
    ('CrVI', 0.9, 0.9, 0.03, 0.26e-3, 1.8e-2, None, None, None, None, 'GI tract', 'IRIS 2024 (OSF lifetime with ADAFs)'),
    ('CrIII', 1500, 1500, None, None, None, None, None, None, None, 'none', 'IRIS'),
    ('Cu', 40, 40, None, None, None, None, None, None, None, 'GI/liver', 'HEAST (EFSA ADI 70)'),
    ('Pb', None, None, None, 8.5e-6, 1.2e-5, 1.5, 'SBP BMDL01 (EFSA 2010)', 0.63, 'CKD BMDL10 (EFSA 2010)', 'kidney/CV/CNS',
     'No RfD; children: neurodevelopment BMDL01 0.50; OEHHA CSF/IUR'),
    ('MeHg', 0.1, 0.1, None, None, None, None, None, None, None, 'nervous system', 'IRIS 2001'),
    ('Ni', 20, 20, 0.09, None, None, None, None, None, None, 'reproductive', 'IRIS soluble salts; ATSDR inhalation MRL (EFSA TDI 13)'),
]
for i, row in enumerate(trv, 2):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        if j > 1: c.fill = INPUT
widths(ws, [10, 12, 12, 11, 13, 12, 12, 26, 12, 22, 14, 70])
NT = len(trv) + 1

# ------------------------------------------------------------------ Water & air
ws = wb.create_sheet('Water_Air')
header(ws, 1, ['Metal key', 'C water (µg/L)', 'C air PM10 (ng/m3)', 'Note'])
wa = [('iAs', 0.990, None, 'Honghu Lake water (assumed all inorganic)'), ('Cd', 0.140, None, ''),
      ('CrVI', 1.630, None, 'upper bound: all Cr as Cr(VI)'), ('CrIII', 1.630, None, 'lower bound: all Cr as Cr(III)'),
      ('Cu', 3.090, None, ''), ('Pb', 3.420, None, '')]
for i, row in enumerate(wa, 2):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        if j in (2, 3): c.fill = INPUT
widths(ws, [10, 14, 16, 45])
NM = len(wa) + 1

# ------------------------------------------------------------------ Food (long format)
ws = wb.create_sheet('Food')
header(ws, 1, ['Metal key', 'Food group', 'C as eaten (µg/kg ww)', 'Intake IR (kg/d)', 'Species fraction S', 'RBA',
               'C·S·RBA·IR (µg/d)'])
fish = {}
for m, i in dict(As=0, Cd=1, Cr=2, Cu=3, Pb=4).items():
    n, mean, sd = combine_groups([(g[0], g[1 + i][0], g[1 + i][1]) for g in D.S15_GROUPS.values()])
    fish[m] = mean * 1000
food = [('iAs', 'fish muscle', fish['As'], 0.05433, 1, 1), ('Cd', 'fish muscle', fish['Cd'], 0.05433, 1, 1),
        ('CrVI', 'fish muscle', fish['Cr'], 0.05433, 1, 1), ('CrIII', 'fish muscle', fish['Cr'], 0.05433, 1, 1),
        ('Cu', 'fish muscle', fish['Cu'], 0.05433, 1, 1), ('Pb', 'fish muscle', fish['Pb'], 0.05433, 1, 1)]
NF = 200
for i in range(2, NF + 2):
    if i - 2 < len(food):
        for j, v in enumerate(food[i - 2], 1):
            ws.cell(row=i, column=j, value=round(v, 4) if isinstance(v, float) else v)
    for j in range(1, 7):
        ws.cell(row=i, column=j).fill = INPUT
    ws.cell(row=i, column=7, value='=IF(A%d="","",C%d*IF(E%d="",1,E%d)*IF(F%d="",1,F%d)*D%d)' % ((i,) * 7))
widths(ws, [10, 22, 18, 14, 14, 8, 16])

# ------------------------------------------------------------------ Results
ws = wb.create_sheet('Results')
cols = ['Metal key', 'D_w (µg/kg/d)', 'D_f (µg/kg/d)', 'EC_a (µg/m3)', 'HQ water', 'HQ food', 'HQ air', 'HI_agg',
        'share water', 'share food', 'share air', 'D_a,eq (µg/kg/d)', 'MOE (BMDL 1)', 'MOE (BMDL 2)', 'ILCR oral',
        'ILCR inhalation', 'ILCR total', 'P_w (allocation)', 'C*_w context water value (µg/L)', 'Target organ']
header(ws, 1, cols)
BW, IRw, EF, ED, ATn, ATc, ET, InhR, RR, Pmin, Pmax = (REF[k] for k in ('BW', 'IRw', 'EF', 'ED', 'AT_nc', 'AT_c', 'ET', 'InhR', 'RR', 'P_min', 'P_max'))
TF = '(%s*%s/%s)' % (EF, ED, ATn)
LTF = '(%s/%s)' % (ATn, ATc)
for r in range(2, NM + 1):
    m = 'Water_Air!A%d' % r
    lk = lambda col: 'INDEX(TRV!$%s$2:$%s$%d,MATCH($A%d,TRV!$A$2:$A$%d,0))' % (col, col, NT, r, NT)
    f = {
        'A': '=%s' % m,
        'B': '=Water_Air!B%d*%s/%s*%s' % (r, IRw, BW, TF),
        'C': '=SUMIF(Food!$A$2:$A$%d,$A%d,Food!$G$2:$G$%d)/%s*%s' % (NF + 1, r, NF + 1, BW, TF),
        'D': '=N(Water_Air!C%d)/1000*%s*%s' % (r, ET, TF),
        'E': '=IF(N(%s)>0,B%d/%s,"")' % (lk('B'), r, lk('B')),
        'F': '=IF(N(%s)>0,C%d/%s,"")' % (lk('C'), r, lk('C')),
        'G': '=IF(N(%s)>0,D%d/%s,"")' % (lk('D'), r, lk('D')),
        'H': '=IF(COUNT(E%d:G%d)=0,"",SUM(E%d:G%d))' % (r, r, r, r),
        'I': '=IF(H%d="","",N(E%d)/H%d)' % (r, r, r),
        'J': '=IF(H%d="","",N(F%d)/H%d)' % (r, r, r),
        'K': '=IF(H%d="","",N(G%d)/H%d)' % (r, r, r),
        'L': '=N(Water_Air!C%d)/1000*%s/%s*%s*%s' % (r, InhR, BW, TF, RR),
        'M': '=IF(N(%s)>0,%s/(B%d+C%d+L%d),"")' % (lk('G'), lk('G'), r, r, r),
        'N': '=IF(N(%s)>0,%s/(B%d+C%d+L%d),"")' % (lk('I'), lk('I'), r, r, r),
        'O': '=(B%d+C%d)*%s*N(%s)' % (r, r, LTF, lk('E')),
        'P': '=D%d*%s*N(%s)' % (r, LTF, lk('F')),
        'Q': '=O%d+P%d' % (r, r),
        'R': '=IF(N(%s)>0,MIN(%s,MAX(%s,1-N(F%d)-N(G%d))),"")' % (lk('B'), Pmax, Pmin, r, r),
        'S': '=IF(R%d="","",%s*%s*R%d/(%s*%s))' % (r, lk('B'), BW, r, IRw, TF),
        'T': '=%s' % lk('K'),
    }
    for col, v in f.items():
        ws['%s%d' % (col, r)] = v
for col in 'BCDEFGHLMNOPQRS':
    for r in range(2, NM + 1):
        ws['%s%d' % (col, r)].number_format = '0.000E+00' if col in 'OPQ' else '0.0000'
for col in 'IJKR':
    for r in range(2, NM + 1):
        ws['%s%d' % (col, r)].number_format = '0%'
widths(ws, [10] + [12] * 18 + [14])
ws.freeze_panes = 'B2'
r0 = NM + 3
ws.cell(row=r0, column=1, value='Cumulative screening HI (all threshold metals, Cr counted once as CrVI):').font = BOLD
ws.cell(row=r0, column=8, value='=SUMPRODUCT((A2:A%d<>"CrIII")*N(+H2:H%d))' % (NM, NM))
ws.cell(row=r0 + 1, column=1, value='Target-organ HI (kidney):').font = BOLD
ws.cell(row=r0 + 1, column=8, value='=SUMIF(T2:T%d,"kidney",H2:H%d)' % (NM, NM))
ws.cell(row=r0 + 3, column=1, value='Read HI for threshold metals; MOE for Pb and iAs (concern if <= 1); ILCR acceptable range 1E-6 to 1E-4. '
        'Shares = empirical source contribution f_m (Eq. 8). C*_w = context-specific drinking-water value (Eq. 9).')
wb.save(OUT)
print('saved', OUT)
