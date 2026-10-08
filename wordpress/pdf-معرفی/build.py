# -*- coding: utf-8 -*-
"""سندِ «معماری متابولیک» را می‌سازد — آماده‌ی تبدیل به PDF.

فونت‌ها: فارسی B Nazanin، لاتین Times New Roman.
هر دو با نام در فایل ثبت می‌شوند؛ ورد روی ویندوز خودش آن‌ها را برمی‌دارد.
"""

import os, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from rtl import (rtl_para, run, cell_shade, borders, cell_borders, table_rtl,
                 cell_valign, no_space, section_rtl, fix_zoom, fix_table,
                 _el, FA, EN)

IMG = os.path.join(HERE, 'img')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'out.docx')

NAVY, NAVY2 = '10161C', '1B2A3A'
GOLD, GOLD_L = '8F6E2C', 'C6A05A'
PAPER, RULE = 'F7F4ED', 'DCD5C7'
INK2, INK3 = '3A4752', '66717B'

W_PAGE = Cm(17.2)          # عرضِ مفیدِ صفحه: ۲۱ − ۱٫۹ − ۱٫۹
LH = 1.35                  # فاصله‌ی خطوط برای متنِ بدنه

FA_DIGITS = str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹')
def fa(n): return str(n).translate(FA_DIGITS)

doc = Document()
fix_zoom(doc)

s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.3)
s.left_margin, s.right_margin = Cm(1.9), Cm(1.9)
s.header_distance, s.footer_distance = Cm(0.8), Cm(0.8)
section_rtl(s)

st = doc.styles['Normal']
st.font.name = EN
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:cs'), FA)
st.element.rPr.rFonts.set(qn('w:eastAsia'), EN)
st.element.rPr.append(_el('w:szCs', val='24'))

# سربرگ
hp = s.header.paragraphs[0]
rtl_para(hp, AL.RIGHT, space_after=0, line=1)
hp.add_run().add_picture(os.path.join(IMG, 'logo-header.png'), width=Cm(4.2))
borders(hp._p, bottom=(RULE, 6))

# پاصفحه
fp = s.footer.paragraphs[0]
rtl_para(fp, AL.CENTER, space_after=0, line=1)
run(fp, 'saadatmehracademy.com', size=9, color=GOLD, bold=True)
run(fp, '   ·   آکادمی تندرستی سعادت‌مهر', size=9, color=INK3)


def H1(title, kicker, sub):
    p = doc.add_paragraph(); rtl_para(p, AL.RIGHT, space_after=1, line=1.1)
    run(p, title, size=23, bold=True, color=NAVY)
    k = doc.add_paragraph(); rtl_para(k, AL.RIGHT, space_after=3, line=1.2)
    run(k, kicker, size=13, bold=True, color=GOLD)
    q = doc.add_paragraph(); rtl_para(q, AL.RIGHT, space_after=8, line=1.3)
    run(q, sub, size=11, color=INK2)
    borders(q._p, bottom=(GOLD, 10))

def H2(num, txt):
    p = doc.add_paragraph()
    rtl_para(p, AL.RIGHT, space_after=4, space_before=11, line=1.15, keep_next=True)
    run(p, f'{num}  ', size=14, bold=True, color=GOLD)
    run(p, txt, size=14, bold=True, color=NAVY)

def body(txt, after=4, size=11.5, align=AL.JUSTIFY):
    p = doc.add_paragraph(); rtl_para(p, align, space_after=after, line=LH)
    run(p, txt, size=size)
    return p

def bullet(txt, after=2, label=None):
    p = doc.add_paragraph()
    rtl_para(p, AL.JUSTIFY, space_after=after, line=LH, indent_r=Cm(0.5))
    p.paragraph_format.first_line_indent = Cm(-0.5)
    run(p, '◗  ', size=10, color=GOLD, bold=True)
    if label: run(p, label, size=11.5, bold=True, color=NAVY2)
    run(p, txt, size=11.5)
    return p

def box(fill=PAPER, width=W_PAGE):
    """یک جعبه‌ی تک‌خانه‌ای با نوارِ طلاییِ کنار."""
    t = doc.add_table(rows=1, cols=1)
    fix_table(t, [width])
    c = t.rows[0].cells[0]; cell_shade(c, fill); no_space(c)
    cell_borders(c, end=(GOLD, 18), top=(RULE, 4), bottom=(RULE, 4), start=(RULE, 4))
    return c

def gap(pt=4):
    doc.add_paragraph().paragraph_format.space_after = Pt(pt)


# ======================= صفحه ۱ =======================
H1('معماری متابولیک',
   'پروتکل ۱۰۰ روزه SMP',
   'راهنمای استراتژیک بازیابی انرژی، تمرکز و پایداری فیزیولوژیک در محیط‌های پرفشار')

ip = doc.add_paragraph(); rtl_para(ip, AL.CENTER, space_after=2, space_before=1, line=1)
ip.add_run().add_picture(os.path.join(IMG, 'cover.png'), width=Cm(7.4))
cap = doc.add_paragraph(); rtl_para(cap, AL.CENTER, space_after=4, line=1.15)
run(cap, 'نقشه‌ی نقاط کلیدی متابولیک: نور، خواب، تغذیه، حرکت و سیستم عصبی',
    size=9, color=INK3)

H2('۱.', 'پارادوکس عملکرد در محیط‌های پرفشار')
body('شما پیچیده‌ترین تصمیم‌های مالی، فنی یا سازمانی را هدایت می‌کنید؛ '
     'اما بیولوژی بدن شما همگام با این بار شناختی پیش نمی‌آید:')
bullet('افت تمرکز بعدازظهر (ساعات ۱۴ تا ۱۷) و تکیه مداوم بر کافئین و قند.')
bullet('بی‌اثر شدن فرمول‌های کلاسیک (تمرینات سنگین، کم‌خوابی و رژیم‌های فرساینده) '
       'در دهه‌های ۳۰ و ۴۰ زندگی.')
bullet('کمبود وقت مطلق برای آشپزی اختصاصی، شمارش کالری یا تمرین‌های روزانه طولانی.',
       after=6)

c = box()
p = c.paragraphs[0]; rtl_para(p, AL.RIGHT, space_after=0, line=1.3)
run(p, 'اصل ماجرا:  ', size=11.5, bold=True, color=GOLD)
run(p, 'مشکل، اراده شما نیست؛ اجرای یک استراتژی منسوخ روی یک بیولوژی مدرن است.',
    size=11.5, bold=True, color=NAVY)

H2('۲.', 'رویکرد حداقل اصطکاک (Zero-Friction)')
body('پروتکل SMP نیازی به خالی کردن وقت در تقویم ندارد؛ بلکه مثل یک لایه نرم‌افزاری، '
     'مستقیم روی سبک زندگی فعلی شما سوار می‌شود:')
for lab, txt in [
    ('نور (هوشیاری پایدار): ', 'تنظیم نور طبیعی صبح و فیلترهای شبانه، برای حذف افت '
     'تمرکز بعدازظهر و خواب عمیق در شب، بدون مصرف کافئین افراطی.'),
    ('خواب (ریکاوری مغزی): ', 'افزایش کیفیت خواب عمیق برای اینکه صبح‌ها بدون حس گیجی '
     '(Brain Fog) بیدار شوید و نیاز به زنگ‌های مکرر آلارم نداشته باشید.'),
    ('غذا (تغذیه گلوکز مغز): ', 'بدون رژیم‌های سخت و زمان‌بر؛ فقط اصلاح زمان‌بندی '
     'وعده‌ها برای جلوگیری از سقوط ناگهانی انرژی و تمایل شدید به قند حین کار.'),
    ('حرکت (مداخلات ریز): ', 'تمرین‌های ۲ تا ۳ دقیقه‌ای کنار میز کار یا حین تماس‌های '
     'صوتی برای تخلیه بار استرس و پیشگیری از قفل شدن گردن و کمر، بدون نیاز به باشگاه.'),
    ('سیستم عصبی (مدیریت استرس): ', 'تکنیک‌های تنفسی و بیولوژیک ۳۰ ثانیه‌ای برای '
     'جلوگیری از ترشح اضافه کورتیزول و حفظ تصمیم‌گیری شفاف در میانه ددلاین‌های بحرانی.'),
]:
    bullet(txt, label=lab, after=3)


# ======================= صفحه ۲ =======================
H2('۳.', 'معماری اجرایی: ۱۰۰ روز در سه فاز متوالی')
body('ورود و خروج از هر فاز مبتنی بر «گیت‌های ارزیابی بیولوژیک» است، نه حدس و گمان:',
     after=6)

PH = [
    ('فاز ۱', 'روز ۱ تا ۳۰', 'Reset', 'حذف اصطکاک',
     'تخلیه استرس سیستمیک، قطع نوسانات قند خون، و بازیابی انرژی پایدار در ساعات کاری.',
     'گیت روز ۳۰: حساسیت انسولینی، کیفیت خواب عمیق، ثبات انرژی.'),
    ('فاز ۲', 'روز ۳۱ تا ۷۰', 'Activation', 'فعال‌سازی',
     'افزایش انعطاف‌پذیری متابولیک (سوخت‌رسانی پایدار به مغز)، ریکاوری سریع‌تر بعد از '
     'روزهای سنگین کاری و بهبود وضوح شناختی.',
     'گیت روز ۷۰: ترکیب بدنی، نرخ اکسیداسیون چربی، توان عملکردی.'),
    ('فاز ۳', 'روز ۷۱ تا ۱۰۰', 'Optimization', 'تثبیت و بهینه‌سازی',
     'سیستمی‌سازی و تثبیت عادات در تقویم شخصی، به‌طوری که حفظ این سطح از انرژی خودکار '
     'و بدون تلاش اضافه باشد.',
     'گیت روز ۱۰۰: شناسنامه متابولیک اختصاصی و استقلال کامل.'),
]
COLW = Cm(5.733)
tb = doc.add_table(rows=2, cols=3)
fix_table(tb, [COLW]*3)
table_rtl(tb)
for i, (ph, days, en, fa_t, desc, gate) in enumerate(PH):
    hc = tb.rows[0].cells[i]
    cell_shade(hc, NAVY); no_space(hc); cell_valign(hc)
    cell_borders(hc, start=('FFFFFF', 12), end=('FFFFFF', 12))
    p = hc.paragraphs[0]; rtl_para(p, AL.CENTER, space_after=0, space_before=3, line=1.1)
    run(p, f'{ph}  ·  {days}', size=10.5, bold=True, color='FFFFFF')
    q = hc.add_paragraph(); rtl_para(q, AL.CENTER, space_after=3, line=1.1)
    run(q, en, size=9, bold=True, color=GOLD_L)

    bc = tb.rows[1].cells[i]
    cell_shade(bc, 'FFFFFF'); no_space(bc)
    cell_borders(bc, start=('FFFFFF', 12), end=('FFFFFF', 12),
                 bottom=(RULE, 6), top=(RULE, 2))
    p = bc.paragraphs[0]; rtl_para(p, AL.RIGHT, space_after=2, space_before=4, line=1.2)
    run(p, fa_t, size=11, bold=True, color=GOLD)
    q = bc.add_paragraph(); rtl_para(q, AL.JUSTIFY, space_after=3, line=1.3)
    run(q, desc, size=10)
    g = bc.add_paragraph(); rtl_para(g, AL.RIGHT, space_after=4, line=1.25)
    run(g, gate, size=9, bold=True, color=INK2)

gap(3)

# ---------- ۴. چک‌لیست ----------
H2('۴.', 'خودارزیابی سریع: وضعیت اصطکاک بیولوژیک شما')
body('هر موردی که وضعیت شما را توصیف می‌کند علامت بزنید:', after=5)

ITEMS = [
    'افت وضوح ذهن، سنگینی یا ولع کافئین/شیرینی بین ساعات ۱۴ تا ۱۷.',
    'بیدار شدن با حس خستگی یا بی‌رمقی با وجود ساعات خواب به ظاهر کافی.',
    'افت شدید تمرکز و تحریک‌پذیری عصبی هنگام فاصله افتادن بین وعده‌ها.',
    'پاسخ ندادن ترکیب بدنی به تمرینات ورزشی و رژیم‌های کلاسیک قبلی.',
    'نفخ، سنگینی مداوم گوارشی یا تحریک‌پذیری ناشی از فشارهای کاری.',
]
WN, WB, WT = Cm(0.9), Cm(0.8), Cm(15.5)
ck = doc.add_table(rows=len(ITEMS), cols=3)
fix_table(ck, [WN, WB, WT])
table_rtl(ck)
for i, txt in enumerate(ITEMS):
    r = ck.rows[i]
    c0 = r.cells[0]; no_space(c0); cell_valign(c0)
    p = c0.paragraphs[0]; rtl_para(p, AL.CENTER, space_after=2, space_before=2, line=1.15)
    run(p, fa(i + 1), size=11, bold=True, color=GOLD)
    # مربعِ تیک: خانه‌ی خالیِ کادردار، نه نویسه‌ی ☐ — آن نویسه در
    # بسیاری از فونت‌های فارسی نیست و به مربعِ «فونت پیدا نشد» بدل می‌شود.
    c1 = r.cells[1]; no_space(c1); cell_valign(c1)
    cell_borders(c1, top=(NAVY2, 8), bottom=(NAVY2, 8),
                 start=(NAVY2, 8), end=(NAVY2, 8))
    p = c1.paragraphs[0]; rtl_para(p, AL.CENTER, space_after=2, space_before=2, line=1)
    run(p, ' ', size=10)
    c2 = r.cells[2]; no_space(c2); cell_valign(c2)
    p = c2.paragraphs[0]
    rtl_para(p, AL.RIGHT, space_after=2, space_before=2, line=1.25, indent_r=Cm(0.22))
    run(p, txt, size=11)

res = doc.add_paragraph()
rtl_para(res, AL.JUSTIFY, space_after=3, space_before=6, line=LH)
run(res, 'نتیجه:  ', size=11.5, bold=True, color=GOLD)
run(res, '۲ تا ۳ علامت یعنی «اصطکاک بیولوژیک آغاز شده»؛ ۴ تا ۵ علامت یعنی '
         '«بدن در فاز دفاعی است و بخشی از ظرفیت تصمیم‌گیری شما خرج مهار تنش می‌شود».',
    size=11.5)

# ---------- ۵. پشتوانه علمی ----------
H2('۵.', 'پشتوانه علمی و اعتبار بالینی')
c = box()
p = c.paragraphs[0]; rtl_para(p, AL.JUSTIFY, space_after=0, line=1.3)
run(p, 'این پروتکل برآمده از کتاب ', size=11)
run(p, 'معماری متابولیک', size=11, bold=True, color=NAVY)
run(p, ' به قلم حسین سعادت‌مهر و یاسمن سعادت‌مهر، با پیشگفتار ', size=11)
run(p, 'دکتر رودیگر کرش', size=11, bold=True, color=NAVY)
run(p, ' (مدیر ارشد سازمان جهانی بهداشت ', size=11)
run(p, 'WHO', size=11, bold=True)
run(p, ')، و مبانی علمی توسعه‌یافته در تعامل با دانشمندان ارشد دانشکده پزشکی هاروارد '
       'است. طراحی‌شده برای افراد پیشرویی که می‌دانند سلامت بیولوژیک، بزرگ‌ترین اهرم '
       'بازدهی و تصمیم‌گیری حرفه‌ای است.', size=11)


# ======================= صفحه ۳ =======================
H2('۶.', 'قدم بعدی: ارزیابی اصطکاک متابولیک')
body('به دلیل سطح تمرکز و نظارت مستقیم در پروتکل SMP، ظرفیت همراهی در هر کوارتر '
     'به ۱۰ نفر محدود است.')
body('اگر دو مورد یا بیشتر از چک‌لیست بالا وضعیت شما را توصیف می‌کند، کافی است شماره '
     'آن موارد را از راه‌های زیر بفرستید. ریشه اصطکاک متابولیک را در روتین روزمره‌تان '
     'تحلیل می‌کنم و یک راهکار اولیه قابل‌اجرا می‌فرستم.', after=8)

# کارتِ معرفیِ وب‌سایت
wt = doc.add_table(rows=1, cols=1)
fix_table(wt, [W_PAGE])
c = wt.rows[0].cells[0]; cell_shade(c, NAVY); no_space(c)
p = c.paragraphs[0]; rtl_para(p, AL.CENTER, space_after=1, space_before=7, line=1.2)
run(p, 'ارزیابی اولیه، پرونده‌های واقعی و شورای علمی را اینجا ببینید',
    size=12, bold=True, color='FFFFFF')
q = c.add_paragraph(); rtl_para(q, AL.CENTER, space_after=2, line=1.15)
run(q, 'saadatmehracademy.com', size=19, bold=True, color=GOLD_L)
r2 = c.add_paragraph(); rtl_para(r2, AL.CENTER, space_after=8, line=1.3)
run(r2, 'پروتکل ۱۰۰ روزه  ·  پرونده‌های تحلیلی  ·  کتاب معماری متابولیک  ·  '
        'فرم ارزیابی اولیه رایگان', size=10, color='C9D2DB')

gap(5)

CONTACT = [
    ('وب‌سایت',  'saadatmehracademy.com'),
    ('رایانامه', 'hossein@saadatmehracademy.com'),
    ('تلفن',     '۰۵۱۳۷۲۶۰۳۵۶'),
    ('همراه',    '۰۹۳۹۵۳۹۷۸۳۳'),
    ('نشانی',    'مشهد، بلوار خیام، خیام ۵، پلاک ۱۲'),
    ('ساعات',    'شنبه تا چهارشنبه، ۹ صبح تا ۴ بعدازظهر'),
]
WL, WV = Cm(2.2), Cm(6.4)
ct = doc.add_table(rows=3, cols=4)
fix_table(ct, [WL, WV, WL, WV])
table_rtl(ct)
for i, (lab, val) in enumerate(CONTACT):
    row, col = i // 2, (i % 2) * 2
    cl = ct.rows[row].cells[col]; no_space(cl); cell_valign(cl)
    p = cl.paragraphs[0]; rtl_para(p, AL.RIGHT, space_after=3, space_before=3, line=1.2)
    run(p, lab, size=10, bold=True, color=GOLD)
    cv = ct.rows[row].cells[col + 1]; no_space(cv); cell_valign(cv)
    p = cv.paragraphs[0]; rtl_para(p, AL.RIGHT, space_after=3, space_before=3, line=1.2)
    run(p, val, size=10, color=INK2)
    cell_borders(cl, bottom=(RULE, 4)); cell_borders(cv, bottom=(RULE, 4))

gap(4)

# عکس و امضا یک واحدند: keep_with_next تا هیچ‌وقت بینشان صفحه نشکند.
ep = doc.add_paragraph()
rtl_para(ep, AL.CENTER, space_after=3, space_before=0, line=1, keep_next=True)
ep.add_run().add_picture(os.path.join(IMG, 'team.jpg'), width=Cm(12.6))

sg = doc.add_paragraph(); rtl_para(sg, AL.CENTER, space_after=1, line=1.2, keep_next=True)
run(sg, 'حسین سعادت‌مهر  ·  یاسمن سعادت‌مهر', size=11.5, bold=True, color=NAVY)
sg2 = doc.add_paragraph(); rtl_para(sg2, AL.CENTER, space_after=2, line=1.25, keep_next=True)
run(sg2, 'Med Nutrition & Sports Medicine Specialist', size=9.5, color=GOLD, bold=True)
sg3 = doc.add_paragraph(); rtl_para(sg3, AL.CENTER, space_after=0, line=1.3)
run(sg3, 'این محتوا جایگزین تشخیص یا درمان پزشکی نیست.', size=8.5, color=INK3)

doc.save(OUT)
print('✓ ساخته شد:', OUT)
